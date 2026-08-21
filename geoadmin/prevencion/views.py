from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.serializers.json import DjangoJSONEncoder
from django.conf import settings
from django.template.loader import get_template
from openpyxl import Workbook
import json
import os
from django.views.decorators.csrf import csrf_exempt
import re
from datetime import datetime, time
from uuid import uuid4
from django.http import HttpResponseForbidden
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils.text import slugify
from django.utils.safestring import mark_safe
from urllib.parse import unquote, urlparse, urlencode
from .models import PrevencionPlantilla
from .forms import FormVigilanciaOpcion, ExportDataPrevencionForm
from core.models import Faena, Genero
from django.http import Http404, HttpResponse, JsonResponse, FileResponse
from django.core.files.storage import FileSystemStorage
from django.core.files.base import File
from xhtml2pdf import pisa
from core.utils import check_and_convert_pdf
from core.decorators import prevencion_riesgo_required, capacitaciones_y_docs_required, admin_or_base_datos_required, prevencion_mantenedor_required
from .models import (
    PrevencionNotificacionAprobacion,
    PrevencionNotificacionDocumento,
    PrevencionHistorialAutorizacion,
    PrevencionDocumento,
    PrevencionDocumentoDifusionTrabajador,
    PrevencionHistorialDifusionTrabajador,
    PrevencionResultadoCurso,
    PrevencionResultadoCursoIntento,
    PrevencionEvidenciaGeneral,
    PrevencionEvidenciaSeccion,
    PrevencionPermisoLlenado,
    VigilanciaArea,
    VigilanciaCargo,
    VigilanciaEvaluacionRiesgo,
    VigilanciaExposicion,
    VigilanciaGes,
    VigilanciaGradoExposicion,
    HigieneCualitativa,
    HigieneCuantitativa,
    Vigilancia,
    VigilanciaMedica,
    Egreso,
    EgresoAdjuntoHistorico,
    VigilanciaEstado,
    VigilanciaNivelRiesgo,
    VigilanciaNivelSeguimiento,
    VigilanciaTipoContrato,
    VigilanciaContrato,
    VigilanciaAdjunto,
    PrevencionEvidenciaItem,
    ValidacionDobleFactor,
)
from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.db.models import Max, Q, Case, When, Value, IntegerField
from django.db.models.fields.files import FieldFile
from user.models import UsuarioProfile, UserInformacionLaboral
from collections import defaultdict
User = get_user_model()
from messenger.utils import notify_group, get_changes_message


def _format_json_to_text(data, indent=0):
    # Convierte JSON (dict/list) en texto plano para exportarlo legible a Excel.
    if not data:
        return ""
    
    # Formato especifico para 'contenido' de prevencion
    if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict) and 'seccion' in data[0]:
        lines = []
        for sec in data:
            seccion_nombre = sec.get('seccion', 'Sin Nombre')
            lines.append(f"[{str(seccion_nombre).upper()}]")
            respuestas = sec.get('respuestas', [])
            if isinstance(respuestas, list):
                for resp in respuestas:
                    if not isinstance(resp, dict): continue
                    pregunta = resp.get('pregunta', '')
                    cumple = resp.get('cumple', '')
                    extras = resp.get('respuestas_extras', [])
                    
                    line = f"- {pregunta}: {cumple}%"
                    if extras and any(extras):
                        extras_str = ", ".join(str(e) for e in extras if e)
                        if extras_str:
                            line += f" ({extras_str})"
                    lines.append(line)
            lines.append("") # empty line between sections
        return "\n".join(lines).strip()
    
    # Generic format
    if isinstance(data, dict):
        lines = []
        for k, v in data.items():
            if isinstance(v, (dict, list)):
                lines.append(" " * indent + f"{k}:")
                lines.append(_format_json_to_text(v, indent + 2))
            else:
                lines.append(" " * indent + f"{k}: {v}")
        return "\n".join(lines)
    elif isinstance(data, list):
        lines = []
        for item in data:
            if isinstance(item, (dict, list)):
                lines.append(_format_json_to_text(item, indent + 2))
                lines.append(" " * indent + "-" * 20)
            else:
                lines.append(" " * indent + f"- {item}")
        return "\n".join(lines)
    else:
        return " " * indent + str(data)

def _serializar_modelo_crudo(instancia, base_url=None):
    # Serializa un registro Django en columnas simples para escribirlo en Excel.
    fila = {}
    for field in instancia._meta.concrete_fields:
        nombre_campo = field.attname if field.is_relation else field.name
        valor = getattr(instancia, nombre_campo)

        # En Documentos, exporta nombres legibles para llaves foráneas clave.
        if isinstance(instancia, PrevencionDocumento) and field.is_relation:
            if field.name == 'plantilla_base':
                nombre_campo = 'plantilla_base'
                valor = str(instancia.plantilla_base) if instancia.plantilla_base else ''
            elif field.name == 'faena':
                nombre_campo = 'faena'
                valor = str(instancia.faena) if instancia.faena else ''
        elif isinstance(instancia, VigilanciaMedica) and field.is_relation:
            # En Vigilancia, no exporta user_id y muestra nombre de faena.
            if field.name == 'user':
                continue
            if field.name == 'faena':
                nombre_campo = 'faena'
                valor = str(instancia.faena) if instancia.faena else ''

        if isinstance(valor, datetime):
            if timezone.is_aware(valor):
                valor = timezone.localtime(valor)
            valor = valor.strftime('%Y-%m-%d %H:%M:%S')
        elif isinstance(valor, FieldFile):
            valor = _nombre_y_url_archivo(valor, base_url=base_url) if valor and valor.name else ''
        elif hasattr(valor, 'isoformat') and valor is not None:
            valor = valor.isoformat()
        elif isinstance(valor, (dict, list)):
            valor = _format_json_to_text(valor)

        # Manejar multiples adjuntos para Documento
        if nombre_campo == 'evidencia_general' and hasattr(instancia, 'evidencias_generales'):
            extra = []
            if valor:
                extra.append(valor)
            for evidencia in instancia.evidencias_generales.all():
                if evidencia.archivo and evidencia.archivo.name:
                    extra.append(_nombre_y_url_archivo(evidencia.archivo, base_url=base_url))
            valor = "\n".join(extra) if extra else ""

        # Manejar multiples adjuntos para Vigilancia Medica
        if nombre_campo == 'documento_adjunto' and hasattr(instancia, 'adjuntos_vigilancia'):
            extra = []
            if valor:
                extra.append(valor)
            for adjunto in instancia.adjuntos_vigilancia.all():
                if adjunto.archivo and adjunto.archivo.name:
                    extra.append(_nombre_y_url_archivo(adjunto.archivo, base_url=base_url))
            valor = "\n".join(extra) if extra else ""

        fila[nombre_campo] = valor
    if isinstance(instancia, VigilanciaMedica):
        # En exportación de Vigilancia, deja los adjuntos al final.
        if 'documento_adjunto' in fila:
            fila['archivos_adjuntos'] = fila.pop('documento_adjunto')
        fila = _ordenar_y_filtrar_columnas_vigilancia(fila)

    return fila


def _normalizar_url_media(url, base_url=None):
    url = str(url or '').strip()
    if not url:
        return ''
    if url.startswith('/') and base_url:
        return f"{str(base_url).rstrip('/')}{url}"
    return url


def _nombre_y_url_archivo(archivo_field, base_url=None):
    if not archivo_field:
        return ''
    nombre = os.path.basename(getattr(archivo_field, 'name', '') or '') or 'archivo'
    url = _normalizar_url_media(getattr(archivo_field, 'url', '') or '', base_url=base_url)
    return f"{nombre} ({url})" if url else nombre


def _es_columna_evidencia_o_adjunto(nombre_columna):
    key = str(nombre_columna or '').strip().lower().replace(' ', '_')
    return key in {
        'evidencia_general',
        'evidencia_por_item',
        'evidencia_seccion',
        'archivos_adjuntos',
        'documento_adjunto',
    } or key.startswith('adjunto_') or key.startswith('evidencia_por_item_') or key.startswith('evidencia_seccion_') or key.startswith('evidencia_general_')


def _extraer_nombre_y_url_desde_texto(valor):
    texto = str(valor or '').strip()
    if not texto or '\n' in texto:
        return None, None
    if texto.startswith(('http://', 'https://', '/')):
        return texto, texto
    # Formato esperado: "archivo.ext (url)"
    idx = texto.rfind(' (')
    if idx == -1:
        idx = texto.rfind('(')
    if idx == -1 or not texto.endswith(')'):
        return None, None
    nombre = texto[:idx].strip()
    url = texto[idx + 1:-1].strip().lstrip('(').strip()
    if not url.startswith(('http://', 'https://', '/')):
        return None, None
    return (nombre or url), url


def _limpiar_texto_multilinea_evidencias(valor):
    # Si hay varios archivos en la misma celda, deja solo nombres (sin URL larga).
    lineas = str(valor or '').splitlines()
    nuevas = []
    for linea in lineas:
        nombre, _ = _extraer_nombre_y_url_desde_texto(linea.strip())
        nuevas.append(nombre or linea.strip())
    return "\n".join([l for l in nuevas if l])


def _aplicar_hipervinculo_si_corresponde(cell, nombre_columna, valor, link_font):
    if not _es_columna_evidencia_o_adjunto(nombre_columna):
        return
    if '\n' in str(valor or ''):
        cell.value = _limpiar_texto_multilinea_evidencias(valor)
        return
    nombre, url = _extraer_nombre_y_url_desde_texto(valor)
    if not url:
        return
    cell.value = nombre
    cell.hyperlink = url
    cell.font = link_font


def _formatear_titulo_columna(texto):
    if str(texto or '').startswith('adjunto_'):
        numero = str(texto).split('_')[-1]
        return f"Adjunto {numero}"
    aliases = {
        'created_at': 'Fecha de creacion',
        'antiguedad_anos': 'Antiguedad años',
        'antiguedad_meses': 'Antiguedad meses',
        'antiguedad_dias': 'Antiguedad dias',
        'rut_trabajador': 'RUT',
        'nombre_completo': 'Nombre completo',
        'genero_texto': 'Genero',
        'fecha_nacimiento': 'Fecha nacimiento',
        'fecha_ingreso': 'Fecha ingreso',
        'fecha_retiro': 'Fecha retiro / desvinculacion',
        'tipo_contrato': 'Tipo de contrato',
        'ges': 'GES',
        'silice_eval_riesgo': 'Silice evaluacion de riesgo',
        'silice_nivel_riesgo': 'Silice nivel de riesgo',
        'silice_grado_expo': 'Silice grado de exposicion',
        'silice_fecha_radio': 'Silice fecha radiografia',
        'silice_fecha_vence': 'Silice fecha vencimiento',
        'silice_vigencia': 'Silice vigencia',
        'silice_obs': 'Silice observaciones',
        'ruido_eval_riesgo': 'Ruido evaluacion de riesgo',
        'ruido_nivel_seguimiento': 'Ruido nivel de seguimiento',
        'ruido_grado_expo': 'Ruido grado de exposicion',
        'ruido_fecha_audio': 'Ruido fecha audiometria',
        'ruido_fecha_vence': 'Ruido fecha vencimiento',
        'ruido_vigencia': 'Ruido vigencia',
        'ruido_obs': 'Ruido observaciones',
        'hipo_exposicion': 'Hipobaria exposicion',
        'hipo_fecha_hemo': 'Hipobaria fecha hemoglobina',
        'hipo_fecha_vence': 'Hipobaria fecha vencimiento',
        'hipo_vigencia': 'Hipobaria vigencia',
        'hipo_obs': 'Hipobaria observaciones',
        'vibra_eval_riesgo': 'Vibraciones evaluacion del riesgo',
        'vibra_exposicion': 'Vibraciones exposicion',
        'rad_exposicion': 'Radiaciones ionizantes exposicion',
        'humos_eval_riesgo': 'Humos metalicos evaluacion del riesgo',
        'humos_exposicion': 'Humos metalicos exposicion',
        'archivos_adjuntos': 'Archivos adjuntos',
    }
    if texto in aliases:
        return aliases[texto]
    texto_limpio = str(texto or '').replace('_', ' ').strip()
    if not texto_limpio:
        return ''
    return texto_limpio[0].upper() + texto_limpio[1:]


def _ordenar_y_filtrar_columnas_vigilancia(fila):
    # Exporta solo columnas funcionales y en el orden de la ficha.
    columnas = [
        'created_at',
        'creado_por',
        'fecha_incidente',
        'rut_trabajador',
        'nombre_completo',
        'genero_texto',
        'fecha_nacimiento',
        'edad',
        'faena',
        'area',
        'ges',
        'cargo',
        'tipo_contrato',
        'fecha_ingreso',
        'antiguedad_anos',
        'antiguedad_meses',
        'antiguedad_dias',
        'fecha_retiro',
        'silice_eval_riesgo',
        'silice_nivel_riesgo',
        'silice_grado_expo',
        'silice_fecha_radio',
        'silice_fecha_vence',
        'silice_vigencia',
        'silice_obs',
        'ruido_eval_riesgo',
        'ruido_nivel_seguimiento',
        'ruido_grado_expo',
        'ruido_fecha_audio',
        'ruido_fecha_vence',
        'ruido_vigencia',
        'ruido_obs',
        'hipo_exposicion',
        'hipo_fecha_hemo',
        'hipo_fecha_vence',
        'hipo_vigencia',
        'hipo_obs',
        'vibra_eval_riesgo',
        'vibra_exposicion',
        'rad_exposicion',
        'humos_eval_riesgo',
        'humos_exposicion',
        'archivos_adjuntos',
    ]
    return {col: fila.get(col, '') for col in columnas}


def _expandir_adjuntos_vigilancia(filas):
    # Convierte "archivos_adjuntos" en columnas: adjunto_1, adjunto_2, ...
    max_adjuntos = 0
    listas_adjuntos = []
    for fila in filas:
        crudo = str(fila.get('archivos_adjuntos', '') or '')
        adjuntos = [linea.strip() for linea in crudo.splitlines() if linea.strip()]
        listas_adjuntos.append(adjuntos)
        if len(adjuntos) > max_adjuntos:
            max_adjuntos = len(adjuntos)

    if max_adjuntos == 0:
        return filas

    filas_expandida = []
    for fila, adjuntos in zip(filas, listas_adjuntos):
        nueva_fila = {k: v for k, v in fila.items() if k != 'archivos_adjuntos'}
        for idx in range(max_adjuntos):
            nueva_fila[f'adjunto_{idx + 1}'] = adjuntos[idx] if idx < len(adjuntos) else ''
        filas_expandida.append(nueva_fila)
    return filas_expandida


def _normalizar_columnas_extras_para_excel(extras):
    # Convierte respuestas_extras a lista de celdas.
    # Si llega texto con separador " - ", lo separa en columnas distintas.
    if not isinstance(extras, list):
        return [str(extras or '').strip()] if str(extras or '').strip() else []

    resultado = []
    for valor in extras:
        texto = str(valor or '').strip()
        if not texto:
            continue
        if ' - ' in texto:
            partes = [p.strip() for p in texto.split(' - ') if p.strip()]
            if partes:
                resultado.extend(partes)
                continue
        resultado.append(texto)
    return resultado


def _construir_filas_documentos_export(documentos):
    filas = []

    def fila_ordenada(base, seccion='', porcentaje='', items='', evidencia_item='', columnas='', evidencia_seccion=''):
        return {
            'creador': base.get('creador', ''),
            'Fecha Creacion': base.get('Fecha Creacion', ''),
            'faena': base.get('faena', ''),
            'plantilla_base': base.get('plantilla_base', ''),
            'seccion': seccion,
            'porcentaje': porcentaje,
            'items': items,
            'evidencia por item': evidencia_item,
            'columnas': columnas,
            'evidencia seccion': evidencia_seccion,
            'observacion_general': base.get('observacion_general', ''),
            'evidencia general': base.get('evidencia general', ''),
        }

    for documento in documentos:
        contenido = documento.contenido if isinstance(documento.contenido, list) else []

        evidencias_seccion_map = defaultdict(list)
        for evidencia in documento.evidencias_secciones.all():
            evidencias_seccion_map[evidencia.indice_seccion].append(_nombre_y_url_archivo(evidencia.archivo))

        evidencias_item_map = defaultdict(list)
        for evidencia in documento.evidencias_items.all():
            evidencias_item_map[(evidencia.indice_seccion, evidencia.indice_fila)].append(
                _nombre_y_url_archivo(evidencia.archivo)
            )

        evidencias_generales = [_nombre_y_url_archivo(evidencia.archivo) for evidencia in documento.evidencias_generales.all()]
        if not evidencias_generales and documento.evidencia_general:
            evidencias_generales = [_nombre_y_url_archivo(documento.evidencia_general)]
        evidencia_general_texto = ' | '.join([e for e in evidencias_generales if e])

        base = {
            'Fecha Creacion': timezone.localtime(documento.fecha_creacion).strftime('%Y-%m-%d %H:%M:%S') if documento.fecha_creacion else '',
            'faena': str(documento.faena) if documento.faena else '',
            'plantilla_base': str(documento.plantilla_base) if documento.plantilla_base else '',
            'observacion_general': documento.observacion_general or '',
            'evidencia general': evidencia_general_texto,
            'creador': documento.creador or '',
        }

        if not contenido:
            filas.append(fila_ordenada(base))
            continue

        for sec_idx, seccion in enumerate(contenido):
            respuestas = seccion.get('respuestas') if isinstance(seccion, dict) else []
            nombre_seccion = (seccion or {}).get('seccion', '') if isinstance(seccion, dict) else ''
            evidencia_seccion_texto = ' | '.join([e for e in evidencias_seccion_map.get(sec_idx, []) if e])

            if not respuestas:
                filas.append(
                    fila_ordenada(
                        base,
                        seccion=nombre_seccion,
                        evidencia_seccion=evidencia_seccion_texto,
                    )
                )
                continue

            for row_idx, respuesta in enumerate(respuestas):
                columnas_extras = respuesta.get('respuestas_extras', []) if isinstance(respuesta, dict) else []
                encabezados_columnas = seccion.get('columnas_extras', []) if isinstance(seccion, dict) else []
                columnas_texto = ''
                if isinstance(columnas_extras, list):
                    pares_columnas = []
                    max_len = max(len(encabezados_columnas), len(columnas_extras))
                    for i in range(max_len):
                        encabezado = ''
                        if i < len(encabezados_columnas):
                            encabezado = str(encabezados_columnas[i] or '').strip()
                        valor_columna = ''
                        if i < len(columnas_extras):
                            valor_columna = str(columnas_extras[i] or '').strip()

                        if not valor_columna:
                            valor_columna = '-'

                        if encabezado:
                            pares_columnas.append(f"{encabezado}: {valor_columna}")
                        else:
                            pares_columnas.append(valor_columna)
                    columnas_texto = ' | '.join(pares_columnas)
                else:
                    columnas_texto = str(columnas_extras or '-')

                evidencia_item_texto = ' | '.join([
                    e for e in evidencias_item_map.get((sec_idx, row_idx), []) if e
                ])

                filas.append(
                    fila_ordenada(
                        base,
                        seccion=nombre_seccion,
                        porcentaje=(respuesta or {}).get('cumple', '') if isinstance(respuesta, dict) else '',
                        items=(respuesta or {}).get('pregunta', '') if isinstance(respuesta, dict) else '',
                        evidencia_item=evidencia_item_texto,
                        columnas=columnas_texto,
                        evidencia_seccion=evidencia_seccion_texto,
                    )
                )

    return filas


def _escribir_hoja_desde_filas(hoja, filas):
    # Escribe encabezados y filas en una hoja Excel, con formato básico.
    if not filas:
        hoja.append(['Sin datos para el rango seleccionado'])
        return

    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    encabezados_clave = list(filas[0].keys())
    encabezados_display = [_formatear_titulo_columna(columna) for columna in encabezados_clave]
    hoja.append(encabezados_display)

    header_font = Font(bold=True)
    link_font = Font(color="0000EE", underline="single")
    # Sin relleno para usar el fondo blanco por defecto de Excel.
    base_fill = PatternFill(fill_type=None)
    border_style = Side(border_style="thin", color="000000")
    border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)
    
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap_alignment = Alignment(wrap_text=True, horizontal="left", vertical="top")
    
    for cell in hoja[1]:
        cell.font = header_font
        cell.alignment = header_alignment
        cell.fill = base_fill
        cell.border = border
    hoja.row_dimensions[1].height = 36

    for fila_idx, fila in enumerate(filas, start=2):
        row_data = [fila.get(columna, '') for columna in encabezados_clave]
        hoja.append(row_data)
        for col_idx, col_name in enumerate(encabezados_clave, start=1):
            cell = hoja.cell(row=fila_idx, column=col_idx)
            _aplicar_hipervinculo_si_corresponde(
                cell,
                col_name,
                fila.get(col_name, ''),
                link_font
            )

    for row in hoja.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_alignment
            cell.fill = base_fill
            cell.border = border

    for col_num, col_name in enumerate(encabezados_clave, 1):
        col_letter = get_column_letter(col_num)
        if col_name in ['contenido', 'evidencia_general', 'observacion_general']:
            hoja.column_dimensions[col_letter].width = 60
        elif col_name in ['archivos_adjuntos', 'documento_adjunto'] or str(col_name).startswith('adjunto_'):
            hoja.column_dimensions[col_letter].width = 55
        elif col_name in ['evidencia por item', 'evidencia seccion', 'evidencia general']:
            hoja.column_dimensions[col_letter].width = 55
        elif col_name in ['silice_obs', 'ruido_obs', 'hipo_obs']:
            hoja.column_dimensions[col_letter].width = 35
        elif col_name in ['ruido_nivel_seguimiento', 'silice_nivel_riesgo']:
            hoja.column_dimensions[col_letter].width = 30
        elif col_name in [
            'vibra_eval_riesgo',
            'vibra_exposicion',
            'rad_exposicion',
            'humos_eval_riesgo',
            'humos_exposicion',
        ]:
            hoja.column_dimensions[col_letter].width = 34
        elif col_name == 'archivos_adjuntos':
            hoja.column_dimensions[col_letter].width = 45
        else:
            # Aumenta ancho para encabezados largos para evitar cortes en Excel.
            titulo = _formatear_titulo_columna(col_name)
            if len(titulo) >= 26:
                hoja.column_dimensions[col_letter].width = 34
            elif len(titulo) >= 20:
                hoja.column_dimensions[col_letter].width = 28
            else:
                hoja.column_dimensions[col_letter].width = 20


def _escribir_hoja_documentos_jerarquico(hoja, documentos_qs, base_url=None):
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    
    header_font = Font(bold=True)
    link_font = Font(color="0000EE", underline="single")
    # Sin relleno para usar el fondo blanco por defecto de Excel.
    base_fill = PatternFill(fill_type=None)
    
    border_style = Side(border_style="thin", color="000000")
    border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)
    alignment_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    alignment_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    documentos = list(documentos_qs)
    columnas_extras_headers = []
    max_ev_item = 0
    max_ev_sec = 0
    max_ev_gral = 0
    for doc_tmp in documentos:
        ev_gral_count = (1 if doc_tmp.evidencia_general else 0) + doc_tmp.evidencias_generales.count()
        max_ev_gral = max(max_ev_gral, ev_gral_count)

        sec_counter = {}
        for es_tmp in doc_tmp.evidencias_secciones.all():
            sec_counter[es_tmp.indice_seccion] = sec_counter.get(es_tmp.indice_seccion, 0) + 1
        if sec_counter:
            max_ev_sec = max(max_ev_sec, max(sec_counter.values()))

        item_counter = {}
        for ei_tmp in doc_tmp.evidencias_items.all():
            key_tmp = (ei_tmp.indice_seccion, ei_tmp.indice_fila)
            item_counter[key_tmp] = item_counter.get(key_tmp, 0) + 1
        if item_counter:
            max_ev_item = max(max_ev_item, max(item_counter.values()))

        contenido_tmp = doc_tmp.contenido if isinstance(doc_tmp.contenido, list) else []
        for sec_tmp in contenido_tmp:
            if not isinstance(sec_tmp, dict):
                continue
            respuestas_tmp = sec_tmp.get('respuestas', [])
            respuestas_tmp = respuestas_tmp if isinstance(respuestas_tmp, list) else []

            for resp_tmp in respuestas_tmp:
                if not isinstance(resp_tmp, dict):
                    continue
                extras_tmp = _normalizar_columnas_extras_para_excel(resp_tmp.get('respuestas_extras', []))
                for idx_extra, _ in enumerate(extras_tmp):
                    while len(columnas_extras_headers) <= idx_extra:
                        columnas_extras_headers.append('columnas')

    # Mantiene al menos 1 subcolumna para "columnas".
    while len(columnas_extras_headers) < 1:
        columnas_extras_headers.append('columnas')

    max_ev_item = max(1, max_ev_item)
    max_ev_sec = max(1, max_ev_sec)
    max_ev_gral = max(1, max_ev_gral)
    headers_ev_item = [f"evidencia por item {i}" for i in range(1, max_ev_item + 1)]
    headers_ev_sec = [f"evidencia seccion {i}" for i in range(1, max_ev_sec + 1)]
    headers_ev_gral = [f"evidencia general {i}" for i in range(1, max_ev_gral + 1)]

    encabezados_columnas_extras = []
    if len(columnas_extras_headers) > 0:
        encabezados_columnas_extras = ["columnas"] + ([""] * (len(columnas_extras_headers) - 1))

    encabezados_base = [
        "creador",
        "Fecha Creacion",
        "faena",
        "plantilla_base",
        "seccion",
        "porcentaje",
        "items",
    ] + headers_ev_item + encabezados_columnas_extras + headers_ev_sec + [
        "observacion_general",
    ] + headers_ev_gral
    encabezados = [_formatear_titulo_columna(columna) for columna in encabezados_base]
    total_columnas = len(encabezados_base)
    total_columnas_extras = len(columnas_extras_headers)
    col_inicio_extras = 8 + len(headers_ev_item)
    col_fin_extras = col_inicio_extras + total_columnas_extras - 1
    col_inicio_evidencia_seccion = col_inicio_extras + total_columnas_extras
    col_observacion_general = col_inicio_evidencia_seccion + len(headers_ev_sec)
    col_inicio_evidencia_general = col_observacion_general + 1
    hoja.append(encabezados)

    for col_idx, cell in enumerate(hoja[1], 1):
        cell.font = header_font
        cell.alignment = alignment_center
        cell.fill = base_fill
        cell.border = border

    if total_columnas_extras > 1:
        hoja.merge_cells(start_row=1, start_column=col_inicio_extras, end_row=1, end_column=col_fin_extras)
        
    current_row = 2

    if not documentos:
        hoja.append(['Sin datos para el rango seleccionado'])
        return

    def _ajustar_recuadros_columnas_por_fila(row_idx, cantidad_fila):
        # No fusiona celdas en filas para mantener recuadros por columna.
        return

    for doc in documentos:
        start_row_doc = current_row
        
        plantilla_nombre = str(doc.plantilla_base.nombre) if doc.plantilla_base else 'Sin Plantilla'
        fecha_creacion = timezone.localtime(doc.fecha_creacion).strftime('%Y-%m-%d %H:%M:%S') if doc.fecha_creacion else ''
        faena_nombre = str(doc.faena.faena) if doc.faena else ''
        obs_gral = doc.observacion_general or ''
        creador = doc.creador or ''
        
        ev_gral_list = []
        if doc.evidencia_general:
            ev_gral_list.append(_nombre_y_url_archivo(doc.evidencia_general, base_url=base_url))
        for eg in doc.evidencias_generales.all():
            if eg.archivo and eg.archivo.name:
                ev_gral_list.append(_nombre_y_url_archivo(eg.archivo, base_url=base_url))
        ev_sec_map = {}
        for es in doc.evidencias_secciones.all():
            if es.indice_seccion not in ev_sec_map:
                ev_sec_map[es.indice_seccion] = []
            if es.archivo and es.archivo.name:
                ev_sec_map[es.indice_seccion].append(_nombre_y_url_archivo(es.archivo, base_url=base_url))
            
        ev_item_map = {}
        for ei in doc.evidencias_items.all():
            key = f"{ei.indice_seccion}_{ei.indice_fila}"
            if key not in ev_item_map:
                ev_item_map[key] = []
            if ei.archivo and ei.archivo.name:
                ev_item_map[key].append(_nombre_y_url_archivo(ei.archivo, base_url=base_url))

        contenido = doc.contenido
        if not isinstance(contenido, list):
            ev_item_cells = [''] * len(headers_ev_item)
            ev_sec_cells = [''] * len(headers_ev_sec)
            ev_gral_cells = ev_gral_list[:len(headers_ev_gral)] + [''] * max(0, len(headers_ev_gral) - len(ev_gral_list))
            fila = [creador, fecha_creacion, faena_nombre, plantilla_nombre, "", "", ""]
            fila.extend(ev_item_cells)
            fila.extend([''] * len(columnas_extras_headers))
            fila.extend(ev_sec_cells)
            fila.extend([obs_gral])
            fila.extend(ev_gral_cells)
            hoja.append(fila)
            for col_idx in range(1, total_columnas + 1):
                cell = hoja.cell(row=current_row, column=col_idx)
                cell.alignment = alignment_left
                cell.fill = base_fill
                cell.border = border
                _aplicar_hipervinculo_si_corresponde(
                    cell,
                    encabezados_base[col_idx - 1],
                    fila[col_idx - 1],
                    link_font
                )
            current_row += 1
            continue

        for sec_idx, sec in enumerate(contenido):
            start_row_sec = current_row
            sec_nombre = sec.get('seccion', '')
            columnas_seccion = sec.get('columnas_extras', []) if isinstance(sec, dict) else []
            columnas_seccion_count = len(columnas_seccion) if isinstance(columnas_seccion, list) else 0
            
            ev_sec_list = ev_sec_map.get(sec_idx, [])
            ev_sec_cells = ev_sec_list[:len(headers_ev_sec)] + [''] * max(0, len(headers_ev_sec) - len(ev_sec_list))
            
            respuestas = sec.get('respuestas', [])
            if not isinstance(respuestas, list) or len(respuestas) == 0:
                ev_item_cells = [''] * len(headers_ev_item)
                ev_gral_cells = ev_gral_list[:len(headers_ev_gral)] + [''] * max(0, len(headers_ev_gral) - len(ev_gral_list))
                fila = [creador, fecha_creacion, faena_nombre, plantilla_nombre, sec_nombre, "", ""]
                fila.extend(ev_item_cells)
                fila.extend([''] * len(columnas_extras_headers))
                fila.extend(ev_sec_cells)
                fila.extend([obs_gral])
                fila.extend(ev_gral_cells)
                hoja.append(fila)
                for col_idx in range(1, total_columnas + 1):
                    cell = hoja.cell(row=current_row, column=col_idx)
                    cell.alignment = alignment_center if col_idx == 6 else alignment_left
                    cell.fill = base_fill
                    cell.border = border
                    _aplicar_hipervinculo_si_corresponde(
                        cell,
                        encabezados_base[col_idx - 1],
                        fila[col_idx - 1],
                        link_font
                    )
                _ajustar_recuadros_columnas_por_fila(current_row, columnas_seccion_count)
                current_row += 1
            else:
                for row_idx, resp in enumerate(respuestas):
                    if not isinstance(resp, dict):
                        continue
                    pregunta = resp.get('pregunta', '')
                    porcentaje = resp.get('cumple', '')
                    extras = _normalizar_columnas_extras_para_excel(resp.get('respuestas_extras', []))
                    
                    ev_item_key = f"{sec_idx}_{row_idx}"
                    ev_item_list = ev_item_map.get(ev_item_key, [])
                    ev_item_cells = ev_item_list[:len(headers_ev_item)] + [''] * max(0, len(headers_ev_item) - len(ev_item_list))
                    ev_gral_cells = ev_gral_list[:len(headers_ev_gral)] + [''] * max(0, len(headers_ev_gral) - len(ev_gral_list))

                    extras_celdas = [str(e or '') for e in extras]
                    if len(extras_celdas) < len(columnas_extras_headers):
                        extras_celdas.extend([''] * (len(columnas_extras_headers) - len(extras_celdas)))
                    elif len(extras_celdas) > len(columnas_extras_headers):
                        extras_celdas = extras_celdas[:len(columnas_extras_headers)]

                    fila = [
                        creador,
                        fecha_creacion,
                        faena_nombre,
                        plantilla_nombre,
                        sec_nombre,
                        porcentaje,
                        pregunta,
                    ] + ev_item_cells + extras_celdas + ev_sec_cells + [obs_gral] + ev_gral_cells
                    hoja.append(fila)
                    for col_idx in range(1, total_columnas + 1):
                        cell = hoja.cell(row=current_row, column=col_idx)
                        cell.alignment = alignment_center if col_idx == 6 else alignment_left
                        cell.fill = base_fill
                        cell.border = border
                        _aplicar_hipervinculo_si_corresponde(
                            cell,
                            encabezados_base[col_idx - 1],
                            fila[col_idx - 1],
                            link_font
                        )
                    _ajustar_recuadros_columnas_por_fila(
                        current_row,
                        len(extras)
                    )
                    current_row += 1
                    
            if current_row - 1 > start_row_sec:
                hoja.merge_cells(start_row=start_row_sec, start_column=5, end_row=current_row-1, end_column=5) # seccion
                for col_idx in range(col_inicio_evidencia_seccion, col_inicio_evidencia_seccion + len(headers_ev_sec)):
                    hoja.merge_cells(
                        start_row=start_row_sec,
                        start_column=col_idx,
                        end_row=current_row-1,
                        end_column=col_idx
                    ) # evidencia seccion

        if current_row - 1 > start_row_doc:
            hoja.merge_cells(start_row=start_row_doc, start_column=1, end_row=current_row-1, end_column=1) # creador
            hoja.merge_cells(start_row=start_row_doc, start_column=2, end_row=current_row-1, end_column=2) # Fecha Creacion
            hoja.merge_cells(start_row=start_row_doc, start_column=3, end_row=current_row-1, end_column=3) # faena
            hoja.merge_cells(start_row=start_row_doc, start_column=4, end_row=current_row-1, end_column=4) # plantilla_base
            hoja.merge_cells(
                start_row=start_row_doc,
                start_column=col_observacion_general,
                end_row=current_row-1,
                end_column=col_observacion_general
            ) # observacion_general
            for col_idx in range(col_inicio_evidencia_general, col_inicio_evidencia_general + len(headers_ev_gral)):
                hoja.merge_cells(
                    start_row=start_row_doc,
                    start_column=col_idx,
                    end_row=current_row-1,
                    end_column=col_idx
                ) # evidencia general

    for col_idx, col_name in enumerate(encabezados_base, 1):
        col_letter = get_column_letter(col_idx)
        if col_name in ('creador', 'Fecha Creacion', 'faena'):
            width = 20
        elif col_name == 'plantilla_base':
            width = 25
        elif col_name == 'seccion':
            width = 20
        elif col_name == 'porcentaje':
            width = 12
        elif col_name == 'items':
            width = 35
        elif str(col_name).startswith('evidencia por item') or str(col_name).startswith('evidencia seccion'):
            width = 55
        elif col_name == 'observacion_general' or str(col_name).startswith('evidencia general'):
            width = 55
        else:
            width = 24
        hoja.column_dimensions[col_letter].width = width


def _get_usuario_profile(user):
    return UsuarioProfile.objects.select_related('faena').filter(user=user).first()


def _normalizar_fecha_para_usuario_profile(fecha_raw):
    fecha = parse_date(str(fecha_raw or '').strip())
    if not fecha:
        return None

    fecha_dt = datetime.combine(fecha, time.min)
    if timezone.is_naive(fecha_dt):
        fecha_dt = timezone.make_aware(fecha_dt, timezone.get_current_timezone())
    return fecha_dt


def _calcular_edad_desde_fecha(fecha_raw):
    # Recalcula edad en backend para no depender del valor enviado por el frontend.
    fecha = fecha_raw if hasattr(fecha_raw, 'year') else parse_date(str(fecha_raw or '').strip())
    if not fecha:
        return None

    hoy = timezone.localdate()
    edad = hoy.year - fecha.year
    if (hoy.month, hoy.day) < (fecha.month, fecha.day):
        edad -= 1
    return max(0, edad)


def _fecha_usuario_profile_como_date(perfil):
    if not perfil or not perfil.fechaNacimiento:
        return None

    fecha = perfil.fechaNacimiento
    if timezone.is_aware(fecha):
        fecha = timezone.localtime(fecha)
    return fecha.date()


def _resolver_genero_por_texto(genero_texto):
    genero_limpio = str(genero_texto or '').strip()
    if not genero_limpio:
        return None
    return Genero.objects.filter(genero__iexact=genero_limpio, status=True).first()


def _sincronizar_datos_personales_usuario(user, genero_id=None, fecha_nacimiento_raw=None, sobrescribir=False):
    if not user:
        return

    perfil, _ = UsuarioProfile.objects.get_or_create(user=user)
    update_fields = []

    if genero_id and (sobrescribir or not perfil.genero_id) and str(perfil.genero_id or '') != str(genero_id):
        perfil.genero_id = genero_id
        update_fields.append('genero')

    fecha_nacimiento_dt = _normalizar_fecha_para_usuario_profile(fecha_nacimiento_raw)
    if fecha_nacimiento_dt and (sobrescribir or not perfil.fechaNacimiento):
        fecha_actual = _fecha_usuario_profile_como_date(perfil)
        if fecha_actual != fecha_nacimiento_dt.date():
            perfil.fechaNacimiento = fecha_nacimiento_dt
            update_fields.append('fechaNacimiento')

    if update_fields:
        perfil.save(update_fields=update_fields)


def _obtener_datos_personales_sincronizados(vigilancia):
    perfil = _get_usuario_profile(vigilancia.user) if vigilancia.user_id else None
    perfil_genero = perfil.genero if perfil and perfil.genero_id else None
    perfil_fecha = _fecha_usuario_profile_como_date(perfil)

    genero_ficha = _resolver_genero_por_texto(vigilancia.genero_texto)
    fecha_ficha = vigilancia.fecha_nacimiento

    genero_obj = perfil_genero or genero_ficha
    fecha_nacimiento = perfil_fecha or fecha_ficha
    update_fields_vigilancia = []

    if perfil_genero and vigilancia.genero_texto != perfil_genero.genero:
        vigilancia.genero_texto = perfil_genero.genero
        update_fields_vigilancia.append('genero_texto')

    if perfil_fecha and vigilancia.fecha_nacimiento != perfil_fecha:
        vigilancia.fecha_nacimiento = perfil_fecha
        update_fields_vigilancia.append('fecha_nacimiento')

    edad_calculada = _calcular_edad_desde_fecha(fecha_nacimiento)
    edad_objetivo = edad_calculada if edad_calculada is not None else 0
    if (vigilancia.edad or 0) != edad_objetivo:
        vigilancia.edad = edad_objetivo
        update_fields_vigilancia.append('edad')

    if update_fields_vigilancia:
        vigilancia.save(update_fields=update_fields_vigilancia)

    if vigilancia.user_id:
        if not perfil_genero and genero_ficha:
            _sincronizar_datos_personales_usuario(vigilancia.user, genero_id=genero_ficha.id, sobrescribir=False)
        if not perfil_fecha and fecha_ficha:
            _sincronizar_datos_personales_usuario(vigilancia.user, fecha_nacimiento_raw=fecha_ficha.isoformat(), sobrescribir=False)

    return {
        'genero_id': genero_obj.id if genero_obj else '',
        'genero_texto': genero_obj.genero if genero_obj else '',
        'fecha_nacimiento': fecha_nacimiento,
    }


def _is_sin_asignar(usuario_profile):
    return bool(
        usuario_profile
        and usuario_profile.faena
        and str(usuario_profile.faena.faena).strip().upper() == "SIN ASIGNAR"
    )


def _faenas_visibles_para_usuario(user):
    faenas_activas = Faena.objects.filter(status=True)
    usuario_profile = _get_usuario_profile(user)

    if _is_sin_asignar(usuario_profile):
        return faenas_activas

    if usuario_profile and usuario_profile.faena_id:
        return faenas_activas.filter(id=usuario_profile.faena_id)

    return faenas_activas.none()


def _faenas_seleccionables_para_usuario(user):
    return _faenas_visibles_para_usuario(user).exclude(faena__iexact='SIN ASIGNAR')


def _filtrar_queryset_por_faena_usuario(queryset, user, campo_faena='faena'):
    usuario_profile = _get_usuario_profile(user)

    if _is_sin_asignar(usuario_profile):
        return queryset

    if usuario_profile and usuario_profile.faena_id:
        return queryset.filter(**{f'{campo_faena}_id': usuario_profile.faena_id})

    return queryset.none()


def _usuario_puede_ver_todas_faenas(user):
    return _is_sin_asignar(_get_usuario_profile(user))


def _ids_documentos_visibles_para_usuario(user, tipo_slug=''):
    # Para capacitaciones, la visibilidad se determina por difusión activa al usuario.
    tipo_slug_norm = str(tipo_slug or '').strip().lower()

    filtros_difusion = {
        'status': True,
        'user': user,
        'documento__status': True,
    }
    if tipo_slug_norm:
        filtros_difusion['documento__plantilla_base__tipo'] = tipo_slug_norm
    return set(
        PrevencionDocumentoDifusionTrabajador.objects
        .filter(**filtros_difusion)
        .values_list('documento_id', flat=True)
    )


def _faena_permitida_para_usuario(user, faena_id):
    if not faena_id:
        return False

    usuario_profile = _get_usuario_profile(user)
    if _is_sin_asignar(usuario_profile):
        return Faena.objects.filter(id=faena_id, status=True).exists()

    if usuario_profile and usuario_profile.faena_id and str(usuario_profile.faena_id) == str(faena_id):
        return Faena.objects.filter(id=usuario_profile.faena_id, status=True).exists()

    return False


def _nombre_usuario_legible(user):
    nombre_usuario = f"{user.first_name} {user.last_name}".strip()
    if not nombre_usuario:
        nombre_usuario = user.username
    return nombre_usuario


def _serializar_notificacion_aprobacion(notificacion, request=None):
    sender_name = _nombre_usuario_legible(notificacion.solicitante)
    review_url = ''
    if request is not None:
        review_url = reverse('view_review_approval_notification', kwargs={'pk': notificacion.id})

    kind = 'sistema' if notificacion.tipo_documento == 'sistema' else 'approval'

    data = {
        'kind': kind,
        'id': notificacion.id,
        'titulo': notificacion.titulo,
        'descripcion': notificacion.descripcion,
        'estado': notificacion.estado,
        'estado_display': notificacion.get_estado_display(),
        'fecha': timezone.localtime(notificacion.fechacreacion).strftime('%d/%m/%Y %H:%M'),
        'solicitante': sender_name,
        'sender': sender_name,
        'plantilla_id': notificacion.plantilla_id,
        'plantilla_nombre': notificacion.plantilla.nombre if getattr(notificacion, 'plantilla', None) else '',
        'tipo_documento': notificacion.tipo_documento or '',
        'motivo_rechazo': notificacion.motivo_rechazo or '',
        'created_ts': int(notificacion.fechacreacion.timestamp()) if notificacion.fechacreacion else 0,
        'fecha_respuesta': (
            timezone.localtime(notificacion.fecha_respuesta).strftime('%d/%m/%Y %H:%M')
            if notificacion.fecha_respuesta else ''
        ),
        'review_url': review_url,
    }
    return data


def _serializar_notificacion_documento(notificacion, request=None):
    documento = getattr(notificacion, 'documento', None)
    review_url = ''
    tipo_documento = str(notificacion.tipo_documento or '').strip().lower()
    if request is not None and documento and documento.pk:
        base_url = reverse('view_documento', kwargs={'pk': documento.pk})
        review_url = f"{base_url}?{urlencode({'notif_doc': notificacion.id, 'origen': 'notificaciones', 'tipo': tipo_documento})}"

    sender_name = 'Sistema'
    return {
        'kind': 'document',
        'id': notificacion.id,
        'documento_id': documento.id if documento else None,
        'titulo': notificacion.titulo,
        'descripcion': notificacion.descripcion,
        'estado': notificacion.estado,
        'estado_display': notificacion.get_estado_display(),
        'fecha': timezone.localtime(notificacion.fechacreacion).strftime('%d/%m/%Y %H:%M'),
        'solicitante': sender_name,
        'sender': sender_name,
        'plantilla_id': documento.plantilla_base_id if documento else None,
        'plantilla_nombre': documento.plantilla_base.nombre if documento and documento.plantilla_base_id else '',
        'tipo_documento': tipo_documento,
        'motivo_rechazo': '',
        'created_ts': int(notificacion.fechacreacion.timestamp()) if notificacion.fechacreacion else 0,
        'fecha_respuesta': (
            timezone.localtime(notificacion.fecha_lectura).strftime('%d/%m/%Y %H:%M')
            if notificacion.fecha_lectura else ''
        ),
        'review_url': review_url,
    }


def _parsear_bool(value):
    if isinstance(value, str):
        return value.strip().lower() in {'1', 'true', 'yes', 'si', 'on'}
    return bool(value)


def _normalizar_rut_chileno(valor):
    # Normaliza un RUT a formato cuerpo+DV sin puntos/guion (ej: 12345678K).
    limpio = re.sub(r'[^0-9kK]', '', str(valor or '')).upper()
    if len(limpio) < 2:
        return ''
    return f"{limpio[:-1].lstrip('0') or '0'}{limpio[-1]}"


def _calcular_dv_rut(cuerpo):
    serie = [2, 3, 4, 5, 6, 7]
    suma = 0
    for idx, digito in enumerate(reversed(str(cuerpo))):
        suma += int(digito) * serie[idx % len(serie)]
    resto = 11 - (suma % 11)
    if resto == 11:
        return '0'
    if resto == 10:
        return 'K'
    return str(resto)


def _rut_chileno_valido(rut_normalizado):
    rut = _normalizar_rut_chileno(rut_normalizado)
    if not re.fullmatch(r'\d{7,8}[0-9K]', rut):
        return False
    cuerpo, dv = rut[:-1], rut[-1]
    return _calcular_dv_rut(cuerpo) == dv


def _formatear_rut_chileno(rut_normalizado):
    rut = _normalizar_rut_chileno(rut_normalizado)
    if not _rut_chileno_valido(rut):
        return str(rut_normalizado or '')
    cuerpo = rut[:-1]
    dv = rut[-1]
    return f"{int(cuerpo):,}".replace(',', '.') + f"-{dv}"


def _extraer_rut_chileno_desde_qr(raw_text):
    # Intenta extraer un RUT válido desde texto QR (texto plano o JSON).
    texto = str(raw_text or '').strip()
    if not texto:
        return ''

    candidatos = []

    if texto.startswith('{') and texto.endswith('}'):
        try:
            payload = json.loads(texto)
            if isinstance(payload, dict):
                for key, value in payload.items():
                    key_normalized = str(key).strip().lower()
                    if 'rut' in key_normalized or key_normalized == 'run':
                        candidatos.append(str(value or ''))
        except Exception:
            pass

    candidatos.extend(re.findall(r'(?<!\d)(\d{7,8}\s*-\s*[0-9kK])(?!\w)', texto))
    candidatos.extend(re.findall(r'(?<!\d)(\d{1,2}(?:\.\d{3}){2}\s*-\s*[0-9kK])(?!\w)', texto, flags=re.IGNORECASE))
    candidatos.extend(re.findall(r'(?<!\d)(\d{7,8}[0-9kK])(?!\w)', texto, flags=re.IGNORECASE))
    candidatos.append(texto)

    for candidato in candidatos:
        rut = _normalizar_rut_chileno(candidato)
        if _rut_chileno_valido(rut):
            return rut
    return ''


def _mapear_tipo_documento_label(tipo_documento):
    tipo = str(tipo_documento or '').strip().lower()
    labels = {
        'charlas': 'Charla',
        'informativos': 'Informativo',
        'difusiones': 'Capacitación',
        'cursos': 'Curso',
    }
    return labels.get(tipo, 'Documento')


def _ordenar_notificaciones_pendientes_primero(queryset, pending_value='pendiente'):
    return queryset.annotate(
        pending_priority=Case(
            When(estado=pending_value, then=Value(0)),
            default=Value(1),
            output_field=IntegerField(),
        )
    ).order_by('pending_priority', '-fechacreacion', '-id')


def _documento_activo_capacitacion_id(request):
    if request is None or not getattr(request, 'user', None):
        return 0
    if not request.user.is_authenticated:
        return 0

    estado_bloqueo = _obtener_estado_bloqueo_capacitacion(request)
    if not estado_bloqueo:
        return 0

    try:
        documento_id = int(estado_bloqueo.get('documento_id') or 0)
    except (TypeError, ValueError):
        documento_id = 0
    return max(0, documento_id)


def _listar_notificaciones_usuario(user, request=None, limit=None):
    approval_qs = (
        PrevencionNotificacionAprobacion.objects
        .filter(destinatario=user)
        .exclude(estado=PrevencionNotificacionAprobacion.Estado.CANCELADA)
        .select_related('solicitante', 'plantilla')
    )
    approval_qs = _ordenar_notificaciones_pendientes_primero(
        approval_qs,
        pending_value=PrevencionNotificacionAprobacion.Estado.PENDIENTE,
    )

    document_qs = (
        PrevencionNotificacionDocumento.objects
        .filter(
            destinatario=user,
            documento__status=True,
            documento__difusion_trabajadores__status=True,
            documento__difusion_trabajadores__user=user,
        )
        .select_related('documento', 'documento__plantilla_base')
        .distinct()
    )
    document_qs = _ordenar_notificaciones_pendientes_primero(
        document_qs,
        pending_value=PrevencionNotificacionDocumento.Estado.PENDIENTE,
    )

    serialized = []
    for item in approval_qs:
        serialized.append(_serializar_notificacion_aprobacion(item, request=request))
    for item in document_qs:
        serialized.append(_serializar_notificacion_documento(item, request=request))

    documento_activo_id = _documento_activo_capacitacion_id(request)
    if documento_activo_id > 0:
        for notificacion in serialized:
            if str(notificacion.get('kind') or '') != 'document':
                continue
            try:
                documento_id = int(notificacion.get('documento_id') or 0)
            except (TypeError, ValueError):
                documento_id = 0
            if documento_id != documento_activo_id:
                continue
            notificacion['estado'] = 'activo'
            notificacion['estado_display'] = 'Activo'

    serialized.sort(
        key=lambda n: (
            0 if str(n.get('estado') or '').lower() == 'pendiente' else 1,
            -(int(n.get('created_ts') or 0)),
            -(int(n.get('id') or 0)),
        )
    )

    if isinstance(limit, int) and limit > 0:
        return serialized[:limit]
    return serialized


def _obtener_estado_detallado_autorizadores_plantilla(plantilla, autorizador_ids=None):
    if not plantilla or not getattr(plantilla, 'id', None):
        return {}

    if autorizador_ids is None:
        autorizador_ids = list(plantilla.autorizadores.values_list('id', flat=True))

    ids_normalizados = []
    ids_vistos = set()
    for raw_id in (autorizador_ids or []):
        try:
            uid = int(raw_id)
        except (TypeError, ValueError):
            continue
        if uid in ids_vistos:
            continue
        ids_vistos.add(uid)
        ids_normalizados.append(uid)

    if not ids_normalizados:
        return {}

    notificaciones_qs = (
        PrevencionNotificacionAprobacion.objects
        .filter(
            plantilla=plantilla,
            destinatario_id__in=ids_normalizados,
        )
        .exclude(estado=PrevencionNotificacionAprobacion.Estado.CANCELADA)
        .order_by('destinatario_id', '-id')
    )

    estados = {}
    for notificacion in notificaciones_qs:
        destinatario_key = str(notificacion.destinatario_id)
        if destinatario_key in estados:
            continue
        estados[destinatario_key] = {
            'notification_id': notificacion.id,
            'estado': notificacion.estado,
            'motivo_rechazo': notificacion.motivo_rechazo or '',
        }
    return estados


def _obtener_estado_autorizadores_plantilla(plantilla, autorizador_ids=None):
    estado_detallado = _obtener_estado_detallado_autorizadores_plantilla(plantilla, autorizador_ids)
    estados = {}
    for destinatario_id, payload in estado_detallado.items():
        try:
            uid = int(destinatario_id)
        except (TypeError, ValueError):
            continue
        estados[uid] = str((payload or {}).get('estado') or '').strip().lower()
    return estados


def _plantilla_autorizada_para_documentos(plantilla):
    if not plantilla or not getattr(plantilla, 'id', None):
        return False

    autorizador_ids = list(plantilla.autorizadores.values_list('id', flat=True))
    if not autorizador_ids:
        # Documentos sin autorizadores no quedan habilitados.
        return False

    estados = _obtener_estado_autorizadores_plantilla(plantilla, autorizador_ids)
    tiene_aprobacion = False
    for autorizador_id in autorizador_ids:
        estado = estados.get(autorizador_id)
        if estado in {
            PrevencionNotificacionAprobacion.Estado.PENDIENTE,
            PrevencionNotificacionAprobacion.Estado.RECHAZADA,
        }:
            return False
        if estado == PrevencionNotificacionAprobacion.Estado.ACEPTADA:
            tiene_aprobacion = True

    # Regla de negocio:
    # - Se habilita con al menos una aprobación.
    # - Si cualquier autorizador entra en pendiente/rechazada, se bloquea.
    return tiene_aprobacion


def _parsear_json_request(request):
    try:
        return json.loads(request.body.decode('utf-8') or '{}')
    except Exception:
        return {}


def _cargo_usuario_id(user):
    # Cargo laboral asociado al usuario autenticado.
    info_laboral = UserInformacionLaboral.objects.select_related('cargo').filter(user=user).first()
    return info_laboral.cargo_id if info_laboral and info_laboral.cargo_id else None


def _usuario_puede_rellenar_documento(user, faena_id, plantilla_id):
    if not faena_id or not plantilla_id:
        return False

    # Admin prevencion siempre puede operar.
    if _es_admin_prevencion(user):
        return True

    permisos_qs = PrevencionPermisoLlenado.objects.filter(
        status=True,
        faena_id=faena_id,
        plantilla_id=plantilla_id,
    )

    # Si no hay reglas cargadas para esa combinacion, se mantiene comportamiento actual.
    if not permisos_qs.exists():
        return True

    cargo_usuario_id = _cargo_usuario_id(user)
    if not cargo_usuario_id:
        return False

    return permisos_qs.filter(cargos__id=cargo_usuario_id).exists()


MAX_ARCHIVOS_POR_INPUT = 5
MAX_TAMANO_ARCHIVO_BYTES = 100 * 1024 * 1024  # 100 MB
TIPOS_DOCUMENTO_PERMITIDOS = {'charlas', 'informativos', 'difusiones', 'cursos'}
TIPOS_CAPACITACION_LIBRE_TRAS_VALIDACION = {
    PrevencionPlantilla.TipoRiesgo.CHARLAS.value,
    PrevencionPlantilla.TipoRiesgo.INFORMATIVOS.value,
    PrevencionPlantilla.TipoRiesgo.DIFUSIONES.value,
}
TIEMPO_ESTIMADO_VALOR_30_SEGUNDOS = 3
TIEMPO_ESTIMADO_SEGUNDOS_CORTO = 30
CURSO_TOTAL_INTENTOS_DEFAULT = 2
CURSO_PORCENTAJE_APROBACION_DEFAULT = 60
CURSO_TOTAL_INTENTOS_DISPONIBLES = (1, 2, 3)
CURSO_PORCENTAJES_APROBACION_DISPONIBLES = (50, 60, 70, 80, 90, 100)
TIPOS_DOCUMENTO_TITULO_APROBACION = {
    'charlas': 'Aprobación pendiente: Charla',
    'informativos': 'Aprobación pendiente: Informativo',
    'difusiones': 'Aprobación pendiente: Difusión',
    'cursos': 'Aprobación pendiente: Curso',
}
CAPACITACION_BLOQUEO_SESSION_KEY = 'capacitacion_bloqueo_activo'
CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY = 'capacitacion_bloqueo_expirado'
CAPACITACION_CURSO_FINALIZADO_SESSION_KEY = 'capacitacion_curso_finalizado'
CAPACITACION_CURSO_RESPUESTAS_BORRADOR_SESSION_KEY = 'capacitacion_curso_respuestas_borrador'
AUTORIZACION_REENVIO_DESCRIPCION_PREFIJO = 'Reenvio automatico por cambios en la plantilla.'


def _construir_titulo_aprobacion(tipo_documento):
    return TIPOS_DOCUMENTO_TITULO_APROBACION.get(str(tipo_documento or '').strip().lower(), '')


def _construir_titulo_vista_documento(tipo_documento):
    tipo_slug = str(tipo_documento or '').strip().lower()
    tipo_label = _mapear_tipo_documento_label(tipo_slug)
    if tipo_label == 'Documento':
        return 'Ver Documento'
    return f"Ver {tipo_label}"


def _construir_titulo_nuevo_documento(tipo_documento):
    tipo_slug = str(tipo_documento or '').strip().lower()
    tipo_label = _mapear_tipo_documento_label(tipo_slug)
    if tipo_label == 'Documento':
        return 'Nuevo Documento'
    prefijo = 'Nueva' if tipo_slug in {'charlas', 'difusiones'} else 'Nuevo'
    return f"{prefijo} {tipo_label}"


def _serializar_estructura_para_comparacion_autorizacion(estructura):
    try:
        return json.dumps(estructura or [], sort_keys=True, cls=DjangoJSONEncoder)
    except Exception:
        return str(estructura or '')


def _obtener_cargo_id_plantilla(plantilla):
    if not plantilla or not getattr(plantilla, 'id', None):
        return None

    permiso = None
    if plantilla.faena_id:
        permiso = PrevencionPermisoLlenado.objects.filter(
            status=True,
            faena_id=plantilla.faena_id,
            plantilla=plantilla,
        ).first()
    if not permiso:
        permiso = PrevencionPermisoLlenado.objects.filter(
            status=True,
            plantilla=plantilla,
        ).order_by('-updated_at', '-id').first()
    if not permiso:
        return None
    return permiso.cargos.values_list('id', flat=True).first()


def _snapshot_plantilla_para_reautorizacion(plantilla):
    if not plantilla or not getattr(plantilla, 'id', None):
        return {}

    return {
        'nombre': str(plantilla.nombre or '').strip(),
        'faena_id': int(plantilla.faena_id or 0),
        'tiempo_estimado_minutos': plantilla.tiempo_estimado_minutos,
        'curso_total_intentos': int(getattr(plantilla, 'curso_total_intentos', 0) or 0),
        'curso_porcentaje_aprobacion': int(getattr(plantilla, 'curso_porcentaje_aprobacion', 0) or 0),
        'modo_contenido': str(plantilla.modo_contenido or PrevencionPlantilla.ModoContenido.ESTRUCTURA),
        'archivo_pdf_nombre': str(getattr(getattr(plantilla, 'archivo_pdf', None), 'name', '') or ''),
        'estructura': _serializar_estructura_para_comparacion_autorizacion(getattr(plantilla, 'estructura', [])),
        'cargo_id': _obtener_cargo_id_plantilla(plantilla),
        'autorizadores_ids': tuple(sorted(plantilla.autorizadores.values_list('id', flat=True))),
    }


def _crear_reautorizaciones_automaticas(plantilla, solicitante):
    if (
        not plantilla
        or not getattr(plantilla, 'id', None)
        or not solicitante
        or not getattr(solicitante, 'is_authenticated', False)
    ):
        return 0

    destinatarios_ids = list(
        plantilla.autorizadores.filter(is_active=True).values_list('id', flat=True)
    )
    if not destinatarios_ids:
        return 0

    now = timezone.now()
    PrevencionNotificacionAprobacion.objects.filter(
        plantilla=plantilla,
        estado=PrevencionNotificacionAprobacion.Estado.PENDIENTE,
    ).update(
        estado=PrevencionNotificacionAprobacion.Estado.CANCELADA,
        leida=True,
        fecha_lectura=now,
        fecha_respuesta=now,
    )

    tipo_documento = str(getattr(plantilla, 'tipo', '') or '').strip().lower()
    tipo_label = _mapear_tipo_documento_label(tipo_documento)
    titulo = (_construir_titulo_aprobacion(tipo_documento) or 'Aprobación pendiente')[:200]
    descripcion_base = (
        f'{_nombre_usuario_legible(solicitante)} solicita tu aprobación para '
        f'{tipo_label.lower()}: "{str(plantilla.nombre or "Sin nombre").strip()}".'
    )
    descripcion = f"{AUTORIZACION_REENVIO_DESCRIPCION_PREFIJO} {descripcion_base}".strip()

    notificaciones = [
        PrevencionNotificacionAprobacion(
            destinatario_id=destinatario_id,
            solicitante=solicitante,
            plantilla=plantilla,
            tipo_documento=tipo_documento,
            titulo=titulo,
            descripcion=descripcion,
            estado=PrevencionNotificacionAprobacion.Estado.PENDIENTE,
            leida=False,
            fecha_lectura=None,
            fecha_respuesta=None,
            motivo_rechazo='',
        )
        for destinatario_id in destinatarios_ids
    ]
    PrevencionNotificacionAprobacion.objects.bulk_create(notificaciones)
    return len(notificaciones)


def _es_notificacion_reenvio_automatico(notificacion):
    descripcion = str(getattr(notificacion, 'descripcion', '') or '').strip().lower()
    prefijo = AUTORIZACION_REENVIO_DESCRIPCION_PREFIJO.lower()
    return bool(descripcion and prefijo and descripcion.startswith(prefijo))


def _formatear_tiempo_restante_capacitacion(segundos_restantes):
    total = max(0, int(segundos_restantes or 0))
    minutos, segundos = divmod(total, 60)
    partes = []
    if minutos:
        partes.append(f"{minutos} minuto{'s' if minutos != 1 else ''}")
    if segundos:
        partes.append(f"{segundos} segundo{'s' if segundos != 1 else ''}")
    if not partes:
        return "0 segundos"
    return " y ".join(partes)


def _tiempo_estimado_a_segundos(tiempo_estimado_minutos):
    try:
        valor = int(tiempo_estimado_minutos or 0)
    except (TypeError, ValueError):
        return 0

    if valor <= 0:
        return 0

    if valor == TIEMPO_ESTIMADO_VALOR_30_SEGUNDOS:
        return TIEMPO_ESTIMADO_SEGUNDOS_CORTO

    return valor * 60


def _configuracion_curso_plantilla(plantilla):
    total_intentos = int(getattr(plantilla, 'curso_total_intentos', 0) or 0)
    if total_intentos not in CURSO_TOTAL_INTENTOS_DISPONIBLES:
        total_intentos = CURSO_TOTAL_INTENTOS_DEFAULT

    porcentaje_requerido = int(getattr(plantilla, 'curso_porcentaje_aprobacion', 0) or 0)
    if porcentaje_requerido not in CURSO_PORCENTAJES_APROBACION_DISPONIBLES:
        porcentaje_requerido = CURSO_PORCENTAJE_APROBACION_DEFAULT

    return {
        'total_intentos': total_intentos,
        'porcentaje_requerido': porcentaje_requerido,
    }


def _normalizar_estado_resultado_curso(documento, resultado=None):
    plantilla = getattr(documento, 'plantilla_base', None)
    config = _configuracion_curso_plantilla(plantilla)
    total_intentos = int(config.get('total_intentos') or CURSO_TOTAL_INTENTOS_DEFAULT)
    porcentaje_requerido = int(
        config.get('porcentaje_requerido') or CURSO_PORCENTAJE_APROBACION_DEFAULT
    )

    if not resultado:
        return {
            'total_intentos': total_intentos,
            'porcentaje_requerido': porcentaje_requerido,
            'intentos_realizados': 0,
            'intentos_restantes': max(total_intentos, 0),
            'aprobado': False,
            'estado_key': 'entregado',
            'estado_label': 'Entregado',
        }

    porcentaje_obtenido = float(getattr(resultado, 'porcentaje_aprobacion', 0.0) or 0.0)
    intentos_guardados = int(getattr(resultado, 'intentos_realizados', 0) or 0)
    tuvo_evaluacion_historica = bool(
        int(getattr(resultado, 'total_preguntas', 0) or 0) > 0
        or porcentaje_obtenido > 0
        or bool(getattr(resultado, 'detalle_respuestas', None))
    )
    intentos_realizados = intentos_guardados if intentos_guardados > 0 else (
        1 if tuvo_evaluacion_historica else 0
    )
    if total_intentos > 0:
        intentos_realizados = min(intentos_realizados, total_intentos)
    intentos_restantes = max(total_intentos - intentos_realizados, 0)

    aprobado = bool(getattr(resultado, 'aprobado', False))
    if not aprobado and intentos_realizados > 0:
        aprobado = porcentaje_obtenido >= float(porcentaje_requerido)

    if aprobado:
        estado_key = 'aprobado'
        estado_label = 'Aprobado'
    elif intentos_realizados > 0:
        estado_key = 'rechazado'
        estado_label = 'Rechazado'
    else:
        estado_key = 'entregado'
        estado_label = 'Entregado'

    return {
        'total_intentos': total_intentos,
        'porcentaje_requerido': porcentaje_requerido,
        'intentos_realizados': intentos_realizados,
        'intentos_restantes': intentos_restantes,
        'aprobado': aprobado,
        'estado_key': estado_key,
        'estado_label': estado_label,
    }


def _documento_capacitacion_validado_para_usuario(documento_id, user):
    if not documento_id or not user or not getattr(user, 'is_authenticated', False):
        return False

    notificacion = (
        PrevencionNotificacionDocumento.objects
        .filter(destinatario=user, documento_id=documento_id)
        .order_by('-fechacreacion', '-id')
        .values('estado', 'leida')
        .first()
    )
    if not notificacion:
        return False

    estado_raw = str(notificacion.get('estado') or '').strip().lower()
    leida = bool(notificacion.get('leida'))
    return leida or estado_raw == PrevencionNotificacionDocumento.Estado.VISTA


def _obtener_estado_bloqueo_capacitacion(request):
    data = request.session.get(CAPACITACION_BLOQUEO_SESSION_KEY)
    if not isinstance(data, dict):
        return None

    try:
        expires_at_ts = int(float(data.get('expires_at_ts') or 0))
    except (TypeError, ValueError):
        request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
        request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
        request.session.modified = True
        return None

    try:
        documento_id = int(data.get('documento_id') or 0)
    except (TypeError, ValueError):
        documento_id = 0

    now_ts = int(timezone.now().timestamp())
    remaining_seconds = max(0, expires_at_ts - now_ts)
    try:
        started_at_ts = int(float(data.get('started_at_ts') or 0))
    except (TypeError, ValueError):
        started_at_ts = 0
        
    if documento_id <= 0:
        request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
        request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
        request.session.modified = True
        return None

    documento_bloqueado = (
        PrevencionDocumento.objects
        .filter(id=documento_id, status=True)
        .select_related('plantilla_base')
        .first()
    )

    # Si el documento fue eliminado/inactivado, el bloqueo queda inválido.
    if not documento_bloqueado:
        request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
        request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
        request.session.modified = True
        return None

    tipo_documento_bloqueado = str(
        getattr(getattr(documento_bloqueado, 'plantilla_base', None), 'tipo', '') or ''
    ).strip().lower()

    if remaining_seconds <= 0:
        if tipo_documento_bloqueado in TIPOS_CAPACITACION_LIBRE_TRAS_VALIDACION:
            if request.session.pop(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY, None) is not None:
                request.session.modified = True
        else:
            if documento_id > 0:
                request.session[CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY] = {
                    'documento_id': documento_id,
                    'tipo_slug': str(data.get('tipo_slug') or '').strip().lower(),
                    'expired_at_ts': now_ts,
                }
            
            if tipo_documento_bloqueado == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
                request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
                request.session.modified = True
            else:
                request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
                request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
                request.session.modified = True
                return None

    if request.user.is_authenticated:
        documentos_visibles_usuario = _ids_documentos_visibles_para_usuario(request.user)
        if documento_id not in documentos_visibles_usuario:
            request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
            request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
            request.session.modified = True
            return None

        notificacion_actual = (
            PrevencionNotificacionDocumento.objects
            .filter(destinatario=request.user, documento_id=documento_id)
            .order_by('-fechacreacion', '-id')
            .values('fechacreacion')
            .first()
        )
        if (
            started_at_ts > 0
            and notificacion_actual
            and notificacion_actual.get('fechacreacion')
            and int(notificacion_actual['fechacreacion'].timestamp()) > started_at_ts
        ):
            # Reasignacion posterior al bloqueo actual: invalida el bloqueo viejo.
            request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
            request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
            request.session.modified = True
            return None

        if (
            tipo_documento_bloqueado in TIPOS_CAPACITACION_LIBRE_TRAS_VALIDACION
            and _documento_capacitacion_validado_para_usuario(documento_id, request.user)
        ):
            # Para tipos no curso, al validar el documento se libera el bloqueo.
            request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
            request.session.pop(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY, None)
            request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
            request.session.modified = True
            return None

    return {
        'documento_id': documento_id,
        'tipo_slug': str(data.get('tipo_slug') or '').strip().lower(),
        'expires_at_ts': expires_at_ts,
        'started_at_ts': started_at_ts,
        'remaining_seconds': remaining_seconds,
    }


def _iniciar_bloqueo_capacitacion(request, documento, forzar_reinicio=False):
    if not documento or not getattr(documento, 'plantilla_base_id', None):
        return None

    plantilla = documento.plantilla_base
    tipo_slug = str(getattr(plantilla, 'tipo', '') or '').strip().lower()
    if tipo_slug not in TIPOS_DOCUMENTO_PERMITIDOS:
        return None

    if (
        tipo_slug in TIPOS_CAPACITACION_LIBRE_TRAS_VALIDACION
        and _documento_capacitacion_validado_para_usuario(documento.id, request.user)
    ):
        return None

    duracion_segundos = _tiempo_estimado_a_segundos(
        getattr(plantilla, 'tiempo_estimado_minutos', 0)
    )

    if duracion_segundos <= 0:
        return None

    estado_actual = _obtener_estado_bloqueo_capacitacion(request)
    if (
        estado_actual
        and estado_actual.get('documento_id') == documento.id
        and not forzar_reinicio
    ):
        return estado_actual

    now_ts = int(timezone.now().timestamp())
    expires_at_ts = now_ts + duracion_segundos
    data = {
        'documento_id': documento.id,
        'tipo_slug': tipo_slug,
        'started_at_ts': now_ts,
        'expires_at_ts': expires_at_ts,
    }
    request.session[CAPACITACION_BLOQUEO_SESSION_KEY] = data
    request.session.pop(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY, None)
    request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
    request.session.modified = True

    data['remaining_seconds'] = expires_at_ts - now_ts
    return data


def _obtener_curso_finalizado_bloqueado(request, estado_bloqueo=None):
    data = request.session.get(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY)
    if not isinstance(data, dict):
        return None

    try:
        documento_id = int(data.get('documento_id') or 0)
    except (TypeError, ValueError):
        documento_id = 0
    tipo_slug = str(data.get('tipo_slug') or '').strip().lower()

    if documento_id <= 0 or tipo_slug != PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
        request.session.modified = True
        return None

    if estado_bloqueo is None:
        estado_bloqueo = _obtener_estado_bloqueo_capacitacion(request)

    if not estado_bloqueo:
        request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
        request.session.modified = True
        return None

    documento_bloqueado_id = int(estado_bloqueo.get('documento_id') or 0)
    segundos_restantes = max(0, int(estado_bloqueo.get('remaining_seconds', 0) or 0))
    if documento_bloqueado_id != documento_id or segundos_restantes <= 0:
        request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
        request.session.modified = True
        return None

    return {
        'documento_id': documento_id,
        'tipo_slug': tipo_slug,
    }


def _obtener_borrador_respuestas_curso(request, documento_id, with_meta=False):
    try:
        doc_key = str(int(documento_id or 0))
    except (TypeError, ValueError):
        return {'respuestas': {}, 'updated_at_ts': 0} if with_meta else {}
    if not doc_key or doc_key == '0':
        return {'respuestas': {}, 'updated_at_ts': 0} if with_meta else {}

    data = request.session.get(CAPACITACION_CURSO_RESPUESTAS_BORRADOR_SESSION_KEY)
    if not isinstance(data, dict):
        return {'respuestas': {}, 'updated_at_ts': 0} if with_meta else {}

    payload = data.get(doc_key)
    if not isinstance(payload, dict):
        return {'respuestas': {}, 'updated_at_ts': 0} if with_meta else {}

    respuestas = _normalizar_respuestas_marcadas_curso(payload.get('respuestas'))
    respuestas_normalizadas = {str(k): int(v) for k, v in respuestas.items()}
    updated_at_ts_raw = payload.get('updated_at_ts')
    try:
        updated_at_ts = int(updated_at_ts_raw or 0)
    except (TypeError, ValueError):
        updated_at_ts = 0

    if with_meta:
        return {
            'respuestas': respuestas_normalizadas,
            'updated_at_ts': updated_at_ts,
        }
    return respuestas_normalizadas


def _guardar_borrador_respuestas_curso(request, documento_id, respuestas_raw):
    try:
        doc_key = str(int(documento_id or 0))
    except (TypeError, ValueError):
        return {}
    if not doc_key or doc_key == '0':
        return {}

    respuestas_normalizadas = _normalizar_respuestas_marcadas_curso(respuestas_raw)
    respuestas_session = {str(k): int(v) for k, v in respuestas_normalizadas.items()}

    data = request.session.get(CAPACITACION_CURSO_RESPUESTAS_BORRADOR_SESSION_KEY)
    if not isinstance(data, dict):
        data = {}
    else:
        data = dict(data)

    data[doc_key] = {
        'respuestas': respuestas_session,
        'updated_at_ts': int(timezone.now().timestamp()),
    }
    request.session[CAPACITACION_CURSO_RESPUESTAS_BORRADOR_SESSION_KEY] = data
    request.session.modified = True
    return respuestas_session


def _limpiar_borrador_respuestas_curso(request, documento_id=None):
    data = request.session.get(CAPACITACION_CURSO_RESPUESTAS_BORRADOR_SESSION_KEY)
    if not isinstance(data, dict):
        return

    if documento_id is None:
        request.session.pop(CAPACITACION_CURSO_RESPUESTAS_BORRADOR_SESSION_KEY, None)
        request.session.modified = True
        return

    try:
        doc_key = str(int(documento_id or 0))
    except (TypeError, ValueError):
        doc_key = ''
    if not doc_key:
        return

    if doc_key in data:
        data = dict(data)
        data.pop(doc_key, None)
        if data:
            request.session[CAPACITACION_CURSO_RESPUESTAS_BORRADOR_SESSION_KEY] = data
        else:
            request.session.pop(CAPACITACION_CURSO_RESPUESTAS_BORRADOR_SESSION_KEY, None)
        request.session.modified = True


def _actualizar_notificacion_estado_curso(documento, usuario, aprobado, fecha=None):
    if not documento or not usuario or not getattr(usuario, 'is_authenticated', False):
        return

    momento = fecha or timezone.now()
    if aprobado:
        PrevencionNotificacionDocumento.objects.filter(
            destinatario=usuario,
            documento=documento,
        ).update(
            estado=PrevencionNotificacionDocumento.Estado.VISTA,
            leida=True,
            fecha_lectura=momento,
        )
    else:
        PrevencionNotificacionDocumento.objects.filter(
            destinatario=usuario,
            documento=documento,
        ).update(
            estado=PrevencionNotificacionDocumento.Estado.PENDIENTE,
            leida=False,
            fecha_lectura=None,
        )


def _auto_finalizar_curso_por_tiempo(request, documento):
    if not documento or not getattr(documento, 'id', None):
        return None
    if not request.user or not getattr(request.user, 'is_authenticated', False):
        return None

    tipo_documento = str(
        getattr(getattr(documento, 'plantilla_base', None), 'tipo', '') or ''
    ).strip().lower()
    if tipo_documento != PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        return None

    respuestas_guardadas = _obtener_borrador_respuestas_curso(request, documento.id)
    momento = timezone.now()

    resultados_qs = (
        PrevencionResultadoCurso.objects
        .filter(documento=documento, trabajador=request.user)
        .order_by('-updated_at', '-id')
    )
    resultado = resultados_qs.first()
    estado_previo = _normalizar_estado_resultado_curso(documento, resultado)

    if bool(estado_previo.get('aprobado')):
        _actualizar_notificacion_estado_curso(
            documento=documento,
            usuario=request.user,
            aprobado=True,
            fecha=momento,
        )
        _limpiar_borrador_respuestas_curso(request, documento.id)
        return {
            'estado_key': 'aprobado',
            'intentos_restantes': int(estado_previo.get('intentos_restantes') or 0),
        }

    intentos_restantes_previos = int(estado_previo.get('intentos_restantes') or 0)
    if intentos_restantes_previos <= 0:
        _actualizar_notificacion_estado_curso(
            documento=documento,
            usuario=request.user,
            aprobado=False,
            fecha=momento,
        )
        _limpiar_borrador_respuestas_curso(request, documento.id)
        return {
            'estado_key': 'rechazado',
            'intentos_restantes': 0,
        }

    evaluacion = _evaluar_resultado_curso(documento, respuestas_guardadas)
    intentos_nuevos = int(estado_previo.get('intentos_realizados') or 0) + 1
    porcentaje_requerido = int(
        estado_previo.get('porcentaje_requerido') or CURSO_PORCENTAJE_APROBACION_DEFAULT
    )
    total_intentos = int(estado_previo.get('total_intentos') or CURSO_TOTAL_INTENTOS_DEFAULT)

    aprobado_actual = float(evaluacion.get('porcentaje_aprobacion') or 0.0) >= float(
        porcentaje_requerido
    )
    intentos_restantes = max(total_intentos - intentos_nuevos, 0)

    if resultado:
        resultados_qs.exclude(id=resultado.id).delete()
        resultado.total_preguntas = int(evaluacion.get('total_preguntas') or 0)
        resultado.total_correctas = int(evaluacion.get('total_correctas') or 0)
        resultado.porcentaje_aprobacion = float(evaluacion.get('porcentaje_aprobacion') or 0.0)
        resultado.intentos_realizados = intentos_nuevos
        resultado.aprobado = aprobado_actual
        resultado.detalle_respuestas = evaluacion.get('detalle') or []
        resultado.save()
    else:
        try:
            resultado = PrevencionResultadoCurso.objects.create(
                documento=documento,
                trabajador=request.user,
                total_preguntas=int(evaluacion.get('total_preguntas') or 0),
                total_correctas=int(evaluacion.get('total_correctas') or 0),
                porcentaje_aprobacion=float(evaluacion.get('porcentaje_aprobacion') or 0.0),
                intentos_realizados=intentos_nuevos,
                aprobado=aprobado_actual,
                detalle_respuestas=evaluacion.get('detalle') or [],
            )
        except IntegrityError:
            resultado = resultados_qs.first()
            if resultado:
                resultado.total_preguntas = int(evaluacion.get('total_preguntas') or 0)
                resultado.total_correctas = int(evaluacion.get('total_correctas') or 0)
                resultado.porcentaje_aprobacion = float(evaluacion.get('porcentaje_aprobacion') or 0.0)
                resultado.intentos_realizados = intentos_nuevos
                resultado.aprobado = aprobado_actual
                resultado.detalle_respuestas = evaluacion.get('detalle') or []
                resultado.save()

    _registrar_historial_intento_curso(
        documento=documento,
        trabajador=request.user,
        numero_intento=intentos_nuevos,
        evaluacion=evaluacion,
        aprobado=aprobado_actual,
        fecha_intento=momento,
    )

    _actualizar_notificacion_estado_curso(
        documento=documento,
        usuario=request.user,
        aprobado=aprobado_actual,
        fecha=momento,
    )
    _limpiar_borrador_respuestas_curso(request, documento.id)
    return {
        'estado_key': 'aprobado' if aprobado_actual else 'rechazado',
        'intentos_restantes': intentos_restantes,
    }


def _registrar_historial_intento_curso(
    documento,
    trabajador,
    numero_intento,
    evaluacion,
    aprobado,
    fecha_intento=None,
):
    if not documento or not trabajador or not getattr(trabajador, 'id', None):
        return None
    try:
        intento_numero = int(numero_intento or 0)
    except (TypeError, ValueError):
        intento_numero = 0
    if intento_numero <= 0:
        return None

    defaults = {
        'total_preguntas': int((evaluacion or {}).get('total_preguntas') or 0),
        'total_correctas': int((evaluacion or {}).get('total_correctas') or 0),
        'porcentaje_aprobacion': float((evaluacion or {}).get('porcentaje_aprobacion') or 0.0),
        'aprobado': bool(aprobado),
        'detalle_respuestas': (evaluacion or {}).get('detalle') or [],
        'fechacreacion': fecha_intento or timezone.now(),
    }
    intento_obj, _ = PrevencionResultadoCursoIntento.objects.update_or_create(
        documento=documento,
        trabajador=trabajador,
        numero_intento=intento_numero,
        defaults=defaults,
    )
    return intento_obj


def _reiniciar_resultado_curso_trabajador(documento, user_id):
    if not documento:
        return 0
    try:
        trabajador_id = int(user_id or 0)
    except (TypeError, ValueError):
        return 0
    if trabajador_id <= 0:
        return 0

    eliminados_resultado = PrevencionResultadoCurso.objects.filter(
        documento=documento,
        trabajador_id=trabajador_id,
    ).delete()[0]
    PrevencionResultadoCursoIntento.objects.filter(
        documento=documento,
        trabajador_id=trabajador_id,
    ).delete()
    return eliminados_resultado


def _limpiar_estado_sesion_curso_usuario_actual(request, documento_id):
    if not request or not getattr(request, 'session', None):
        return
    try:
        doc_id = int(documento_id or 0)
    except (TypeError, ValueError):
        doc_id = 0
    if doc_id <= 0:
        return

    session_modificada = False

    bloqueo = request.session.get(CAPACITACION_BLOQUEO_SESSION_KEY)
    if isinstance(bloqueo, dict):
        try:
            bloqueo_doc_id = int(bloqueo.get('documento_id') or 0)
        except (TypeError, ValueError):
            bloqueo_doc_id = 0
        if bloqueo_doc_id == doc_id:
            request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
            request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
            session_modificada = True

    bloqueo_expirado = request.session.get(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY)
    if isinstance(bloqueo_expirado, dict):
        try:
            bloqueo_exp_doc_id = int(bloqueo_expirado.get('documento_id') or 0)
        except (TypeError, ValueError):
            bloqueo_exp_doc_id = 0
        if bloqueo_exp_doc_id == doc_id:
            request.session.pop(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY, None)
            session_modificada = True

    _limpiar_borrador_respuestas_curso(request, doc_id)
    if session_modificada:
        request.session.modified = True


def _indice_opcion_a_letra(indice):
    try:
        idx = int(indice)
    except (TypeError, ValueError):
        return ''

    if idx < 0:
        return ''
    if idx < 26:
        return chr(65 + idx)
    return f"O{idx + 1}"


def _indice_desde_letra_opcion(letra):
    texto = str(letra or '').strip().upper()
    if not texto:
        return None
    if len(texto) == 1 and 'A' <= texto <= 'Z':
        return ord(texto) - ord('A')
    return None


def _extraer_preguntas_curso_desde_documento(documento):
    plantilla = getattr(documento, 'plantilla_base', None)
    estructura = []
    if plantilla and isinstance(getattr(plantilla, 'estructura', None), list):
        estructura = plantilla.estructura

    preguntas = []
    for idx, seccion in enumerate(estructura):
        if not isinstance(seccion, dict):
            continue

        titulo = str(seccion.get('titulo') or '').strip() or f"Pregunta {idx + 1}"
        enunciado = str(seccion.get('enunciado') or '').strip()
        imagen_url = str(seccion.get('imagen_url') or '').strip()

        opciones = []
        opciones_raw = seccion.get('opciones')
        if isinstance(opciones_raw, list):
            for opcion in opciones_raw:
                texto_opcion = str(opcion or '').strip()
                if texto_opcion:
                    opciones.append(texto_opcion)

        respuesta_correcta = str(seccion.get('respuesta_correcta') or '').strip().upper()
        justificacion = str(seccion.get('justificacion') or '').strip()

        filas = seccion.get('filas')
        if isinstance(filas, list) and filas:
            if not enunciado:
                primer_valor = str(filas[0] or '').strip()
                es_opcion = bool(re.match(r'^[a-zA-Z]\s*[-.)]\s*(.+)$', primer_valor))
                es_respuesta = primer_valor.lower().startswith('respuesta')
                es_justificacion = primer_valor.lower().startswith('justificaci')
                if primer_valor and not es_opcion and not es_respuesta and not es_justificacion:
                    enunciado = primer_valor

            if not opciones:
                for fila in filas:
                    texto = str(fila or '').strip()
                    opcion_match = re.match(r'^[a-zA-Z]\s*[-.)]\s*(.+)$', texto)
                    if opcion_match:
                        opciones.append(opcion_match.group(1).strip())

            if not respuesta_correcta:
                linea_respuesta = next(
                    (
                        str(fila or '').strip()
                        for fila in filas
                        if str(fila or '').strip().lower().startswith('respuesta')
                    ),
                    '',
                )
                if linea_respuesta:
                    match = re.search(r':\s*([A-Za-z])', linea_respuesta)
                    if match:
                        respuesta_correcta = match.group(1).upper()

            if not justificacion:
                linea_justificacion = next(
                    (
                        str(fila or '').strip()
                        for fila in filas
                        if str(fila or '').strip().lower().startswith('justificaci')
                    ),
                    '',
                )
                if linea_justificacion:
                    match = re.search(r':\s*(.+)$', linea_justificacion)
                    if match:
                        justificacion = match.group(1).strip()

        preguntas.append({
            'indice': idx,
            'titulo': titulo,
            'enunciado': enunciado,
            'imagen_url': imagen_url,
            'opciones': opciones,
            'respuesta_correcta': respuesta_correcta,
            'justificacion': justificacion,
        })

    return preguntas


def _normalizar_respuestas_marcadas_curso(raw):
    respuestas = {}
    if isinstance(raw, dict):
        items = raw.items()
    elif isinstance(raw, list):
        items = enumerate(raw)
    else:
        return respuestas

    for key, value in items:
        try:
            pregunta_idx = int(key)
            opcion_idx = int(value)
        except (TypeError, ValueError):
            continue
        if pregunta_idx < 0 or opcion_idx < 0:
            continue
        respuestas[pregunta_idx] = opcion_idx
    return respuestas


def _evaluar_resultado_curso(documento, respuestas_marcadas_raw):
    preguntas = _extraer_preguntas_curso_desde_documento(documento)
    respuestas_marcadas = _normalizar_respuestas_marcadas_curso(respuestas_marcadas_raw)

    detalle = []
    total_correctas = 0
    total_preguntas = len(preguntas)

    for pregunta in preguntas:
        idx = int(pregunta.get('indice') or 0)
        opciones = list(pregunta.get('opciones') or [])
        respuesta_correcta_letra = str(pregunta.get('respuesta_correcta') or '').strip().upper()
        respuesta_correcta_idx = _indice_desde_letra_opcion(respuesta_correcta_letra)

        seleccion_idx = respuestas_marcadas.get(idx)
        seleccion_valida = isinstance(seleccion_idx, int) and 0 <= seleccion_idx < len(opciones)
        seleccion_letra = _indice_opcion_a_letra(seleccion_idx) if seleccion_valida else ''
        seleccion_texto = opciones[seleccion_idx] if seleccion_valida else ''

        es_correcta = (
            respuesta_correcta_idx is not None
            and seleccion_valida
            and seleccion_idx == respuesta_correcta_idx
        )
        if es_correcta:
            total_correctas += 1

        opciones_detalle = []
        for opcion_idx, opcion_texto in enumerate(opciones):
            opciones_detalle.append({
                'indice': opcion_idx,
                'letra': _indice_opcion_a_letra(opcion_idx),
                'texto': opcion_texto,
                'es_marcada': bool(seleccion_valida and opcion_idx == seleccion_idx),
                'es_correcta': bool(
                    respuesta_correcta_idx is not None and opcion_idx == respuesta_correcta_idx
                ),
            })

        detalle.append({
            'indice': idx,
            'numero': idx + 1,
            'titulo': pregunta.get('titulo') or f"Pregunta {idx + 1}",
            'enunciado': pregunta.get('enunciado') or '',
            'imagen_url': pregunta.get('imagen_url') or '',
            'opciones': opciones_detalle,
            'respuesta_correcta_letra': respuesta_correcta_letra or '-',
            'respuesta_correcta_texto': (
                opciones[respuesta_correcta_idx]
                if isinstance(respuesta_correcta_idx, int) and 0 <= respuesta_correcta_idx < len(opciones)
                else ''
            ),
            'justificacion': pregunta.get('justificacion') or '',
            'respondida': bool(seleccion_valida),
            'seleccion_letra': seleccion_letra or '-',
            'seleccion_texto': seleccion_texto or 'Sin respuesta',
            'es_correcta': es_correcta,
        })

    porcentaje_aprobacion = round(
        (float(total_correctas) * 100.0 / float(total_preguntas)),
        2
    ) if total_preguntas > 0 else 0.0

    return {
        'total_preguntas': total_preguntas,
        'total_correctas': total_correctas,
        'porcentaje_aprobacion': porcentaje_aprobacion,
        'detalle': detalle,
    }


def _obtener_historial_autorizaciones_plantilla(plantilla, limite_notificaciones=100, limite_cambios=100, limite_total=150):
    if not plantilla or not getattr(plantilla, 'id', None):
        return []

    eventos_historial = []

    historial_notificaciones_qs = (
        PrevencionNotificacionAprobacion.objects
        .filter(plantilla=plantilla)
        .select_related('solicitante', 'destinatario')
        .order_by('-fechacreacion', '-id')
    )
    for notificacion in historial_notificaciones_qs[:limite_notificaciones]:
        estado_notificacion = notificacion.estado
        estado_display = notificacion.get_estado_display()
        motivo_rechazo = (notificacion.motivo_rechazo or '').strip()
        evento_notificacion = 'Solicitud de aprobación'
        if _es_notificacion_reenvio_automatico(notificacion):
            evento_notificacion = 'Cambio de plantilla y reenvío de solicitud'

        # Compatibilidad histórica:
        # versiones anteriores podían convertir rechazos ya resueltos a "cancelada"
        # al sincronizar autorizadores. Si existe motivo de rechazo, se preserva como rechazada.
        if (
            estado_notificacion == PrevencionNotificacionAprobacion.Estado.CANCELADA
            and motivo_rechazo
        ):
            estado_notificacion = PrevencionNotificacionAprobacion.Estado.RECHAZADA
            estado_display = str(PrevencionNotificacionAprobacion.Estado.RECHAZADA.label)

        eventos_historial.append({
            'fecha_raw': notificacion.fechacreacion,
            'fecha': timezone.localtime(notificacion.fechacreacion).strftime('%d/%m/%Y %H:%M'),
            'evento': evento_notificacion,
            'solicitante': _nombre_usuario_legible(notificacion.solicitante),
            'destinatario': _nombre_usuario_legible(notificacion.destinatario),
            'estado': estado_notificacion,
            'estado_display': estado_display,
            'fecha_respuesta': (
                timezone.localtime(notificacion.fecha_respuesta).strftime('%d/%m/%Y %H:%M')
                if notificacion.fecha_respuesta else '-'
            ),
            'motivo_rechazo': motivo_rechazo,
        })

    historial_cambios_qs = (
        PrevencionHistorialAutorizacion.objects
        .filter(plantilla=plantilla)
        .select_related('actor', 'destinatario')
        .order_by('-fechacreacion', '-id')
    )
    for cambio in historial_cambios_qs[:limite_cambios]:
        fecha_cambio = timezone.localtime(cambio.fechacreacion).strftime('%d/%m/%Y %H:%M')
        mostrar_fecha_respuesta = cambio.accion in {
            PrevencionHistorialAutorizacion.Accion.AGREGADO,
            PrevencionHistorialAutorizacion.Accion.ELIMINADO,
        }
        eventos_historial.append({
            'fecha_raw': cambio.fechacreacion,
            'fecha': fecha_cambio,
            'evento': 'Cambio de autorizadores',
            'solicitante': _nombre_usuario_legible(cambio.actor) if cambio.actor else '-',
            'destinatario': _nombre_usuario_legible(cambio.destinatario) if cambio.destinatario else '-',
            'estado': cambio.accion,
            'estado_display': cambio.get_accion_display(),
            'fecha_respuesta': fecha_cambio if mostrar_fecha_respuesta else '-',
        })

    historial_ordenado = sorted(
        eventos_historial,
        key=lambda item: item.get('fecha_raw') or timezone.now(),
        reverse=True,
    )[:limite_total]

    for item in historial_ordenado:
        item.pop('fecha_raw', None)
    return historial_ordenado


VIGILANCIA_OPCIONES = {
    'ges': {
        'model': VigilanciaGes,
        'titulo': 'GES',
        'texto_boton_nuevo': 'Nuevo GES',
        'mensaje_exito': 'GES creado con exito!',
        'es_numerico': False,
        'menu_key': 'v-ges',
    },
    'area': {
        'model': VigilanciaArea,
        'titulo': 'Area',
        'texto_boton_nuevo': 'Nueva Area',
        'mensaje_exito': 'Area creada con exito!',
        'es_numerico': False,
        'menu_key': 'v-area',
    },
    'cargo': {
        'model': VigilanciaCargo,
        'titulo': 'Cargo',
        'texto_boton_nuevo': 'Nuevo Cargo',
        'mensaje_exito': 'Cargo creado con exito!',
        'es_numerico': False,
        'menu_key': 'v-cargo',
    },
    'tipo-contrato': {
        'model': VigilanciaTipoContrato,
        'titulo': 'Tipo de Contrato',
        'texto_boton_nuevo': 'Nuevo Tipo de Contrato',
        'mensaje_exito': 'Tipo de Contrato creado con exito!',
        'es_numerico': False,
        'menu_key': 'v-tipo-contrato',
    },
    'contrato': {
        'model': VigilanciaContrato,
        'titulo': 'Contrato',
        'texto_boton_nuevo': 'Nuevo Contrato',
        'mensaje_exito': 'Contrato creado con exito!',
        'es_numerico': False,
        'menu_key': 'v-contrato',
    },
    'evaluacion-riesgo': {
        'model': VigilanciaEvaluacionRiesgo,
        'titulo': 'Evaluacion de Riesgo',
        'texto_boton_nuevo': 'Nueva Evaluacion de Riesgo',
        'mensaje_exito': 'Evaluacion de Riesgo creada con exito!',
        'es_numerico': False,
        'menu_key': 'v-evaluacion-riesgo',
    },
    'exposicion': {
        'model': VigilanciaExposicion,
        'titulo': 'Exposicion',
        'texto_boton_nuevo': 'Nueva Exposicion',
        'mensaje_exito': 'Exposicion creada con exito!',
        'es_numerico': False,
        'menu_key': 'v-exposicion',
    },
    'nivel-riesgo': {
        'model': VigilanciaNivelRiesgo,
        'titulo': 'Nivel de Riesgo',
        'texto_boton_nuevo': 'Nuevo Nivel de Riesgo',
        'mensaje_exito': 'Nivel de Riesgo creado con exito!',
        'es_numerico': False,
        'menu_key': 'v-nivel-riesgo',
    },
    'grado-exposicion': {
        'model': VigilanciaGradoExposicion,
        'titulo': 'Grado de Exposicion',
        'texto_boton_nuevo': 'Nuevo Grado de Exposicion',
        'mensaje_exito': 'Grado de Exposicion creado con exito!',
        'es_numerico': False,
        'menu_key': 'v-grado-exposicion',
    },
    'nivel-seguimiento': {
        'model': VigilanciaNivelSeguimiento,
        'titulo': 'Nivel de Seguimiento',
        'texto_boton_nuevo': 'Nuevo Nivel de Seguimiento',
        'mensaje_exito': 'Nivel de Seguimiento creado con exito!',
        'es_numerico': False,
        'menu_key': 'v-nivel-seguimiento',
    },
    'estado': {
        'model': VigilanciaEstado,
        'titulo': 'Estado',
        'texto_boton_nuevo': 'Nuevo Estado',
        'mensaje_exito': 'Estado creado con exito!',
        'es_numerico': False,
        'menu_key': 'v-estado',
    },
}


def _es_admin_prevencion(user):
    if not user or not user.is_authenticated:
        return False
    perfil = _get_usuario_profile(user)
    return bool(perfil and perfil.seccionPrevencion == 'ADMINISTRADOR')

def _usuario_autorizado_opcion_vigilancia(user, tipo_slug):
    if not user or not user.is_authenticated:
        return False
    perfil = _get_usuario_profile(user)
    if not perfil:
        return False
    rol = perfil.seccionPrevencion
    if rol == 'SIN ASIGNAR':
        return False
    if tipo_slug in ['ges', 'area', 'cargo', 'tipo-contrato', 'contrato']:
        return rol in ['ADMINISTRADOR', 'BASE DATOS']
    return True


def _get_vigilancia_opcion_config(tipo_slug):
    return VIGILANCIA_OPCIONES.get(tipo_slug)


def _construir_titulo_detalle_documento(documento):
    tipo = 'Sin tipo'
    faena = 'Sin faena'

    if documento.plantilla_base_id:
        tipo = documento.plantilla_base.get_tipo_display()

    if documento.faena_id:
        faena = documento.faena.faena

    return f"Detalle - {tipo} - {faena}"


def _obtener_nombre_archivo_desde_url(url):
    if not url:
        return 'Archivo'
    return unquote(str(url).split('?')[0].split('#')[0].rstrip('/').split('/')[-1]) or 'Archivo'


def _formatear_texto_para_pdf(texto, longitud_bloque=10):
    valor = str(texto or '').strip()
    if not valor:
        return '-'

    resultado = []
    for i, caracter in enumerate(valor):
        if i > 0 and i % longitud_bloque == 0:
            resultado.append(' ')
        resultado.append(caracter)

    return ''.join(resultado)


def _formatear_texto_para_pdf_html(texto, longitud_bloque=10):
    valor = _formatear_texto_para_pdf(texto, longitud_bloque=longitud_bloque)
    if valor == '-':
        return valor
    return mark_safe(valor.replace(' ', '<br/>'))


def _es_extension_imagen_previsualizable(extension):
    return extension.lower() in {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp', '.svg'}


def _resolver_src_pdf(uri, rel):
    if not uri:
        return uri

    if uri.startswith('http://') or uri.startswith('https://'):
        return uri

    if os.path.isabs(uri) and os.path.exists(uri):
        return uri

    if uri.startswith(settings.MEDIA_URL):
        path = os.path.join(settings.MEDIA_ROOT, uri.replace(settings.MEDIA_URL, '', 1))
        return path

    if uri.startswith(settings.STATIC_URL):
        path = os.path.join(settings.STATIC_ROOT, uri.replace(settings.STATIC_URL, '', 1))
        if os.path.exists(path):
            return path

    return uri


def _resolver_ruta_local_adjunto_pdf(url_archivo):
    # Convierte URL de media a ruta local para poder previsualizar/convertir PDFs.
    url_archivo = str(url_archivo or '').strip()
    if not url_archivo:
        return None

    # Caso directo en disco.
    if os.path.isabs(url_archivo) and os.path.exists(url_archivo):
        return url_archivo

    parsed = urlparse(url_archivo)
    ruta = parsed.path or url_archivo

    media_url = str(getattr(settings, 'MEDIA_URL', '') or '')
    if media_url and ruta.startswith(media_url):
        relativo = ruta.replace(media_url, '', 1).lstrip('/')
        candidato = os.path.join(settings.MEDIA_ROOT, relativo)
        return candidato if os.path.exists(candidato) else None

    # Fallback por nombre dentro de MEDIA_ROOT.
    nombre = _obtener_nombre_archivo_desde_url(url_archivo)
    if nombre:
        candidato = os.path.join(settings.MEDIA_ROOT, nombre)
        if os.path.exists(candidato):
            return candidato

    return None


def _ruta_local_a_media_url(ruta_local):
    ruta_local = str(ruta_local or '').strip()
    if not ruta_local:
        return ''

    media_root = os.path.abspath(str(getattr(settings, 'MEDIA_ROOT', '') or ''))
    media_url = str(getattr(settings, 'MEDIA_URL', '/media/') or '/media/')
    abs_path = os.path.abspath(ruta_local)

    if media_root and abs_path.startswith(media_root):
        relativo = os.path.relpath(abs_path, media_root).replace(os.sep, '/')
        return f"{media_url.rstrip('/')}/{relativo}"

    marcador = f"{os.sep}media{os.sep}"
    if marcador in abs_path:
        relativo = abs_path.split(marcador, 1)[1].replace(os.sep, '/')
        return f"{media_url.rstrip('/')}/{relativo}"

    return ''


def _construir_archivo_evidencia(archivo_field, request, es_pdf_view=False):
    if not archivo_field:
        return None

    try:
        archivo_url = archivo_field.url
    except Exception:
        return None

    nombre_archivo = _obtener_nombre_archivo_desde_url(archivo_url)
    extension = os.path.splitext(nombre_archivo)[1].lower()
    es_imagen = _es_extension_imagen_previsualizable(extension)
    es_pdf_mismo = extension == '.pdf'
    url_descarga = request.build_absolute_uri(archivo_url) if request else archivo_url

    archivo_data = {
        'nombre_archivo': nombre_archivo,
        'nombre_archivo_html': _formatear_texto_para_pdf_html(nombre_archivo, longitud_bloque=50),
        'extension': extension or 'sin extension',
        'url_descarga': url_descarga,
        'url_descarga_html': _formatear_texto_para_pdf_html(url_descarga, longitud_bloque=65),
    }

    if es_imagen:
        archivo_data['tipo_visualizacion'] = 'imagen'
        try:
            archivo_data['archivo_src'] = archivo_field.path
        except Exception:
            archivo_data['tipo_visualizacion'] = 'no_previsualizable'
            archivo_data['mensaje_visualizacion'] = (
                'Este archivo no es compartible para vista previa en el PDF. '
                'Usa el enlace de descarga para revisarlo.'
            )
    elif es_pdf_mismo and es_pdf_view:
        try:
            paginas = check_and_convert_pdf(archivo_field.path)
            archivo_data['tipo_visualizacion'] = 'pdf_multiple'
            archivo_data['archivo_src_multiple'] = paginas
        except Exception as e:
            print("Error al convertir PDF en Prevencion:", e)
            archivo_data['tipo_visualizacion'] = 'no_previsualizable'
            archivo_data['mensaje_visualizacion'] = (
                'Ocurrió un error al procesar las páginas de este PDF.'
            )
    else:
        archivo_data['tipo_visualizacion'] = 'no_previsualizable'
        archivo_data['mensaje_visualizacion'] = (
            'Este archivo no es compartible para vista previa en el PDF. '
            'Usa el enlace de descarga para revisarlo.'
        )

    return archivo_data


def _construir_anexo_evidencias(documento, contenido_documento, request, es_pdf_view=False):
    anexos = []
    indice = 1
    evidencias_secciones = list(documento.evidencias_secciones.all().order_by('id'))
    evidencias_items = list(documento.evidencias_items.all().order_by('id'))
    evidencias_generales = list(documento.evidencias_generales.all().order_by('id'))

    def agregar_grupo(titulo, origen, archivos):
        nonlocal indice
        if not archivos:
            return

        anexos.append({
            'indice': indice,
            'titulo': titulo,
            'origen': origen,
            'archivos': archivos,
        })
        indice += 1

    for sec_idx, seccion in enumerate(contenido_documento):
        archivos_seccion = [
            _construir_archivo_evidencia(evidencia.archivo, request, es_pdf_view)
            for evidencia in evidencias_secciones
            if evidencia.indice_seccion == sec_idx
        ]
        agregar_grupo(
            f"Evidencia de Sección {sec_idx + 1}: {seccion.get('seccion') or 'Sin nombre'}",
            'Sección',
            [archivo for archivo in archivos_seccion if archivo],
        )

        for row_idx, fila in enumerate(seccion.get('respuestas', [])):
            archivos_item = [
                _construir_archivo_evidencia(evidencia.archivo, request, es_pdf_view)
                for evidencia in evidencias_items
                if evidencia.indice_seccion == sec_idx and evidencia.indice_fila == row_idx
            ]
            agregar_grupo(
                f"Evidencia Item {sec_idx + 1}.{row_idx + 1}: {fila.get('pregunta') or 'Sin pregunta'}",
                'Item',
                [archivo for archivo in archivos_item if archivo],
            )

    if evidencias_generales:
        archivos_generales = [
            _construir_archivo_evidencia(evidencia.archivo, request, es_pdf_view)
            for evidencia in evidencias_generales
        ]
        agregar_grupo(
            'Evidencia General',
            'Cierre del Reporte',
            [archivo for archivo in archivos_generales if archivo],
        )
    elif documento.evidencia_general:
        archivo_legacy = _construir_archivo_evidencia(documento.evidencia_general, request, es_pdf_view)
        agregar_grupo(
            'Evidencia General',
            'Cierre del Reporte',
            [archivo_legacy] if archivo_legacy else [],
        )

    return anexos


def _preparar_documento_para_visualizacion(documento):
    evidencias_seccion_map = defaultdict(list)
    for evidencia in documento.evidencias_secciones.all().order_by('id'):
        evidencias_seccion_map[evidencia.indice_seccion].append(evidencia.archivo.url)

    item_evidencias_map = defaultdict(list)
    for evidencia in documento.evidencias_items.all().order_by('id'):
        item_evidencias_map[(evidencia.indice_seccion, evidencia.indice_fila)].append(evidencia.archivo.url)

    contenido = json.loads(json.dumps(documento.contenido or []))
    for sec_idx, seccion in enumerate(contenido):
        seccion['evidencias_urls'] = evidencias_seccion_map.get(sec_idx, [])
        seccion['evidencias_nombres'] = [
            _formatear_texto_para_pdf_html(_obtener_nombre_archivo_desde_url(url), longitud_bloque=40)
            for url in seccion['evidencias_urls']
        ]
        if 'respuestas' in seccion:
            for row_idx, fila in enumerate(seccion['respuestas']):
                fila['evidencia_urls'] = item_evidencias_map.get((sec_idx, row_idx), [])
                fila['evidencia_nombres'] = [
                    _formatear_texto_para_pdf_html(_obtener_nombre_archivo_desde_url(url))
                    for url in fila['evidencia_urls']
                ]

    return contenido


def _opciones_vigilancia(modelo):
    return modelo.objects.filter(status=True).order_by('correlativo', 'id')


def _catalogos_vigilancia_formulario():
    return {
        'opciones_ges': _opciones_vigilancia(VigilanciaGes),
        'opciones_area': _opciones_vigilancia(VigilanciaArea),
        'opciones_cargo': _opciones_vigilancia(VigilanciaCargo),
        'opciones_tipo_contrato': _opciones_vigilancia(VigilanciaTipoContrato),
        'opciones_contrato': _opciones_vigilancia(VigilanciaContrato),
        'opciones_eval_riesgo': _opciones_vigilancia(VigilanciaEvaluacionRiesgo),
        'opciones_nivel_riesgo': _opciones_vigilancia(VigilanciaNivelRiesgo),
        'opciones_grado_exposicion': _opciones_vigilancia(VigilanciaGradoExposicion),
        'opciones_nivel_seguimiento': _opciones_vigilancia(VigilanciaNivelSeguimiento),
        'opciones_exposicion': _opciones_vigilancia(VigilanciaExposicion),
    }


def _normalizar_valor_vigilancia(valor, es_numerico=False):
    valor_limpio = str(valor or '').strip()
    if not valor_limpio:
        raise ValueError('Debes ingresar un valor.')

    if not es_numerico:
        return valor_limpio

    valor_numerico = valor_limpio.replace(',', '.')
    try:
        numero = float(valor_numerico)
    except (TypeError, ValueError):
        raise ValueError('Este campo debe ser numerico.')

    if numero.is_integer():
        return str(int(numero))

    return str(numero).rstrip('0').rstrip('.')


def _obtener_archivos_subidos(request):
    archivos = []
    for key in request.FILES:
        archivos.extend(request.FILES.getlist(key))
    return archivos


def _validar_archivos_subidos(request):
    for key in request.FILES:
        archivos = request.FILES.getlist(key)
        if len(archivos) > MAX_ARCHIVOS_POR_INPUT:
            return False, f"No puedes subir mas de {MAX_ARCHIVOS_POR_INPUT} archivos por cada evidencia."

        for archivo in archivos:
            if getattr(archivo, 'size', 0) > MAX_TAMANO_ARCHIVO_BYTES:
                return False, "No puede excederse de los 100 MB por archivo."

    return True, None


def _contar_archivos_existentes_por_input(documento, input_name):
    if input_name == 'evidencia_general':
        existentes = documento.evidencias_generales.count()
        if existentes == 0 and documento.evidencia_general:
            existentes = 1
        return existentes

    if input_name.startswith('file_seccion_'):
        try:
            indice = int(input_name.split('_')[-1])
        except (TypeError, ValueError):
            return 0
        return documento.evidencias_secciones.filter(indice_seccion=indice).count()

    if input_name.startswith('file_item_'):
        parts = input_name.split('_')
        if len(parts) < 4:
            return 0
        try:
            sec_idx = int(parts[2])
            row_idx = int(parts[3])
        except (TypeError, ValueError):
            return 0
        return documento.evidencias_items.filter(
            indice_seccion=sec_idx,
            indice_fila=row_idx
        ).count()

    return 0


def _validar_total_archivos_por_item_en_edicion(documento, request):
    for key in request.FILES:
        archivos_nuevos = request.FILES.getlist(key)
        existentes = _contar_archivos_existentes_por_input(documento, key)

        if existentes + len(archivos_nuevos) > MAX_ARCHIVOS_POR_INPUT:
            restantes = max(0, MAX_ARCHIVOS_POR_INPUT - existentes)
            if restantes == 0:
                return False, f"Este registro ya tiene {MAX_ARCHIVOS_POR_INPUT} archivos. Primero elimina uno para subir otro."
            return False, (
                f"Solo puedes subir {restantes} archivo(s) mas en esta evidencia "
                f"(maximo {MAX_ARCHIVOS_POR_INPUT} en total)."
            )

    return True, None


def _contar_adjuntos_vigilancia(vigilancia):
    total = 1 if vigilancia.documento_adjunto else 0
    prefetched_adjuntos = getattr(vigilancia, '_prefetched_objects_cache', {}).get('adjuntos_vigilancia')
    if prefetched_adjuntos is not None:
        return total + len(prefetched_adjuntos)
    return total + vigilancia.adjuntos_vigilancia.count()


def _validar_total_adjuntos_vigilancia_en_edicion(vigilancia, request):
    archivos_nuevos = request.FILES.getlist('documento')
    if not archivos_nuevos:
        return True, None

    existentes = _contar_adjuntos_vigilancia(vigilancia)
    if existentes + len(archivos_nuevos) > MAX_ARCHIVOS_POR_INPUT:
        restantes = max(0, MAX_ARCHIVOS_POR_INPUT - existentes)
        if restantes == 0:
            return False, f"Esta ficha ya tiene {MAX_ARCHIVOS_POR_INPUT} archivos. Primero elimina uno para subir otro."
        return False, (
            f"Solo puedes subir {restantes} archivo(s) mas en esta ficha "
            f"(maximo {MAX_ARCHIVOS_POR_INPUT} en total)."
        )

    return True, None


def _serializar_adjuntos_vigilancia(vigilancia):
    archivos = []

    if vigilancia.documento_adjunto:
        archivos.append({
            'id': f'legacy-{vigilancia.pk}',
            'url': vigilancia.documento_adjunto.url,
            'name': os.path.basename(vigilancia.documento_adjunto.name or '') or 'Archivo',
            'delete_url': reverse('delete_vigilancia_doc', args=[vigilancia.pk]),
        })

    for adjunto in vigilancia.adjuntos_vigilancia.all():
        if not adjunto.archivo:
            continue
        archivos.append({
            'id': adjunto.id,
            'url': adjunto.archivo.url,
            'name': os.path.basename(adjunto.archivo.name or '') or 'Archivo',
            'delete_url': reverse('delete_vigilancia_adjunto', args=[adjunto.pk]),
        })

    return archivos


def _to_float_or_none(value):
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text:
        return None

    text = text.replace('%', '').replace(',', '.')
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def _calcular_porcentaje_promedio_documento(contenido):
    if not isinstance(contenido, list):
        return 0.0

    valores = []
    for seccion in contenido:
        if not isinstance(seccion, dict):
            continue

        respuestas = seccion.get('respuestas', [])
        if not isinstance(respuestas, list):
            continue

        for fila in respuestas:
            if not isinstance(fila, dict):
                continue

            cumple = _to_float_or_none(fila.get('cumple'))
            if cumple is None:
                continue

            cumple = max(0.0, min(100.0, cumple))
            valores.append(cumple)

    if not valores:
        return 0.0

    return round(sum(valores) / len(valores), 1)


def _mensaje_actualizacion_prevencion(tipo_slug):
    mensajes = {
        'condiciones': 'Condicion general actualizada correctamente.',
        'silice': 'Silice actualizada correctamente.',
        'ruido': 'Ruido actualizado correctamente.',
        'hipobaria': 'Hipobaria actualizada correctamente.',
        'psicosocial': 'Psicosocial actualizado correctamente.',
        'tmert': 'TMERT actualizado correctamente.',
        'mmc': 'MMC actualizado correctamente.',
        'charlas': 'Charla actualizada correctamente.',
        'informativos': 'Informativo actualizado correctamente.',
        'difusiones': 'Difusion actualizada correctamente.',
        'cursos': 'Curso actualizado correctamente.',
    }
    return mensajes.get(tipo_slug, 'Plantilla actualizada correctamente.')

@login_required
@prevencion_mantenedor_required
def manage_prevencion(request, tipo_slug):

    if tipo_slug not in [v.value for v in PrevencionPlantilla.TipoRiesgo]:
        messages.error(request, "Tipo de prevención no válido.")
        return redirect('dashboardPrevencion')

    tipos_map = {v.value: v.label for v in PrevencionPlantilla.TipoRiesgo}
    boton_nuevo_texto_map = {
        'condiciones': 'Nueva Condicion General',
        'silice': 'Nuevo Silice',
        'ruido': 'Nuevo Ruido',
        'hipobaria': 'Nueva Hipobaria',
        'psicosocial': 'Nuevo Psicosocial',
        'tmert': 'Nuevo TMERT',
        'mmc': 'Nuevo MMC',
        'charlas': 'Nueva Charla',
        'informativos': 'Nuevo Informativo',
        'difusiones': 'Nueva Difusion',
        'cursos': 'Nuevo Curso',
    }
    boton_nuevo_texto = boton_nuevo_texto_map.get(
        tipo_slug,
        f"Nuevo {tipos_map.get(tipo_slug, tipo_slug.title())}"
    )

    plantillas_qs = PrevencionPlantilla.objects.filter(
        tipo=tipo_slug,
        status=True
    ).select_related('faena').order_by('-fechacreacion')

    plantillas_eliminadas_qs = PrevencionPlantilla.objects.filter(
        tipo=tipo_slug,
        status=False
    ).select_related('faena').order_by('-fechacreacion')

    plantillas = list(plantillas_qs)
    plantillas_eliminadas = list(plantillas_eliminadas_qs)
    
    cargo_por_plantilla = {}
    plantillas_con_documentos_ids = set()

    if plantillas and tipo_slug in TIPOS_DOCUMENTO_PERMITIDOS:
        plantillas_con_documentos_ids = set(
            PrevencionDocumento.objects.filter(
                plantilla_base_id__in=[p.id for p in plantillas]
            ).values_list('plantilla_base_id', flat=True).distinct()
        )

    permisos_qs = PrevencionPermisoLlenado.objects.filter(
        status=True,
        plantilla_id__in=[p.id for p in plantillas]
    ).prefetch_related('cargos')

    for permiso in permisos_qs:
        cargos = list(
            permiso.cargos.order_by('correlativo', 'id').values_list('valor', flat=True)
        )
        cargo_por_plantilla[permiso.plantilla_id] = ', '.join(cargos) if cargos else 'Todos'

    for plantilla in plantillas:
        plantilla.cargo_display = cargo_por_plantilla.get(plantilla.id, 'Todos')
        plantilla.edicion_bloqueada = plantilla.id in plantillas_con_documentos_ids
        
    for p_elim in plantillas_eliminadas:
        p_elim.cargo_display = 'Todos'
    
    active_tab = (request.GET.get('tab') or '').strip().lower()
    if active_tab not in {'activos', 'eliminados'}:
        active_tab = 'activos'
    
    context = {
        'titulo_categoria': tipos_map.get(tipo_slug, tipo_slug.title()),
        'slug_categoria': tipo_slug,
        'plantillas': plantillas,
        'plantillas_eliminadas': plantillas_eliminadas,
        'boton_nuevo_texto': boton_nuevo_texto,
        'mostrar_boton_nuevo': True,
        'habilitar_duplicado_docs': tipo_slug in TIPOS_DOCUMENTO_PERMITIDOS,
        'active_tab': active_tab,
        'sidebar': 'mantenedor_prev',
        'sidebarmenu': f't-{tipo_slug}'
    }
    return render(request, 'pages/prevencion/manage_prevencion.html', context)

@login_required
@capacitaciones_y_docs_required
def manage_capacitaciones(request, tipo_slug):
    tipo_slug = str(tipo_slug or '').strip().lower()
    if tipo_slug not in TIPOS_DOCUMENTO_PERMITIDOS:
        messages.error(request, "Tipo de capacitación no válido.")
        return redirect('dashboardPrevencion')

    tipos_map = {v.value: v.label for v in PrevencionPlantilla.TipoRiesgo}
    documentos_visibles_ids = _ids_documentos_visibles_para_usuario(
        user=request.user,
        tipo_slug=tipo_slug,
    )

    documentos_qs = (
        PrevencionDocumento.objects
        .filter(
            status=True,
            plantilla_base__tipo=tipo_slug,
            id__in=documentos_visibles_ids,
        )
        .filter(
            difusion_trabajadores__status=True,
            difusion_trabajadores__user=request.user,
        )
        .select_related('faena', 'plantilla_base')
        .order_by('-fecha_creacion', '-id')
        .distinct()
    )

    estado_leido_capacitacion_label = (
        'Validado'
        if tipo_slug in {
            PrevencionPlantilla.TipoRiesgo.CHARLAS.value,
            PrevencionPlantilla.TipoRiesgo.INFORMATIVOS.value,
            PrevencionPlantilla.TipoRiesgo.DIFUSIONES.value,
        }
        else 'Leído'
    )

    documentos_ids = list(documentos_qs.values_list('id', flat=True))
    ultima_difusion_por_documento = {}
    estado_notificacion_por_documento = {}
    estado_curso_por_documento = {}
    
    if documentos_ids:
        difusiones_qs = (
            PrevencionDocumentoDifusionTrabajador.objects
            .filter(status=True, documento_id__in=documentos_ids, user=request.user)
            .order_by('documento_id', '-fechacreacion', '-id')
            .values('documento_id', 'fechacreacion', 'creado_por', 'motivo_rechazo')
        )
        for registro_difusion in difusiones_qs:
            documento_id = int(registro_difusion.get('documento_id') or 0)
            if documento_id and documento_id not in ultima_difusion_por_documento:
                ultima_difusion_por_documento[documento_id] = registro_difusion

        notificaciones_qs = (
            PrevencionNotificacionDocumento.objects
            .filter(destinatario=request.user, documento_id__in=documentos_ids)
            .order_by('documento_id', '-fechacreacion', '-id')
            .values('documento_id', 'estado', 'leida')
        )
        for notificacion in notificaciones_qs:
            documento_id = int(notificacion.get('documento_id') or 0)
            if not documento_id or documento_id in estado_notificacion_por_documento:
                continue
            estado_raw = str(notificacion.get('estado') or '').strip().lower()
            leida = bool(notificacion.get('leida'))
            if leida or estado_raw == PrevencionNotificacionDocumento.Estado.VISTA:
                estado_notificacion_por_documento[documento_id] = {
                    'estado_label': estado_leido_capacitacion_label,
                    'estado_key': 'leido',
                }
            else:
                estado_notificacion_por_documento[documento_id] = {
                    'estado_label': 'Entregado',
                    'estado_key': 'entregado',
                }

        if tipo_slug == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
            resultados_qs = (
                PrevencionResultadoCurso.objects
                .filter(documento_id__in=documentos_ids, trabajador=request.user)
                .select_related('documento__plantilla_base')
                .order_by('documento_id', '-updated_at', '-id')
            )
            for resultado in resultados_qs:
                documento_id = int(getattr(resultado, 'documento_id', 0) or 0)
                if not documento_id or documento_id in estado_curso_por_documento:
                    continue
                documento_resultado = getattr(resultado, 'documento', None)
                if not documento_resultado:
                    continue
                estado_curso_por_documento[documento_id] = _normalizar_estado_resultado_curso(
                    documento_resultado,
                    resultado,
                )

    estado_bloqueo = _obtener_estado_bloqueo_capacitacion(request)
    curso_finalizado_bloqueado = _obtener_curso_finalizado_bloqueado(request, estado_bloqueo)
    documento_curso_finalizado_id = int(curso_finalizado_bloqueado.get('documento_id') or 0) if curso_finalizado_bloqueado else 0
    documento_activo_id = int(estado_bloqueo.get('documento_id') or 0) if estado_bloqueo else 0
    segundos_restantes_activo = max(0, int(estado_bloqueo.get('remaining_seconds', 0) or 0)) if estado_bloqueo else 0
    tipo_activo_bloqueo = str(estado_bloqueo.get('tipo_slug') or '').strip().lower() if estado_bloqueo else ''
    
    mantener_activo_sin_tiempo = bool(
        estado_bloqueo
        and segundos_restantes_activo <= 0
        and (
            tipo_activo_bloqueo in TIPOS_CAPACITACION_LIBRE_TRAS_VALIDACION
            or tipo_activo_bloqueo == PrevencionPlantilla.TipoRiesgo.CURSOS.value
        )
    )
    bloqueo_vigente = bool(
        estado_bloqueo
        and (segundos_restantes_activo > 0 or mantener_activo_sin_tiempo)
    )

    documentos_difundidos = []
    cargo_por_documento = {}
    for documento in documentos_qs:
        if documento.id not in cargo_por_documento:
            cargo_display = 'Todos'
            if not documento.plantilla_base_id:
                cargo_display = '-'
            else:
                permiso = _permiso_llenado_para_documento(documento)
                if permiso:
                    cargos_permiso = list(
                        permiso.cargos.filter(status=True)
                        .order_by('correlativo', 'id')
                        .values_list('valor', flat=True)
                    )
                    cargo_display = ', '.join(cargos_permiso) if cargos_permiso else 'Todos'
            cargo_por_documento[documento.id] = cargo_display

        registro_difusion = ultima_difusion_por_documento.get(documento.id) or {}
        fecha_difusion = registro_difusion.get('fechacreacion') or documento.fecha_creacion
        difundido_por = _resolver_nombre_difusor(
            registro_difusion.get('creado_por') or documento.creador
        ) or '-'
        
        motivo_rechazo = registro_difusion.get('motivo_rechazo') or ''
        estado_notificacion = estado_notificacion_por_documento.get(documento.id) or {
            'estado_label': 'Entregado',
            'estado_key': 'entregado',
        }
        
        estado_curso = None
        intentos_restantes_curso = None
        bloqueo_apertura_motivo = ''
        puede_abrir = True

        if estado_notificacion.get('estado_key') == 'leido' and tipo_slug != PrevencionPlantilla.TipoRiesgo.CURSOS.value:
            puede_abrir = False
            bloqueo_apertura_motivo = 'Ya validaste esta capacitación.'

        if tipo_slug == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
            estado_curso = estado_curso_por_documento.get(documento.id)
            if not estado_curso:
                estado_curso = _normalizar_estado_resultado_curso(documento, None)

            intentos_restantes_curso = int(estado_curso.get('intentos_restantes') or 0)
            if estado_curso.get('estado_key') in {'aprobado', 'rechazado'}:
                estado_notificacion = {
                    'estado_label': estado_curso.get('estado_label') or 'Entregado',
                    'estado_key': estado_curso.get('estado_key') or 'entregado',
                }

            if bool(estado_curso.get('aprobado')):
                puede_abrir = False
                bloqueo_apertura_motivo = 'No puedes volver a realizar este curso porque ya está aprobado.'
            elif (
                int(estado_curso.get('intentos_realizados') or 0) > 0
                and intentos_restantes_curso <= 0
            ):
                puede_abrir = False
                bloqueo_apertura_motivo = 'No puedes volver a realizar este curso porque ya agotaste tus intentos.'

        if motivo_rechazo and tipo_slug != PrevencionPlantilla.TipoRiesgo.CURSOS.value:
            puede_abrir = False
            bloqueo_apertura_motivo = 'No puedes abrir este documento porque tu validación de identidad fue rechazada.'
            estado_notificacion = {
                'estado_label': 'Rechazado',
                'estado_key': 'rechazado',
            }

        if bloqueo_vigente and documento.id == documento_activo_id:
            if (
                tipo_slug == PrevencionPlantilla.TipoRiesgo.CURSOS.value
                and documento.id == documento_curso_finalizado_id
            ):
                if estado_curso and estado_curso.get('estado_key') in {'aprobado', 'rechazado'}:
                    estado_notificacion = {
                        'estado_label': estado_curso.get('estado_label') or 'Entregado',
                        'estado_key': estado_curso.get('estado_key') or 'entregado',
                    }
                else:
                    estado_notificacion = {
                        'estado_label': estado_leido_capacitacion_label,
                        'estado_key': 'leido',
                    }
            else:
                if not (estado_curso and bool(estado_curso.get('aprobado'))):
                    estado_notificacion = {
                        'estado_label': 'Activo',
                        'estado_key': 'activo',
                    }
                    puede_abrir = True  
                    bloqueo_apertura_motivo = ''

        documentos_difundidos.append({
            'registro_id': None,
            'documento_id': documento.id,
            'fecha_difusion': fecha_difusion,
            'tipo_label': documento.plantilla_base.get_tipo_display() if documento.plantilla_base_id else '-',
            'faena_nombre': documento.faena.faena if documento.faena_id else '-',
            'cargo_display': cargo_por_documento.get(documento.id, '-'),
            'plantilla_nombre': documento.plantilla_base.nombre if documento.plantilla_base_id else 'Plantilla eliminada',
            'difundido_por': difundido_por,
            'estado_notificacion': estado_notificacion['estado_label'],
            'estado_notificacion_key': estado_notificacion['estado_key'],
            'curso_intentos_restantes': intentos_restantes_curso,
            'puede_abrir': puede_abrir,
            'bloqueo_apertura_motivo': bloqueo_apertura_motivo,
            'motivo_rechazo': motivo_rechazo,
        })

    alerta_capacitacion_activa = None
    if bloqueo_vigente:
        documento_activo_id = int(estado_bloqueo.get('documento_id') or 0)
        documentos_visibles_usuario = _ids_documentos_visibles_para_usuario(request.user)
        if documento_activo_id and documento_activo_id in documentos_visibles_usuario:
            documento_activo = (
                PrevencionDocumento.objects
                .filter(id=documento_activo_id, status=True)
                .select_related('plantilla_base')
                .first()
            )
            if documento_activo:
                tipo_activo = str(
                    getattr(documento_activo.plantilla_base, 'tipo', '') or ''
                ).strip().lower()
                params_volver = {'origen': 'capacitaciones'}
                if tipo_activo in TIPOS_DOCUMENTO_PERMITIDOS:
                    params_volver['tipo'] = tipo_activo
                params_volver['next'] = request.get_full_path()

                alerta_capacitacion_activa = {
                    'documento_id': documento_activo.id,
                    'remaining_seconds': max(0, int(estado_bloqueo.get('remaining_seconds', 0) or 0)),
                    'curso_finalizado': bool(
                        tipo_activo == PrevencionPlantilla.TipoRiesgo.CURSOS.value
                        and documento_activo.id == documento_curso_finalizado_id
                    ),
                    'nombre_documento_activo': (
                        documento_activo.plantilla_base.nombre
                        if documento_activo.plantilla_base_id else
                        f"Documento #{documento_activo.id}"
                    ),
                    'tipo_activo_slug': tipo_activo,
                    'url_volver_activo': (
                        f"{reverse('view_documento', kwargs={'pk': documento_activo.id})}"
                        f"?{urlencode(params_volver)}"
                    ),
                }

    context = {
        'titulo_categoria': tipos_map.get(tipo_slug, tipo_slug.title()),
        'slug_categoria': tipo_slug,
        'documentos_difundidos': documentos_difundidos,
        'alerta_capacitacion_activa': alerta_capacitacion_activa,
        'sidebarmain': 'capacitaciones',
        'sidebar': 'historial_doc',
        'sidebarmenu': f't-{tipo_slug}',
    }
    return render(request, 'pages/prevencion/manage_capacitaciones.html', context)

@login_required
@prevencion_mantenedor_required
def edit_prevencion(request, tipo_slug, pk=None):
    def _to_int_or_none(value):
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    tipos_map = {v.value: v.label for v in PrevencionPlantilla.TipoRiesgo}
    tipos_solo_items = {
        PrevencionPlantilla.TipoRiesgo.CHARLAS.value,
        PrevencionPlantilla.TipoRiesgo.INFORMATIVOS.value,
        PrevencionPlantilla.TipoRiesgo.DIFUSIONES.value,
    }
    tipos_con_modo_pdf = tipos_solo_items
    permitir_columnas_extras = tipo_slug not in tipos_solo_items
    habilitar_modo_pdf = tipo_slug in tipos_con_modo_pdf
    es_curso = tipo_slug == PrevencionPlantilla.TipoRiesgo.CURSOS.value
    duplicate_source_id = ''
    duplicate_source = None
    selected_cargo_id_post = None
    selected_autorizadores_ids_post = None
    tiempos_estimados_disponibles = [3] + list(range(10, 181, 10))
    curso_intentos_disponibles = list(CURSO_TOTAL_INTENTOS_DISPONIBLES)
    curso_porcentajes_aprobacion_disponibles = list(CURSO_PORCENTAJES_APROBACION_DISPONIBLES)

    if not pk:
        duplicate_source_id = (request.POST.get('duplicate_from') or request.GET.get('duplicate_from') or '').strip()
        if duplicate_source_id:
            duplicate_source = get_object_or_404(
                PrevencionPlantilla,
                pk=duplicate_source_id,
                status=True,
                tipo=tipo_slug
            )

    duplicate_source_cargo_id = None
    permiso_origen = None
    if duplicate_source and duplicate_source.faena_id:
        permiso_origen = PrevencionPermisoLlenado.objects.filter(
            status=True,
            faena_id=duplicate_source.faena_id,
            plantilla=duplicate_source
        ).first()
    if duplicate_source and not permiso_origen:
        permiso_origen = PrevencionPermisoLlenado.objects.filter(
            status=True,
            plantilla=duplicate_source
        ).order_by('-updated_at', '-id').first()
    if permiso_origen:
        duplicate_source_cargo_id = permiso_origen.cargos.values_list('id', flat=True).first()
    if pk:
        # === SEGURIDAD ===
        # Agregamos status=True a la consulta. 
        # Si el usuario pone un ID de una plantilla eliminada (status=False),
        # get_object_or_404 lanzará un error 404 en lugar de cargarla.
        # También validamos que el tipo coincida con la URL para evitar mezclas.
        plantilla = get_object_or_404(PrevencionPlantilla, pk=pk, status=True, tipo=tipo_slug)
        titulo = f"Editar Plantilla: {plantilla.nombre}"
    else:
        # Crear nueva
        plantilla = PrevencionPlantilla(tipo=tipo_slug)
        if duplicate_source:
            titulo = f"Duplicar Plantilla: {duplicate_source.nombre}"
            if request.method != 'POST':
                plantilla.nombre = duplicate_source.nombre
                plantilla.faena_id = duplicate_source.faena_id
                plantilla.tiempo_estimado_minutos = duplicate_source.tiempo_estimado_minutos
                plantilla.curso_total_intentos = duplicate_source.curso_total_intentos
                plantilla.curso_porcentaje_aprobacion = duplicate_source.curso_porcentaje_aprobacion
                if habilitar_modo_pdf:
                    plantilla.modo_contenido = duplicate_source.modo_contenido
                    if duplicate_source.archivo_pdf:
                        plantilla.archivo_pdf = duplicate_source.archivo_pdf
                plantilla.estructura = json.loads(
                    json.dumps(duplicate_source.estructura or [], cls=DjangoJSONEncoder)
                )
        else:
            titulo = f"Nueva Plantilla: {tipos_map.get(tipo_slug, 'Prevención')}"

    if es_curso and request.method != 'POST':
        # Solo aplicar default automático en cursos ya existentes o duplicados.
        # En creación nueva se mantiene "Seleccione..." para obligar selección explícita.
        aplicar_default_curso = bool(plantilla.id or duplicate_source)
        if aplicar_default_curso:
            if plantilla.curso_total_intentos not in curso_intentos_disponibles:
                plantilla.curso_total_intentos = CURSO_TOTAL_INTENTOS_DEFAULT
            if plantilla.curso_porcentaje_aprobacion not in curso_porcentajes_aprobacion_disponibles:
                plantilla.curso_porcentaje_aprobacion = CURSO_PORCENTAJE_APROBACION_DEFAULT

    snapshot_pre_reautorizacion = (
        _snapshot_plantilla_para_reautorizacion(plantilla)
        if pk and tipo_slug in TIPOS_DOCUMENTO_PERMITIDOS
        else {}
    )

    # Bloque POST
    if request.method == 'POST':
        try:
            plantilla.nombre = (request.POST.get('nombre') or '').strip()
            plantilla.faena_id = _to_int_or_none(request.POST.get('faena'))
            plantilla.tiempo_estimado_minutos = _to_int_or_none(request.POST.get('tiempo_estimado_minutos'))
            if es_curso:
                plantilla.curso_total_intentos = _to_int_or_none(request.POST.get('curso_total_intentos'))
                plantilla.curso_porcentaje_aprobacion = _to_int_or_none(request.POST.get('curso_porcentaje_aprobacion'))
            else:
                plantilla.curso_total_intentos = None
                plantilla.curso_porcentaje_aprobacion = None
            cargo_id = _to_int_or_none(request.POST.get('cargo'))
            selected_cargo_id_post = cargo_id
            if not plantilla.creador:
                plantilla.creador = request.user.username 
            
            modo_contenido_post = (request.POST.get('modo_contenido') or '').strip()
            if habilitar_modo_pdf and modo_contenido_post == PrevencionPlantilla.ModoContenido.PDF:
                plantilla.modo_contenido = PrevencionPlantilla.ModoContenido.PDF
            else:
                plantilla.modo_contenido = PrevencionPlantilla.ModoContenido.ESTRUCTURA

            plantilla.autorizacion_incompleta = _parsear_bool(
                request.POST.get('has_unselected_autorizadores', False)
            )

            error_validacion = False
            if (
                plantilla.tiempo_estimado_minutos is not None
                and plantilla.tiempo_estimado_minutos not in tiempos_estimados_disponibles
            ):
                messages.error(request, "El tiempo estimado seleccionado no es válido.")
                error_validacion = True
            if es_curso:
                if plantilla.curso_total_intentos not in curso_intentos_disponibles:
                    messages.error(request, "El total de intentos seleccionado no es válido.")
                    error_validacion = True
                if plantilla.curso_porcentaje_aprobacion not in curso_porcentajes_aprobacion_disponibles:
                    messages.error(request, "El porcentaje de aprobación seleccionado no es válido.")
                    error_validacion = True

            autorizadores_ids = []
            autorizadores_validos_ids = []
            autorizadores_raw = (request.POST.get('autorizador_ids') or '').strip()
            if autorizadores_raw:
                seen_autorizadores = set()
                for raw_id in autorizadores_raw.split(','):
                    uid = _to_int_or_none((raw_id or '').strip())
                    if uid and uid not in seen_autorizadores:
                        seen_autorizadores.add(uid)
                        autorizadores_ids.append(uid)
            selected_autorizadores_ids_post = list(autorizadores_ids)

            if len(autorizadores_ids) > 5:
                messages.error(request, "Solo se permiten hasta 5 autorizadores en Autorización.")
                error_validacion = True

            if autorizadores_ids:
                autorizadores_validos_ids = list(
                    User.objects.filter(
                        id__in=autorizadores_ids, 
                        is_active=True,
                        usuarioprofile__seccionPrevencion='ADMINISTRADOR'
                    ).values_list('id', flat=True)
                )
                if len(autorizadores_validos_ids) != len(autorizadores_ids):
                    messages.error(request, "Uno o más autorizadores ya no están disponibles.")
                    error_validacion = True

            archivo_pdf = request.FILES.get('archivo_pdf')
            eliminar_pdf_actual = (request.POST.get('remove_existing_pdf') or '').strip() == '1'
            if (
                duplicate_source
                and not pk
                and habilitar_modo_pdf
                and plantilla.modo_contenido == PrevencionPlantilla.ModoContenido.PDF
                and not archivo_pdf
                and not eliminar_pdf_actual
                and duplicate_source.archivo_pdf
            ):
                # En duplicado (POST), conserva el PDF del origen mientras no se quite ni se reemplace.
                plantilla.archivo_pdf = duplicate_source.archivo_pdf
            pdf_actual = plantilla.archivo_pdf if (plantilla.archivo_pdf and plantilla.archivo_pdf.name) else None
            tiene_pdf_actual = bool(pdf_actual) and not eliminar_pdf_actual

            if plantilla.modo_contenido == PrevencionPlantilla.ModoContenido.PDF:
                if archivo_pdf:
                    nombre_pdf = (archivo_pdf.name or '').lower()
                    if not nombre_pdf.endswith('.pdf'):
                        messages.error(request, "Solo se permiten archivos PDF.")
                        error_validacion = True
                    else:
                        plantilla.archivo_pdf = archivo_pdf
                elif not tiene_pdf_actual:
                    messages.error(request, "Debes adjuntar un PDF para continuar.")
                    error_validacion = True

            estructura_final = []
            if plantilla.modo_contenido != PrevencionPlantilla.ModoContenido.PDF:
                if es_curso:
                    question_indices_str = request.POST.get('question_indices', '')
                    question_ids = [q for q in question_indices_str.split(',') if q]

                    for idx, q_id in enumerate(question_ids, start=1):
                        titulo_pregunta = (request.POST.get(f'pregunta_titulo_{q_id}') or '').strip() or f"pregunta {idx}"
                        enunciado = (request.POST.get(f'pregunta_texto_{q_id}') or '').strip()
                        justificacion = (request.POST.get(f'justificacion_{q_id}') or '').strip()
                        respuesta_correcta = (request.POST.get(f'respuesta_{q_id}') or '').strip().upper()
                        imagen_url = (request.POST.get(f'imagen_actual_{q_id}') or '').strip()
                        imagen_pregunta = request.FILES.get(f'pregunta_imagen_{q_id}')

                        option_indices = [
                            o for o in (request.POST.get(f'option_indices_{q_id}', '') or '').split(',')
                            if o
                        ]
                        opciones = []
                        for o_idx in option_indices:
                            opcion_valor = (request.POST.get(f'opcion_{q_id}_{o_idx}') or '').strip()
                            if opcion_valor:
                                opciones.append(opcion_valor)

                        if not enunciado:
                            messages.error(request, f"La pregunta {idx} debe tener enunciado.")
                            error_validacion = True

                        if len(opciones) < 2:
                            messages.error(request, f"La pregunta {idx} debe tener al menos dos opciones.")
                            error_validacion = True

                        letras_validas = [chr(65 + i) for i in range(len(opciones))]
                        if not respuesta_correcta:
                            messages.error(request, f"La pregunta {idx} debe tener respuesta correcta.")
                            error_validacion = True
                        elif respuesta_correcta not in letras_validas:
                            messages.error(
                                request,
                                f"La respuesta correcta de la pregunta {idx} no coincide con sus opciones."
                            )
                            error_validacion = True

                        if not justificacion:
                            messages.error(request, f"La pregunta {idx} debe tener justificación.")
                            error_validacion = True

                        if imagen_pregunta:
                            nombre_imagen = (imagen_pregunta.name or '').lower()
                            es_imagen_tipo = (getattr(imagen_pregunta, 'content_type', '') or '').lower().startswith('image/')
                            es_imagen_ext = nombre_imagen.endswith(('.jpg', '.jpeg', '.png', '.gif', '.webp', '.bmp'))
                            if not (es_imagen_tipo or es_imagen_ext):
                                messages.error(request, f"La imagen de la pregunta {idx} no es válida. Solo se permiten fotos.")
                                error_validacion = True
                            else:
                                fs = FileSystemStorage()
                                nombre_base, ext = os.path.splitext(os.path.basename(imagen_pregunta.name))
                                ext = (ext or '').lower()
                                if ext not in ('.jpg', '.jpeg', '.png', '.gif', '.webp', '.bmp'):
                                    ext = '.jpg'
                                nombre_seguro = slugify(nombre_base) or 'pregunta'
                                ruta_guardado = f"prevencion/cursos_preguntas/{nombre_seguro}-{uuid4().hex}{ext}"
                                nombre_archivo = fs.save(ruta_guardado, imagen_pregunta)
                                imagen_url = fs.url(nombre_archivo)

                        filas = []
                        if enunciado:
                            filas.append(enunciado)
                        for o_num, opcion in enumerate(opciones):
                            letra = chr(97 + o_num) if o_num < 26 else f"o{o_num + 1}"
                            filas.append(f"{letra} - {opcion}")
                        if respuesta_correcta:
                            filas.append(f"Respuesta correcta: {respuesta_correcta}")
                        if justificacion:
                            filas.append(f"Justificación: {justificacion}")

                        if not filas:
                            filas = ["Item sin título"]

                        estructura_final.append({
                            'titulo': titulo_pregunta,
                            'enunciado': enunciado,
                            'imagen_url': imagen_url,
                            'opciones': opciones,
                            'respuesta_correcta': respuesta_correcta,
                            'justificacion': justificacion,
                            'columnas_extras': [],
                            'filas': filas,
                        })

                    if not estructura_final:
                        estructura_final = [{
                            'titulo': 'pregunta 1',
                            'enunciado': '',
                            'imagen_url': '',
                            'opciones': [],
                            'respuesta_correcta': '',
                            'justificacion': '',
                            'columnas_extras': [],
                            'filas': ['Item sin título'],
                        }]
                else:
                    item_indices_str = request.POST.get('item_indices', '')
                    if item_indices_str:
                        items_ids = item_indices_str.split(',')
                        for i_id in items_ids:
                            if not i_id:
                                continue
                            
                            titulo_seccion = request.POST.get(f'nombre_item_{i_id}')
                            
                            columnas_extras = []
                            if permitir_columnas_extras:
                                head_indices_str = request.POST.get(f'head_indices_{i_id}', '')
                                if head_indices_str:
                                    for h_id in head_indices_str.split(','):
                                        col_val = request.POST.get(f'nombre_head_{i_id}_{h_id}')
                                        if col_val:
                                            columnas_extras.append(col_val)

                            row_indices_str = request.POST.get(f'row_indices_{i_id}', '')
                            filas = []
                            if row_indices_str:
                                for r_id in row_indices_str.split(','):
                                    # Guardar el texto real del item.
                                    item_texto = request.POST.get(f'nombre_row_{i_id}_{r_id}')
                                    # Si el usuario escribió algo, se guarda; si no, usa un valor por defecto.
                                    filas.append(item_texto if item_texto else "Item sin título")

                            estructura_final.append({
                                'titulo': titulo_seccion,
                                'columnas_extras': columnas_extras,
                                'filas': filas
                            })

                    # En tipos "solo items", la estructura mínima obligatoria es:
                    # al menos 1 sección y al menos 1 ítem (1.1) en la primera sección.
                    if not permitir_columnas_extras:
                        if not estructura_final:
                            estructura_final = [{
                                'titulo': '',
                                'columnas_extras': [],
                                'filas': ['Item sin título']
                            }]
                        elif not estructura_final[0].get('filas'):
                            estructura_final[0]['filas'] = ['Item sin título']
            
            plantilla.estructura = estructura_final

            mismo_nombre = False
            misma_faena = False
            mismo_cargo = False
            mismo_tiempo_estimado = False
            misma_configuracion_curso = True
            if duplicate_source:
                mismo_nombre = plantilla.nombre == (duplicate_source.nombre or '').strip()
                misma_faena = plantilla.faena_id == duplicate_source.faena_id
                mismo_cargo = cargo_id == duplicate_source_cargo_id
                mismo_tiempo_estimado = (
                    plantilla.tiempo_estimado_minutos == duplicate_source.tiempo_estimado_minutos
                )
                if es_curso:
                    misma_configuracion_curso = (
                        int(plantilla.curso_total_intentos or 0) == int(duplicate_source.curso_total_intentos or 0)
                        and int(plantilla.curso_porcentaje_aprobacion or 0) == int(duplicate_source.curso_porcentaje_aprobacion or 0)
                    )

            if error_validacion:
                pass
            elif (
                duplicate_source
                and mismo_nombre
                and misma_faena
                and mismo_cargo
                and mismo_tiempo_estimado
                and misma_configuracion_curso
            ):
                messages.error(
                    request,
                    "Para duplicar debes cambiar al menos la faena, el nombre, el cargo o el tiempo estimado."
                )
            else:
                if eliminar_pdf_actual and pdf_actual:
                    pdf_actual.delete(save=False)
                    if not archivo_pdf:
                        plantilla.archivo_pdf = None
                plantilla.save()
                plantilla.autorizadores.set(autorizadores_validos_ids)

                # Mantiene una única configuración de permisos por plantilla/faena.
                PrevencionPermisoLlenado.objects.filter(plantilla=plantilla).exclude(
                    faena_id=plantilla.faena_id
                ).delete()

                if cargo_id:
                    if not VigilanciaCargo.objects.filter(status=True, id=cargo_id).exists():
                        messages.error(request, "El cargo seleccionado no es válido.")
                        return redirect('edit_prevencion', tipo_slug=tipo_slug, pk=plantilla.id)

                    permiso, _ = PrevencionPermisoLlenado.objects.get_or_create(
                        faena_id=plantilla.faena_id,
                        plantilla=plantilla,
                        defaults={
                            'status': True,
                            'creador': request.user.username
                        }
                    )
                    if not permiso.status:
                        permiso.status = True
                    if not permiso.creador:
                        permiso.creador = request.user.username
                    permiso.save()
                    permiso.cargos.set([cargo_id])
                else:
                    # Sin cargo seleccionado: elimina regla para conservar acceso abierto.
                    PrevencionPermisoLlenado.objects.filter(
                        faena_id=plantilla.faena_id,
                        plantilla=plantilla
                    ).delete()

                total_reautorizaciones = 0
                if pk and tipo_slug in TIPOS_DOCUMENTO_PERMITIDOS:
                    snapshot_post_reautorizacion = _snapshot_plantilla_para_reautorizacion(plantilla)
                    if snapshot_pre_reautorizacion != snapshot_post_reautorizacion:
                        total_reautorizaciones = _crear_reautorizaciones_automaticas(
                            plantilla=plantilla,
                            solicitante=request.user,
                        )

                if duplicate_source and not pk:
                    messages.success(request, "Plantilla duplicada correctamente.")
                else:
                    messages.success(request, _mensaje_actualizacion_prevencion(tipo_slug))
                    if total_reautorizaciones > 0:
                        messages.info(
                            request,
                            "Se reenviaron automáticamente las solicitudes de autorización por cambios en la plantilla.",
                        )
                if tipo_slug in TIPOS_DOCUMENTO_PERMITIDOS:
                    tipo_label = _mapear_tipo_documento_label(tipo_slug)
                    user_name = request.user.get_full_name() or request.user.username
                    es_femenino = tipo_slug in {'charlas', 'difusiones'}
                    if not pk:
                        prefijo = 'Nueva' if es_femenino else 'Nuevo'
                        participio = 'creada' if es_femenino else 'creado'
                        notify_group(
                            'prevencion_general',
                            f'{prefijo} {tipo_label.lower()}: {plantilla.nombre}',
                            f'{prefijo} {tipo_label.lower()} {participio}: {plantilla.nombre}\n{participio.capitalize()} por: {user_name}'
                        )
                    else:
                        participio = 'actualizada' if es_femenino else 'actualizado'
                        notify_group(
                            'prevencion_general',
                            f'{tipo_label} {participio}: {plantilla.nombre}',
                            f'{tipo_label} {participio}: {plantilla.nombre}\n{participio.capitalize()} por: {user_name}'
                        )
                # Flujo en 2 pasos para documentos: guardar primero y luego autorización.
                if (not pk) and (tipo_slug in TIPOS_DOCUMENTO_PERMITIDOS):
                    edit_url = reverse('edit_prevencion', kwargs={'tipo_slug': tipo_slug, 'pk': plantilla.id})
                    return redirect(f"{edit_url}?tab=autorizacion")
                return redirect('manage_prevencion', tipo_slug=tipo_slug)
            
        except Exception as e:
            messages.error(request, f"Error al guardar: {str(e)}")
    
    faenas = Faena.objects.filter(status=True)
    cargos = VigilanciaCargo.objects.filter(status=True).order_by('correlativo', 'id')
    selected_cargo_id = None

    if selected_cargo_id_post is not None:
        selected_cargo_id = selected_cargo_id_post
    elif duplicate_source:
        selected_cargo_id = duplicate_source_cargo_id
    elif plantilla.id:
        permiso_actual = None
        if plantilla.faena_id:
            permiso_actual = PrevencionPermisoLlenado.objects.filter(
                status=True,
                faena_id=plantilla.faena_id,
                plantilla=plantilla
            ).first()
        if not permiso_actual:
            permiso_actual = PrevencionPermisoLlenado.objects.filter(
                status=True,
                plantilla=plantilla
            ).order_by('-updated_at', '-id').first()
        if permiso_actual:
            selected_cargo_id = permiso_actual.cargos.values_list('id', flat=True).first()

    trabajadores_autorizacion = []
    usuarios_autorizacion_qs = (
        User.objects.filter(
            is_active=True, 
            usuarioprofile__seccionPrevencion='ADMINISTRADOR'
        ).order_by('first_name', 'last_name', 'username')
    )
    for usuario in usuarios_autorizacion_qs:
        nombre_completo = f"{(usuario.first_name or '').strip()} {(usuario.last_name or '').strip()}".strip()
        if not nombre_completo:
            nombre_completo = usuario.username
        trabajadores_autorizacion.append({
            'id': usuario.id,
            'texto': nombre_completo,
        })

    if selected_autorizadores_ids_post is not None:
        selected_autorizadores_ids = list(selected_autorizadores_ids_post)
    elif duplicate_source:
        selected_autorizadores_ids = []
    elif plantilla.id:
        selected_autorizadores_ids = list(plantilla.autorizadores.values_list('id', flat=True))
    else:
        selected_autorizadores_ids = []

    autorizacion_ids_validos = {item['id'] for item in trabajadores_autorizacion}
    selected_autorizadores_ids = [uid for uid in selected_autorizadores_ids if uid in autorizacion_ids_validos][:5]

    autorizadores_estado = {}
    if plantilla.id and selected_autorizadores_ids:
        autorizadores_estado = _obtener_estado_detallado_autorizadores_plantilla(
            plantilla,
            selected_autorizadores_ids,
        )

    historial_autorizaciones = _obtener_historial_autorizaciones_plantilla(plantilla)

    modo_contenido_actual = plantilla.modo_contenido or PrevencionPlantilla.ModoContenido.ESTRUCTURA
    if not habilitar_modo_pdf:
        modo_contenido_actual = PrevencionPlantilla.ModoContenido.ESTRUCTURA
    archivo_pdf_url = plantilla.archivo_pdf.url if plantilla.archivo_pdf else ''
    archivo_pdf_nombre = os.path.basename(plantilla.archivo_pdf.name) if plantilla.archivo_pdf else ''

    # JSON seguro para el frontend
    estructura_json = json.dumps(plantilla.estructura, cls=DjangoJSONEncoder) if plantilla.estructura else '[]'

    autorizacion_habilitada = bool(plantilla.id)
    tab_request = (request.GET.get('tab') or '').strip().lower()
    tabs_disponibles = {'principal'}
    if autorizacion_habilitada:
        tabs_disponibles.update({'autorizacion', 'historial'})
    active_tab = tab_request if tab_request in tabs_disponibles else 'principal'

    context = {
        'titulo_categoria': tipos_map.get(tipo_slug, 'Prevención'),
        'titulo_pagina': titulo,
        'slug_categoria': tipo_slug,
        'plantilla': plantilla,
        'faenas': faenas,
        'cargos': cargos,
        'tiempos_estimados_disponibles': tiempos_estimados_disponibles,
        'curso_intentos_disponibles': curso_intentos_disponibles,
        'curso_porcentajes_aprobacion_disponibles': curso_porcentajes_aprobacion_disponibles,
        'selected_cargo_id': selected_cargo_id,
        'modo_duplicado': bool(duplicate_source),
        'duplicate_from_id': duplicate_source.id if duplicate_source else '',
        'es_curso': es_curso,
        'permitir_columnas_extras': permitir_columnas_extras,
        'habilitar_modo_pdf': habilitar_modo_pdf,
        'modo_contenido_actual': modo_contenido_actual,
        'archivo_pdf_url': archivo_pdf_url,
        'archivo_pdf_nombre': archivo_pdf_nombre,
        'estructura_json': estructura_json,
        'trabajadores_autorizacion_json': json.dumps(trabajadores_autorizacion, cls=DjangoJSONEncoder),
        'selected_autorizadores_ids_json': json.dumps(selected_autorizadores_ids, cls=DjangoJSONEncoder),
        'autorizadores_estado_json': json.dumps(autorizadores_estado, cls=DjangoJSONEncoder),
        'autorizacion_incompleta': bool(getattr(plantilla, 'autorizacion_incompleta', False)),
        'autorizacion_habilitada': autorizacion_habilitada,
        'historial_autorizaciones': historial_autorizaciones,
        'active_tab': active_tab,
        'sidebar': 'mantenedor_prev',
        'sidebarmenu': f't-{tipo_slug}'
    }
    return render(request, 'pages/prevencion/edit_prevencion.html', context)

@login_required
@prevencion_mantenedor_required
def delete_prevencion(request, pk):
    # === SEGURIDAD ===
    # También protegemos el borrado. Solo se puede borrar algo que esté activo.
    plantilla = get_object_or_404(PrevencionPlantilla, pk=pk, status=True)
    
    tipo = plantilla.tipo
    plantilla.status = False
    plantilla.save()

    if request.headers.get('x-requested-with') == 'XMLHttpRequest':
        return JsonResponse({
            'success': True,
            'message': 'Plantilla eliminada.',
            'tipo_slug': tipo,
            'redirect_url': reverse('manage_prevencion', kwargs={'tipo_slug': tipo}),
        })
    
    messages.success(request, "Plantilla eliminada.")
    return redirect('manage_prevencion', tipo_slug=tipo)


# === GESTIÓN DE DOCUMENTOS (PREVENCIONISTA) ===

# 1. LISTADO HISTORIAL (Datatable)
@login_required
@prevencion_riesgo_required
def history_prevencion(request):
    def _preparar_documentos_historial(queryset):
        documentos_lista = []
        for documento in queryset:
            documento.porcentaje_promedio = _calcular_porcentaje_promedio_documento(documento.contenido)

            if not documento.plantilla_base_id:
                documento.cargo_display = '-'
            else:
                permiso = _permiso_llenado_para_documento(documento)
                if not permiso:
                    documento.cargo_display = 'Todos'
                else:
                    cargos = list(
                        permiso.cargos.filter(status=True)
                        .order_by('correlativo', 'id')
                        .values_list('valor', flat=True)
                    )
                    documento.cargo_display = ', '.join(cargos) if cargos else 'Todos'

            documentos_lista.append(documento)
        return documentos_lista

    documentos_aprobados_qs = PrevencionDocumento.objects.filter(status=True).select_related('faena', 'plantilla_base')
    documentos_aprobados_qs = _filtrar_queryset_por_faena_usuario(
        documentos_aprobados_qs,
        request.user
    ).order_by('-fecha_creacion')

    documentos_eliminados_qs = PrevencionDocumento.objects.filter(status=False).select_related('faena', 'plantilla_base')
    documentos_eliminados_qs = _filtrar_queryset_por_faena_usuario(
        documentos_eliminados_qs,
        request.user
    ).order_by('-fecha_creacion')

    active_tab = (request.GET.get('tab') or '').strip().lower()
    if active_tab not in {'aprobados', 'eliminados'}:
        active_tab = 'aprobados'

    context = {
        'documentos_aprobados': _preparar_documentos_historial(documentos_aprobados_qs),
        'documentos_eliminados': _preparar_documentos_historial(documentos_eliminados_qs),
        'active_tab': active_tab,
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo', # Mantiene expandido el menú padre
        'sidebar': 'historial_doc'          # Marca activo el sub-item Historial
    }
    return render(request, 'pages/prevencion/history_prevencion.html', context)


def _permiso_llenado_para_documento(documento):
    if not documento or not documento.plantilla_base_id:
        return None

    permiso = PrevencionPermisoLlenado.objects.filter(
        status=True,
        faena_id=documento.faena_id,
        plantilla_id=documento.plantilla_base_id,
    ).first()

    if permiso:
        return permiso

    return (
        PrevencionPermisoLlenado.objects.filter(
            status=True,
            plantilla_id=documento.plantilla_base_id,
        )
        .order_by('-updated_at', '-id')
        .first()
    )


def _catalogo_trabajadores_para_difusion_documento(documento, aplicar_filtro_cargo=True):
    permiso = _permiso_llenado_para_documento(documento)
    cargos_permitidos = []
    cargo_ids = []

    if permiso:
        cargos_permitidos = list(
            permiso.cargos.filter(status=True).order_by('correlativo', 'id').values_list('valor', flat=True)
        )
        cargo_ids = list(
            permiso.cargos.filter(status=True).values_list('id', flat=True)
        )

    trabajadores_qs = (
        User.objects.filter(
            is_active=True,
            usuarioprofile__isnull=False,
            usuarioprofile__faena_id=documento.faena_id,
        )
        .select_related(
            'usuarioprofile__faena',
            'informacion_laboral__cargo',
            'informacion_laboral__area',
            'informacion_laboral__ges',
        )
        .order_by('first_name', 'last_name', 'username')
    )

    if aplicar_filtro_cargo and cargo_ids:
        trabajadores_qs = trabajadores_qs.filter(informacion_laboral__cargo_id__in=cargo_ids)

    trabajadores = []
    for usuario in trabajadores_qs:
        info = getattr(usuario, 'informacion_laboral', None)
        perfil = getattr(usuario, 'usuarioprofile', None)
        nombre = f"{(usuario.first_name or '').strip()} {(usuario.last_name or '').strip()}".strip() or usuario.username
        trabajadores.append({
            'user': usuario,
            'username': usuario.username,
            'nombre': nombre,
            'cargo': getattr(getattr(info, 'cargo', None), 'valor', '') or '',
            'faena_id': getattr(perfil, 'faena_id', None) or documento.faena_id,
            'faena_nombre': getattr(getattr(perfil, 'faena', None), 'faena', '') or (documento.faena.faena if documento.faena_id else ''),
            'area': getattr(getattr(info, 'area', None), 'valor', '') or '',
            'ges': getattr(getattr(info, 'ges', None), 'valor', '') or '',
        })

    return trabajadores, cargos_permitidos


def _catalogo_todos_los_trabajadores_para_agregar(excluir_ids=None):
    ids_excluir = {1}
    if excluir_ids:
        for valor in excluir_ids:
            try:
                ids_excluir.add(int(valor))
            except (TypeError, ValueError):
                continue

    usuarios_qs = (
        User.objects
        .exclude(id__in=ids_excluir)
        .select_related(
            'usuarioprofile__faena',
            'informacion_laboral__cargo',
            'informacion_laboral__area',
            'informacion_laboral__ges',
        )
        .order_by('first_name', 'last_name', 'username')
    )

    catalogo = []
    for usuario in usuarios_qs:
        info = getattr(usuario, 'informacion_laboral', None)
        perfil = getattr(usuario, 'usuarioprofile', None)
        nombre = f"{(usuario.first_name or '').strip()} {(usuario.last_name or '').strip()}".strip() or usuario.username

        catalogo.append({
            'user': usuario,
            'username': usuario.username,
            'nombre': nombre,
            'cargo': getattr(getattr(info, 'cargo', None), 'valor', '') or '',
            'faena_id': getattr(perfil, 'faena_id', None) or '',
            'faena_nombre': getattr(getattr(perfil, 'faena', None), 'faena', '') or '',
            'area': getattr(getattr(info, 'area', None), 'valor', '') or '',
            'ges': getattr(getattr(info, 'ges', None), 'valor', '') or '',
        })

    return catalogo


def _nombre_completo_usuario(usuario):
    if not usuario:
        return ''
    return f"{(usuario.first_name or '').strip()} {(usuario.last_name or '').strip()}".strip() or usuario.username


def _resolver_nombre_difusor(valor):
    texto = str(valor or '').strip()
    if not texto:
        return ''

    usuario = User.objects.filter(username__iexact=texto).only('first_name', 'last_name', 'username').first()
    if usuario:
        nombre_resuelto = _nombre_completo_usuario(usuario)
        if nombre_resuelto and nombre_resuelto.strip().lower() != texto.lower():
            return nombre_resuelto

    # Fallback para registros antiguos que guardaron solo el nombre corto ("eden").
    candidatos = list(
        User.objects.filter(first_name__iexact=texto)
        .exclude(last_name__isnull=True)
        .exclude(last_name__exact='')
        .only('first_name', 'last_name', 'username')
        .order_by('id')[:2]
    )
    if len(candidatos) == 1:
        return _nombre_completo_usuario(candidatos[0])

    return texto


def _nombre_difusor_desde_request(user):
    if not user or not user.is_authenticated:
        return ''
    return _nombre_completo_usuario(user)


def _registrar_historial_difusion_documento(
    documento,
    accion,
    actor=None,
    trabajador=None,
    trabajador_rut='',
    trabajador_nombre='',
):
    accion_normalizada = str(accion or '').strip().lower()
    acciones_validas = {
        PrevencionHistorialDifusionTrabajador.Accion.AGREGADO,
        PrevencionHistorialDifusionTrabajador.Accion.QUITADO,
    }
    if accion_normalizada not in acciones_validas:
        return None

    rut = str(trabajador_rut or '').strip()
    nombre = str(trabajador_nombre or '').strip()
    if trabajador:
        if not rut:
            rut = str(getattr(trabajador, 'username', '') or '').strip()
        if not nombre:
            nombre = _nombre_usuario_legible(trabajador)

    actor_nombre = _nombre_usuario_legible(actor) if actor else ''
    if not actor_nombre:
        actor_nombre = _resolver_nombre_difusor(getattr(documento, 'creador', '')) or ''

    return PrevencionHistorialDifusionTrabajador.objects.create(
        documento=documento,
        trabajador=trabajador if trabajador and getattr(trabajador, 'id', None) else None,
        trabajador_rut=rut,
        trabajador_nombre=nombre,
        actor=actor if actor and getattr(actor, 'id', None) else None,
        actor_nombre=actor_nombre,
        accion=accion_normalizada,
    )


def _crear_registro_difusion_documento(documento, trabajador, creado_por='', status=True):
    usuario = trabajador.get('user')
    if not usuario or not usuario.id:
        return None

    defaults = {
        'rut_trabajador': trabajador.get('username') or usuario.username,
        'nombre_completo': trabajador.get('nombre') or usuario.username,
        'cargo': trabajador.get('cargo') or '',
        'faena_id': trabajador.get('faena_id') or documento.faena_id,
        'area': trabajador.get('area') or '',
        'ges': trabajador.get('ges') or '',
        'creado_por': _resolver_nombre_difusor(creado_por) or '',
        'status': status,
    }

    registro, created = PrevencionDocumentoDifusionTrabajador.objects.get_or_create(
        documento=documento,
        user=usuario,
        defaults=defaults,
    )

    if created:
        return registro

    cambios = []
    for campo, valor in defaults.items():
        if getattr(registro, campo) != valor:
            setattr(registro, campo, valor)
            cambios.append(campo)
    if cambios:
        registro.save(update_fields=cambios)

    return registro


def _inicializar_snapshot_difusion_documento(documento, user=None):
    if PrevencionDocumentoDifusionTrabajador.objects.filter(documento=documento).exists():
        return 0

    creador = _resolver_nombre_difusor(getattr(documento, 'creador', ''))
    if not creador:
        creador = _nombre_difusor_desde_request(user)
    total_creados = 0

    with transaction.atomic():
        trabajadores_base, _ = _catalogo_trabajadores_para_difusion_documento(
            documento=documento,
            aplicar_filtro_cargo=True,
        )
        for trabajador in trabajadores_base:
            registro = _crear_registro_difusion_documento(
                documento=documento,
                trabajador=trabajador,
                creado_por=creador,
                status=True,
            )
            if registro:
                total_creados += 1
                _registrar_historial_difusion_documento(
                    documento=documento,
                    accion=PrevencionHistorialDifusionTrabajador.Accion.AGREGADO,
                    actor=user,
                    trabajador=getattr(registro, 'user', None),
                    trabajador_rut=getattr(registro, 'rut_trabajador', '') or trabajador.get('username', ''),
                    trabajador_nombre=getattr(registro, 'nombre_completo', '') or trabajador.get('nombre', ''),
                )

    return total_creados


def _obtener_trabajadores_para_difusion(documento):
    trabajadores, cargos_permitidos = _catalogo_trabajadores_para_difusion_documento(
        documento=documento,
        aplicar_filtro_cargo=True,
    )
    return [
        {
            'rut': fila.get('username') or '-',
            'nombre': fila.get('nombre') or '-',
            'cargo': fila.get('cargo') or '-',
            'faena': fila.get('faena_nombre') or '-',
            'area': fila.get('area') or '-',
            'ges': fila.get('ges') or '-',
        }
        for fila in trabajadores
    ], cargos_permitidos


def _asegurar_notificacion_documento_disponible(documento, destinatario, forzar_pendiente=False):
    if not documento or not destinatario or not getattr(destinatario, 'id', None):
        return None

    tipo_documento = ''
    nombre_plantilla = 'Documento'
    if documento.plantilla_base_id:
        tipo_documento = str(documento.plantilla_base.tipo or '').strip().lower()
        nombre_plantilla = documento.plantilla_base.nombre or nombre_plantilla

    tipo_label = _mapear_tipo_documento_label(tipo_documento)
    titulo = f'Documento disponible: {tipo_label}'
    descripcion = f'Tienes disponible para revisión {tipo_label.lower()}: "{nombre_plantilla}".'

    defaults_base = {
        'tipo_documento': tipo_documento,
        'titulo': titulo[:200],
        'descripcion': descripcion,
    }

    if forzar_pendiente:
        notificacion, _ = PrevencionNotificacionDocumento.objects.update_or_create(
            destinatario=destinatario,
            documento=documento,
            defaults={
                **defaults_base,
                'estado': PrevencionNotificacionDocumento.Estado.PENDIENTE,
                'leida': False,
                'fechacreacion': timezone.now(),
                'fecha_inicio_lectura': None,
                'fecha_lectura': None,
            },
        )
        return notificacion

    notificacion, creada = PrevencionNotificacionDocumento.objects.get_or_create(
        destinatario=destinatario,
        documento=documento,
        defaults={
            **defaults_base,
            'estado': PrevencionNotificacionDocumento.Estado.PENDIENTE,
            'leida': False,
            'fechacreacion': timezone.now(),
            'fecha_inicio_lectura': None,
            'fecha_lectura': None,
        }
    )
    if not creada:
        update_fields = []
        if notificacion.tipo_documento != tipo_documento:
            notificacion.tipo_documento = tipo_documento
            update_fields.append('tipo_documento')
        if notificacion.titulo != titulo[:200]:
            notificacion.titulo = titulo[:200]
            update_fields.append('titulo')
        if notificacion.descripcion != descripcion:
            notificacion.descripcion = descripcion
            update_fields.append('descripcion')
        if update_fields:
            notificacion.save(update_fields=update_fields)

    return notificacion


def _crear_notificacion_documento_disponible(documento, destinatario):
    return _asegurar_notificacion_documento_disponible(
        documento=documento,
        destinatario=destinatario,
        forzar_pendiente=True,
    )


def _notificar_documento_a_trabajadores_difusion(documento):
    if not documento:
        return 0

    registros = (
        PrevencionDocumentoDifusionTrabajador.objects
        .filter(documento=documento, status=True, user__isnull=False)
        .select_related('user')
    )
    enviados = 0
    for registro in registros:
        if _crear_notificacion_documento_disponible(documento, registro.user):
            enviados += 1
    return enviados

# 2. VER DETALLE (Solo Lectura)
@login_required
@capacitaciones_y_docs_required
def view_documento(request, pk):
    origen = str(request.GET.get('origen') or '').strip().lower()
    notif_doc_id = str(request.GET.get('notif_doc') or '').strip()
    confirmar_cambio_activo = _parsear_bool(request.GET.get('confirmar_cambio'))
    acceso_desde_capacitaciones = origen in {'capacitaciones', 'notificaciones'} or notif_doc_id.isdigit()

    if acceso_desde_capacitaciones:
        documentos_visibles_ids = _ids_documentos_visibles_para_usuario(request.user)
        documentos = PrevencionDocumento.objects.filter(
            status=True,
            id__in=documentos_visibles_ids,
        )
    else:
        documentos = _filtrar_queryset_por_faena_usuario(
            PrevencionDocumento.objects.filter(status=True),
            request.user
        )
    documento = get_object_or_404(documentos, pk=pk)
    plantilla = documento.plantilla_base

    tipo_documento = str(getattr(plantilla, 'tipo', '') or '').strip().lower()
    tipo_origen = str(request.GET.get('tipo') or '').strip().lower()
    next_raw = str(request.GET.get('next') or '').strip()
    back_url_next = ''
    if next_raw and url_has_allowed_host_and_scheme(
        url=next_raw,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        next_parts = urlparse(next_raw)
        back_url_next = next_parts.path or '/'
        if next_parts.query:
            back_url_next = f"{back_url_next}?{next_parts.query}"
        if next_parts.fragment:
            back_url_next = f"{back_url_next}#{next_parts.fragment}"

    es_origen_capacitaciones = origen == 'capacitaciones'
    es_origen_notificaciones = origen == 'notificaciones'

    if es_origen_notificaciones and tipo_origen not in TIPOS_DOCUMENTO_PERMITIDOS:
        tipo_origen = tipo_documento if tipo_documento in TIPOS_DOCUMENTO_PERMITIDOS else ''

    if es_origen_capacitaciones:
        if tipo_origen not in TIPOS_DOCUMENTO_PERMITIDOS:
            tipo_origen = tipo_documento if tipo_documento in TIPOS_DOCUMENTO_PERMITIDOS else ''
        es_origen_capacitaciones = tipo_origen in TIPOS_DOCUMENTO_PERMITIDOS

    back_url = reverse('history_prevencion')
    sidebarmain = 'prevencion_riesgo'
    sidebarmenu = ''
    if es_origen_capacitaciones:
        back_url = reverse('manage_capacitaciones', kwargs={'tipo_slug': tipo_origen})
        sidebarmain = 'capacitaciones'
        sidebarmenu = f't-{tipo_origen}'
    elif es_origen_notificaciones:
        back_url = reverse('view_all_approval_notifications')
        if tipo_origen in TIPOS_DOCUMENTO_PERMITIDOS:
            sidebarmain = 'capacitaciones'
            sidebarmenu = f't-{tipo_origen}'
    if back_url_next:
        back_url = back_url_next

    es_flujo_capacitacion = (
        tipo_documento in TIPOS_DOCUMENTO_PERMITIDOS
        and acceso_desde_capacitaciones
    )

    if es_flujo_capacitacion:
        difusion = PrevencionDocumentoDifusionTrabajador.objects.filter(
            documento=documento, user=request.user
        ).first()
        es_curso = tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value
        
        if difusion and difusion.motivo_rechazo and not es_curso:
            messages.error(request, "Acceso bloqueado: Falló la validación de identidad o el proceso fue cancelado.")
            return redirect(back_url)

    if es_flujo_capacitacion:
        notificacion_inicio = (
            PrevencionNotificacionDocumento.objects
            .filter(destinatario=request.user, documento=documento)
            .first()
        )
        if notificacion_inicio and not notificacion_inicio.fecha_inicio_lectura:
            notificacion_inicio.fecha_inicio_lectura = timezone.now()
            notificacion_inicio.save(update_fields=['fecha_inicio_lectura'])

        if tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
            session_key = f'inicio_intento_{documento.id}'
            if session_key not in request.session:
                request.session[session_key] = timezone.now().isoformat()
                request.session.modified = True

    documento_ya_validado = _documento_capacitacion_validado_para_usuario(documento.id, request.user)
    estado_curso_usuario = None
    if es_flujo_capacitacion and tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        resultado_curso_usuario = (
            PrevencionResultadoCurso.objects
            .filter(documento=documento, trabajador=request.user)
            .order_by('-updated_at', '-id')
            .first()
        )
        estado_curso_usuario = _normalizar_estado_resultado_curso(
            documento,
            resultado_curso_usuario,
        )
        if bool(estado_curso_usuario.get('aprobado')):
            _limpiar_borrador_respuestas_curso(request, documento.id)
            messages.info(request, "Este curso ya está aprobado y no se puede volver a seleccionar.")
            return redirect(back_url)
        if (
            int(estado_curso_usuario.get('intentos_realizados') or 0) > 0
            and int(estado_curso_usuario.get('intentos_restantes') or 0) <= 0
        ):
            _limpiar_borrador_respuestas_curso(request, documento.id)
            messages.error(request, "No te quedan intentos disponibles para este curso.")
            return redirect(back_url)

    aplica_bloqueo_capacitacion = (
        es_flujo_capacitacion
        and not (
            tipo_documento in TIPOS_CAPACITACION_LIBRE_TRAS_VALIDACION
            and documento_ya_validado
        )
    )
    estado_bloqueo_actual = None
    alerta_capacitacion_activa = None
    tiempo_expirado_al_ingresar = False
 
    if aplica_bloqueo_capacitacion:
        estado_bloqueo = _obtener_estado_bloqueo_capacitacion(request)
        curso_finalizado_bloqueado = _obtener_curso_finalizado_bloqueado(request, estado_bloqueo)
        if estado_bloqueo:
            documento_bloqueado_id = int(estado_bloqueo.get('documento_id') or 0)
            segundos_restantes_bloqueo = max(
                0,
                int(estado_bloqueo.get('remaining_seconds', 0) or 0)
            )
            if documento_bloqueado_id != documento.id and segundos_restantes_bloqueo > 0:
                restante = _formatear_tiempo_restante_capacitacion(
                    segundos_restantes_bloqueo
                )
                messages.add_message(
                    request,
                    messages.ERROR,
                    "Aún no puedes abrir otro archivo",
                    extra_tags=f"Tiempo restante: {restante}. __countdown_seconds={segundos_restantes_bloqueo}__",
                )
                return redirect(back_url)

            if (
                documento_bloqueado_id == documento.id
                and segundos_restantes_bloqueo > 0
                and tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value
                and curso_finalizado_bloqueado
                and int(curso_finalizado_bloqueado.get('documento_id') or 0) == documento.id
            ):
                restante = _formatear_tiempo_restante_capacitacion(segundos_restantes_bloqueo)
                messages.add_message(
                    request,
                    messages.ERROR,
                    "Este curso ya fue finalizado",
                    extra_tags=f"Podrás volver a abrirlo cuando termine el tiempo restante: {restante}. __countdown_seconds={segundos_restantes_bloqueo}__",
                )
                return redirect(back_url)

            if documento_bloqueado_id != documento.id:
                if confirmar_cambio_activo:
                    estado_bloqueo_actual = _iniciar_bloqueo_capacitacion(
                        request,
                        documento,
                        forzar_reinicio=True,
                    )
                    if not estado_bloqueo_actual:
                        request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
                        request.session.modified = True
                else:
                    documento_activo = (
                        PrevencionDocumento.objects
                        .filter(id=documento_bloqueado_id, status=True)
                        .select_related('plantilla_base')
                        .first()
                    )

                    nombre_documento_activo = 'Documento activo'
                    tipo_documento_activo = ''
                    if documento_activo and documento_activo.plantilla_base_id:
                        nombre_documento_activo = (
                            documento_activo.plantilla_base.nombre
                            or nombre_documento_activo
                        )
                        tipo_documento_activo = str(
                            documento_activo.plantilla_base.tipo or ''
                        ).strip().lower()

                    query_continuar = request.GET.copy()
                    query_continuar['confirmar_cambio'] = '1'
                    query_string_continuar = query_continuar.urlencode()
                    url_continuar = request.path
                    if query_string_continuar:
                        url_continuar = f"{request.path}?{query_string_continuar}"

                    params_volver = {'origen': 'capacitaciones'}
                    if tipo_documento_activo in TIPOS_DOCUMENTO_PERMITIDOS:
                        params_volver['tipo'] = tipo_documento_activo
                    if back_url:
                        params_volver['next'] = back_url

                    url_volver = back_url
                    if documento_activo and documento_activo.id:
                        url_volver = (
                            f"{reverse('view_documento', kwargs={'pk': documento_activo.id})}"
                            f"?{urlencode(params_volver)}"
                        )

                    alerta_capacitacion_activa = {
                        'nombre_documento_activo': nombre_documento_activo,
                        'tipo_activo_slug': tipo_documento_activo,
                        'tipo_destino_slug': tipo_documento,
                        'url_continuar': url_continuar,
                        'url_volver_activo': url_volver,
                    }
            else:
                # Mismo documento: permitido, sin reiniciar contador.
                estado_bloqueo_actual = estado_bloqueo
        else:
            bloqueo_expirado = request.session.get(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY)
            documento_expirado_id = 0
            if isinstance(bloqueo_expirado, dict):
                try:
                    documento_expirado_id = int(bloqueo_expirado.get('documento_id') or 0)
                except (TypeError, ValueError):
                    documento_expirado_id = 0

            if documento_expirado_id == documento.id:
                request.session.pop(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY, None)
                request.session.modified = True

            estado_bloqueo_actual = _iniciar_bloqueo_capacitacion(request, documento)

    aviso_tiempo_bloqueo_capacitacion = ''
    aviso_minutos_bloqueo_capacitacion = 0
    contador_segundos_restantes = 0
    if estado_bloqueo_actual:
        segundos_restantes = max(0, int(estado_bloqueo_actual.get('remaining_seconds', 0) or 0))
        contador_segundos_restantes = segundos_restantes
        aviso_tiempo_bloqueo_capacitacion = _formatear_tiempo_restante_capacitacion(
            segundos_restantes
        )
        if segundos_restantes > 0:
            aviso_minutos_bloqueo_capacitacion = max(1, (segundos_restantes + 59) // 60)

    contenido_documento = _preparar_documento_para_visualizacion(documento)

    if plantilla and not _plantilla_autorizada_para_documentos(plantilla):
        messages.error(
            request,
            "Este documento aún no está habilitado: falta aprobación de todos los autorizadores."
        )
        return redirect(back_url)

    tipo_slug = ''
    tipo_label = 'Sin tipo'
    modo_contenido = PrevencionPlantilla.ModoContenido.ESTRUCTURA
    estructura_lectura = []

    if plantilla:
        tipo_slug = plantilla.tipo
        tipo_label = plantilla.get_tipo_display()
        modo_contenido = plantilla.modo_contenido or PrevencionPlantilla.ModoContenido.ESTRUCTURA
        estructura_lectura = plantilla.estructura or []

    tipo_titulo_slug = tipo_slug if tipo_slug in TIPOS_DOCUMENTO_PERMITIDOS else ''
    if not tipo_titulo_slug and tipo_origen in TIPOS_DOCUMENTO_PERMITIDOS:
        tipo_titulo_slug = tipo_origen
    titulo_vista_documento = _construir_titulo_vista_documento(tipo_titulo_slug)

    if not estructura_lectura:
        for seccion in (documento.contenido or []):
            filas = []
            for fila in (seccion.get('respuestas') or []):
                filas.append({
                    'item': fila.get('pregunta', ''),
                    'respuestas_extras': fila.get('respuestas_extras') or [],
                })

            estructura_lectura.append({
                'titulo': seccion.get('seccion', ''),
                'columnas_extras': seccion.get('columnas_extras') or [],
                'filas': filas,
            })
    
    evidencias_generales = list(documento.evidencias_generales.all().order_by('id'))
    evidencia_general_legacy_url = None
    if not evidencias_generales and documento.evidencia_general:
        evidencia_general_legacy_url = documento.evidencia_general.url

    trabajadores_difusion, cargos_difusion = _obtener_trabajadores_para_difusion(documento)
    permiso_documento = _permiso_llenado_para_documento(documento)
    cargo_display_documento = 'Todos'
    if permiso_documento:
        cargos_permiso = list(
            permiso_documento.cargos.filter(status=True)
            .order_by('correlativo', 'id')
            .values_list('valor', flat=True)
        )
        if cargos_permiso:
            cargo_display_documento = ', '.join(cargos_permiso)

    tiempo_estimado_minutos = 0
    tiempo_estimado_documento = 'Sin tiempo estimado'
    if plantilla:
        try:
            tiempo_estimado_minutos = int(getattr(plantilla, 'tiempo_estimado_minutos', 0) or 0)
        except (TypeError, ValueError):
            tiempo_estimado_minutos = 0
        tiempo_estimado_segundos = _tiempo_estimado_a_segundos(tiempo_estimado_minutos)
        if tiempo_estimado_segundos > 0:
            tiempo_estimado_documento = _formatear_tiempo_restante_capacitacion(
                tiempo_estimado_segundos
            )

    curso_intentos_restantes = None
    curso_porcentaje_exigencia = None
    respuestas_curso_iniciales = {}
    if tipo_slug == PrevencionPlantilla.TipoRiesgo.CURSOS.value and plantilla:
        if estado_curso_usuario is None:
            resultado_curso_usuario = (
                PrevencionResultadoCurso.objects
                .filter(documento=documento, trabajador=request.user)
                .order_by('-updated_at', '-id')
                .first()
            )
            estado_curso_usuario = _normalizar_estado_resultado_curso(
                documento,
                resultado_curso_usuario,
            )
        curso_intentos_restantes = int(
            (estado_curso_usuario or {}).get('intentos_restantes') or 0
        )
        curso_porcentaje_exigencia = int(
            (estado_curso_usuario or {}).get('porcentaje_requerido')
            or CURSO_PORCENTAJE_APROBACION_DEFAULT
        )
        borrador_curso = _obtener_borrador_respuestas_curso(
            request,
            documento.id,
            with_meta=True,
        )
        respuestas_curso_iniciales = dict(borrador_curso.get('respuestas') or {})
        borrador_updated_at_ts = int(borrador_curso.get('updated_at_ts') or 0)

        if respuestas_curso_iniciales:
            fecha_difusion_actual = (
                PrevencionDocumentoDifusionTrabajador.objects
                .filter(
                    documento=documento,
                    user_id=request.user.id,
                    status=True,
                )
                .order_by('-fechacreacion', '-id')
                .values_list('fechacreacion', flat=True)
                .first()
            )
            try:
                fecha_difusion_actual_ts = int(fecha_difusion_actual.timestamp()) if fecha_difusion_actual else 0
            except Exception:
                fecha_difusion_actual_ts = 0

            if fecha_difusion_actual_ts > 0 and borrador_updated_at_ts < fecha_difusion_actual_ts:
                _limpiar_borrador_respuestas_curso(request, documento.id)
                respuestas_curso_iniciales = {}

    pdf_view_url = reverse('documento_pdf_view', args=[documento.pk])
    if (
        plantilla
        and modo_contenido == PrevencionPlantilla.ModoContenido.PDF
        and plantilla.archivo_pdf
    ):
        pdf_view_url = reverse('ajax_preview_plantilla_pdf', args=[plantilla.id])

    exigir_tiempo_finalizado = bool(
        estado_bloqueo_actual
        and int(estado_bloqueo_actual.get('remaining_seconds', 0) or 0) > 0
    )

    ocultar_boton_finalizado_por_validacion = bool(
        tipo_documento in TIPOS_CAPACITACION_LIBRE_TRAS_VALIDACION
        and documento_ya_validado
    )
    mostrar_boton_finalizado = bool(
        es_flujo_capacitacion
        and (es_origen_capacitaciones or es_origen_notificaciones)
        and not ocultar_boton_finalizado_por_validacion
    )
    url_finalizar_capacitacion = ''
    if mostrar_boton_finalizado:
        params_finalizar = {
            'origen': 'capacitaciones' if es_origen_capacitaciones else 'notificaciones',
        }
        if tipo_origen in TIPOS_DOCUMENTO_PERMITIDOS:
            params_finalizar['tipo'] = tipo_origen
        if back_url:
            params_finalizar['next'] = back_url
        if notif_doc_id.isdigit():
            params_finalizar['notif_doc'] = notif_doc_id
        url_finalizar_capacitacion = (
            f"{reverse('finalizar_capacitacion_documento', kwargs={'pk': documento.pk})}"
            f"?{urlencode(params_finalizar)}"
        )
    url_guardar_resultado_curso = ''
    url_guardar_borrador_curso = ''
    if mostrar_boton_finalizado and tipo_slug == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        url_guardar_resultado_curso = reverse(
            'guardar_resultado_capacitacion_documento',
            kwargs={'pk': documento.pk},
        )
        url_guardar_borrador_curso = reverse(
            'guardar_borrador_capacitacion_documento',
            kwargs={'pk': documento.pk},
        )

    context = {
        'documento': documento,
        'contenido_documento': contenido_documento,
        'titulo_detalle_documento': _construir_titulo_detalle_documento(documento),
        'titulo_vista_documento': titulo_vista_documento,
        'cargo_display_documento': cargo_display_documento,
        'tiempo_estimado_documento': tiempo_estimado_documento,
        'curso_intentos_restantes': curso_intentos_restantes,
        'curso_porcentaje_exigencia': curso_porcentaje_exigencia,
        'tipo_slug': tipo_slug,
        'tipo_label': tipo_label,
        'modo_contenido': modo_contenido,
        'estructura_lectura': estructura_lectura,
        'pdf_view_url': pdf_view_url,
        'trabajadores_difusion': trabajadores_difusion,
        'cargos_difusion': cargos_difusion,
        'evidencias_generales': evidencias_generales,
        'evidencia_general_legacy_url': evidencia_general_legacy_url,
        'sidebarmain': sidebarmain,
        'sidebarmenu': sidebarmenu,
        'sidebar': 'historial_doc',
        'back_url': back_url,
        'mostrar_boton_finalizado': mostrar_boton_finalizado,
        'url_finalizar_capacitacion': url_finalizar_capacitacion,
        'url_guardar_resultado_curso': url_guardar_resultado_curso,
        'url_guardar_borrador_curso': url_guardar_borrador_curso,
        'respuestas_curso_iniciales_json': json.dumps(
            respuestas_curso_iniciales or {},
            cls=DjangoJSONEncoder,
        ),
        'aviso_tiempo_bloqueo_capacitacion': aviso_tiempo_bloqueo_capacitacion,
        'aviso_minutos_bloqueo_capacitacion': aviso_minutos_bloqueo_capacitacion,
        'contador_segundos_restantes': contador_segundos_restantes,
        'exigir_tiempo_finalizado': exigir_tiempo_finalizado,
        'tiempo_expirado_al_ingresar': tiempo_expirado_al_ingresar,
        'alerta_capacitacion_activa': alerta_capacitacion_activa,
        'es_flujo_capacitacion': es_flujo_capacitacion,
    }
    return render(request, 'pages/prevencion/view_documento.html', context)


@login_required
@capacitaciones_y_docs_required
def guardar_borrador_capacitacion_documento(request, pk):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido.'}, status=405)

    documentos_visibles_ids = _ids_documentos_visibles_para_usuario(request.user)
    documentos = (
        PrevencionDocumento.objects
        .filter(status=True, id__in=documentos_visibles_ids)
        .select_related('plantilla_base')
    )
    documento = get_object_or_404(documentos, pk=pk)

    tipo_documento = str(getattr(getattr(documento, 'plantilla_base', None), 'tipo', '') or '').strip().lower()
    if tipo_documento != PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        return JsonResponse({'ok': False, 'error': 'El documento no corresponde a un curso.'}, status=400)

    payload = _parsear_json_request(request)
    respuestas_raw = payload.get('respuestas') if isinstance(payload, dict) else {}
    respuestas_guardadas = _guardar_borrador_respuestas_curso(
        request,
        documento.id,
        respuestas_raw,
    )

    bloqueo = request.session.get(CAPACITACION_BLOQUEO_SESSION_KEY) 
    remaining_seconds = 0
    
    if (
        isinstance(bloqueo, dict)
        and int(bloqueo.get('documento_id') or 0) == documento.id
    ):
        now_ts = int(timezone.now().timestamp())
        expires_at_ts = int(float(bloqueo.get('expires_at_ts') or 0))
        remaining_seconds = max(0, expires_at_ts - now_ts)

    return JsonResponse({
        'ok': True,
        'respuestas': respuestas_guardadas,
        'total_respuestas': len(respuestas_guardadas),
        'remaining_seconds': remaining_seconds,
    })


@login_required
@capacitaciones_y_docs_required
def finalizar_capacitacion_documento(request, pk):
    origen = str(request.GET.get('origen') or '').strip().lower()
    notif_doc_id = str(request.GET.get('notif_doc') or '').strip()
    acceso_desde_capacitaciones = origen in {'capacitaciones', 'notificaciones'} or notif_doc_id.isdigit()

    if acceso_desde_capacitaciones:
        documentos_visibles_ids = _ids_documentos_visibles_para_usuario(request.user)
        documentos = PrevencionDocumento.objects.filter(
            status=True,
            id__in=documentos_visibles_ids,
        )
    else:
        documentos = _filtrar_queryset_por_faena_usuario(
            PrevencionDocumento.objects.filter(status=True),
            request.user
        )

    documento = get_object_or_404(documentos, pk=pk)
    plantilla = documento.plantilla_base
    tipo_documento = str(getattr(plantilla, 'tipo', '') or '').strip().lower()
    tipo_origen = str(request.GET.get('tipo') or '').strip().lower()
    next_raw = str(request.GET.get('next') or '').strip()

    es_origen_capacitaciones = origen == 'capacitaciones'
    es_origen_notificaciones = origen == 'notificaciones'

    if es_origen_notificaciones and tipo_origen not in TIPOS_DOCUMENTO_PERMITIDOS:
        tipo_origen = tipo_documento if tipo_documento in TIPOS_DOCUMENTO_PERMITIDOS else ''

    if es_origen_capacitaciones:
        if tipo_origen not in TIPOS_DOCUMENTO_PERMITIDOS:
            tipo_origen = tipo_documento if tipo_documento in TIPOS_DOCUMENTO_PERMITIDOS else ''
        es_origen_capacitaciones = tipo_origen in TIPOS_DOCUMENTO_PERMITIDOS

    back_url = reverse('history_prevencion')
    if es_origen_capacitaciones:
        back_url = reverse('manage_capacitaciones', kwargs={'tipo_slug': tipo_origen})
    elif es_origen_notificaciones:
        back_url = reverse('view_all_approval_notifications')

    if next_raw and url_has_allowed_host_and_scheme(
        url=next_raw,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        next_parts = urlparse(next_raw)
        back_url = next_parts.path or '/'
        if next_parts.query:
            back_url = f"{back_url}?{next_parts.query}"
        if next_parts.fragment:
            back_url = f"{back_url}#{next_parts.fragment}"

    params_view = {
        'origen': 'capacitaciones' if es_origen_capacitaciones else 'notificaciones',
    }
    if tipo_origen in TIPOS_DOCUMENTO_PERMITIDOS:
        params_view['tipo'] = tipo_origen
    if back_url:
        params_view['next'] = back_url
    if notif_doc_id.isdigit():
        params_view['notif_doc'] = notif_doc_id
    url_volver_documento = (
        f"{reverse('view_documento', kwargs={'pk': documento.pk})}"
        f"?{urlencode(params_view)}"
    )

    aplica_bloqueo_capacitacion = (
        tipo_documento in TIPOS_DOCUMENTO_PERMITIDOS
        and acceso_desde_capacitaciones
    )
    estado_bloqueo = _obtener_estado_bloqueo_capacitacion(request)

    if aplica_bloqueo_capacitacion and estado_bloqueo:
        documento_bloqueado_id = int(estado_bloqueo.get('documento_id') or 0)
        if documento_bloqueado_id != documento.id:
            messages.error(
                request,
                "Hay otra capacitación activa sin finalizar. Vuelve a ese documento para completarla."
            )
            return redirect(url_volver_documento)

    segundos_restantes_bloqueo = max(0, int(estado_bloqueo.get('remaining_seconds', 0) or 0)) if estado_bloqueo else 0

    if (
        tipo_documento in TIPOS_CAPACITACION_LIBRE_TRAS_VALIDACION
        and aplica_bloqueo_capacitacion
        and estado_bloqueo
        and int(estado_bloqueo.get('documento_id') or 0) == documento.id
        and segundos_restantes_bloqueo > 0
    ):
        messages.error(
            request,
            "Debes esperar a que termine el tiempo de lectura antes de validar esta capacitación."
        )
        return redirect(url_volver_documento)

    now = timezone.now()
    estado_curso = None
    if tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        resultado_curso = (
            PrevencionResultadoCurso.objects
            .filter(documento=documento, trabajador=request.user)
            .order_by('-updated_at', '-id')
            .first()
        )
        estado_curso = _normalizar_estado_resultado_curso(documento, resultado_curso)
        if not resultado_curso or int(estado_curso.get('intentos_realizados') or 0) <= 0:
            messages.error(
                request,
                "Debes completar el curso antes de finalizarlo."
            )
            return redirect(url_volver_documento)

    if tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        if bool((estado_curso or {}).get('aprobado')):
            PrevencionNotificacionDocumento.objects.filter(
                destinatario=request.user,
                documento=documento,
            ).update(
                estado=PrevencionNotificacionDocumento.Estado.VISTA,
                leida=True,
                fecha_lectura=now,
            )
        else:
            PrevencionNotificacionDocumento.objects.filter(
                destinatario=request.user,
                documento=documento,
            ).update(
                estado=PrevencionNotificacionDocumento.Estado.PENDIENTE,
                leida=False,
                fecha_lectura=None,
            )
    else:
        PrevencionNotificacionDocumento.objects.filter(
            destinatario=request.user,
            documento=documento,
        ).update(
            estado=PrevencionNotificacionDocumento.Estado.VISTA,
            leida=True,
            fecha_lectura=now,
        )

    if tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        _limpiar_borrador_respuestas_curso(request, documento.id)

    if (
        tipo_documento in TIPOS_CAPACITACION_LIBRE_TRAS_VALIDACION
        and aplica_bloqueo_capacitacion
        and estado_bloqueo
        and int(estado_bloqueo.get('documento_id') or 0) == documento.id
    ):
        # En tipos no curso, al validar se libera de inmediato la restricción.
        request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
        request.session.pop(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY, None)
        request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
        request.session.modified = True

    # No liberar el bloqueo al finalizar manualmente:
    # la restricción de tiempo debe mantenerse hasta que expire el contador.
    if (
        tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value
        and aplica_bloqueo_capacitacion
        and estado_bloqueo
        and int(estado_bloqueo.get('documento_id') or 0) == documento.id
        and segundos_restantes_bloqueo > 0
        and not bool((estado_curso or {}).get('aprobado'))
    ):
        request.session[CAPACITACION_CURSO_FINALIZADO_SESSION_KEY] = {
            'documento_id': documento.id,
            'tipo_slug': tipo_documento,
            'finalizado_at_ts': int(now.timestamp()),
        }
        request.session.modified = True
    else:
        request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)

    bloqueo_expirado = request.session.get(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY)
    if isinstance(bloqueo_expirado, dict):
        try:
            documento_expirado_id = int(bloqueo_expirado.get('documento_id') or 0)
        except (TypeError, ValueError):
            documento_expirado_id = 0
        if documento_expirado_id == documento.id:
            request.session.pop(CAPACITACION_BLOQUEO_EXPIRADO_SESSION_KEY, None)
            request.session.modified = True

    if tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        intentos_restantes = int((estado_curso or {}).get('intentos_restantes') or 0)
        if bool((estado_curso or {}).get('aprobado')):
            messages.success(request, "Curso aprobado.")
        elif intentos_restantes > 0:
            messages.error(
                request,
                f"Curso rechazado. Te quedan {intentos_restantes} intento{'s' if intentos_restantes != 1 else ''}."
            )
        else:
            messages.error(request, "Curso rechazado. No quedan intentos disponibles.")
    else:
        messages.success(request, "Capacitación validada.")
    return redirect(back_url)


@login_required
@capacitaciones_y_docs_required
def guardar_resultado_capacitacion_documento(request, pk):
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'Método no permitido.'}, status=405)

    documentos_visibles_ids = _ids_documentos_visibles_para_usuario(request.user)
    documentos = (
        PrevencionDocumento.objects
        .filter(status=True, id__in=documentos_visibles_ids)
        .select_related('plantilla_base')
    )
    documento = get_object_or_404(documentos, pk=pk)

    tipo_documento = str(getattr(getattr(documento, 'plantilla_base', None), 'tipo', '') or '').strip().lower()
    if tipo_documento != PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        return JsonResponse({'ok': False, 'error': 'El documento no corresponde a un curso.'}, status=400)

    if request.session.get(CAPACITACION_BLOQUEO_SESSION_KEY):
        request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
        request.session.modified = True

    payload = _parsear_json_request(request)
    if not isinstance(payload, dict):
        payload = {}
    
    forzar_rechazo = payload.get('forzar_rechazo') is True
    
    if forzar_rechazo:
        respuestas_marcadas = {}
        _limpiar_borrador_respuestas_curso(request, documento.id)
        
        difusion = PrevencionDocumentoDifusionTrabajador.objects.filter(
            documento=documento, user=request.user
        ).first()
        if difusion:
            motivo_actual = str(difusion.motivo_rechazo or '').lower()
            if 'identidad' not in motivo_actual and 'carnet' not in motivo_actual:
                difusion.motivo_rechazo = 'cancel'
                difusion.save(update_fields=['motivo_rechazo'])
    else:
        respuestas_brutas = payload.get('respuestas')
        respuestas_marcadas = respuestas_brutas if isinstance(respuestas_brutas, dict) else {}
        
        difusion = PrevencionDocumentoDifusionTrabajador.objects.filter(
            documento=documento, user=request.user
        ).first()
        if difusion and difusion.motivo_rechazo:
            difusion.motivo_rechazo = None
            difusion.save(update_fields=['motivo_rechazo'])
        
    evaluacion = _evaluar_resultado_curso(documento, respuestas_marcadas)
    
    estado_base = _normalizar_estado_resultado_curso(documento, None)
    total_intentos = int(estado_base.get('total_intentos') or CURSO_TOTAL_INTENTOS_DEFAULT)
    porcentaje_requerido = int(estado_base.get('porcentaje_requerido') or CURSO_PORCENTAJE_APROBACION_DEFAULT)

    resultados_qs = (
        PrevencionResultadoCurso.objects
        .filter(documento=documento, trabajador=request.user)
        .order_by('-updated_at', '-id')
    )
    resultado = resultados_qs.first()
    estado_previo = _normalizar_estado_resultado_curso(documento, resultado)

    if bool(estado_previo.get('aprobado')):
        _limpiar_borrador_respuestas_curso(request, documento.id)
        return JsonResponse(
            {
                'ok': False,
                'error': 'Este curso ya está aprobado y no se puede volver a intentar.',
                'estado_curso': 'aprobado',
                'intentos_restantes': int(estado_previo.get('intentos_restantes') or 0),
                'total_intentos': total_intentos,
                'porcentaje_requerido': porcentaje_requerido,
            },
            status=409,
        )

    if int(estado_previo.get('intentos_restantes') or 0) <= 0:
        _limpiar_borrador_respuestas_curso(request, documento.id)
        return JsonResponse(
            {
                'ok': False,
                'error': 'No quedan intentos disponibles para este curso.',
                'estado_curso': 'rechazado',
                'intentos_restantes': 0,
                'total_intentos': total_intentos,
                'porcentaje_requerido': porcentaje_requerido,
            },
            status=409,
        )

    intentos_nuevos = int(estado_previo.get('intentos_realizados') or 0) + 1
    aprobado_actual = float(evaluacion.get('porcentaje_aprobacion') or 0.0) >= float(porcentaje_requerido)
    intentos_restantes = max(total_intentos - intentos_nuevos, 0)

    if resultado:
        # Si quedaron duplicados históricos, conserva solo el último.
        resultados_qs.exclude(id=resultado.id).delete()
        resultado.total_preguntas = int(evaluacion.get('total_preguntas') or 0)
        resultado.total_correctas = int(evaluacion.get('total_correctas') or 0)
        resultado.porcentaje_aprobacion = float(evaluacion.get('porcentaje_aprobacion') or 0.0)
        resultado.intentos_realizados = intentos_nuevos
        resultado.aprobado = aprobado_actual
        resultado.detalle_respuestas = evaluacion.get('detalle') or []
        # save() completo para que auto_now actualice updated_at.
        resultado.save()
    else:
        resultado = PrevencionResultadoCurso.objects.create(
            documento=documento,
            trabajador=request.user,
            total_preguntas=int(evaluacion.get('total_preguntas') or 0),
            total_correctas=int(evaluacion.get('total_correctas') or 0),
            porcentaje_aprobacion=float(evaluacion.get('porcentaje_aprobacion') or 0.0),
            intentos_realizados=intentos_nuevos,
            aprobado=aprobado_actual,
            detalle_respuestas=evaluacion.get('detalle') or [],
        )

    _registrar_historial_intento_curso(
        documento=documento,
        trabajador=request.user,
        numero_intento=intentos_nuevos,
        evaluacion=evaluacion,
        aprobado=aprobado_actual,
        fecha_intento=timezone.now(),
    )

    _limpiar_borrador_respuestas_curso(request, documento.id)

    session_key = f'inicio_intento_{documento.id}'
    segundos_empleados = 0
    
    if session_key in request.session:
        try:
            inicio_str = request.session[session_key]
            inicio_dt = datetime.fromisoformat(inicio_str)
            segundos_empleados = max(int((timezone.now() - inicio_dt).total_seconds()), 0)
            del request.session[session_key]
            request.session.modified = True
        except Exception:
            pass
            
    ultimo_intento_guardado = PrevencionResultadoCursoIntento.objects.filter(
        documento=documento, trabajador=request.user
    ).order_by('-id').first()
    
    if ultimo_intento_guardado:
        ultimo_intento_guardado.tiempo_segundos = segundos_empleados
        ultimo_intento_guardado.save(update_fields=['tiempo_segundos'])

    return JsonResponse({
        'ok': True,
        'resultado_id': resultado.id,
        'total_preguntas': resultado.total_preguntas,
        'total_correctas': resultado.total_correctas,
        'porcentaje_aprobacion': resultado.porcentaje_aprobacion,
        'estado_curso': 'aprobado' if aprobado_actual else 'rechazado',
        'aprobado': aprobado_actual,
        'porcentaje_requerido': porcentaje_requerido,
        'intentos_realizados': intentos_nuevos,
        'total_intentos': total_intentos,
        'intentos_restantes': intentos_restantes,
    })

@login_required
@capacitaciones_y_docs_required
def view_resultado_capacitacion_documento(request, documento_id, user_id):
    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True).select_related('faena', 'plantilla_base'),
        request.user
    )
    documento = get_object_or_404(documentos, pk=documento_id)

    tipo_documento = str(getattr(getattr(documento, 'plantilla_base', None), 'tipo', '') or '').strip().lower()
    if tipo_documento != PrevencionPlantilla.TipoRiesgo.CURSOS.value:
        messages.error(request, "Esta vista de resultado está disponible solo para cursos.")
        return redirect('documento_difusion_trabajadores', documento_id=documento.id)

    registro = get_object_or_404(
        PrevencionDocumentoDifusionTrabajador.objects.select_related('user').filter(
            documento=documento,
            user_id=user_id,
            status=True,
        )
    )

    resultado = (
        PrevencionResultadoCurso.objects
        .filter(documento=documento, trabajador_id=user_id)
        .order_by('-updated_at', '-id')
        .first()
    )
    intentos_qs = (
        PrevencionResultadoCursoIntento.objects
        .filter(documento=documento, trabajador_id=user_id)
        .order_by('-numero_intento', '-id')
    )
    intento_numero_solicitado = 0
    try:
        intento_numero_solicitado = int(request.GET.get('intento') or 0)
    except (TypeError, ValueError):
        intento_numero_solicitado = 0
    intento_seleccionado = None
    if intento_numero_solicitado > 0:
        intento_seleccionado = intentos_qs.filter(numero_intento=intento_numero_solicitado).first()
    if not intento_seleccionado:
        intento_seleccionado = intentos_qs.first()

    evaluacion_vacia = _evaluar_resultado_curso(documento, {})
    detalle_respuestas = evaluacion_vacia.get('detalle') or []
    total_preguntas = int(evaluacion_vacia.get('total_preguntas') or 0)
    total_correctas = 0
    porcentaje_aprobacion = 0.0
    fecha_resultado = None
    intento_actual_numero = None

    if intento_seleccionado:
        if isinstance(intento_seleccionado.detalle_respuestas, list) and intento_seleccionado.detalle_respuestas:
            detalle_respuestas = intento_seleccionado.detalle_respuestas
        total_preguntas = int(intento_seleccionado.total_preguntas or total_preguntas)
        total_correctas = int(intento_seleccionado.total_correctas or 0)
        porcentaje_aprobacion = float(intento_seleccionado.porcentaje_aprobacion or 0.0)
        fecha_resultado = intento_seleccionado.fechacreacion
        intento_actual_numero = int(intento_seleccionado.numero_intento or 0)

    elif resultado:
        if isinstance(resultado.detalle_respuestas, list) and resultado.detalle_respuestas:
            detalle_respuestas = resultado.detalle_respuestas
        total_preguntas = int(resultado.total_preguntas or total_preguntas)
        total_correctas = int(resultado.total_correctas or 0)
        porcentaje_aprobacion = float(resultado.porcentaje_aprobacion or 0.0)
        fecha_resultado = resultado.updated_at
        intento_actual_numero = int(getattr(resultado, 'intentos_realizados', 0) or 0) or None

    detalle_respuestas_normalizado = []
    for idx, pregunta in enumerate(detalle_respuestas):
        if not isinstance(pregunta, dict):
            continue
        numero_pregunta = int(pregunta.get('numero') or (idx + 1))
        titulo_raw = str(pregunta.get('titulo') or '').strip()
        titulo_mostrar = titulo_raw
        if titulo_raw.lower() == f"pregunta {numero_pregunta}".lower():
            titulo_mostrar = ''

        pregunta_normalizada = dict(pregunta)
        pregunta_normalizada['numero'] = numero_pregunta
        pregunta_normalizada['titulo_mostrar'] = titulo_mostrar
        detalle_respuestas_normalizado.append(pregunta_normalizada)

    detalle_respuestas = detalle_respuestas_normalizado

    trabajador = registro.user
    nombre_trabajador = (
        registro.nombre_completo
        or (trabajador.get_full_name().strip() if trabajador else '')
        or (trabajador.username if trabajador else registro.rut_trabajador or '-')
    )
    base_url_resultado = reverse(
        'view_resultado_capacitacion_documento',
        kwargs={'documento_id': documento.id, 'user_id': user_id},
    )
    lista_intentos = []
    intentos_map = {}
    for intento in (
        PrevencionResultadoCursoIntento.objects
        .filter(documento=documento, trabajador_id=user_id)
        .order_by('numero_intento', 'id')
    ):
        numero = int(getattr(intento, 'numero_intento', 0) or 0)
        if numero <= 0:
            continue
        porcentaje = float(getattr(intento, 'porcentaje_aprobacion', 0.0) or 0.0)
        intentos_map[numero] = {
            'numero': numero,
            'porcentaje': porcentaje,
            'fecha': getattr(intento, 'fechacreacion', None),
            'url': f"{base_url_resultado}?intento={numero}",
        }

    intentos_esperados = int(getattr(resultado, 'intentos_realizados', 0) or 0) if resultado else 0
    if intentos_map:
        intentos_esperados = max(intentos_esperados, max(intentos_map.keys()))
    if intentos_esperados <= 0 and resultado:
        intentos_esperados = 1

    for numero in range(1, intentos_esperados + 1):
        if numero in intentos_map:
            lista_intentos.append(intentos_map[numero])
        else:
            lista_intentos.append({
                'numero': numero,
                'porcentaje': None,
                'fecha': None,
                'url': '',
            })

    if resultado and intentos_esperados > 0 and (intentos_esperados not in intentos_map):
        lista_intentos[-1] = {
            'numero': intentos_esperados,
            'porcentaje': float(getattr(resultado, 'porcentaje_aprobacion', 0.0) or 0.0),
            'fecha': getattr(resultado, 'updated_at', None),
            'url': base_url_resultado,
        }

    context = {
        'documento': documento,
        'registro': registro,
        'nombre_trabajador': nombre_trabajador,
        'rut_trabajador': registro.rut_trabajador or (trabajador.username if trabajador else '-'),
        'cargo_trabajador': registro.cargo or '-',
        'total_preguntas': total_preguntas,
        'total_correctas': total_correctas,
        'porcentaje_aprobacion': porcentaje_aprobacion,
        'detalle_respuestas': detalle_respuestas,
        'fecha_resultado': fecha_resultado,
        'intento_actual_numero': intento_actual_numero,
        'lista_intentos': lista_intentos,
        'titulo_detalle_documento': _construir_titulo_detalle_documento(documento),
        'back_url': reverse('documento_difusion_trabajadores', kwargs={'documento_id': documento.id}),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'historial_doc',
    }
    return render(request, 'pages/prevencion/view_resultado_curso.html', context)


@login_required
@prevencion_riesgo_required
def documento_difusion_trabajadores(request, documento_id):
    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True).select_related('faena', 'plantilla_base'),
        request.user
    )
    documento = get_object_or_404(documentos, pk=documento_id)
    tipo_documento = str(getattr(getattr(documento, 'plantilla_base', None), 'tipo', '') or '').strip().lower()
    es_documento_curso = tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value
    estado_leido_capacitacion_label = (
        'Validado'
        if tipo_documento in {
            PrevencionPlantilla.TipoRiesgo.CHARLAS.value,
            PrevencionPlantilla.TipoRiesgo.INFORMATIVOS.value,
            PrevencionPlantilla.TipoRiesgo.DIFUSIONES.value,
        }
        else 'Leído'
    )
    if documento.plantilla_base and not _plantilla_autorizada_para_documentos(documento.plantilla_base):
        messages.error(
            request,
            "Este documento aún no está habilitado para difusión: falta aprobación de todos los autorizadores."
        )
        return redirect('history_prevencion')
    _inicializar_snapshot_difusion_documento(documento, request.user)

    filas_qs = (
        PrevencionDocumentoDifusionTrabajador.objects.filter(
            documento=documento,
            status=True,
        )
        .select_related('faena', 'user')
        .order_by('nombre_completo', 'rut_trabajador', 'id')
    )

    def _formatear_fecha_hora_detalle(fecha):
        if not fecha:
            return '-'
        try:
            valor = timezone.localtime(fecha) if timezone.is_aware(fecha) else fecha
            return valor.strftime('%d/%m/%Y %H:%M')
        except Exception:
            return '-'

    def _formatear_porcentaje_detalle(valor):
        try:
            numero = float(valor or 0.0)
        except (TypeError, ValueError):
            numero = 0.0
        return f"{numero:.2f}%".replace('.', ',')

    filas = []
    cache_nombres_difusor = {}
    user_ids_en_tabla = set()
    for registro in filas_qs:
        nombre_usuario = ''
        if registro.user:
            nombre_usuario = registro.user.get_full_name().strip() or registro.user.username

        difundido_por_raw = str(registro.creado_por or '').strip()
        if not difundido_por_raw:
            difundido_por_raw = str(getattr(documento, 'creador', '') or '').strip()
        if difundido_por_raw in cache_nombres_difusor:
            difundido_por = cache_nombres_difusor[difundido_por_raw]
        else:
            difundido_por = _resolver_nombre_difusor(difundido_por_raw) or '-'
            cache_nombres_difusor[difundido_por_raw] = difundido_por
  
        if difundido_por_raw and difundido_por != '-' and difundido_por != difundido_por_raw:
            registro.creado_por = difundido_por
            registro.save(update_fields=['creado_por'])

        filas.append({
            'registro': registro,
            'user': registro.user,
            'username': registro.rut_trabajador or (registro.user.username if registro.user else '-'),
            'nombre': registro.nombre_completo or nombre_usuario or '-',
            'cargo': registro.cargo or '-',
            'faena_nombre': (
                registro.faena.faena
                if registro.faena_id else
                (documento.faena.faena if documento.faena_id else '-')
             ),
            'difundido_por': difundido_por or '-',
            'area': registro.area or '-',
            'ges': registro.ges or '-',
            'motivo_rechazo': registro.motivo_rechazo,
        })
        if registro.user_id:
            user_ids_en_tabla.add(registro.user_id)

    estado_notificacion_por_user_id = {}
    detalle_notificacion_por_user_id = {}
    if user_ids_en_tabla:
        notificaciones_qs = (
            PrevencionNotificacionDocumento.objects
            .filter(documento=documento, destinatario_id__in=user_ids_en_tabla)
            .order_by('destinatario_id', '-fechacreacion', '-id')
            .values(
                'destinatario_id',
                'estado',
                'leida',
                'fechacreacion',
                'fecha_inicio_lectura',
                'fecha_lectura',
            )
        )
        for notificacion in notificaciones_qs:
            user_id = int(notificacion.get('destinatario_id') or 0)
            if not user_id or user_id in estado_notificacion_por_user_id:
                continue
            estado_raw = str(notificacion.get('estado') or '').strip().lower()
            leida = bool(notificacion.get('leida'))
            if leida or estado_raw == PrevencionNotificacionDocumento.Estado.VISTA:
                estado_notificacion_por_user_id[user_id] = {
                    'estado_label': estado_leido_capacitacion_label,
                    'estado_key': 'leido',
                }
            else:
                estado_notificacion_por_user_id[user_id] = {
                    'estado_label': 'Entregado',
                    'estado_key': 'entregado',
                }
            detalle_notificacion_por_user_id[user_id] = {
                'fecha_entregado': notificacion.get('fechacreacion'),
                'fecha_inicio_lectura': notificacion.get('fecha_inicio_lectura'),
                'fecha_validado': notificacion.get('fecha_lectura'),
                'estado_raw': estado_raw,
                'leida': leida,
            }

    resultados_curso_por_user_id = {}
    estado_curso_por_user_id = {}
    intentos_historial_por_user_id = defaultdict(list)
    if es_documento_curso and user_ids_en_tabla:
        resultados_qs = (
            PrevencionResultadoCurso.objects
            .filter(documento=documento, trabajador_id__in=user_ids_en_tabla)
            .select_related('documento__plantilla_base')
            .order_by('trabajador_id', '-updated_at', '-id')
        )
        for resultado in resultados_qs:
            uid = int(getattr(resultado, 'trabajador_id', 0) or 0)
            if not uid:
                continue
            if uid not in estado_curso_por_user_id:
                estado_curso_por_user_id[uid] = _normalizar_estado_resultado_curso(
                    documento,
                    resultado,
                )
            if uid not in resultados_curso_por_user_id:
                resultados_curso_por_user_id[uid] = {
                    'total_correctas': int(getattr(resultado, 'total_correctas', 0) or 0),
                    'total_preguntas': int(getattr(resultado, 'total_preguntas', 0) or 0),
                    'porcentaje_aprobacion': float(getattr(resultado, 'porcentaje_aprobacion', 0.0) or 0.0),
                }
        intentos_qs = (
            PrevencionResultadoCursoIntento.objects
            .filter(documento=documento, trabajador_id__in=user_ids_en_tabla)
            .order_by('trabajador_id', 'numero_intento', 'id')
        )
        for intento in intentos_qs:
            uid = int(getattr(intento, 'trabajador_id', 0) or 0)
            if uid <= 0:
                continue
            intento_numero = int(getattr(intento, 'numero_intento', 0) or 0)
            if intento_numero <= 0:
                continue
            intentos_historial_por_user_id[uid].append({
                'numero': intento_numero,
                'porcentaje': float(getattr(intento, 'porcentaje_aprobacion', 0.0) or 0.0),
                'fecha_raw': getattr(intento, 'fechacreacion', None),
                'fecha': _formatear_fecha_hora_detalle(getattr(intento, 'fechacreacion', None)),
                'motivo_rechazo': getattr(intento, 'motivo_rechazo', None),
                'tiempo_segundos': getattr(intento, 'tiempo_segundos', 0),
            })

    for fila in filas:
        user = fila.get('user')
        if not user or not getattr(user, 'id', None):
            fila['estado_notificacion'] = '-'
            fila['estado_notificacion_key'] = 'sin_estado'
            fila['resultado_disponible'] = False
            fila['url_resultado_curso'] = ''
            continue
        estado_notificacion = estado_notificacion_por_user_id.get(user.id) or {
            'estado_label': 'Entregado',
            'estado_key': 'entregado',
        }
        if es_documento_curso:
            estado_curso = estado_curso_por_user_id.get(user.id)
            if not estado_curso:
                estado_curso = _normalizar_estado_resultado_curso(documento, None)
            if estado_curso.get('estado_key') in {'aprobado', 'rechazado'}:
                estado_notificacion = {
                    'estado_label': estado_curso.get('estado_label') or 'Entregado',
                    'estado_key': estado_curso.get('estado_key') or 'entregado',
                }
        fila['estado_notificacion'] = estado_notificacion['estado_label']
        fila['estado_notificacion_key'] = estado_notificacion['estado_key']
        resultado_user = resultados_curso_por_user_id.get(user.id)
        fila['resultado_disponible'] = bool(resultado_user)
        fila['resultado_total_correctas'] = int((resultado_user or {}).get('total_correctas') or 0)
        fila['resultado_total_preguntas'] = int((resultado_user or {}).get('total_preguntas') or 0)
        fila['resultado_porcentaje_aprobacion'] = float((resultado_user or {}).get('porcentaje_aprobacion') or 0.0)
        fila['url_resultado_curso'] = reverse(
            'view_resultado_capacitacion_documento',
            kwargs={'documento_id': documento.id, 'user_id': user.id},
        ) if es_documento_curso else ''

        registro = fila.get('registro')
        detalle_notificacion = detalle_notificacion_por_user_id.get(user.id) or {}
        fecha_entregado = detalle_notificacion.get('fecha_entregado') or getattr(registro, 'fechacreacion', None)
        fecha_activo = detalle_notificacion.get('fecha_inicio_lectura')
        fecha_validado = detalle_notificacion.get('fecha_validado')

        motivo_rechazo = registro.motivo_rechazo if registro else None
        fecha_rechazo = registro.updated_at if (registro and motivo_rechazo) else None
        
        es_rechazado_curso = es_documento_curso and estado_notificacion.get('estado_key') == 'rechazado'
        intentos_user = intentos_historial_por_user_id.get(user.id, [])
        
        if es_documento_curso:
            if estado_notificacion.get('estado_key') == 'aprobado':
                motivo_rechazo = None
                fecha_rechazo = None
                if registro and registro.motivo_rechazo:
                    registro.motivo_rechazo = None
                    registro.save(update_fields=['motivo_rechazo'])
            else:
                if intentos_user:
                    ultimo_intento = intentos_user[-1]
                    motivo_intento = ultimo_intento.get('motivo_rechazo')
                    
                    if motivo_intento:
                        motivo_rechazo = motivo_intento
                    else:
                        motivo_rechazo = registro.motivo_rechazo if registro else None
                        
                    fecha_rechazo = ultimo_intento.get('fecha_raw')
                    
                    if registro and motivo_intento and registro.motivo_rechazo != motivo_intento:
                        registro.motivo_rechazo = motivo_intento
                        registro.save(update_fields=['motivo_rechazo'])

        fila['motivo_rechazo'] = motivo_rechazo

        tiempo_lectura = '-'

        if es_documento_curso and intentos_user:
            tiempos_intentos = []
            
            for intento in intentos_user:
                segundos_bd = intento.get('tiempo_segundos', 0)
                end_time = intento.get('fecha_raw')
                numero = int(intento.get('numero') or 0)
                
                if segundos_bd > 0:
                    tiempos_intentos.append(_formatear_tiempo_restante_capacitacion(segundos_bd))
                elif numero == 1 and fecha_activo and end_time:
                    try:
                        segundos = max(int((end_time - fecha_activo).total_seconds()), 0)
                        tiempos_intentos.append(_formatear_tiempo_restante_capacitacion(segundos))
                    except Exception:
                        tiempos_intentos.append('-')
                else:
                    tiempos_intentos.append('-')
                
            tiempo_lectura = ", ".join(tiempos_intentos)
            
        else:
            if (motivo_rechazo or es_rechazado_curso) and fecha_activo and fecha_rechazo:
                try:
                    segundos = max(int((fecha_rechazo - fecha_activo).total_seconds()), 0)
                except Exception:
                    segundos = 0
                tiempo_lectura = _formatear_tiempo_restante_capacitacion(segundos)
            elif fecha_activo and fecha_validado:
                try:
                    segundos = max(int((fecha_validado - fecha_activo).total_seconds()), 0)
                except Exception:
                    segundos = 0
                tiempo_lectura = _formatear_tiempo_restante_capacitacion(segundos)
                
            elif fecha_activo and not fecha_validado:
                tiempo_lectura = 'Pendiente de validación'

        fila['info_boton_activo'] = bool(user and getattr(user, 'id', None))
        fila['info_fecha_entregado'] = _formatear_fecha_hora_detalle(fecha_entregado)
        fila['info_fecha_activo'] = _formatear_fecha_hora_detalle(fecha_activo)
        fila['info_fecha_validado'] = _formatear_fecha_hora_detalle(fecha_validado)
        fila['info_fecha_rechazo'] = _formatear_fecha_hora_detalle(fecha_rechazo) if fecha_rechazo else '-'
        fila['info_tiempo_lectura'] = tiempo_lectura
        fila['info_es_curso'] = bool(es_documento_curso)
        fila['info_intentos'] = []
        fila['info_intentos_script_id'] = f"info-intentos-{documento.id}-{user.id}"

        if es_documento_curso:
            intentos_resumen = []
            for intento_item in intentos_historial_por_user_id.get(user.id, []):
                intento_num = int(intento_item.get('numero') or 0)
                if intento_num <= 0:
                    continue
                intentos_resumen.append({
                    'numero': intento_num,
                    'porcentaje': _formatear_porcentaje_detalle(intento_item.get('porcentaje')),
                    'fecha': str(intento_item.get('fecha') or '-'),
                    'url': (
                        f"{fila['url_resultado_curso']}?intento={intento_num}"
                        if fila.get('url_resultado_curso') else '#'
                    ),
                })
            intentos_esperados = int(
                (estado_curso_por_user_id.get(user.id, {}) or {}).get('intentos_realizados') or 0
            )
            if intentos_esperados <= 0 and intentos_resumen:
                intentos_esperados = max(item.get('numero', 0) for item in intentos_resumen)
            if intentos_esperados <= 0 and fila.get('resultado_disponible'):
                intentos_esperados = 1

            intentos_por_numero = {
                int(item.get('numero') or 0): item
                for item in intentos_resumen
                if int(item.get('numero') or 0) > 0
            }
            intentos_resumen = []
            
            for numero in range(1, intentos_esperados + 1):
                if numero in intentos_por_numero:
                    intentos_resumen.append(intentos_por_numero[numero])
                else:
                    intentos_resumen.append({
                        'numero': numero,
                        'porcentaje': '-',
                        'fecha': '-',
                        'url': '#',
                    })

            if fila.get('resultado_disponible') and intentos_esperados > 0 and intentos_esperados not in intentos_por_numero:
                intentos_resumen[-1] = {
                    'numero': intentos_esperados,
                    'porcentaje': _formatear_porcentaje_detalle(
                        fila.get('resultado_porcentaje_aprobacion')
                    ),
                    'fecha': _formatear_fecha_hora_detalle(fecha_validado),
                    'url': fila.get('url_resultado_curso') or '#',
                }
            fila['info_intentos'] = intentos_resumen

    claves_estado_validadas = {'aprobado'} if es_documento_curso else {'leido'}
    total_trabajadores = len(filas)
    total_validados = sum(
        1
        for fila in filas
        if str(fila.get('estado_notificacion_key') or '').strip().lower() in claves_estado_validadas
    )
    total_entregados = sum(
        1
        for fila in filas
        if str(fila.get('estado_notificacion_key') or '').strip().lower() == 'entregado'
    )
    porcentaje_validacion = 0.0
    if total_trabajadores > 0:
        porcentaje_validacion = (float(total_validados) / float(total_trabajadores)) * 100.0
    promedio_porcentaje_aprobacion = 0.0
    if es_documento_curso:
        porcentajes_curso_aprobados = [
            float((resultado or {}).get('porcentaje_aprobacion') or 0.0)
            for uid, resultado in resultados_curso_por_user_id.items()
            if bool((estado_curso_por_user_id.get(uid) or {}).get('aprobado'))
        ]
        if porcentajes_curso_aprobados:
            promedio_porcentaje_aprobacion = (
                sum(porcentajes_curso_aprobados) / float(len(porcentajes_curso_aprobados))
            )
    resumen_trabajadores = {
        'name_documento': (
            documento.plantilla_base.nombre
            if documento.plantilla_base_id else
            f"Documento #{documento.id}"
        ),
        'tipo_documento': (
            documento.plantilla_base.get_tipo_display()
            if documento.plantilla_base_id else
            'Documento'
        ),
        'faena_nombre': documento.faena.faena if documento.faena_id else '-',
        'total': total_trabajadores,
        'validados': total_validados,
        'porcentaje_validacion': porcentaje_validacion,
        'entregados': total_entregados,
        'promedio_porcentaje_aprobacion': promedio_porcentaje_aprobacion,
    }

    _, cargos_permitidos = _catalogo_trabajadores_para_difusion_documento(
        documento=documento,
        aplicar_filtro_cargo=True,
    )
    trabajadores_catalogo = _catalogo_todos_los_trabajadores_para_agregar()
    trabajadores_disponibles = []
    dropdown_user_ids = set()

    for trabajador in trabajadores_catalogo:
        usuario = trabajador.get('user')
        if not usuario or not usuario.id:
            continue
        if usuario.id in user_ids_en_tabla or usuario.id in dropdown_user_ids:
            continue
        trabajadores_disponibles.append(trabajador)
        dropdown_user_ids.add(usuario.id)

    inactivos_qs = (
        PrevencionDocumentoDifusionTrabajador.objects.filter(
            documento=documento,
            status=False,
            user__isnull=False,
        )
        .select_related('faena', 'user')
        .order_by('nombre_completo', 'rut_trabajador', 'id')
    )

    for registro in inactivos_qs:
        if not registro.user_id or registro.user_id in user_ids_en_tabla or registro.user_id in dropdown_user_ids:
            continue
        nombre_usuario = registro.user.get_full_name().strip() if registro.user else ''
        trabajadores_disponibles.append({
            'user': registro.user,
            'username': registro.rut_trabajador or (registro.user.username if registro.user else '-'),
            'nombre': registro.nombre_completo or nombre_usuario or (registro.user.username if registro.user else '-'),
            'cargo': registro.cargo or '',
            'faena_id': registro.faena_id or documento.faena_id,
            'faena_nombre': registro.faena.faena if registro.faena_id else (documento.faena.faena if documento.faena_id else ''),
            'area': registro.area or '',
            'ges': registro.ges or '',
        })
        dropdown_user_ids.add(registro.user_id)

    historial_movimientos_trabajadores = []
    historial_qs = (
        PrevencionHistorialDifusionTrabajador.objects
        .filter(documento=documento)
        .select_related('trabajador', 'actor')
        .order_by('-fechacreacion', '-id')[:200]
    )
    for movimiento in historial_qs:
        fecha_movimiento = movimiento.fechacreacion
        if fecha_movimiento and timezone.is_aware(fecha_movimiento):
            fecha_movimiento = timezone.localtime(fecha_movimiento)
        fecha_order = int(fecha_movimiento.timestamp()) if fecha_movimiento else 0

        trabajador_nombre = str(movimiento.trabajador_nombre or '').strip()
        if not trabajador_nombre and movimiento.trabajador:
            trabajador_nombre = _nombre_usuario_legible(movimiento.trabajador)
        if not trabajador_nombre:
            trabajador_nombre = '-'

        trabajador_rut = str(movimiento.trabajador_rut or '').strip()
        if not trabajador_rut and movimiento.trabajador:
            trabajador_rut = str(movimiento.trabajador.username or '').strip()
        if not trabajador_rut:
            trabajador_rut = '-'

        actor_nombre = str(movimiento.actor_nombre or '').strip()
        if not actor_nombre and movimiento.actor:
            actor_nombre = _nombre_usuario_legible(movimiento.actor)
        if not actor_nombre:
            actor_nombre = '-'

        historial_movimientos_trabajadores.append({
            'fecha': fecha_movimiento.strftime('%d/%m/%Y %H:%M') if fecha_movimiento else '-',
            'fecha_order': fecha_order,
            'accion_key': movimiento.accion,
            'accion_label': movimiento.get_accion_display(),
            'trabajador_nombre': trabajador_nombre,
            'trabajador_rut': trabajador_rut,
            'actor_nombre': actor_nombre,
        })

    context = {
        'documento': documento,
        'titulo_detalle_documento': _construir_titulo_detalle_documento(documento),
        'filas_trabajadores': filas,
        'resumen_trabajadores': resumen_trabajadores,
        'trabajadores_disponibles': trabajadores_disponibles,
        'historial_movimientos_trabajadores': historial_movimientos_trabajadores,
        'cargos_difusion': cargos_permitidos,
        'es_documento_curso': es_documento_curso,
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'historial_doc',
    }
    return render(request, 'pages/prevencion/documento_difusion_trabajadores.html', context)


@login_required
@prevencion_riesgo_required
def add_documento_difusion_trabajador(request, documento_id):
    if request.method != 'POST':
        return redirect('documento_difusion_trabajadores', documento_id=documento_id)

    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True).select_related('faena', 'plantilla_base'),
        request.user
    )
    documento = get_object_or_404(documentos, pk=documento_id)
    tipo_documento = str(getattr(getattr(documento, 'plantilla_base', None), 'tipo', '') or '').strip().lower()
    es_documento_curso = tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value
    if documento.plantilla_base and not _plantilla_autorizada_para_documentos(documento.plantilla_base):
        messages.error(
            request,
            "Este documento aún no está habilitado para difusión: falta aprobación de todos los autorizadores."
        )
        return redirect('history_prevencion')
    user_id = str(request.POST.get('user_id') or '').strip()
    if not user_id:
        messages.info(request, "Selecciona un trabajador para agregar.")
        return redirect('documento_difusion_trabajadores', documento_id=documento.id)

    trabajador_activo = PrevencionDocumentoDifusionTrabajador.objects.filter(
        documento=documento,
        user_id=user_id,
        status=True,
    ).exists()
    if trabajador_activo:
        messages.info(request, "El trabajador ya está agregado en la lista.")
        return redirect('documento_difusion_trabajadores', documento_id=documento.id)

    registro_inactivo = (
        PrevencionDocumentoDifusionTrabajador.objects.filter(
            documento=documento,
            user_id=user_id,
            status=False,
        )
        .order_by('-id')
        .first()
    )
    if registro_inactivo:
        nombre_difusor = _nombre_difusor_desde_request(request.user)
        registro_inactivo.status = True
        registro_inactivo.creado_por = nombre_difusor or request.user.username
        registro_inactivo.fechacreacion = timezone.now()
        registro_inactivo.save(update_fields=['status', 'creado_por', 'fechacreacion'])
        if registro_inactivo.user_id:
            if es_documento_curso:
                _reiniciar_resultado_curso_trabajador(documento, registro_inactivo.user_id)
                if registro_inactivo.user_id == request.user.id:
                    _limpiar_estado_sesion_curso_usuario_actual(request, documento.id)
            _crear_notificacion_documento_disponible(documento, registro_inactivo.user)
        _registrar_historial_difusion_documento(
            documento=documento,
            accion=PrevencionHistorialDifusionTrabajador.Accion.AGREGADO,
            actor=request.user,
            trabajador=registro_inactivo.user,
            trabajador_rut=registro_inactivo.rut_trabajador,
            trabajador_nombre=registro_inactivo.nombre_completo,
        )
        messages.success(request, "Trabajador agregado a la lista correctamente.")
        return redirect('documento_difusion_trabajadores', documento_id=documento.id)

    trabajadores_catalogo = _catalogo_todos_los_trabajadores_para_agregar()
    trabajadores_por_id = {
        str(item['user'].id): item
        for item in trabajadores_catalogo
        if item.get('user') and item['user'].id
    }
    trabajador = trabajadores_por_id.get(user_id)
    if not trabajador:
        messages.info(request, "El trabajador seleccionado no está disponible para agregar.")
        return redirect('documento_difusion_trabajadores', documento_id=documento.id)

    nombre_difusor = _nombre_difusor_desde_request(request.user)
    registro_creado = _crear_registro_difusion_documento(
        documento=documento,
        trabajador=trabajador,
        creado_por=nombre_difusor or request.user.username,
        status=True,
    )
    if registro_creado and registro_creado.user_id:
        if es_documento_curso:
            _reiniciar_resultado_curso_trabajador(documento, registro_creado.user_id)
            if registro_creado.user_id == request.user.id:
                _limpiar_estado_sesion_curso_usuario_actual(request, documento.id)
        _crear_notificacion_documento_disponible(documento, registro_creado.user)
    if registro_creado:
        _registrar_historial_difusion_documento(
            documento=documento,
            accion=PrevencionHistorialDifusionTrabajador.Accion.AGREGADO,
            actor=request.user,
            trabajador=getattr(registro_creado, 'user', None),
            trabajador_rut=getattr(registro_creado, 'rut_trabajador', ''),
            trabajador_nombre=getattr(registro_creado, 'nombre_completo', ''),
        )
    messages.success(request, "Trabajador agregado a la lista correctamente.")
    return redirect('documento_difusion_trabajadores', documento_id=documento.id)


@login_required
@prevencion_riesgo_required
def remove_documento_difusion_trabajador(request, documento_id):
    if request.method != 'POST':
        return redirect('documento_difusion_trabajadores', documento_id=documento_id)

    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True).select_related('faena', 'plantilla_base'),
        request.user
    )
    documento = get_object_or_404(documentos, pk=documento_id)
    tipo_documento = str(getattr(getattr(documento, 'plantilla_base', None), 'tipo', '') or '').strip().lower()
    es_documento_curso = tipo_documento == PrevencionPlantilla.TipoRiesgo.CURSOS.value
    if documento.plantilla_base and not _plantilla_autorizada_para_documentos(documento.plantilla_base):
        messages.error(
            request,
            "Este documento aún no está habilitado para difusión: falta aprobación de todos los autorizadores."
        )
        return redirect('history_prevencion')
    user_id = str(request.POST.get('user_id') or '').strip()
    registro_id = str(request.POST.get('registro_id') or '').strip()
    if not user_id and not registro_id:
        messages.error(request, "No se recibió el trabajador a remover.")
        return redirect('documento_difusion_trabajadores', documento_id=documento.id)

    registros_qs = PrevencionDocumentoDifusionTrabajador.objects.filter(
        documento=documento,
        status=True,
    )
    if registro_id:
        registro = registros_qs.filter(id=registro_id).first()
    else:
        registro = registros_qs.filter(user_id=user_id).order_by('-id').first()
    if not registro:
        messages.info(request, "El trabajador no está activo en la lista.")
        return redirect('documento_difusion_trabajadores', documento_id=documento.id)

    registro.status = False
    registro.save(update_fields=['status'])
    if registro.user_id:
        if es_documento_curso:
            _reiniciar_resultado_curso_trabajador(documento, registro.user_id)
        PrevencionNotificacionDocumento.objects.filter(
            destinatario_id=registro.user_id,
            documento=documento,
        ).delete()
        if es_documento_curso and registro.user_id == request.user.id:
            _limpiar_estado_sesion_curso_usuario_actual(request, documento.id)
    _registrar_historial_difusion_documento(
        documento=documento,
        accion=PrevencionHistorialDifusionTrabajador.Accion.QUITADO,
        actor=request.user,
        trabajador=registro.user,
        trabajador_rut=registro.rut_trabajador,
        trabajador_nombre=registro.nombre_completo,
    )
    messages.success(request, "Trabajador eliminado de la lista.")
    return redirect('documento_difusion_trabajadores', documento_id=documento.id)


@login_required
@capacitaciones_y_docs_required
def documento_pdf_view(request, pk):
    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True),
        request.user
    )
    documento = get_object_or_404(documentos, pk=pk)
    if documento.plantilla_base and not _plantilla_autorizada_para_documentos(documento.plantilla_base):
        messages.error(
            request,
            "Este documento aún no está habilitado: falta aprobación de todos los autorizadores."
        )
        return redirect('history_prevencion')

    plantilla = documento.plantilla_base
    if (
        plantilla
        and (plantilla.modo_contenido == PrevencionPlantilla.ModoContenido.PDF)
        and plantilla.archivo_pdf
    ):
        try:
            plantilla.archivo_pdf.open('rb')
        except Exception:
            return HttpResponse(status=404)

        nombre_archivo = os.path.basename(plantilla.archivo_pdf.name or '') or f"documento-{documento.id}.pdf"
        response = FileResponse(plantilla.archivo_pdf, content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="{nombre_archivo}"'
        return response

    contenido_documento = _preparar_documento_para_visualizacion(documento)
    anexo_evidencias = _construir_anexo_evidencias(documento, contenido_documento, request, es_pdf_view=True)
    evidencias_generales = list(documento.evidencias_generales.all().order_by('id'))
    evidencia_general_legacy_url = None
    evidencias_generales_nombres = []
    if not evidencias_generales and documento.evidencia_general:
        evidencia_general_legacy_url = documento.evidencia_general.url
    else:
        evidencias_generales_nombres = [
            _formatear_texto_para_pdf_html(_obtener_nombre_archivo_desde_url(evidencia.archivo.url), longitud_bloque=40)
            for evidencia in evidencias_generales
        ]

    titulo_detalle_documento = _construir_titulo_detalle_documento(documento)
    template = get_template('pages/pdfs/prevencion_documento_pdf.html')
    html = template.render({
        'documento': documento,
        'contenido_documento': contenido_documento,
        'titulo_detalle_documento': titulo_detalle_documento,
        'evidencias_generales': evidencias_generales,
        'evidencias_generales_nombres': evidencias_generales_nombres,
        'evidencia_general_legacy_url': evidencia_general_legacy_url,
        'evidencia_general_legacy_nombre': _formatear_texto_para_pdf_html(_obtener_nombre_archivo_desde_url(evidencia_general_legacy_url), longitud_bloque=40) if evidencia_general_legacy_url else '',
        'observacion_general_pdf': _formatear_texto_para_pdf_html(documento.observacion_general, longitud_bloque=60),
        'fecha_descarga': timezone.localtime(timezone.now()),
        'anexo_evidencias': anexo_evidencias,
    })

    nombre_tipo = 'reporte'
    if documento.plantilla_base_id:
        nombre_tipo = documento.plantilla_base.get_tipo_display()

    nombre_faena = documento.faena.faena if documento.faena_id else 'sin-faena'
    filename = f"{slugify(nombre_tipo)}-{slugify(nombre_faena)}-{documento.id}.pdf"

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    pisa_status = pisa.CreatePDF(html, dest=response, link_callback=_resolver_src_pdf)

    if pisa_status.err:
        return HttpResponse('Error al generar el PDF.', status=500)

    return response

# 3. CREAR DOCUMENTO (Actualizado con Archivos)
@login_required
@prevencion_riesgo_required
def new_documento(request):
    if request.method == 'POST':
        try:
            archivos_validos, msg_archivos = _validar_archivos_subidos(request)
            if not archivos_validos:
                messages.error(request, msg_archivos)
                return redirect('new_documento')

            plantilla_id = request.POST.get('plantilla')
            if not plantilla_id:
                messages.error(request, "Debes seleccionar una plantilla en Configuración de Acceso.")
                return redirect('new_documento')

            plantilla = PrevencionPlantilla.objects.filter(
                id=plantilla_id,
                status=True
            ).first()
            if not plantilla:
                messages.error(request, "La plantilla seleccionada no está disponible.")
                return redirect('new_documento')

            if not _plantilla_autorizada_para_documentos(plantilla):
                messages.error(
                    request,
                    "La plantilla seleccionada aún no está aprobada por todos los autorizadores."
                )
                return redirect('new_documento')

            # Solo se admite una faena por documento.
            faenas_raw = request.POST.get('faenas_ids') or ''
            faenas_destino_ids = []
            seen_faenas = set()
            for raw_id in str(faenas_raw).split(','):
                fid = str(raw_id).strip()
                if fid and fid not in seen_faenas:
                    seen_faenas.add(fid)
                    faenas_destino_ids.append(fid)

            if len(faenas_destino_ids) > 1:
                messages.error(request, "Solo se permite seleccionar una faena.")
                return redirect('new_documento')

            faena_destino_id = faenas_destino_ids[0] if faenas_destino_ids else str(plantilla.faena_id)

            if str(faena_destino_id) != str(plantilla.faena_id):
                messages.error(request, "La plantilla seleccionada no coincide con la faena elegida.")
                return redirect('new_documento')

            if not _faena_permitida_para_usuario(request.user, faena_destino_id):
                messages.error(request, "No tienes permisos para crear documentos en la faena seleccionada.")
                return redirect('new_documento')

            if not _usuario_puede_rellenar_documento(request.user, faena_destino_id, plantilla.id):
                messages.error(
                    request,
                    "Tu cargo no tiene permisos para rellenar esta plantilla en la faena seleccionada."
                )
                return redirect('new_documento')

            contenido_data = request.POST.get('contenido_data') or '[]'
            contenido_json = json.loads(contenido_data)
            observacion = (request.POST.get('observacion') or '').strip()
            creador = _nombre_usuario_legible(request.user)

            with transaction.atomic():
                doc = PrevencionDocumento(
                    plantilla_base=plantilla,
                    faena_id=faena_destino_id,
                    contenido=contenido_json,
                    observacion_general=observacion,
                    creador=creador,
                )
                doc.save()

                for archivo in request.FILES.getlist('evidencia_general'):
                    if hasattr(archivo, 'seek'):
                        archivo.seek(0)
                    PrevencionEvidenciaGeneral.objects.create(documento=doc, archivo=archivo)

                for key in request.FILES:
                    if key == 'evidencia_general':
                        continue

                    archivos = request.FILES.getlist(key)

                    if key.startswith('file_seccion_'):
                        try:
                            idx = int(key.split('_')[-1])
                        except (TypeError, ValueError):
                            continue

                        for archivo in archivos:
                            if hasattr(archivo, 'seek'):
                                archivo.seek(0)
                            PrevencionEvidenciaSeccion.objects.create(
                                documento=doc,
                                indice_seccion=idx,
                                archivo=archivo
                            )

                    elif key.startswith('file_item_'):
                        parts = key.split('_')
                        if len(parts) < 4:
                            continue

                        try:
                            sec_idx = int(parts[2])
                            row_idx = int(parts[3])
                        except (TypeError, ValueError):
                            continue

                        for archivo in archivos:
                            if hasattr(archivo, 'seek'):
                                archivo.seek(0)
                            PrevencionEvidenciaItem.objects.create(
                                documento=doc,
                                indice_seccion=sec_idx,
                                indice_fila=row_idx,
                                archivo=archivo
                            )

                _inicializar_snapshot_difusion_documento(doc, request.user)
                _notificar_documento_a_trabajadores_difusion(doc)

            tipo_slug = str(getattr(plantilla, 'tipo', '') or '').strip().lower()
            tipo_label = _mapear_tipo_documento_label(tipo_slug)
            es_femenino = tipo_slug in {'charlas', 'difusiones'}
            participio = 'creada' if es_femenino else 'creado'
            
            messages.success(request, f"{tipo_label} {participio} correctamente.")

            return redirect('history_prevencion')

        except json.JSONDecodeError:
            messages.error(request, "No se pudo procesar la estructura del formulario. Intenta nuevamente.")
        except Exception as e:
            messages.error(request, f"Error al guardar: {e}")

    faenas = _faenas_seleccionables_para_usuario(request.user)
    tipos_plantilla = [
        (value, label)
        for value, label in PrevencionPlantilla.TipoRiesgo.choices
        if value in TIPOS_DOCUMENTO_PERMITIDOS
    ]
    tipo_documento_inicial = str(request.GET.get('tipo') or '').strip().lower()
    if tipo_documento_inicial not in TIPOS_DOCUMENTO_PERMITIDOS:
        tipo_documento_inicial = ''
    
    context = {
        'faenas': faenas,
        'tipos_plantilla': tipos_plantilla,
        'tipo_documento_inicial': tipo_documento_inicial,
        'titulo_vista_nuevo_documento': _construir_titulo_nuevo_documento(tipo_documento_inicial),
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo',
        # Mantiene resaltado "Documentos" al entrar en Crear.
        'sidebar': 'historial_doc'
    }
    return render(request, 'pages/prevencion/new_documento.html', context)

# 4. EDITAR DOCUMENTO
@login_required
@prevencion_riesgo_required
def edit_documento(request, pk):
    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True),
        request.user
    )
    doc = get_object_or_404(documentos, pk=pk)

    messages.info(request, "La edición de reportes fue deshabilitada. Usa la opción Trabajadores (solo lectura).")
    return redirect('view_documento', pk=doc.pk)

    if doc.plantilla_base_id and not _usuario_puede_rellenar_documento(request.user, doc.faena_id, doc.plantilla_base_id):
        messages.error(request, "Tu cargo no tiene permisos para editar este documento.")
        return redirect('history_prevencion')

    if request.method == 'POST':
        try:
            archivos_validos, msg_archivos = _validar_archivos_subidos(request)
            if not archivos_validos:
                messages.error(request, msg_archivos)
                return redirect('edit_documento', pk=pk)

            archivos_totales_validos, msg_totales = _validar_total_archivos_por_item_en_edicion(doc, request)
            if not archivos_totales_validos:
                messages.error(request, msg_totales)
                return redirect('edit_documento', pk=pk)

            # A. Actualizar Datos Texto
            doc.contenido = json.loads(request.POST.get('contenido_data'))
            doc.observacion_general = request.POST.get('observacion')

            doc.save()

            for archivo in request.FILES.getlist('evidencia_general'):
                PrevencionEvidenciaGeneral.objects.create(documento=doc, archivo=archivo)

            for key in request.FILES:
                if key == 'evidencia_general':
                    continue

                archivos = request.FILES.getlist(key)

                if key.startswith('file_seccion_'):
                    try:
                        idx = int(key.split('_')[-1])
                    except (TypeError, ValueError):
                        continue

                    for archivo in archivos:
                        PrevencionEvidenciaSeccion.objects.create(
                            documento=doc,
                            indice_seccion=idx,
                            archivo=archivo
                        )

                elif key.startswith('file_item_'):
                    parts = key.split('_')
                    if len(parts) < 4:
                        continue

                    try:
                        sec_idx = int(parts[2])
                        row_idx = int(parts[3])
                    except (TypeError, ValueError):
                        continue

                    for archivo in archivos:
                        PrevencionEvidenciaItem.objects.create(
                            documento=doc,
                            indice_seccion=sec_idx,
                            indice_fila=row_idx,
                            archivo=archivo
                        )

            messages.success(request, "Reporte actualizado correctamente.")
            return redirect('history_prevencion')

        except Exception as e:
            messages.error(request, f"Error al actualizar: {e}")

        # PREPARAR DATOS PARA EL TEMPLATE
    evidencias_secciones = defaultdict(list)
    for evidencia in doc.evidencias_secciones.all().order_by('id'):
        evidencias_secciones[str(evidencia.indice_seccion)].append({
            'id': evidencia.id,
            'url': evidencia.archivo.url
        })

    evidencias_items = defaultdict(list)
    for evidencia in doc.evidencias_items.all().order_by('id'):
        key = f"{evidencia.indice_seccion}_{evidencia.indice_fila}"
        evidencias_items[key].append({
            'id': evidencia.id,
            'url': evidencia.archivo.url
        })

    contenido_json = json.dumps(doc.contenido, cls=DjangoJSONEncoder)
    evidencias_secciones_json = json.dumps(evidencias_secciones, cls=DjangoJSONEncoder)
    evidencias_items_json = json.dumps(evidencias_items, cls=DjangoJSONEncoder)
    evidencias_generales = list(doc.evidencias_generales.all().order_by('id'))
    evidencia_general_legacy_url = None
    if not evidencias_generales and doc.evidencia_general:
        evidencia_general_legacy_url = doc.evidencia_general.url

    context = {
        'documento': doc,
        'titulo_detalle_documento': _construir_titulo_detalle_documento(doc),
        'contenido_json': contenido_json,
        'evidencias_secciones_json': evidencias_secciones_json,
        'evidencias_items_json': evidencias_items_json,
        'evidencias_generales': evidencias_generales,
        'evidencia_general_legacy_url': evidencia_general_legacy_url,
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'historial_doc'
    }
    return render(request, 'pages/prevencion/edit_documento.html', context)

# 5. ELIMINAR DOCUMENTO (Soft Delete)
@login_required
@prevencion_riesgo_required
def delete_documento(request, pk):
    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True),
        request.user
    )
    doc = get_object_or_404(documentos, pk=pk)

    # Si el documento eliminado es el bloqueado en capacitación, liberar bloqueo inmediato.
    estado_bloqueo = _obtener_estado_bloqueo_capacitacion(request)
    if estado_bloqueo and int(estado_bloqueo.get('documento_id') or 0) == doc.id:
        request.session.pop(CAPACITACION_BLOQUEO_SESSION_KEY, None)
        request.session.pop(CAPACITACION_CURSO_FINALIZADO_SESSION_KEY, None)
        request.session.modified = True

    doc.status = False # Borrado lógico
    doc.save()
    messages.success(request, "Reporte eliminado correctamente.")
    return redirect('history_prevencion')


# === AJAX HELPERS ===

# A. AJAX: Cargar plantillas por Faena
@login_required
@prevencion_mantenedor_required
def ajax_load_plantillas(request):
    faena_id = request.GET.get('faena_id')
    faena_ids = request.GET.getlist('faena_ids[]') or request.GET.getlist('faena_ids')
    tipo = (request.GET.get('tipo') or '').strip()

    if faena_id:
        faena_consulta = str(faena_id).strip()
    else:
        faenas_consulta = [str(valor).strip() for valor in faena_ids if str(valor).strip()]
        faenas_consulta = list(dict.fromkeys(faenas_consulta))
        if len(faenas_consulta) != 1:
            return JsonResponse([], safe=False)
        faena_consulta = faenas_consulta[0]

    if not faena_consulta:
        return JsonResponse([], safe=False)

    if not _faena_permitida_para_usuario(request.user, faena_consulta):
        return JsonResponse([], safe=False)

    plantillas_qs = PrevencionPlantilla.objects.filter(
        faena_id=faena_consulta,
        status=True,
        tipo__in=TIPOS_DOCUMENTO_PERMITIDOS,
    )

    if tipo:
        if tipo not in TIPOS_DOCUMENTO_PERMITIDOS:
            return JsonResponse([], safe=False)
        plantillas_qs = plantillas_qs.filter(tipo=tipo)

    plantillas = list(plantillas_qs.select_related('faena').order_by('nombre', 'id'))
    plantillas = [p for p in plantillas if _plantilla_autorizada_para_documentos(p)]
    permisos_qs = PrevencionPermisoLlenado.objects.filter(
        status=True,
        faena_id=faena_consulta,
        plantilla_id__in=[p.id for p in plantillas]
    ).prefetch_related('cargos')

    cargos_por_plantilla = {}
    cargo_values_por_plantilla = {}
    cargo_todos_por_plantilla = {}
    for permiso in permisos_qs:
        cargos = list(
            permiso.cargos.order_by('correlativo', 'id').values_list('valor', flat=True)
        )
        cargos = list(dict.fromkeys([str(valor).strip() for valor in cargos if str(valor).strip()]))
        if cargos:
            cargos_por_plantilla[permiso.plantilla_id] = ', '.join(cargos)
            cargo_values_por_plantilla[permiso.plantilla_id] = cargos
            cargo_todos_por_plantilla[permiso.plantilla_id] = False
        else:
            cargos_por_plantilla[permiso.plantilla_id] = 'Todos'
            cargo_values_por_plantilla[permiso.plantilla_id] = ['Todos']
            cargo_todos_por_plantilla[permiso.plantilla_id] = True

    payload = []
    for plantilla in plantillas:
        cargo_values = cargo_values_por_plantilla.get(plantilla.id, ['Todos'])
        cargo_todos = cargo_todos_por_plantilla.get(plantilla.id, True)
        payload.append({
            'id': plantilla.id,
            'nombre': plantilla.nombre,
            'tipo': plantilla.tipo,
            'faena_id': plantilla.faena_id,
            'faena__faena': plantilla.faena.faena if plantilla.faena_id else '',
            'cargo_display': cargos_por_plantilla.get(plantilla.id, 'Todos'),
            'cargo_values': cargo_values,
            'cargo_todos': cargo_todos,
        })
    return JsonResponse(payload, safe=False)

# B. AJAX: Obtener estructura de una plantilla específica
@login_required
@prevencion_mantenedor_required
def ajax_get_structure(request):
    plantilla_id = request.GET.get('plantilla_id')
    plantilla = get_object_or_404(PrevencionPlantilla, pk=plantilla_id, status=True)
    if not _faena_permitida_para_usuario(request.user, plantilla.faena_id):
        return JsonResponse([], safe=False)
    if not _plantilla_autorizada_para_documentos(plantilla):
        return JsonResponse(
            {'error': 'La plantilla aún no está aprobada por todos los autorizadores.'},
            status=403
        )

    preview_images = []
    if plantilla.archivo_pdf and (plantilla.modo_contenido == PrevencionPlantilla.ModoContenido.PDF):
        try:
            paginas = check_and_convert_pdf(plantilla.archivo_pdf.path)
            for ruta in paginas:
                media_url = _ruta_local_a_media_url(ruta)
                if media_url:
                    preview_images.append(media_url)
        except Exception:
            preview_images = []

    payload = {
        'modo_contenido': plantilla.modo_contenido or PrevencionPlantilla.ModoContenido.ESTRUCTURA,
        'estructura': plantilla.estructura or [],
        'archivo_pdf_url': reverse('ajax_preview_plantilla_pdf', args=[plantilla.id]) if plantilla.archivo_pdf else '',
        'archivo_pdf_nombre': os.path.basename(plantilla.archivo_pdf.name) if plantilla.archivo_pdf else '',
        'pdf_preview_images': preview_images,
    }
    return JsonResponse(payload)


@login_required
@prevencion_mantenedor_required
def ajax_preview_plantilla_pdf(request, plantilla_id):
    plantilla = get_object_or_404(PrevencionPlantilla, pk=plantilla_id, status=True)
    if not _faena_permitida_para_usuario(request.user, plantilla.faena_id):
        return HttpResponse(status=403)
    if not _plantilla_autorizada_para_documentos(plantilla):
        return HttpResponse(status=403)

    if not plantilla.archivo_pdf:
        return HttpResponse(status=404)

    try:
        plantilla.archivo_pdf.open('rb')
    except Exception:
        return HttpResponse(status=404)

    nombre_archivo = os.path.basename(plantilla.archivo_pdf.name or '') or 'plantilla.pdf'
    response = FileResponse(plantilla.archivo_pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{nombre_archivo}"'
    return response


@login_required
@prevencion_mantenedor_required
def ajax_create_approval_notification(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Método no permitido.'}, status=405)

    payload = _parsear_json_request(request)
    destinatario_id = str(payload.get('destinatario_id') or '').strip()
    if not destinatario_id.isdigit():
        return JsonResponse({'success': False, 'error': 'Destinatario inválido.'}, status=400)

    destinatario = User.objects.filter(id=int(destinatario_id), is_active=True).first()
    if not destinatario:
        return JsonResponse({'success': False, 'error': 'Usuario destinatario no encontrado.'}, status=404)

    perfil_destinatario = getattr(destinatario, 'usuarioprofile', None)
    if not perfil_destinatario or perfil_destinatario.seccionPrevencion != 'ADMINISTRADOR':
        return JsonResponse({'success': False, 'error': 'El destinatario debe ser administrador en Prevención.'}, status=400)

    plantilla = None
    plantilla_id_raw = str(payload.get('plantilla_id') or '').strip()
    if plantilla_id_raw:
        if not plantilla_id_raw.isdigit():
            return JsonResponse({'success': False, 'error': 'Plantilla inválida.'}, status=400)
        plantilla = PrevencionPlantilla.objects.filter(id=int(plantilla_id_raw), status=True).first()
        if not plantilla:
            return JsonResponse({'success': False, 'error': 'Plantilla no encontrada.'}, status=404)

    tipo_documento = str(payload.get('tipo_documento') or '').strip().lower()
    if tipo_documento and tipo_documento not in TIPOS_DOCUMENTO_PERMITIDOS:
        tipo_documento = ''

    titulo_payload = str(payload.get('titulo') or '').strip()
    titulo_generado = _construir_titulo_aprobacion(tipo_documento)
    titulo = (titulo_generado or titulo_payload or 'Aprobación pendiente').strip()[:200]
    descripcion = str(payload.get('descripcion') or '').strip()
    if not descripcion:
        descripcion = f'{_nombre_usuario_legible(request.user)} solicita tu aprobación.'

    notificacion_qs = PrevencionNotificacionAprobacion.objects.filter(
        destinatario=destinatario,
        solicitante=request.user,
        estado=PrevencionNotificacionAprobacion.Estado.PENDIENTE,
        tipo_documento=tipo_documento,
        titulo=titulo,
        descripcion=descripcion,
    )
    if plantilla:
        notificacion_qs = notificacion_qs.filter(plantilla=plantilla)
    else:
        notificacion_qs = notificacion_qs.filter(plantilla__isnull=True)

    notificacion = notificacion_qs.order_by('-id').first()
    created = False
    if not notificacion:
        notificacion = PrevencionNotificacionAprobacion.objects.create(
            destinatario=destinatario,
            solicitante=request.user,
            plantilla=plantilla,
            tipo_documento=tipo_documento,
            titulo=titulo,
            descripcion=descripcion,
        )
        created = True

    return JsonResponse({
        'success': True,
        'created': created,
        'notification': _serializar_notificacion_aprobacion(notificacion, request=request),
    })


@login_required
@prevencion_mantenedor_required
def ajax_sync_autorizadores_prevencion(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Método no permitido.'}, status=405)

    payload = _parsear_json_request(request)
    plantilla_id_raw = str(payload.get('plantilla_id') or '').strip()
    if not plantilla_id_raw.isdigit():
        return JsonResponse({'success': False, 'error': 'Plantilla inválida.'}, status=400)

    plantilla = PrevencionPlantilla.objects.filter(id=int(plantilla_id_raw), status=True).first()
    if not plantilla:
        return JsonResponse({'success': False, 'error': 'Plantilla no encontrada.'}, status=404)
    if plantilla.tipo not in TIPOS_DOCUMENTO_PERMITIDOS:
        return JsonResponse({'success': False, 'error': 'La autorización solo aplica a documentos.'}, status=400)

    ids_payload = payload.get('autorizador_ids', [])
    if isinstance(ids_payload, str):
        ids_payload = [segment.strip() for segment in ids_payload.split(',') if segment.strip()]
    if not isinstance(ids_payload, list):
        return JsonResponse({'success': False, 'error': 'Formato de autorizadores inválido.'}, status=400)

    autorizadores_ids = []
    seen = set()
    for raw_id in ids_payload:
        uid = str(raw_id or '').strip()
        if not uid:
            continue
        if not uid.isdigit():
            return JsonResponse({'success': False, 'error': 'ID de autorizador inválido.'}, status=400)
        uid_int = int(uid)
        if uid_int in seen:
            continue
        seen.add(uid_int)
        autorizadores_ids.append(uid_int)

    if len(autorizadores_ids) > 5:
        return JsonResponse({'success': False, 'error': 'Solo se permiten hasta 5 autorizadores.'}, status=400)

    has_unselected_rows_raw = payload.get('has_unselected_rows', False)
    if isinstance(has_unselected_rows_raw, str):
        has_unselected_rows = has_unselected_rows_raw.strip().lower() in {'1', 'true', 'yes', 'si', 'on'}
    else:
        has_unselected_rows = bool(has_unselected_rows_raw)

    autorizadores_validos_ids = list(
        User.objects.filter(
            id__in=autorizadores_ids, 
            is_active=True, 
            usuarioprofile__seccionPrevencion='ADMINISTRADOR'
        ).values_list('id', flat=True)
    )
    if len(autorizadores_validos_ids) != len(autorizadores_ids):
        return JsonResponse({'success': False, 'error': 'Uno o más autorizadores no son válidos.'}, status=400)

    autorizadores_previos = set(plantilla.autorizadores.values_list('id', flat=True))
    autorizadores_actuales = set(autorizadores_validos_ids)
    autorizadores_agregados = list(autorizadores_actuales - autorizadores_previos)
    autorizadores_eliminados = list(autorizadores_previos - autorizadores_actuales)

    plantilla.autorizadores.set(autorizadores_validos_ids)
    if plantilla.autorizacion_incompleta != has_unselected_rows:
        plantilla.autorizacion_incompleta = has_unselected_rows
        plantilla.save(update_fields=['autorizacion_incompleta'])

    if autorizadores_eliminados:
        now = timezone.now()
        PrevencionNotificacionAprobacion.objects.filter(
            plantilla=plantilla,
            destinatario_id__in=autorizadores_eliminados,
            estado=PrevencionNotificacionAprobacion.Estado.PENDIENTE,
        ).update(
            estado=PrevencionNotificacionAprobacion.Estado.CANCELADA,
            leida=True,
            fecha_lectura=now,
            fecha_respuesta=now,
        )

    eventos_historial = [
        PrevencionHistorialAutorizacion(
            plantilla=plantilla,
            actor=request.user,
            destinatario_id=destinatario_id,
            accion=PrevencionHistorialAutorizacion.Accion.AGREGADO,
        )
        for destinatario_id in autorizadores_agregados
    ] + [
        PrevencionHistorialAutorizacion(
            plantilla=plantilla,
            actor=request.user,
            destinatario_id=destinatario_id,
            accion=PrevencionHistorialAutorizacion.Accion.ELIMINADO,
        )
        for destinatario_id in autorizadores_eliminados
    ]
    if eventos_historial:
        PrevencionHistorialAutorizacion.objects.bulk_create(eventos_historial)

    return JsonResponse({
        'success': True,
        'autorizador_ids': autorizadores_validos_ids,
        'added_ids': autorizadores_agregados,
        'removed_ids': autorizadores_eliminados,
        'has_unselected_rows': has_unselected_rows,
    })


@login_required
@prevencion_mantenedor_required
def ajax_list_approval_history_prevencion(request, plantilla_id):
    plantilla = PrevencionPlantilla.objects.filter(id=plantilla_id, status=True).first()
    if not plantilla:
        return JsonResponse({'success': False, 'error': 'Plantilla no encontrada.'}, status=404)

    historial = _obtener_historial_autorizaciones_plantilla(plantilla)
    return JsonResponse({
        'success': True,
        'history': historial,
    })


@login_required
@prevencion_mantenedor_required
def ajax_list_approval_status_prevencion(request, plantilla_id):
    plantilla = PrevencionPlantilla.objects.filter(id=plantilla_id, status=True).first()
    if not plantilla:
        return JsonResponse({'success': False, 'error': 'Plantilla no encontrada.'}, status=404)

    estados = _obtener_estado_detallado_autorizadores_plantilla(plantilla)
    return JsonResponse({
        'success': True,
        'status_by_user': estados,
    })


@login_required
@capacitaciones_y_docs_required
def ajax_list_approval_notifications(request):
    notificaciones = _listar_notificaciones_usuario(request.user, request=request, limit=4)
    documento_activo_id = _documento_activo_capacitacion_id(request)
    unread_approval_count = PrevencionNotificacionAprobacion.objects.filter(
        destinatario=request.user,
        estado=PrevencionNotificacionAprobacion.Estado.PENDIENTE,
        leida=False,
    ).count()
    unread_document_qs = PrevencionNotificacionDocumento.objects.filter(
        destinatario=request.user,
        estado=PrevencionNotificacionDocumento.Estado.PENDIENTE,
        leida=False,
        documento__status=True,
        documento__difusion_trabajadores__status=True,
        documento__difusion_trabajadores__user=request.user,
    )
    if documento_activo_id > 0:
        unread_document_qs = unread_document_qs.exclude(documento_id=documento_activo_id)
    unread_document_count = unread_document_qs.distinct().count()

    return JsonResponse({
        'success': True,
        'unread_count': unread_approval_count + unread_document_count,
        'notifications': notificaciones,
    })


@login_required
@capacitaciones_y_docs_required
def view_all_approval_notifications(request):
    notificaciones = _listar_notificaciones_usuario(request.user, request=request)

    context = {
        'notificaciones': notificaciones,
        'sidebar': 'dashboard',
    }
    return render(request, 'pages/prevencion/all_approval_notifications.html', context)


@login_required
@capacitaciones_y_docs_required
def view_review_approval_notification(request, pk):
    notificacion = get_object_or_404(
        PrevencionNotificacionAprobacion.objects.select_related('solicitante', 'plantilla', 'plantilla__faena'),
        pk=pk,
        destinatario=request.user,
    )

    if notificacion.estado == PrevencionNotificacionAprobacion.Estado.PENDIENTE and not notificacion.leida:
        now = timezone.now()
        notificacion.leida = True
        if not notificacion.fecha_lectura:
            notificacion.fecha_lectura = now
            notificacion.save(update_fields=['leida', 'fecha_lectura'])
        else:
            notificacion.save(update_fields=['leida'])

    plantilla = notificacion.plantilla
    estructura = []
    if plantilla and isinstance(plantilla.estructura, list):
        estructura = plantilla.estructura
    rut_destinatario_normalizado = _normalizar_rut_chileno(notificacion.destinatario.username)

    context = {
        'notificacion': notificacion,
        'plantilla': plantilla,
        'estructura': estructura,
        'tipo_documento_slug': str(notificacion.tipo_documento or '').strip().lower(),
        'tipo_documento_label': _mapear_tipo_documento_label(notificacion.tipo_documento),
        'can_resolve': notificacion.estado == PrevencionNotificacionAprobacion.Estado.PENDIENTE,
        'resolve_url': reverse('ajax_resolve_approval_notification', kwargs={'pk': notificacion.id}),
        'rut_destinatario_qr_normalizado': rut_destinatario_normalizado,
        'rut_destinatario_qr_display': _formatear_rut_chileno(rut_destinatario_normalizado),
        'sidebar': 'dashboard',
    }
    return render(request, 'pages/prevencion/review_approval_notification.html', context)


@login_required
@capacitaciones_y_docs_required
def ajax_resolve_approval_notification(request, pk):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Método no permitido.'}, status=405)

    notificacion = get_object_or_404(
        PrevencionNotificacionAprobacion,
        pk=pk,
        destinatario=request.user,
    )
    if notificacion.estado != PrevencionNotificacionAprobacion.Estado.PENDIENTE:
        return JsonResponse({'success': False, 'error': 'La solicitud ya fue procesada.'}, status=409)

    payload = _parsear_json_request(request)
    accion = str(payload.get('accion') or '').strip().lower()

    confirmo_revision = _parsear_bool(payload.get('confirmar_revision', False))
    motivo_rechazo = str(payload.get('motivo_rechazo') or '').strip()
    # qr_payload = str(payload.get('qr_payload') or '').strip()
    # QR temporalmente desactivado (flujo clásico sin escaneo).

    is_sistema = notificacion.tipo_documento == 'sistema'

    if accion == 'aceptar':
        nuevo_estado = PrevencionNotificacionAprobacion.Estado.ACEPTADA
    elif accion == 'rechazar':
        nuevo_estado = PrevencionNotificacionAprobacion.Estado.RECHAZADA
    else:
        return JsonResponse({'success': False, 'error': 'Acción inválida.'}, status=400)

    if not is_sistema and not confirmo_revision:
        return JsonResponse({'success': False, 'error': 'Debes revisar el contenido completo antes de responder.'}, status=400)
    if nuevo_estado == PrevencionNotificacionAprobacion.Estado.RECHAZADA and not motivo_rechazo:
        return JsonResponse({'success': False, 'error': 'Debes escribir el motivo de rechazo.'}, status=400)
    if len(motivo_rechazo) > 1000:
        return JsonResponse({'success': False, 'error': 'El motivo de rechazo no puede superar 1000 caracteres.'}, status=400)

    # if nuevo_estado == PrevencionNotificacionAprobacion.Estado.ACEPTADA:
    #     rut_esperado = _normalizar_rut_chileno(notificacion.destinatario.username)
    #     if not _rut_chileno_valido(rut_esperado):
    #         return JsonResponse(
    #             {'success': False, 'error': 'No se pudo validar tu RUT de usuario para aprobar esta solicitud.'},
    #             status=400,
    #         )
    #
    #     rut_escaneado = _extraer_rut_chileno_desde_qr(qr_payload)
    #     if not rut_escaneado:
    #         return JsonResponse(
    #             {'success': False, 'error': 'No se pudo leer un RUT válido desde el código QR.'},
    #             status=400,
    #         )
    #     if rut_escaneado != rut_esperado:
    #         return JsonResponse(
    #             {'success': False, 'error': 'El RUT del código QR no coincide con tu usuario.'},
    #             status=403,
    #         )

    now = timezone.now()
    notificacion.estado = nuevo_estado
    notificacion.leida = True
    if not notificacion.fecha_lectura:
        notificacion.fecha_lectura = now
    notificacion.fecha_respuesta = now
    notificacion.motivo_rechazo = motivo_rechazo if nuevo_estado == PrevencionNotificacionAprobacion.Estado.RECHAZADA else ''
    notificacion.save(update_fields=['estado', 'leida', 'fecha_lectura', 'fecha_respuesta', 'motivo_rechazo'])

    if notificacion.solicitante != request.user:
        estado_texto = 'completada' if nuevo_estado == PrevencionNotificacionAprobacion.Estado.ACEPTADA else 'rechazada'
        desc = f'Autorización {estado_texto} para {notificacion.plantilla.nombre if notificacion.plantilla else ""} por: {request.user.get_full_name() or request.user.username}'
        if motivo_rechazo:
            desc += f'\nMotivo: {motivo_rechazo}'
        PrevencionNotificacionAprobacion.objects.create(
            destinatario=notificacion.solicitante,
            solicitante=request.user,
            plantilla=notificacion.plantilla,
            tipo_documento='sistema',
            titulo=f'Autorización {estado_texto}',
            descripcion=desc,
        )

    return JsonResponse({
        'success': True,
        'notification': _serializar_notificacion_aprobacion(notificacion, request=request),
    })


@login_required
@prevencion_riesgo_required
def ajax_cancel_approval_notification(request, pk):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Método no permitido.'}, status=405)

    notificacion = get_object_or_404(
        PrevencionNotificacionAprobacion,
        pk=pk,
        solicitante=request.user,
    )

    if notificacion.estado != PrevencionNotificacionAprobacion.Estado.PENDIENTE:
        return JsonResponse({'success': True, 'notification': _serializar_notificacion_aprobacion(notificacion, request=request)})

    now = timezone.now()
    notificacion.estado = PrevencionNotificacionAprobacion.Estado.CANCELADA
    notificacion.leida = True
    if not notificacion.fecha_lectura:
        notificacion.fecha_lectura = now
    notificacion.fecha_respuesta = now
    notificacion.save(update_fields=['estado', 'leida', 'fecha_lectura', 'fecha_respuesta'])

    return JsonResponse({
        'success': True,
        'notification': _serializar_notificacion_aprobacion(notificacion, request=request),
    })


@login_required
@admin_or_base_datos_required
def manage_vigilancia_option(request, tipo_slug):
    if not _usuario_autorizado_opcion_vigilancia(request.user, tipo_slug):
        return redirect('dashboardPrevencion')

    config = _get_vigilancia_opcion_config(tipo_slug)
    if not config:
        messages.error(request, 'Tipo de opcion de vigilancia no valido.')
        return redirect('dashboardPrevencion')

    storage = messages.get_messages(request)
    storage.used = True

    opciones = config['model'].objects.all().order_by('correlativo', 'id')
    context = {
        'opciones': opciones,
        'tipo_slug': tipo_slug,
        'titulo_opcion': config['titulo'],
        'texto_boton_nuevo': config['texto_boton_nuevo'],
        'sidebarmain': 'manage_system',
        'sidebar': 'mantenedor_prev',
        'sidebarmenu': config['menu_key'],
    }
    return render(request, 'pages/prevencion/manage_vigilancia_opcion.html', context)


@login_required
@admin_or_base_datos_required
def new_vigilancia_option(request, tipo_slug):
    if not _usuario_autorizado_opcion_vigilancia(request.user, tipo_slug):
        return redirect('dashboardPrevencion')

    config = _get_vigilancia_opcion_config(tipo_slug)
    if not config:
        messages.error(request, 'Tipo de opcion de vigilancia no valido.')
        return redirect('dashboardPrevencion')

    form = FormVigilanciaOpcion()
    form.fields['valor'].label = config['titulo']
    if config['es_numerico']:
        form.fields['valor'].help_text = 'Solo valores numericos.'

    context = {
        'formopcion': form,
        'tipo_slug': tipo_slug,
        'titulo_opcion': config['titulo'],
        'titulo_nuevo': config['texto_boton_nuevo'],
        'mensaje_exito': config['mensaje_exito'],
        'es_numerico': config['es_numerico'],
        'sidebarmain': 'manage_system',
        'sidebar': 'mantenedor_prev',
        'sidebarmenu': config['menu_key'],
    }
    return render(request, 'pages/prevencion/new_vigilancia_opcion.html', context)


@login_required
@admin_or_base_datos_required
def save_new_vigilancia_option(request, tipo_slug):
    if not _usuario_autorizado_opcion_vigilancia(request.user, tipo_slug):
        return JsonResponse({'success': False, 'error': 'No autorizado.'}, status=403)

    config = _get_vigilancia_opcion_config(tipo_slug)
    if not config:
        return JsonResponse({'success': False, 'error': 'Tipo de opcion no valido.'}, status=400)

    if request.method != 'POST':
        return redirect('new_vigilancia_option', tipo_slug=tipo_slug)

    formulario = FormVigilanciaOpcion(data=request.POST)
    if not formulario.is_valid():
        return JsonResponse({'success': False, 'error': 'Formulario invalido.'}, status=400)

    try:
        valor = _normalizar_valor_vigilancia(
            formulario.cleaned_data['valor'],
            es_numerico=config['es_numerico']
        )
    except ValueError as exc:
        return JsonResponse({'success': False, 'error': str(exc)}, status=400)

    modelo = config['model']
    existe = modelo.objects.filter(valor__iexact=valor).exists()
    if existe:
        return JsonResponse({'success': False, 'error': 'Este valor ya existe.'}, status=400)

    nombre_usuario = f"{request.user.first_name} {request.user.last_name}".strip()
    if not nombre_usuario:
        nombre_usuario = request.user.username

    try:
        with transaction.atomic():
            ultimo_correlativo = (
                modelo.objects
                .select_for_update()
                .aggregate(max_correlativo=Max('correlativo'))
                .get('max_correlativo') or 0
            )

            modelo.objects.create(
                correlativo=ultimo_correlativo + 1,
                valor=valor,
                status=True,
                creador=nombre_usuario,
            )
    except IntegrityError:
        return JsonResponse({'success': False, 'error': 'No se pudo generar el correlativo. Intenta nuevamente.'}, status=409)

    user_name = request.user.get_full_name() or request.user.username
    notify_group('prevencion_general', f'Opción de vigilancia creada: {valor}', f'Opción de Vigilancia ({config["titulo"]}): {valor}\nCreada por: {user_name}')
    return JsonResponse({'success': True})


@login_required
@admin_or_base_datos_required
def status_vigilancia_option(request, tipo_slug):
    if not _usuario_autorizado_opcion_vigilancia(request.user, tipo_slug):
        return redirect('dashboardPrevencion')

    config = _get_vigilancia_opcion_config(tipo_slug)
    if not config:
        messages.error(request, 'Tipo de opcion de vigilancia no valido.')
        return redirect('dashboardPrevencion')

    if request.method == 'POST':
        opcion = get_object_or_404(
            config['model'],
            id=request.POST.get('id'),
        )
        opcion.status = not opcion.status
        opcion.save(update_fields=['status'])

        accion = 'Deshabilitada' if not opcion.status else 'Habilitada'
        messages.success(request, f'Opcion {accion} Correctamente')
        user_name = request.user.get_full_name() or request.user.username
        descripcion = getattr(opcion, 'valor', str(opcion))
        notify_group('prevencion_general', f'Opción {accion.lower()}: {descripcion}', f'Opción de Vigilancia ({config["titulo"]}): {descripcion}\n{accion} por: {user_name}')

    return redirect('manage_vigilancia_option', tipo_slug=tipo_slug)



# =========================================================
#  VIGILANCIA MÉDICA (CORREGIDO)
# =========================================================

@login_required
@prevencion_riesgo_required
def manage_vigilancia(request):
    vigilancias = VigilanciaMedica.objects.filter(status=True).select_related('faena').prefetch_related('adjuntos_vigilancia')
    vigilancias = _filtrar_queryset_por_faena_usuario(vigilancias, request.user).order_by('-fecha_incidente', '-created_at')
    context = {
        'vigilancias': vigilancias,
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/manage_vigilancia.html', context)

@login_required
def new_vigilancia(request):
    faenas_visibles = _faenas_visibles_para_usuario(request.user)

    if request.method == 'POST':
        try:
            archivos_validos, msg_archivos = _validar_archivos_subidos(request)
            if not archivos_validos:
                messages.error(request, msg_archivos)
                return redirect('new_vigilancia')

            faena_id = request.POST.get('faena')
            if not _faena_permitida_para_usuario(request.user, faena_id):
                messages.error(request, "No tienes permisos para registrar una ficha en esa faena.")
                return redirect('new_vigilancia')

            # 1. Búsqueda de usuario por RUT.
            rut_raw = request.POST.get('rut')
            rut_limpio = rut_raw.replace('.', '').replace('-', '').upper()
            rut_guion = f"{rut_limpio[:-1]}-{rut_limpio[-1]}" if len(rut_limpio) > 1 else rut_limpio

            usuario_asociado = User.objects.filter(
                Q(username__iexact=rut_limpio) | Q(username__iexact=rut_guion) | Q(username__iexact=rut_raw)
            ).first()
            
            # 2. Género: guardar texto legible desde el ID seleccionado.
            genero_id = request.POST.get('genero')
            genero_obj = Genero.objects.filter(pk=genero_id).first()
            genero_txt = genero_obj.genero if genero_obj else None  # Guardamos "Masculino", no "1"
            fecha_nacimiento_raw = request.POST.get('fecha_nacimiento') or None
            edad_calculada = _calcular_edad_desde_fecha(fecha_nacimiento_raw)

            vigilancia = VigilanciaMedica(
                user=usuario_asociado,
                # 1. PERSONALES
                rut_trabajador=rut_raw,
                nombre_completo=request.POST.get('nombre_completo'),
                area=request.POST.get('area'),
                cargo=request.POST.get('cargo'),
                faena_id=faena_id,
                
                genero_texto=genero_txt,  # Aquí se guarda el texto del género.
                
                fecha_nacimiento=fecha_nacimiento_raw,
                edad=edad_calculada if edad_calculada is not None else 0,
                fecha_ingreso=request.POST.get('fecha_ingreso') or None,
                antiguedad_anos=request.POST.get('antiguedad_anos') or 0,
                antiguedad_meses=request.POST.get('antiguedad_meses') or 0,
                antiguedad_dias=request.POST.get('antiguedad_dias') or 0,
                tipo_contrato=request.POST.get('tipo_contrato'),
                fecha_retiro=request.POST.get('fecha_retiro') or None,
                ges=request.POST.get('ges'),

                # 2. Sílice
                silice_eval_riesgo=request.POST.get('silice_eval_riesgo'),
                silice_nivel_riesgo=request.POST.get('silice_nivel_riesgo'),
                silice_grado_expo=request.POST.get('silice_grado_expo'),
                silice_fecha_radio=request.POST.get('silice_fecha_radio') or None,
                silice_fecha_vence=request.POST.get('silice_fecha_vence') or None,
                silice_vigencia=request.POST.get('silice_vigencia'),
                silice_obs=request.POST.get('silice_obs'),

                # 3. RUIDO
                ruido_eval_riesgo=request.POST.get('ruido_eval_riesgo'),
                ruido_nivel_seguimiento=request.POST.get('ruido_nivel_seguimiento'),
                ruido_grado_expo=request.POST.get('ruido_grado_expo'),
                ruido_fecha_audio=request.POST.get('ruido_fecha_audio') or None,
                ruido_fecha_vence=request.POST.get('ruido_fecha_vence') or None,
                ruido_vigencia=request.POST.get('ruido_vigencia'),
                ruido_obs=request.POST.get('ruido_obs'),

                # 4. HIPOBARIA
                hipo_exposicion=request.POST.get('hipo_exposicion'),
                hipo_fecha_hemo=request.POST.get('hipo_fecha_hemo') or None,
                hipo_fecha_vence=request.POST.get('hipo_fecha_vence') or None,
                hipo_vigencia=request.POST.get('hipo_vigencia'),
                hipo_obs=request.POST.get('hipo_obs'),

                # 5. OTROS
                vibra_eval_riesgo=request.POST.get('vibra_eval_riesgo'),
                vibra_exposicion=request.POST.get('vibra_exposicion'),
                rad_exposicion=request.POST.get('rad_exposicion'),
                humos_eval_riesgo=request.POST.get('humos_eval_riesgo'),
                humos_exposicion=request.POST.get('humos_exposicion'),

                # Meta (se usa 'fecha_incidente' porque así se llama el input en HTML).
                fecha_incidente=request.POST.get('fecha_incidente') or None,
                creado_por=request.user.username
            )
            
            with transaction.atomic():
                vigilancia.save()
                _sincronizar_datos_personales_usuario(
                    usuario_asociado,
                    genero_id=genero_id,
                    fecha_nacimiento_raw=fecha_nacimiento_raw,
                    sobrescribir=True
                )
                for archivo in request.FILES.getlist('documento'):
                    VigilanciaAdjunto.objects.create(vigilancia=vigilancia, archivo=archivo)
            messages.success(request, "Vigilancia guardada correctamente.")
            user_name = request.user.get_full_name() or request.user.username
            notify_group('prevencion_general', f'Nueva ficha vigilancia: {vigilancia.nombre_completo}', f'Vigilancia Médica: {vigilancia.nombre_completo}\nCreada por: {user_name}')
            return redirect('manage_vigilancia')
            
        except Exception as e:
            messages.error(request, f"Error al guardar: {e}")

    context = {
        'faenas': _faenas_seleccionables_para_usuario(request.user),
        'generos': Genero.objects.filter(status=True), 
        **_catalogos_vigilancia_formulario(),
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo', 
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/new_vigilancia.html', context)


@login_required
@prevencion_riesgo_required
def edit_vigilancia(request, pk):
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.filter(status=True).prefetch_related('adjuntos_vigilancia'),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)

    if vigilancia.egresado:
        return _redireccion_egreso_desde_ficha(
            request,
            vigilancia,
            mensaje="No se puede modificar una ficha médica que ya se encuentra egresada.",
        )

    datos_personales_sincronizados = _obtener_datos_personales_sincronizados(vigilancia)
    
    if request.method == 'POST':
        old_vigilancia = VigilanciaMedica.objects.get(pk=vigilancia.pk)
        try:
            archivos_validos, msg_archivos = _validar_archivos_subidos(request)
            if not archivos_validos:
                messages.error(request, msg_archivos)
                return redirect('edit_vigilancia', pk=pk)

            total_adjuntos_valido, total_adjuntos_msg = _validar_total_adjuntos_vigilancia_en_edicion(vigilancia, request)
            if not total_adjuntos_valido:
                messages.error(request, total_adjuntos_msg)
                return redirect('edit_vigilancia', pk=pk)

            # 1. PERSONALES
            vigilancia.rut_trabajador = request.POST.get('rut')
            vigilancia.nombre_completo = request.POST.get('nombre_completo')
            vigilancia.area = request.POST.get('area')
            vigilancia.cargo = request.POST.get('cargo')
            nueva_faena_id = request.POST.get('faena')
            if not _faena_permitida_para_usuario(request.user, nueva_faena_id):
                messages.error(request, "No tienes permisos para asignar esa faena.")
                return redirect('edit_vigilancia', pk=pk)
            vigilancia.faena_id = nueva_faena_id
            
            # 2. Género: guardar texto legible desde el ID seleccionado.
            genero_id = request.POST.get('genero')
            genero_obj = Genero.objects.filter(pk=genero_id).first()
            vigilancia.genero_texto = genero_obj.genero if genero_obj else None  # Aquí se guarda el texto del género.
            
            fecha_nacimiento_raw = request.POST.get('fecha_nacimiento') or None
            edad_calculada = _calcular_edad_desde_fecha(fecha_nacimiento_raw)
            vigilancia.fecha_nacimiento = fecha_nacimiento_raw
            vigilancia.edad = edad_calculada if edad_calculada is not None else 0
            
            vigilancia.fecha_ingreso = request.POST.get('fecha_ingreso') or None
            vigilancia.antiguedad_anos = request.POST.get('antiguedad_anos') or 0
            vigilancia.antiguedad_meses = request.POST.get('antiguedad_meses') or 0
            vigilancia.antiguedad_dias = request.POST.get('antiguedad_dias') or 0
            vigilancia.tipo_contrato = request.POST.get('tipo_contrato')
            vigilancia.fecha_retiro = request.POST.get('fecha_retiro') or None
            vigilancia.ges = request.POST.get('ges')

            # 2. Sílice
            vigilancia.silice_eval_riesgo = request.POST.get('silice_eval_riesgo')
            vigilancia.silice_nivel_riesgo = request.POST.get('silice_nivel_riesgo')
            vigilancia.silice_grado_expo = request.POST.get('silice_grado_expo')
            vigilancia.silice_fecha_radio = request.POST.get('silice_fecha_radio') or None
            vigilancia.silice_fecha_vence = request.POST.get('silice_fecha_vence') or None
            vigilancia.silice_vigencia = request.POST.get('silice_vigencia')
            vigilancia.silice_obs = request.POST.get('silice_obs')

            # 3. RUIDO
            vigilancia.ruido_eval_riesgo = request.POST.get('ruido_eval_riesgo')
            vigilancia.ruido_nivel_seguimiento = request.POST.get('ruido_nivel_seguimiento')
            vigilancia.ruido_grado_expo = request.POST.get('ruido_grado_expo')
            vigilancia.ruido_fecha_audio = request.POST.get('ruido_fecha_audio') or None
            vigilancia.ruido_fecha_vence = request.POST.get('ruido_fecha_vence') or None
            vigilancia.ruido_vigencia = request.POST.get('ruido_vigencia')
            vigilancia.ruido_obs = request.POST.get('ruido_obs')

            # 4. HIPOBARIA
            vigilancia.hipo_exposicion = request.POST.get('hipo_exposicion')
            vigilancia.hipo_fecha_hemo = request.POST.get('hipo_fecha_hemo') or None
            vigilancia.hipo_fecha_vence = request.POST.get('hipo_fecha_vence') or None
            vigilancia.hipo_vigencia = request.POST.get('hipo_vigencia')
            vigilancia.hipo_obs = request.POST.get('hipo_obs')

            # 5. OTROS
            vigilancia.vibra_eval_riesgo = request.POST.get('vibra_eval_riesgo')
            vigilancia.vibra_exposicion = request.POST.get('vibra_exposicion')
            vigilancia.rad_exposicion = request.POST.get('rad_exposicion')
            vigilancia.humos_eval_riesgo = request.POST.get('humos_eval_riesgo')
            vigilancia.humos_exposicion = request.POST.get('humos_exposicion')

            # META (Usamos 'fecha_incidente' para que coincida con el HTML)
            vigilancia.fecha_incidente = request.POST.get('fecha_incidente') or None

            with transaction.atomic():
                vigilancia.save()
                _sincronizar_datos_personales_usuario(
                    vigilancia.user,
                    genero_id=genero_id,
                    fecha_nacimiento_raw=fecha_nacimiento_raw,
                    sobrescribir=True
                )
                for archivo in request.FILES.getlist('documento'):
                    VigilanciaAdjunto.objects.create(vigilancia=vigilancia, archivo=archivo)
            user_name = request.user.get_full_name() or request.user.username
            changes = get_changes_message(old_vigilancia, vigilancia)
            if changes:
                notify_group('prevencion_general', f'Ficha vigilancia actualizada: {vigilancia.nombre_completo}', f'Vigilancia Médica: {vigilancia.nombre_completo}\nActualizada por: {user_name}{changes}')
            messages.success(request, "Vigilancia actualizada correctamente.")
            return redirect('manage_vigilancia')
            
        except Exception as e:
            messages.error(request, f"Error al actualizar: {e}")

    context = {
        'vigilancia': vigilancia,
        'archivos_adjuntos': _serializar_adjuntos_vigilancia(vigilancia),
        'genero_seleccionado_id': datos_personales_sincronizados['genero_id'],
        'fecha_nacimiento_mostrada': datos_personales_sincronizados['fecha_nacimiento'],
        'faenas': _faenas_seleccionables_para_usuario(request.user),
        'generos': Genero.objects.filter(status=True),
        **_catalogos_vigilancia_formulario(),
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/edit_vigilancia.html', context)

@login_required
@prevencion_riesgo_required
def view_vigilancia(request, pk):
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.filter(status=True).prefetch_related('adjuntos_vigilancia'),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)
    _obtener_datos_personales_sincronizados(vigilancia)
    context = {
        'vigilancia': vigilancia,
        'archivos_adjuntos': _serializar_adjuntos_vigilancia(vigilancia),
        'sidebarmain': 'prevencion_riesgo', 
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/view_vigilancia.html', context)


@login_required
@prevencion_riesgo_required
def vigilancia_pdf_view(request, pk):
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.filter(status=True).prefetch_related('adjuntos_vigilancia'),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)
    # Prepara cada adjunto con metadatos de visualización en PDF (imagen/pdf/otro).
    archivos_adjuntos_raw = _serializar_adjuntos_vigilancia(vigilancia)
    archivos_adjuntos = []
    for archivo in archivos_adjuntos_raw:
        nombre = str((archivo or {}).get('name', '') or '')
        url = str((archivo or {}).get('url', '') or '')
        nombre_lower = nombre.lower()
        url_lower = url.lower().split('?', 1)[0]
        registro = {
            'id': (archivo or {}).get('id'),
            'name': nombre,
            'url': url,
            'paginas_pdf': [],
        }

        if nombre_lower.endswith(('.png', '.jpg', '.jpeg', '.gif', '.webp')) or url_lower.endswith(('.png', '.jpg', '.jpeg', '.gif', '.webp')):
            registro['tipo'] = 'imagen'
        elif nombre_lower.endswith('.pdf') or url_lower.endswith('.pdf'):
            # Si el adjunto es PDF, lo convierte a imágenes para respetar el orden de adjuntos.
            registro['tipo'] = 'pdf'
            ruta_local = _resolver_ruta_local_adjunto_pdf(url)
            if ruta_local:
                try:
                    registro['paginas_pdf'] = check_and_convert_pdf(ruta_local) or []
                except Exception:
                    registro['paginas_pdf'] = []
        else:
            registro['tipo'] = 'otro'

        archivos_adjuntos.append(registro)

    template = get_template('pages/pdfs/vigilancia_medica_pdf.html')
    html = template.render({
        'vigilancia': vigilancia,
        'archivos_adjuntos': archivos_adjuntos,
        'fecha_descarga': timezone.localtime(timezone.now()),
    })

    filename = f"vigilancia-medica-{vigilancia.id}.pdf"
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename=\"{filename}\"'
    pisa_status = pisa.CreatePDF(html, dest=response, link_callback=_resolver_src_pdf)

    if pisa_status.err:
        return HttpResponse('Error al generar el PDF.', status=500)

    return response

@login_required
@prevencion_riesgo_required
def ajax_search_user_by_rut(request):
    rut_raw = request.GET.get('rut', '').strip()
    if not rut_raw: return JsonResponse({'found': False})
    
    rut_limpio = rut_raw.replace('.', '').replace('-', '').upper()
    rut_guion = f"{rut_limpio[:-1]}-{rut_limpio[-1]}" if len(rut_limpio) > 1 else ""

    try:
        user = User.objects.filter(Q(username__iexact=rut_limpio) | Q(username__iexact=rut_guion) | Q(username__iexact=rut_raw)).first()
        if user:
            nombre = f"{user.first_name} {user.last_name}".strip()
            cargo = ''
            genero_id = ''
            faena_id = ''
            fecha_nacimiento = ''
            area = ''
            ges = ''
            tipo_contrato = ''
            fecha_ingreso = ''
            fecha_retiro = ''
            
            if hasattr(user, 'usuarioprofile'):
                p = user.usuarioprofile
                if p.genero: genero_id = p.genero.id
                if p.faena and _faena_permitida_para_usuario(request.user, p.faena.id):
                    faena_id = p.faena.id
                if p.fechaNacimiento:
                    fecha_nacimiento = p.fechaNacimiento.strftime('%Y-%m-%d')

            # Trae datos laborales desde la nueva tabla user_informacion_laboral.
            info_laboral = UserInformacionLaboral.objects.filter(user=user).select_related(
                'area',
                'ges',
                'cargo',
                'tipo_contrato'
            ).first()
            if info_laboral:
                area = info_laboral.area.valor if info_laboral.area else ''
                ges = info_laboral.ges.valor if info_laboral.ges else ''
                cargo = info_laboral.cargo.valor if info_laboral.cargo else ''
                tipo_contrato = info_laboral.tipo_contrato.valor if info_laboral.tipo_contrato else ''
                fecha_ingreso = info_laboral.fechaIngreso.strftime('%Y-%m-%d') if info_laboral.fechaIngreso else ''
                fecha_retiro = (
                    info_laboral.fechaDesvinculacion.strftime('%Y-%m-%d')
                    if info_laboral.fechaDesvinculacion else ''
                )

            data = {
                'found': True,
                'nombre': nombre,
                'cargo': cargo,
                'genero_id': genero_id,
                'faena_id': faena_id,
                'fecha_nacimiento': fecha_nacimiento,
                'area': area,
                'ges': ges,
                'tipo_contrato': tipo_contrato,
                'fecha_ingreso': fecha_ingreso,
                'fecha_retiro': fecha_retiro,
            }
        else:
            data = {'found': False}
    except Exception as e:
        data = {'found': False, 'error': str(e)}
    return JsonResponse(data)

@login_required
@prevencion_riesgo_required
def delete_vigilancia_doc(request, pk):
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.filter(status=True),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)

    if vigilancia.egresado:
        return _redireccion_egreso_desde_ficha(
            request,
            vigilancia,
            mensaje="No se puede modificar una ficha médica que ya se encuentra egresada.",
        )
    
    if vigilancia.documento_adjunto:
        # Borra el archivo del sistema de archivos y limpia el campo
        vigilancia.documento_adjunto.delete(save=False)
        vigilancia.documento_adjunto = None
        vigilancia.save()
        messages.success(request, "Archivo adjunto eliminado correctamente.")
        
    return redirect('edit_vigilancia', pk=pk)


@login_required
@prevencion_riesgo_required
def delete_vigilancia_adjunto(request, pk):
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.filter(status=True),
        request.user
    )
    adjunto = get_object_or_404(
        VigilanciaAdjunto.objects.select_related('vigilancia'),
        pk=pk,
        vigilancia__in=vigilancias
    )

    vigilancia_id = adjunto.vigilancia_id
    adjunto.archivo.delete(save=False)
    adjunto.delete()
    messages.success(request, "Archivo adjunto eliminado correctamente.")
    return redirect('edit_vigilancia', pk=vigilancia_id)


@login_required
@prevencion_riesgo_required
def delete_documento_evidencia(request, pk):
    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True),
        request.user
    )
    doc = get_object_or_404(documentos, pk=pk)

    evidence_id_raw = request.GET.get('evidence_id')
    if evidence_id_raw:
        try:
            evidence_id = int(evidence_id_raw)
        except (TypeError, ValueError):
            messages.error(request, "Identificador de evidencia no valido.")
            return redirect('edit_documento', pk=pk)

        evidencia = doc.evidencias_generales.filter(id=evidence_id).first()
        if evidencia:
            evidencia.archivo.delete(save=False)
            evidencia.delete()
            messages.success(request, "Evidencia general eliminada.")
        else:
            messages.error(request, "No se encontro la evidencia seleccionada.")
        return redirect('edit_documento', pk=pk)

    eliminadas = 0
    for evidencia in doc.evidencias_generales.all():
        evidencia.archivo.delete(save=False)
        evidencia.delete()
        eliminadas += 1

    # Compatibilidad con registros antiguos (campo legacy)
    if doc.evidencia_general:
        doc.evidencia_general.delete(save=False)
        doc.evidencia_general = None
        doc.save(update_fields=['evidencia_general'])
        eliminadas += 1

    if eliminadas:
        messages.success(request, "Evidencia general eliminada correctamente.")

    return redirect('edit_documento', pk=pk)

@login_required
@prevencion_riesgo_required
def delete_documento_evidencia_seccion(request, pk, indice):
    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True),
        request.user
    )
    doc = get_object_or_404(documentos, pk=pk)

    evidence_id_raw = request.GET.get('evidence_id')
    if evidence_id_raw:
        try:
            evidence_id = int(evidence_id_raw)
        except (TypeError, ValueError):
            messages.error(request, "Identificador de evidencia no valido.")
            return redirect('edit_documento', pk=pk)

        evidencia = PrevencionEvidenciaSeccion.objects.filter(
            id=evidence_id,
            documento=doc,
            indice_seccion=indice
        ).first()
        if evidencia:
            evidencia.archivo.delete(save=False)
            evidencia.delete()
            messages.success(request, f"Evidencia de la seccion {indice + 1} eliminada.")
        else:
            messages.error(request, "No se encontro la evidencia seleccionada.")
        return redirect('edit_documento', pk=pk)

    evidencias = PrevencionEvidenciaSeccion.objects.filter(documento=doc, indice_seccion=indice)
    eliminadas = 0
    for evidencia in evidencias:
        evidencia.archivo.delete(save=False)
        evidencia.delete()
        eliminadas += 1

    if eliminadas:
        messages.success(request, f"Evidencias de la seccion {indice + 1} eliminadas.")

    return redirect('edit_documento', pk=pk)


@login_required
@prevencion_riesgo_required
def delete_documento_evidencia_item(request, pk, sec_idx, row_idx):
    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True),
        request.user
    )
    doc = get_object_or_404(documentos, pk=pk)

    evidence_id_raw = request.GET.get('evidence_id')
    if evidence_id_raw:
        try:
            evidence_id = int(evidence_id_raw)
        except (TypeError, ValueError):
            messages.error(request, "Identificador de evidencia no valido.")
            return redirect('edit_documento', pk=pk)

        evidencia = PrevencionEvidenciaItem.objects.filter(
            id=evidence_id,
            documento=doc,
            indice_seccion=sec_idx,
            indice_fila=row_idx
        ).first()
        if evidencia:
            evidencia.archivo.delete(save=False)
            evidencia.delete()
            messages.success(request, f"Evidencia del item {sec_idx+1}.{row_idx+1} eliminada.")
        else:
            messages.error(request, "No se encontro la evidencia seleccionada.")
        return redirect('edit_documento', pk=pk)

    evidencias = PrevencionEvidenciaItem.objects.filter(
        documento=doc,
        indice_seccion=sec_idx,
        indice_fila=row_idx
    )
    eliminadas = 0
    for evidencia in evidencias:
        evidencia.archivo.delete(save=False)
        evidencia.delete()
        eliminadas += 1

    if eliminadas:
        messages.success(request, f"Evidencias del item {sec_idx+1}.{row_idx+1} eliminadas.")

    return redirect('edit_documento', pk=pk)


@login_required
@prevencion_riesgo_required
def export_data_prevencion_view(request):
    # Genera un Excel con datos crudos de Documentos y Vigilancia Médica por rango de fechas.
    form = ExportDataPrevencionForm(request.POST or None)

    if request.method == 'POST':
        if not form.is_valid():
            messages.error(request, "Debes ingresar un rango de fechas valido.")
            return render(request, 'pages/prevencion/export_data.html', {
                'form': form,
                'sidebar': 'prevencion_export_data',
                'sidebarmain': 'prevencion_reportes',
            })

        fecha_inicio = form.cleaned_data['fecha_inicio']
        fecha_final = form.cleaned_data['fecha_final']

        if fecha_final < fecha_inicio:
            messages.error(request, "La fecha inicial no puede ser mayor a la fecha final.")
            return render(request, 'pages/prevencion/export_data.html', {
                'form': form,
                'sidebar': 'prevencion_export_data',
                'sidebarmain': 'prevencion_reportes',
            })

        # Rango del día completo para evitar pérdidas por hora/zona horaria.
        fecha_inicio_dt = datetime.combine(fecha_inicio, time.min)
        fecha_final_dt = datetime.combine(fecha_final, time.max)
        if timezone.is_naive(fecha_inicio_dt):
            fecha_inicio_dt = timezone.make_aware(fecha_inicio_dt, timezone.get_current_timezone())
        if timezone.is_naive(fecha_final_dt):
            fecha_final_dt = timezone.make_aware(fecha_final_dt, timezone.get_current_timezone())

        # Hoja 1: Documentos.
        documentos_qs = PrevencionDocumento.objects.filter(
            status=True,
            fecha_creacion__gte=fecha_inicio_dt,
            fecha_creacion__lte=fecha_final_dt,
        ).order_by('id').select_related('plantilla_base', 'faena').prefetch_related(
            'evidencias_generales',
            'evidencias_secciones',
            'evidencias_items',
        )
        documentos_qs = _filtrar_queryset_por_faena_usuario(documentos_qs, request.user)

        # Hoja 2: Vigilancia Médica.
        vigilancias_qs = VigilanciaMedica.objects.filter(
            status=True,
            fecha_incidente__gte=fecha_inicio,
            fecha_incidente__lte=fecha_final,
        ).order_by('id').prefetch_related('adjuntos_vigilancia')
        vigilancias_qs = _filtrar_queryset_por_faena_usuario(vigilancias_qs, request.user)

        base_url = request.build_absolute_uri('/').rstrip('/')
        filas_vigilancia = [_serializar_modelo_crudo(vig, base_url=base_url) for vig in vigilancias_qs]
        filas_vigilancia = _expandir_adjuntos_vigilancia(filas_vigilancia)

        workbook = Workbook()
        hoja_documentos = workbook.active
        hoja_documentos.title = 'Documentos'
        _escribir_hoja_documentos_jerarquico(hoja_documentos, documentos_qs, base_url=base_url)

        hoja_vigilancia = workbook.create_sheet(title='Vigilancia medica')
        _escribir_hoja_desde_filas(hoja_vigilancia, filas_vigilancia)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = (
            f'attachment; filename="prevencion-datos-{fecha_inicio.isoformat()}-a-{fecha_final.isoformat()}.xlsx"'
        )
        workbook.save(response)
        return response

    return render(request, 'pages/prevencion/export_data.html', {
        'form': form,
        'sidebar': 'prevencion_export_data',
        'sidebarmain': 'prevencion_reportes',
    })


# === BLOQUE PORTADO DESDE GEOATACAMA (VIGILANCIA RENOVADA) ===

def _conteo_trabajadores_por_filtros(faena_id, area_valor, ges_valor):
    # Cuenta trabajadores activos por combinación faena + area + ges.
    if not faena_id or not area_valor or not ges_valor:
        return 0

    return (
        UserInformacionLaboral.objects.filter(
            user__is_active=True,
            user__usuarioprofile__faena_id=faena_id,
            area__valor__iexact=str(area_valor).strip(),
            ges__valor__iexact=str(ges_valor).strip(),
        )
        .exclude(user__role=User.Role.SIN_ASIGNAR)
        .values('user_id')
        .distinct()
        .count()
    )


def _contratos_por_filtros_para_usuario(user, faena_id, area_valor, ges_valor, trabajadores_qs=None):
    # Obtiene contratos visibles por combinación faena + area + ges.
    if not faena_id or not area_valor or not ges_valor:
        return []

    base_qs = trabajadores_qs if trabajadores_qs is not None else _trabajadores_visibles_para_usuario(user)
    contratos_qs = (
        base_qs.filter(
            usuarioprofile__faena_id=faena_id,
            informacion_laboral__area__valor__iexact=str(area_valor).strip(),
            informacion_laboral__ges__valor__iexact=str(ges_valor).strip(),
        )
        .order_by()
        .values_list('informacion_laboral__tipo_contrato__valor', flat=True)
        .distinct()
    )

    contratos = sorted({
        str(contrato).strip()
        for contrato in contratos_qs
        if str(contrato or '').strip()
    })
    return contratos


def _resolver_agente_y_detalle_desde_post(post_data, agent_input_map):
    # Valida seleccion de agente en cuantitativa: exactamente uno y con detalle completo.
    selected_agents = [agent for agent in agent_input_map.keys() if post_data.get(f'{agent}_habilitado')]
    if len(selected_agents) > 1:
        return None, None, "Solo se puede habilitar un agente."
    if not selected_agents:
        return None, None, "Debes seleccionar un agente para la evaluacion cuantitativa."

    agente = selected_agents[0]
    agente_detalle = {}
    faltantes = []
    detalle_labels = getattr(HigieneCuantitativa, 'DETALLE_LABELS', {})

    for detail_key, input_name in agent_input_map[agente]:
        value = (post_data.get(input_name) or '').strip()
        if value:
            agente_detalle[detail_key] = value
            continue
        faltantes.append(detalle_labels.get(detail_key, detail_key.replace('_', ' ').title()))

    if faltantes:
        faltantes_txt = ', '.join(faltantes)
        return None, None, f"Completa los campos del agente seleccionado: {faltantes_txt}."

    return agente, agente_detalle, None


def _trabajadores_visibles_para_usuario(user):
    # Base comun de trabajadores visibles para construir lista por cuantitativa.
    usuarios = (
        User.objects.filter(is_active=True, usuarioprofile__isnull=False)
        .select_related(
            'usuarioprofile',
            'informacion_laboral__area',
            'informacion_laboral__ges',
            'informacion_laboral__tipo_contrato',
        )
        .order_by('first_name', 'last_name', 'username')
    )
    return _filtrar_queryset_por_faena_usuario(usuarios, user, campo_faena='usuarioprofile__faena')


def _normalizar_texto_filtro(valor):
    return ' '.join(str(valor or '').strip().upper().split())


def _cuantitativas_visibles_para_usuario(user):
    return _filtrar_queryset_por_faena_usuario(
        HigieneCuantitativa.objects.filter(status=True).select_related('faena', 'cualitativa'),
        user
    )


def _vigilancias_visibles_para_usuario(user):
    return _filtrar_queryset_por_faena_usuario(
        Vigilancia.objects.filter(status=True).select_related('faena', 'cuantitativa', 'cuantitativa__cualitativa'),
        user
    )


def _egresos_visibles_para_usuario(user):
    # En egreso se permite visibilidad por snapshot (egreso.faena) y por su vigilancia asociada.
    # Esto evita falsos 404 cuando el snapshot histórico quedó con faena distinta o vacía.
    usuario_profile = _get_usuario_profile(user)
    egresos = Egreso.objects.select_related('cuantitativa', 'faena', 'vigilancia').prefetch_related('adjuntos_historicos')

    if _is_sin_asignar(usuario_profile):
        return egresos

    if usuario_profile and usuario_profile.faena_id:
        return egresos.filter(
            Q(faena_id=usuario_profile.faena_id) |
            Q(vigilancia__faena_id=usuario_profile.faena_id)
        )

    return egresos.none()


def _catalogo_trabajadores_para_cuantitativa(user, cuantitativa):
    # Catalogo filtrado solo por faena de la cuantitativa.
    trabajadores_visibles = _trabajadores_visibles_para_usuario(user)
    if cuantitativa.faena_id:
        trabajadores_visibles = trabajadores_visibles.filter(
            usuarioprofile__faena_id=cuantitativa.faena_id
        )
    catalogo = []

    for trabajador in trabajadores_visibles:
        try:
            info_laboral = trabajador.informacion_laboral
        except UserInformacionLaboral.DoesNotExist:
            info_laboral = None

        area_trabajador = info_laboral.area.valor if info_laboral and info_laboral.area else ''
        ges_trabajador = info_laboral.ges.valor if info_laboral and info_laboral.ges else ''

        nombre = f"{trabajador.first_name} {trabajador.last_name}".strip() or trabajador.username
        faena_nombre = ''
        if hasattr(trabajador, 'usuarioprofile') and getattr(trabajador.usuarioprofile, 'faena', None):
            faena_nombre = trabajador.usuarioprofile.faena.faena

        catalogo.append({
            'user': trabajador,
            'username': trabajador.username,
            'nombre': nombre,
            'faena_id': trabajador.usuarioprofile.faena_id if hasattr(trabajador, 'usuarioprofile') else '',
            'faena_nombre': faena_nombre,
            'area': area_trabajador,
            'ges': ges_trabajador,
            'cargo': info_laboral.cargo.valor if info_laboral and info_laboral.cargo else '',
        })

    return catalogo


def _trabajador_coincide_filtro_automatico(contexto_filtro, trabajador):
    # Define quienes entran automaticamente a la tabla (faena ya viene filtrada en el catalogo).
    area_objetivo = _normalizar_texto_filtro(getattr(contexto_filtro, 'area', ''))
    ges_objetivo = _normalizar_texto_filtro(getattr(contexto_filtro, 'ges', ''))
    area_trabajador = _normalizar_texto_filtro(trabajador.get('area'))
    ges_trabajador = _normalizar_texto_filtro(trabajador.get('ges'))

    if area_objetivo and area_trabajador != area_objetivo:
        return False
    if ges_objetivo and ges_trabajador != ges_objetivo:
        return False
    return True

def _reemplazar_archivo_vigilancia(vigilancia, archivo_nuevo):
    # Mantiene un solo archivo por ficha: al subir uno nuevo se reemplaza el anterior.
    if not archivo_nuevo:
        return

    if vigilancia.documento_adjunto:
        vigilancia.documento_adjunto.delete(save=False)
        vigilancia.documento_adjunto = None
        vigilancia.save(update_fields=['documento_adjunto'])

    for adjunto in vigilancia.adjuntos_vigilancia.all():
        if adjunto.archivo:
            adjunto.archivo.delete(save=False)
        adjunto.delete()

    VigilanciaAdjunto.objects.create(vigilancia=vigilancia, archivo=archivo_nuevo)


def _archivo_principal_vigilancia(vigilancia):
    # Prioriza el adjunto nuevo; fallback al campo legacy de la ficha.
    adjunto_principal = vigilancia.adjuntos_vigilancia.order_by('-id').first()
    if adjunto_principal and adjunto_principal.archivo:
        return adjunto_principal.archivo
    if vigilancia.documento_adjunto:
        return vigilancia.documento_adjunto
    return None


def _serializar_adjuntos_vigilancia(vigilancia):
    # Vista funcional: expone solo el archivo vigente/principal.
    adjunto_principal = vigilancia.adjuntos_vigilancia.order_by('-id').first()
    if adjunto_principal and adjunto_principal.archivo:
        return [{
            'id': adjunto_principal.id,
            'url': adjunto_principal.archivo.url,
            'name': os.path.basename(adjunto_principal.archivo.name or '') or 'Archivo',
            'delete_url': reverse('delete_vigilancia_adjunto', args=[adjunto_principal.pk]),
        }]

    if vigilancia.documento_adjunto:
        return [{
            'id': f'legacy-{vigilancia.pk}',
            'url': vigilancia.documento_adjunto.url,
            'name': os.path.basename(vigilancia.documento_adjunto.name or '') or 'Archivo',
            'delete_url': reverse('delete_vigilancia_doc', args=[vigilancia.pk]),
        }]

    return []


def _serializar_archivos_egreso(egreso):
    if not egreso:
        return []

    filas = []
    if egreso.documento_adjunto:
        filas.append({
            'id': f'egreso-actual-{egreso.pk}',
            'orden': 0,
            'es_actual': True,
            'fecha_carga': egreso.documento_adjunto_cargado_at or egreso.fecha_egreso or egreso.created_at,
            'url': egreso.documento_adjunto.url,
            'name': os.path.basename(egreso.documento_adjunto.name or '') or 'Archivo',
        })

    for historico in egreso.adjuntos_historicos.all():
        if not historico.archivo:
            continue
        filas.append({
            'id': f'egreso-historico-{historico.pk}',
            'orden': 1,
            'es_actual': False,
            'fecha_carga': historico.created_at,
            'url': historico.archivo.url,
            'name': os.path.basename(historico.archivo.name or '') or 'Archivo',
        })

    return filas


def _archivo_actual_egreso(egreso):
    if not egreso or not egreso.documento_adjunto:
        return None
    return {
        'id': f'egreso-actual-{egreso.pk}',
        'url': egreso.documento_adjunto.url,
        'name': os.path.basename(egreso.documento_adjunto.name or '') or 'Archivo',
        'fecha_carga': egreso.documento_adjunto_cargado_at or egreso.fecha_egreso or egreso.created_at,
    }


def _mover_adjunto_actual_a_historico_egreso(egreso):
    if not egreso or not egreso.documento_adjunto:
        return None

    archivo_actual = egreso.documento_adjunto
    nombre_archivo = os.path.basename(archivo_actual.name or '') or f"egreso-{egreso.pk}"
    fecha_original = (
        egreso.documento_adjunto_cargado_at
        or egreso.fecha_egreso
        or egreso.created_at
        or timezone.now()
    )

    historico = EgresoAdjuntoHistorico.objects.create(egreso=egreso)
    try:
        archivo_actual.open('rb')
        historico.archivo.save(nombre_archivo, File(archivo_actual), save=True)
    finally:
        try:
            archivo_actual.close()
        except Exception:
            pass

    EgresoAdjuntoHistorico.objects.filter(pk=historico.pk).update(created_at=fecha_original)
    historico.created_at = fecha_original

    archivo_actual.delete(save=False)
    egreso.documento_adjunto = None
    egreso.documento_adjunto_cargado_at = None
    egreso.save(update_fields=['documento_adjunto', 'documento_adjunto_cargado_at'])
    return historico


def _guardar_adjunto_actual_egreso(egreso, archivo_nuevo, fecha_carga=None):
    if not egreso or not archivo_nuevo:
        return

    nombre_archivo = os.path.basename(getattr(archivo_nuevo, 'name', '') or '') or f"egreso-{egreso.pk}"
    egreso.documento_adjunto_cargado_at = fecha_carga or timezone.now()
    egreso.documento_adjunto.save(nombre_archivo, archivo_nuevo, save=False)
    egreso.save(update_fields=['documento_adjunto', 'documento_adjunto_cargado_at'])


def _reemplazar_archivo_egreso(egreso, archivo_nuevo):
    if not egreso or not archivo_nuevo:
        return
    _mover_adjunto_actual_a_historico_egreso(egreso)
    _guardar_adjunto_actual_egreso(egreso, archivo_nuevo, fecha_carga=timezone.now())

@login_required
@prevencion_riesgo_required
def higiene_cualitativa(request):
    registros = HigieneCualitativa.objects.select_related('faena')
    registros = _filtrar_queryset_por_faena_usuario(registros, request.user).order_by('-fecha', '-id')
    context = {
        'registros': registros,
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'higiene_cualitativa',
    }
    return render(request, 'pages/prevencion/higiene_cualitativa.html', context)


@login_required
@prevencion_riesgo_required
def new_higiene_cualitativa(request):
    faenas = _faenas_seleccionables_para_usuario(request.user)
    estados = VigilanciaEstado.objects.filter(status=True).order_by('correlativo', 'id')
    contratos = VigilanciaContrato.objects.filter(status=True).order_by('correlativo', 'id')

    if request.method == 'POST':
        faena_id = request.POST.get('faena')
        codigo = (request.POST.get('codigo') or '').strip()
        nombre = (request.POST.get('nombre') or '').strip()
        fecha = request.POST.get('fecha')
        estado = (request.POST.get('estado') or '').strip()
        contrato = (request.POST.get('contrato') or '').strip()
        documento = request.FILES.get('documento')

        if not _faena_permitida_para_usuario(request.user, faena_id):
            messages.error(request, "No tienes permisos para registrar en esa faena.")
            return redirect('new_higiene_cualitativa')

        if not all([faena_id, codigo, nombre, fecha, estado, contrato, documento]):
            messages.error(request, "Todos los campos son obligatorios.")
            return redirect('new_higiene_cualitativa')

        ext = os.path.splitext((documento.name or '').lower())[1]
        if ext not in {'.pdf', '.doc', '.docx'}:
            messages.error(request, "El documento debe ser PDF o Word (.doc, .docx).")
            return redirect('new_higiene_cualitativa')

        nombre_usuario = f"{request.user.first_name} {request.user.last_name}".strip()
        if not nombre_usuario:
            nombre_usuario = request.user.username

        informe_tecnico = os.path.basename(documento.name or '').strip() or 'Informe tecnico'

        try:
            HigieneCualitativa.objects.create(
                faena_id=faena_id,
                codigo=codigo,
                nombre=nombre,
                fecha=fecha,
                estado=estado,
                contrato=contrato,
                informe_tecnico=informe_tecnico,
                documento=documento,
                creador=nombre_usuario,
                status=True,
            )
            messages.success(request, "Evaluacion cualitativa creada correctamente.")
            user_name = request.user.get_full_name() or request.user.username
            notify_group('prevencion_general', f'Higiene Cualitativa creada: {codigo}', f'Higiene Ocupacional Cualitativa: {codigo} - {nombre}\nCreada por: {user_name}')
            return redirect('higiene_cualitativa')
        except IntegrityError:
            messages.error(request, "El codigo ya existe, usa otro distinto.")
            return redirect('new_higiene_cualitativa')
        except Exception as e:
            messages.error(request, f"No se pudo guardar: {e}")
            return redirect('new_higiene_cualitativa')

    context = {
        'faenas': faenas,
        'estados': estados,
        'contratos': contratos,
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'higiene_cualitativa',
    }
    return render(request, 'pages/prevencion/new_higiene_cualitativa.html', context)


@login_required
@prevencion_riesgo_required
def view_higiene_cualitativa(request, pk):
    registro = get_object_or_404(HigieneCualitativa, id=pk)
    if not _faena_permitida_para_usuario(request.user, registro.faena_id):
        messages.error(request, "No tienes permisos para ver este registro.")
        return redirect('higiene_cualitativa')

    estados = VigilanciaEstado.objects.filter(status=True).order_by('correlativo', 'id')
    context = {
        'registro': registro,
        'estados': estados,
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'higiene_cualitativa',
    }
    return render(request, 'pages/prevencion/view_higiene_cualitativa.html', context)


@login_required
@prevencion_riesgo_required
def edit_higiene_cualitativa(request, pk):
    registro = get_object_or_404(HigieneCualitativa, id=pk)
    faenas = _faenas_seleccionables_para_usuario(request.user)
    estados = VigilanciaEstado.objects.filter(status=True).order_by('correlativo', 'id')
    contratos = VigilanciaContrato.objects.filter(status=True).order_by('correlativo', 'id')

    if not _faena_permitida_para_usuario(request.user, registro.faena_id):
        messages.error(request, "No tienes permisos para editar este registro.")
        return redirect('higiene_cualitativa')

    if request.method == 'POST':
        faena_id = request.POST.get('faena')
        codigo = (request.POST.get('codigo') or '').strip()
        nombre = (request.POST.get('nombre') or '').strip()
        fecha = request.POST.get('fecha')
        estado = (request.POST.get('estado') or '').strip()
        contrato = (request.POST.get('contrato') or '').strip()
        documento = request.FILES.get('documento')

        if not _faena_permitida_para_usuario(request.user, faena_id):
            messages.error(request, "No tienes permisos para editar este registro.")
            return redirect('higiene_cualitativa')

        if not all([faena_id, codigo, nombre, fecha, estado, contrato]):
            messages.error(request, "Todos los campos son obligatorios, excepto el archivo.")
            return redirect('edit_higiene_cualitativa', pk=pk)

        old = HigieneCualitativa.objects.get(pk=registro.pk)

        if not all([faena_id, codigo, nombre, fecha, estado, contrato]):
            messages.error(request, "Todos los campos son obligatorios, excepto el archivo.")
            return redirect('edit_higiene_cualitativa', pk=pk)

        if documento:
            ext = os.path.splitext((documento.name or '').lower())[1]
            if ext not in {'.pdf', '.doc', '.docx'}:
                messages.error(request, "El documento debe ser PDF o Word (.doc, .docx).")
                return redirect('edit_higiene_cualitativa', pk=pk)
            registro.informe_tecnico = os.path.basename(documento.name or '').strip() or registro.informe_tecnico
            registro.documento = documento

        try:
            registro.faena_id = faena_id
            registro.codigo = codigo
            registro.nombre = nombre
            registro.fecha = fecha
            registro.estado = estado
            registro.contrato = contrato
            registro.save()

            messages.success(request, "Evaluacion cualitativa actualizada correctamente.")
            user_name = request.user.get_full_name() or request.user.username
            changes = get_changes_message(old, registro)
            if changes:
                notify_group('prevencion_general', f'Higiene Cualitativa actualizada: {registro.codigo}', f'Higiene Ocupacional Cualitativa: {registro.codigo} - {registro.nombre}\nActualizada por: {user_name}{changes}')
            return redirect('higiene_cualitativa')
        except IntegrityError:
            messages.error(request, "El codigo ya existe, usa otro distinto.")
            return redirect('edit_higiene_cualitativa', pk=pk)
        except Exception as e:
            messages.error(request, f"No se pudo actualizar: {e}")
            return redirect('edit_higiene_cualitativa', pk=pk)

    context = {
        'registro': registro,
        'faenas': faenas,
        'estados': estados,
        'contratos': contratos,
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'higiene_cualitativa',
    }
    return render(request, 'pages/prevencion/edit_higiene_cualitativa.html', context)


@login_required
@prevencion_riesgo_required
def disable_higiene_cualitativa(request, pk):
    if request.method != 'POST':
        return redirect('higiene_cualitativa')

    if not _es_admin_prevencion(request.user):
        messages.error(request, "No tienes permisos para deshabilitar registros.")
        return redirect('higiene_cualitativa')

    registro = get_object_or_404(HigieneCualitativa, id=pk)
    registro.status = not registro.status
    registro.save(update_fields=['status'])
    accion = "habilitada" if registro.status else "deshabilitada"
    messages.success(request, f"Evaluacion cualitativa {accion} correctamente.")
    user_name = request.user.get_full_name() or request.user.username
    notify_group('prevencion_general', f'Higiene Cualitativa {accion}: {registro.codigo}', f'Higiene Ocupacional Cualitativa: {registro.codigo} - {registro.nombre}\n{accion.capitalize()} por: {user_name}')
    return redirect('higiene_cualitativa')


@login_required
@prevencion_riesgo_required
def higiene_cuantitativa(request):
    registros_qs = HigieneCuantitativa.objects.select_related('faena', 'cualitativa')
    registros_qs = _filtrar_queryset_por_faena_usuario(registros_qs, request.user).order_by('-fecha', '-id')
    registros = list(registros_qs)

    trabajadores_visibles = _trabajadores_visibles_para_usuario(request.user)
    cache_contratos = {}
    for registro in registros:
        clave = (
            registro.faena_id,
            _normalizar_texto_filtro(registro.area),
            _normalizar_texto_filtro(registro.ges),
        )
        if clave not in cache_contratos:
            cache_contratos[clave] = _contratos_por_filtros_para_usuario(
                user=request.user,
                faena_id=registro.faena_id,
                area_valor=registro.area,
                ges_valor=registro.ges,
                trabajadores_qs=trabajadores_visibles,
            )
        contratos = cache_contratos[clave]
        registro.contrato_filtro_texto = ' | '.join(contratos) if contratos else '-'

    context = {
        'registros': registros,
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'higiene_cuantitativa',
    }
    return render(request, 'pages/prevencion/higiene_cuantitativa.html', context)


@login_required
@prevencion_riesgo_required
def new_higiene_cuantitativa(request):
    agent_input_map = {
        'silice': [
            ('evaluacion_riesgo', 'silice_eval_riesgo'),
            ('nivel_riesgo', 'silice_nivel_riesgo'),
            ('grado_exposicion', 'silice_grado_exposicion'),
        ],
        'ruido': [
            ('evaluacion_riesgo', 'ruido_eval_riesgo'),
            ('nivel_seguimiento', 'ruido_nivel_seguimiento'),
            ('grado_exposicion', 'ruido_grado_exposicion'),
        ],
        'hipobaria': [
            ('exposicion', 'hipobaria_exposicion'),
        ],
        'vibraciones': [
            ('evaluacion_riesgo', 'vibraciones_eval_riesgo'),
            ('exposicion', 'vibraciones_exposicion'),
        ],
        'radiaciones': [
            ('exposicion', 'radiaciones_exposicion'),
        ],
        'humos': [
            ('evaluacion_riesgo', 'humos_eval_riesgo'),
            ('exposicion', 'humos_exposicion'),
        ],
    }

    cualitativas = HigieneCualitativa.objects.filter(status=True).select_related('faena')
    cualitativas = _filtrar_queryset_por_faena_usuario(cualitativas, request.user)
    cualitativas = cualitativas.order_by('-fechacreacion', '-id')
    catalogos_vigilancia = _catalogos_vigilancia_formulario()
    areas = catalogos_vigilancia['opciones_area']
    ges_opciones = catalogos_vigilancia['opciones_ges']

    if request.method == 'POST':
        cualitativa_id = request.POST.get('cualitativa')
        codigo = (request.POST.get('codigo') or '').strip()
        nombre = (request.POST.get('nombre') or '').strip()
        fecha = request.POST.get('fecha')
        area = (request.POST.get('area') or '').strip()
        ges = (request.POST.get('ges') or '').strip()
        documento = request.FILES.get('documento')

        agente, agente_detalle, msg_agente = _resolver_agente_y_detalle_desde_post(request.POST, agent_input_map)
        if msg_agente:
            messages.error(request, msg_agente)
            return redirect('new_higiene_cuantitativa')

        if not all([cualitativa_id, codigo, nombre, fecha, area, ges, documento]):
            messages.error(request, "Todos los campos son obligatorios.")
            return redirect('new_higiene_cuantitativa')

        ext = os.path.splitext((documento.name or '').lower())[1]
        if ext not in {'.pdf', '.doc', '.docx'}:
            messages.error(request, "El documento debe ser PDF o Word (.doc, .docx).")
            return redirect('new_higiene_cuantitativa')

        cualitativa = cualitativas.filter(id=cualitativa_id).first()
        if not cualitativa:
            messages.error(request, "La evaluacion cualitativa seleccionada no esta disponible.")
            return redirect('new_higiene_cuantitativa')

        if not _faena_permitida_para_usuario(request.user, cualitativa.faena_id):
            messages.error(request, "No tienes permisos para registrar en esa faena.")
            return redirect('new_higiene_cuantitativa')

        nombre_usuario = f"{request.user.first_name} {request.user.last_name}".strip()
        if not nombre_usuario:
            nombre_usuario = request.user.username

        try:
            informe_tecnico = os.path.basename(documento.name or '').strip() or 'Informe tecnico'
            HigieneCuantitativa.objects.create(
                faena_id=cualitativa.faena_id,
                cualitativa=cualitativa,
                codigo=codigo,
                nombre=nombre,
                fecha=fecha,
                area=area,
                ges=ges,
                informe_tecnico=informe_tecnico,
                documento=documento,
                agente=agente,
                agente_detalle=agente_detalle,
                creador=nombre_usuario,
                status=True,
            )
            messages.success(request, "Evaluacion cuantitativa creada correctamente.")
            user_name = request.user.get_full_name() or request.user.username
            notify_group('prevencion_general', f'Higiene Cuantitativa creada: {codigo}', f'Higiene Ocupacional Cuantitativa: {codigo} - {nombre}\nCreada por: {user_name}')
            return redirect('higiene_cuantitativa')
        except IntegrityError:
            messages.error(request, "El codigo de cuantitativa ya existe, usa otro distinto.")
            return redirect('new_higiene_cuantitativa')
        except Exception as e:
            messages.error(request, f"No se pudo guardar: {e}")
            return redirect('new_higiene_cuantitativa')

    context = {
        'cualitativas': cualitativas.order_by('-fechacreacion', '-id'),
        'areas': areas,
        'ges_opciones': ges_opciones,
        'opciones_eval_riesgo': catalogos_vigilancia['opciones_eval_riesgo'],
        'opciones_nivel_riesgo': catalogos_vigilancia['opciones_nivel_riesgo'],
        'opciones_nivel_seguimiento': catalogos_vigilancia['opciones_nivel_seguimiento'],
        'opciones_grado_exposicion': catalogos_vigilancia['opciones_grado_exposicion'],
        'opciones_exposicion': catalogos_vigilancia['opciones_exposicion'],
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'higiene_cuantitativa',
    }
    return render(request, 'pages/prevencion/new_higiene_cuantitativa.html', context)


@login_required
@prevencion_riesgo_required
def ajax_conteo_trabajadores(request):
    faena_id = (request.GET.get('faena_id') or '').strip()
    area_valor = (request.GET.get('area') or '').strip()
    ges_valor = (request.GET.get('ges') or '').strip()

    if not faena_id or not area_valor or not ges_valor:
        return JsonResponse({'count': 0})

    if not _faena_permitida_para_usuario(request.user, faena_id):
        return JsonResponse({'count': 0})

    conteo = _conteo_trabajadores_por_filtros(
        faena_id=faena_id,
        area_valor=area_valor,
        ges_valor=ges_valor,
    )
    return JsonResponse({'count': conteo})


@login_required
@prevencion_riesgo_required
def view_higiene_cuantitativa(request, pk):
    registro = get_object_or_404(HigieneCuantitativa, id=pk)
    if not _faena_permitida_para_usuario(request.user, registro.faena_id):
        messages.error(request, "No tienes permisos para ver este registro.")
        return redirect('higiene_cuantitativa')

    detalle = registro.agente_detalle if isinstance(registro.agente_detalle, dict) else {}
    conteo_trabajadores = _conteo_trabajadores_por_filtros(
        faena_id=registro.faena_id,
        area_valor=registro.area,
        ges_valor=registro.ges,
    )
    context = {
        'registro': registro,
        'detalle': detalle,
        'conteo_trabajadores': conteo_trabajadores,
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'higiene_cuantitativa',
    }
    return render(request, 'pages/prevencion/view_higiene_cuantitativa.html', context)


@login_required
@prevencion_riesgo_required
def edit_higiene_cuantitativa(request, pk):
    registro = get_object_or_404(HigieneCuantitativa, id=pk)
    if not _faena_permitida_para_usuario(request.user, registro.faena_id):
        messages.error(request, "No tienes permisos para editar este registro.")
        return redirect('higiene_cuantitativa')

    agent_input_map = {
        'silice': [
            ('evaluacion_riesgo', 'silice_eval_riesgo'),
            ('nivel_riesgo', 'silice_nivel_riesgo'),
            ('grado_exposicion', 'silice_grado_exposicion'),
        ],
        'ruido': [
            ('evaluacion_riesgo', 'ruido_eval_riesgo'),
            ('nivel_seguimiento', 'ruido_nivel_seguimiento'),
            ('grado_exposicion', 'ruido_grado_exposicion'),
        ],
        'hipobaria': [
            ('exposicion', 'hipobaria_exposicion'),
        ],
        'vibraciones': [
            ('evaluacion_riesgo', 'vibraciones_eval_riesgo'),
            ('exposicion', 'vibraciones_exposicion'),
        ],
        'radiaciones': [
            ('exposicion', 'radiaciones_exposicion'),
        ],
        'humos': [
            ('evaluacion_riesgo', 'humos_eval_riesgo'),
            ('exposicion', 'humos_exposicion'),
        ],
    }

    cualitativas = HigieneCualitativa.objects.filter(status=True).select_related('faena')
    cualitativas = _filtrar_queryset_por_faena_usuario(cualitativas, request.user)
    catalogos_vigilancia = _catalogos_vigilancia_formulario()

    if request.method == 'POST':
        cualitativa_id = request.POST.get('cualitativa')
        codigo = (request.POST.get('codigo') or '').strip()
        nombre = (request.POST.get('nombre') or '').strip()
        fecha = request.POST.get('fecha')
        area = (request.POST.get('area') or '').strip()
        ges = (request.POST.get('ges') or '').strip()
        documento = request.FILES.get('documento')

        if not all([cualitativa_id, codigo, nombre, fecha, area, ges]):
            messages.error(request, "Todos los campos principales son obligatorios.")
            return redirect('edit_higiene_cuantitativa', pk=pk)

        old = HigieneCuantitativa.objects.get(pk=registro.pk)

        agente, agente_detalle, msg_agente = _resolver_agente_y_detalle_desde_post(request.POST, agent_input_map)
        if msg_agente:
            messages.error(request, msg_agente)
            return redirect('edit_higiene_cuantitativa', pk=pk)

        cualitativa = cualitativas.filter(id=cualitativa_id).first()
        if not cualitativa:
            messages.error(request, "La evaluacion cualitativa seleccionada no esta disponible.")
            return redirect('edit_higiene_cuantitativa', pk=pk)

        if not _faena_permitida_para_usuario(request.user, cualitativa.faena_id):
            messages.error(request, "No tienes permisos para asignar esa faena.")
            return redirect('edit_higiene_cuantitativa', pk=pk)

        if documento:
            ext = os.path.splitext((documento.name or '').lower())[1]
            if ext not in {'.pdf', '.doc', '.docx'}:
                messages.error(request, "El documento debe ser PDF o Word (.doc, .docx).")
                return redirect('edit_higiene_cuantitativa', pk=pk)

        try:
            registro.faena_id = cualitativa.faena_id
            registro.cualitativa = cualitativa
            registro.codigo = codigo
            registro.nombre = nombre
            registro.fecha = fecha
            registro.area = area
            registro.ges = ges
            if documento:
                registro.informe_tecnico = os.path.basename(documento.name or '').strip() or registro.informe_tecnico
                registro.documento = documento
            registro.agente = agente
            registro.agente_detalle = agente_detalle
            registro.save()

            messages.success(request, "Evaluacion cuantitativa actualizada correctamente.")
            user_name = request.user.get_full_name() or request.user.username
            changes = get_changes_message(old, registro)
            if changes:
                notify_group('prevencion_general', f'Higiene Cuantitativa actualizada: {registro.codigo}', f'Higiene Ocupacional Cuantitativa: {registro.codigo} - {registro.nombre}\nActualizada por: {user_name}{changes}')
            return redirect('higiene_cuantitativa')
        except IntegrityError:
            messages.error(request, "El codigo de cuantitativa ya existe, usa otro distinto.")
            return redirect('edit_higiene_cuantitativa', pk=pk)
        except Exception as e:
            messages.error(request, f"No se pudo actualizar: {e}")
            return redirect('edit_higiene_cuantitativa', pk=pk)

    detalle = registro.agente_detalle if isinstance(registro.agente_detalle, dict) else {}
    context = {
        'registro': registro,
        'detalle': detalle,
        'cualitativas': cualitativas.order_by('-fechacreacion', '-id'),
        'areas': catalogos_vigilancia['opciones_area'],
        'ges_opciones': catalogos_vigilancia['opciones_ges'],
        'opciones_eval_riesgo': catalogos_vigilancia['opciones_eval_riesgo'],
        'opciones_nivel_riesgo': catalogos_vigilancia['opciones_nivel_riesgo'],
        'opciones_nivel_seguimiento': catalogos_vigilancia['opciones_nivel_seguimiento'],
        'opciones_grado_exposicion': catalogos_vigilancia['opciones_grado_exposicion'],
        'opciones_exposicion': catalogos_vigilancia['opciones_exposicion'],
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'higiene_cuantitativa',
    }
    return render(request, 'pages/prevencion/edit_higiene_cuantitativa.html', context)


@login_required
@prevencion_riesgo_required
def disable_higiene_cuantitativa(request, pk):
    if request.method != 'POST':
        return redirect('higiene_cuantitativa')

    if not _es_admin_prevencion(request.user):
        messages.error(request, "No tienes permisos para deshabilitar registros.")
        return redirect('higiene_cuantitativa')

    registro = get_object_or_404(HigieneCuantitativa, id=pk)
    registro.status = not registro.status
    registro.save(update_fields=['status'])
    accion = "habilitada" if registro.status else "deshabilitada"
    messages.success(request, f"Evaluacion cuantitativa {accion} correctamente.")
    user_name = request.user.get_full_name() or request.user.username
    notify_group('prevencion_general', f'Higiene Cuantitativa {accion}: {registro.codigo}', f'Higiene Ocupacional Cuantitativa: {registro.codigo} - {registro.nombre}\n{accion.capitalize()} por: {user_name}')
    return redirect('higiene_cuantitativa')


@login_required
@prevencion_riesgo_required
def manage_vigilancia(request):
    vigilancias = _vigilancias_con_total_trabajadores_para_usuario(request.user)
    trabajadores_visibles = _trabajadores_visibles_para_usuario(request.user)
    cache_contratos = {}

    for vigilancia in vigilancias:
        clave = (
            vigilancia.faena_id,
            _normalizar_texto_filtro(vigilancia.area),
            _normalizar_texto_filtro(vigilancia.ges),
        )
        if clave not in cache_contratos:
            cache_contratos[clave] = _contratos_por_filtros_para_usuario(
                user=request.user,
                faena_id=vigilancia.faena_id,
                area_valor=vigilancia.area,
                ges_valor=vigilancia.ges,
                trabajadores_qs=trabajadores_visibles,
            )
        contratos = cache_contratos[clave]
        vigilancia.contrato_filtro_texto = ' | '.join(contratos) if contratos else '-'

    context = {
        'vigilancias': vigilancias,
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/manage_vigilancia.html', context)


def _campos_snapshot_egreso():
    # Campos que se congelan al egresar para preservar historial.
    return [
        'user_id',
        'cuantitativa_id',
        'vigilancia_id',
        'rut_trabajador',
        'nombre_completo',
        'genero_texto',
        'fecha_nacimiento',
        'edad',
        'area',
        'cargo',
        'faena_id',
        'fecha_ingreso',
        'antiguedad_anos',
        'antiguedad_meses',
        'antiguedad_dias',
        'tipo_contrato',
        'fecha_retiro',
        'ges',
        'silice_eval_riesgo',
        'silice_nivel_riesgo',
        'silice_grado_expo',
        'silice_fecha_radio',
        'silice_fecha_vence',
        'silice_vigencia',
        'silice_obs',
        'ruido_eval_riesgo',
        'ruido_nivel_seguimiento',
        'ruido_grado_expo',
        'ruido_fecha_audio',
        'ruido_fecha_vence',
        'ruido_vigencia',
        'ruido_obs',
        'hipo_exposicion',
        'hipo_fecha_hemo',
        'hipo_fecha_vence',
        'hipo_vigencia',
        'hipo_obs',
        'vibra_eval_riesgo',
        'vibra_exposicion',
        'vibra_fecha_vence',
        'vibra_vigencia',
        'rad_exposicion',
        'rad_fecha_vence',
        'rad_vigencia',
        'humos_eval_riesgo',
        'humos_exposicion',
        'humos_fecha_vence',
        'humos_vigencia',
        'fecha_incidente',
        'organismo_administrador',
        'clinica',
        'observacion_general',
        'creado_por',
    ]


def _egreso_reutilizable_desde_ficha(ficha):
    if not ficha:
        return None

    egreso = Egreso.objects.filter(ficha_origen_id=ficha.id).order_by('-fecha_egreso', '-id').first()
    if egreso:
        return egreso

    if ficha.vigilancia_id and ficha.user_id:
        return (
            Egreso.objects.filter(vigilancia_id=ficha.vigilancia_id, user_id=ficha.user_id)
            .order_by('-fecha_egreso', '-id')
            .first()
        )
    return None


def _crear_egreso_desde_ficha(ficha, fecha_egreso=None):
    # Crea (o reutiliza) una fila historica en prevencion_egreso.
    fecha_evento = fecha_egreso or timezone.now()
    payload = {campo: getattr(ficha, campo) for campo in _campos_snapshot_egreso()}
    payload['ficha_origen_id'] = ficha.id
    payload['fecha_egreso'] = fecha_evento

    archivo_origen = _archivo_principal_vigilancia(ficha)
    tiene_adjunto_origen = bool(archivo_origen and getattr(archivo_origen, 'name', None))

    egreso = _egreso_reutilizable_desde_ficha(ficha)
    if egreso:
        if tiene_adjunto_origen:
            _mover_adjunto_actual_a_historico_egreso(egreso)
        for campo, valor in payload.items():
            setattr(egreso, campo, valor)
        egreso.save()
    else:
        egreso = Egreso.objects.create(**payload)

    if tiene_adjunto_origen:
        nombre_archivo = os.path.basename(archivo_origen.name) or f"egreso-{egreso.pk}"
        try:
            archivo_origen.open('rb')
            _guardar_adjunto_actual_egreso(
                egreso,
                File(archivo_origen, name=nombre_archivo),
                fecha_carga=fecha_evento,
            )
        finally:
            try:
                archivo_origen.close()
            except Exception:
                pass

    return egreso


def _egreso_mas_reciente_desde_ficha(ficha):
    if not ficha:
        return None
    egreso = Egreso.objects.filter(ficha_origen_id=ficha.id).order_by('-fecha_egreso', '-id').first()
    if egreso:
        return egreso
    if ficha.vigilancia_id and ficha.user_id:
        return (
            Egreso.objects.filter(vigilancia_id=ficha.vigilancia_id, user_id=ficha.user_id)
            .order_by('-fecha_egreso', '-id')
            .first()
        )
    return None


def _calcular_vigencia_desde_fecha(fecha_raw):
    if not fecha_raw:
        return 'PENDIENTE'
    fecha_obj = parse_date(fecha_raw) if isinstance(fecha_raw, str) else fecha_raw
    if not fecha_obj:
        return 'PENDIENTE'
    return 'VIGENTE' if fecha_obj >= timezone.localdate() else 'VENCIDO'


def _redireccion_egreso_desde_ficha(request, ficha, mensaje=None):
    egreso = _egreso_mas_reciente_desde_ficha(ficha)
    if egreso:
        if mensaje:
            messages.info(request, mensaje)
        return redirect('view_vigilancia_egreso', pk=egreso.id)
    if mensaje:
        messages.info(request, mensaje)
    return redirect('egreso')


def _resumen_fichas_por_usuario_en_vigilancia(vigilancia, user_ids=None):
    # Resumen por usuario: ficha activa, ultima ficha y existencia de egresos.
    fichas = VigilanciaMedica.objects.filter(vigilancia=vigilancia).select_related('user', 'faena')
    if user_ids is not None:
        fichas = fichas.filter(user_id__in=user_ids)

    resumen = {}
    for ficha in fichas.order_by('-created_at', '-id'):
        if not ficha.user_id:
            continue
        data = resumen.setdefault(ficha.user_id, {'activa': None, 'ultima': None, 'tiene_egreso': False})
        if data['ultima'] is None:
            data['ultima'] = ficha
        if ficha.status and data['activa'] is None:
            data['activa'] = ficha

    user_ids_objetivo = set(resumen.keys())
    if user_ids is not None:
        user_ids_objetivo |= set(user_ids)
    if user_ids_objetivo:
        user_ids_con_egreso = set(
            Egreso.objects.filter(
                vigilancia=vigilancia,
                user_id__in=user_ids_objetivo,
            ).values_list('user_id', flat=True)
        )
        for uid in user_ids_objetivo:
            data = resumen.setdefault(uid, {'activa': None, 'ultima': None, 'tiene_egreso': False})
            data['tiene_egreso'] = uid in user_ids_con_egreso

    return resumen


def _trabajador_oculto_segun_resumen(resumen_usuario):
    if not resumen_usuario:
        return False
    if resumen_usuario.get('activa'):
        return False
    # Solo ocultar cuando ya existe egreso; los eliminados deben poder reingresar.
    return bool(resumen_usuario.get('tiene_egreso'))


def _trabajador_eliminado_segun_resumen(resumen_usuario):
    if not resumen_usuario:
        return False
    if resumen_usuario.get('activa'):
        return False
    if resumen_usuario.get('tiene_egreso'):
        return False
    # Ultima ficha inactiva y sin egreso: eliminado de la lista actual.
    return resumen_usuario.get('ultima') is not None


def _prefill_agente_desde_cuantitativa(cuantitativa):
    agente = (cuantitativa.agente or '').strip() if cuantitativa else ''
    agente_detalle = cuantitativa.agente_detalle if cuantitativa and isinstance(cuantitativa.agente_detalle, dict) else {}

    if agente == 'silice':
        return {
            'silice_eval_riesgo': agente_detalle.get('evaluacion_riesgo', ''),
            'silice_nivel_riesgo': agente_detalle.get('nivel_riesgo', ''),
            'silice_grado_expo': agente_detalle.get('grado_exposicion', ''),
        }
    if agente == 'ruido':
        return {
            'ruido_eval_riesgo': agente_detalle.get('evaluacion_riesgo', ''),
            'ruido_nivel_seguimiento': agente_detalle.get('nivel_seguimiento', ''),
            'ruido_grado_expo': agente_detalle.get('grado_exposicion', ''),
        }
    if agente == 'hipobaria':
        return {
            'hipo_exposicion': agente_detalle.get('exposicion', ''),
        }
    if agente == 'vibraciones':
        return {
            'vibra_eval_riesgo': agente_detalle.get('evaluacion_riesgo', ''),
            'vibra_exposicion': agente_detalle.get('exposicion', ''),
        }
    if agente == 'radiaciones':
        return {
            'rad_exposicion': agente_detalle.get('exposicion', ''),
        }
    if agente == 'humos':
        return {
            'humos_eval_riesgo': agente_detalle.get('evaluacion_riesgo', ''),
            'humos_exposicion': agente_detalle.get('exposicion', ''),
        }
    return {}


def _datos_laborales_actuales_de_usuario(usuario):
    if not usuario:
        return {}

    try:
        info_laboral = usuario.informacion_laboral
    except UserInformacionLaboral.DoesNotExist:
        info_laboral = None

    perfil = getattr(usuario, 'usuarioprofile', None)
    nombre = f"{(usuario.first_name or '').strip()} {(usuario.last_name or '').strip()}".strip() or usuario.username

    return {
        'username': usuario.username,
        'nombre': nombre,
        'cargo': info_laboral.cargo.valor if info_laboral and info_laboral.cargo else '',
        'faena_id': getattr(perfil, 'faena_id', None) if perfil else None,
        'faena_nombre': getattr(getattr(perfil, 'faena', None), 'faena', '') if perfil else '',
        'area': info_laboral.area.valor if info_laboral and info_laboral.area else '',
        'ges': info_laboral.ges.valor if info_laboral and info_laboral.ges else '',
    }


def _crear_ficha_vigilancia_desde_trabajador(
    vigilancia,
    trabajador,
    creado_por='',
    status=True,
    prefill_agente=None,
):
    # Crea una ficha con snapshot de datos laborales actuales del trabajador.
    cuantitativa = vigilancia.cuantitativa
    usuario = trabajador.get('user')
    if not usuario:
        return None

    datos_actuales = _datos_laborales_actuales_de_usuario(usuario)

    if prefill_agente is None:
        prefill_agente = _prefill_agente_desde_cuantitativa(cuantitativa)

    rut_trabajador = trabajador.get('username') or datos_actuales.get('username') or usuario.username
    nombre_completo = trabajador.get('nombre') or datos_actuales.get('nombre') or usuario.username
    cargo = trabajador.get('cargo') or datos_actuales.get('cargo') or ''
    faena_id = trabajador.get('faena_id') or datos_actuales.get('faena_id') or None
    area = trabajador.get('area') or datos_actuales.get('area') or ''
    ges = trabajador.get('ges') or datos_actuales.get('ges') or ''

    return VigilanciaMedica.objects.create(
        user=usuario,
        cuantitativa=cuantitativa,
        vigilancia=vigilancia,
        rut_trabajador=rut_trabajador,
        nombre_completo=nombre_completo,
        area=area,
        cargo=cargo,
        faena_id=faena_id,
        tipo_contrato='',
        ges=ges,
        silice_eval_riesgo=prefill_agente.get('silice_eval_riesgo', ''),
        silice_nivel_riesgo=prefill_agente.get('silice_nivel_riesgo', ''),
        silice_grado_expo=prefill_agente.get('silice_grado_expo', ''),
        ruido_eval_riesgo=prefill_agente.get('ruido_eval_riesgo', ''),
        ruido_nivel_seguimiento=prefill_agente.get('ruido_nivel_seguimiento', ''),
        ruido_grado_expo=prefill_agente.get('ruido_grado_expo', ''),
        hipo_exposicion=prefill_agente.get('hipo_exposicion', ''),
        vibra_eval_riesgo=prefill_agente.get('vibra_eval_riesgo', ''),
        vibra_exposicion=prefill_agente.get('vibra_exposicion', ''),
        rad_exposicion=prefill_agente.get('rad_exposicion', ''),
        humos_eval_riesgo=prefill_agente.get('humos_eval_riesgo', ''),
        humos_exposicion=prefill_agente.get('humos_exposicion', ''),
        creado_por=creado_por or '',
        status=status,
    )


def _inicializar_snapshot_vigilancia(vigilancia, user, creado_por=''):
    # Congela una sola vez el set automatico inicial de trabajadores de la vigilancia.
    if getattr(vigilancia, 'snapshot_inicializado', False):
        return 0

    cuantitativa = vigilancia.cuantitativa
    if not cuantitativa:
        vigilancia.snapshot_inicializado = True
        vigilancia.save(update_fields=['snapshot_inicializado'])
        return 0

    catalogo = _catalogo_trabajadores_para_cuantitativa(user, cuantitativa)
    user_ids_catalogo = [
        item['user'].id
        for item in catalogo
        if item.get('user') and item['user'].id
    ]
    resumen_fichas = _resumen_fichas_por_usuario_en_vigilancia(vigilancia, user_ids=user_ids_catalogo)
    prefill_agente = _prefill_agente_desde_cuantitativa(cuantitativa)
    creador = creado_por or (user.username if user and user.is_authenticated else '') or (vigilancia.creador or '')
    total_creadas = 0

    with transaction.atomic():
        for trabajador in catalogo:
            usuario = trabajador.get('user')
            if not usuario or not usuario.id:
                continue

            resumen_usuario = resumen_fichas.get(usuario.id)
            if _trabajador_oculto_segun_resumen(resumen_usuario):
                continue
            if _trabajador_eliminado_segun_resumen(resumen_usuario):
                continue
            if resumen_usuario and resumen_usuario.get('ultima'):
                continue
            if not _trabajador_coincide_filtro_automatico(vigilancia, trabajador):
                continue

            _crear_ficha_vigilancia_desde_trabajador(
                vigilancia=vigilancia,
                trabajador=trabajador,
                creado_por=creador,
                status=True,
                prefill_agente=prefill_agente,
            )
            total_creadas += 1

        vigilancia.snapshot_inicializado = True
        vigilancia.save(update_fields=['snapshot_inicializado'])

    return total_creadas


def _vigilancias_con_total_trabajadores_para_usuario(user):
    vigilancias = list(_vigilancias_visibles_para_usuario(user).order_by('-fechacreacion', '-id'))

    # La columna "Trabajadores" debe reflejar la cantidad visible en la Lista de Trabajadores
    # (automaticos por area+ges + agregados manualmente, excluyendo removidos/egresados).
    for vigilancia in vigilancias:
        if not getattr(vigilancia, 'snapshot_inicializado', False):
            _inicializar_snapshot_vigilancia(
                vigilancia=vigilancia,
                user=user,
                creado_por=user.username if user and user.is_authenticated else '',
            )

        if getattr(vigilancia, 'snapshot_inicializado', False):
            vigilancia.total_trabajadores = VigilanciaMedica.objects.filter(
                vigilancia=vigilancia,
                status=True,
            ).count()
            continue

        catalogo = _catalogo_trabajadores_para_cuantitativa(user, vigilancia.cuantitativa)
        user_ids_catalogo = {item['user'].id for item in catalogo}

        # Debe usar el mismo criterio de la grilla para que el numero coincida.
        resumen_fichas = _resumen_fichas_por_usuario_en_vigilancia(vigilancia)
        total_visibles = 0
        for trabajador in catalogo:
            user_id = trabajador['user'].id
            resumen_usuario = resumen_fichas.get(user_id)
            ficha = (resumen_usuario or {}).get('activa')
            coincide_automatico = _trabajador_coincide_filtro_automatico(vigilancia, trabajador)
            eliminado = _trabajador_eliminado_segun_resumen(resumen_usuario)

            if _trabajador_oculto_segun_resumen(resumen_usuario) and not ficha:
                continue
            if eliminado and not ficha:
                continue
            if ficha or coincide_automatico:
                total_visibles += 1

        # Suma fichas activas historicas aunque el trabajador ya no coincida con el catalogo actual.
        fichas_activas_fuera_catalogo = VigilanciaMedica.objects.filter(
            vigilancia=vigilancia,
            status=True,
        )
        if vigilancia.faena_id:
            fichas_activas_fuera_catalogo = fichas_activas_fuera_catalogo.filter(faena_id=vigilancia.faena_id)
        if user_ids_catalogo:
            fichas_activas_fuera_catalogo = fichas_activas_fuera_catalogo.exclude(
                user_id__in=user_ids_catalogo
            )
        total_visibles += fichas_activas_fuera_catalogo.count()

        vigilancia.total_trabajadores = total_visibles

    return vigilancias


def _vigilancias_con_total_egresados_para_usuario(user):
    # Total para pantalla de egreso: historial fijo por vigilancia (sin depender del estado actual del trabajador).
    vigilancias = list(_vigilancias_visibles_para_usuario(user).order_by('-fechacreacion', '-id'))

    for vigilancia in vigilancias:
        vigilancia.total_trabajadores = Egreso.objects.filter(vigilancia=vigilancia).count()

    return vigilancias


def _validar_ficha_para_egreso(ficha):
    # Para egreso solo exigimos fechas de seguimiento del agente activo.
    # Los datos laborales (area/faena/ges/cargo) se mantienen como snapshot historico,
    # pero no bloquean el egreso.
    campos_obligatorios = []

    agente = (ficha.cuantitativa.agente if ficha.cuantitativa else '') or ''
    if agente == 'silice':
        campos_obligatorios.extend([
            ('Fecha Radiografia', ficha.silice_fecha_radio),
            ('Fecha Vencimiento', ficha.silice_fecha_vence),
        ])
    elif agente == 'ruido':
        campos_obligatorios.extend([
            ('Fecha Audiometria', ficha.ruido_fecha_audio),
            ('Fecha Vencimiento', ficha.ruido_fecha_vence),
        ])
    elif agente == 'hipobaria':
        campos_obligatorios.extend([
            ('Fecha Hemoglobina', ficha.hipo_fecha_hemo),
            ('Fecha Vencimiento', ficha.hipo_fecha_vence),
        ])
    elif agente == 'vibraciones':
        campos_obligatorios.extend([
            ('Fecha Vencimiento', ficha.vibra_fecha_vence),
        ])
    elif agente == 'radiaciones':
        campos_obligatorios.extend([
            ('Fecha Vencimiento', ficha.rad_fecha_vence),
        ])
    elif agente == 'humos':
        campos_obligatorios.extend([
            ('Fecha Vencimiento', ficha.humos_fecha_vence),
        ])

    faltantes = [nombre for nombre, valor in campos_obligatorios if valor in (None, '')]
    return len(faltantes) == 0, faltantes


@login_required
@prevencion_riesgo_required
def egreso(request):
    vigilancias = _vigilancias_con_total_egresados_para_usuario(request.user)
    trabajadores_visibles = _trabajadores_visibles_para_usuario(request.user)
    cache_contratos = {}

    for vigilancia in vigilancias:
        clave = (
            vigilancia.faena_id,
            _normalizar_texto_filtro(vigilancia.area),
            _normalizar_texto_filtro(vigilancia.ges),
        )
        if clave not in cache_contratos:
            cache_contratos[clave] = _contratos_por_filtros_para_usuario(
                user=request.user,
                faena_id=vigilancia.faena_id,
                area_valor=vigilancia.area,
                ges_valor=vigilancia.ges,
                trabajadores_qs=trabajadores_visibles,
            )
        contratos = cache_contratos[clave]
        vigilancia.contrato_filtro_texto = ' | '.join(contratos) if contratos else '-'

    context = {
        'vigilancias': vigilancias,
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'egreso',
    }
    return render(request, 'pages/prevencion/egreso.html', context)


@login_required
@prevencion_riesgo_required
def egreso_detalle(request, vigilancia_id):
    vigilancia = get_object_or_404(
        _vigilancias_visibles_para_usuario(request.user),
        id=vigilancia_id
    )
    egresos = (
        Egreso.objects.filter(vigilancia=vigilancia)
        .select_related('faena')
        .order_by('-fecha_incidente', '-fecha_egreso', '-id')
    )

    context = {
        'vigilancia': vigilancia,
        'egresos': egresos,
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'egreso',
    }
    return render(request, 'pages/prevencion/egreso_detalle.html', context)


@login_required
@prevencion_riesgo_required
def reintegrar_vigilancia_trabajador(request, vigilancia_id, egreso_id):
    # Reintegra trabajador desde egreso a la lista activa de vigilancia.
    if request.method != 'POST':
        return redirect('egreso_detalle', vigilancia_id=vigilancia_id)

    vigilancia = get_object_or_404(
        _vigilancias_visibles_para_usuario(request.user),
        id=vigilancia_id
    )
    egreso = get_object_or_404(
        Egreso.objects.select_related('ficha_origen'),
        id=egreso_id,
        vigilancia=vigilancia,
    )

    ficha = egreso.ficha_origen
    if ficha is None and egreso.user_id:
        ficha = (
            VigilanciaMedica.objects.filter(vigilancia=vigilancia, user_id=egreso.user_id)
            .order_by('-id')
            .first()
        )

    if ficha is None:
        messages.error(request, "No se encontró la ficha de vigilancia para reintegrar.")
        return redirect('egreso_detalle', vigilancia_id=vigilancia.id)

    with transaction.atomic():
        ficha.egresado = False
        ficha.fecha_egreso = None
        ficha.status = True
        ficha.fecha_retiro = None
        ficha.save(update_fields=['egresado', 'fecha_egreso', 'status', 'fecha_retiro'])
        if egreso.ficha_origen_id != ficha.id:
            egreso.ficha_origen = ficha
            egreso.save(update_fields=['ficha_origen'])

    messages.success(request, "Trabajador reintegrado a Vigilancia correctamente (historial conservado).")
    user_name = request.user.get_full_name() or request.user.username
    notify_group('prevencion_general', f'Trabajador reintegrado: {ficha.nombre_completo}', f'Vigilancia Médica: {ficha.nombre_completo}\nReintegrado por: {user_name}')
    return redirect('egreso_detalle', vigilancia_id=vigilancia.id)


@login_required
@prevencion_riesgo_required
def new_vigilancia(request):
    cuantitativas_visibles = _cuantitativas_visibles_para_usuario(request.user).order_by('-fecha', '-id')

    if request.method == 'POST':
        cuantitativa_id = request.POST.get('cuantitativa_id')
        codigo = (request.POST.get('codigo') or '').strip()
        nombre = (request.POST.get('nombre') or '').strip()
        cuantitativa = cuantitativas_visibles.filter(id=cuantitativa_id).first()
        if not cuantitativa:
            messages.error(request, "Debes seleccionar una evaluación cuantitativa válida.")
            return redirect('new_vigilancia')
        if not codigo or not nombre:
            messages.error(request, "Debes ingresar código y nombre para la nueva vigilancia.")
            return redirect('new_vigilancia')
        if Vigilancia.objects.filter(codigo__iexact=codigo).exists():
            messages.error(request, "El código de vigilancia ya existe. Debe ser único.")
            return redirect('new_vigilancia')

        nombre_usuario = f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username

        try:
            with transaction.atomic():
                vigilancia = Vigilancia.objects.create(
                    cuantitativa=cuantitativa,
                    faena_id=cuantitativa.faena_id,
                    area=cuantitativa.area or '',
                    ges=cuantitativa.ges or '',
                    codigo=codigo,
                    nombre=nombre,
                    creador=nombre_usuario,
                    status=True,
                )
                _inicializar_snapshot_vigilancia(
                    vigilancia=vigilancia,
                    user=request.user,
                    creado_por=request.user.username,
                )
        except IntegrityError:
            messages.error(request, "El código de vigilancia ya existe. Debe ser único.")
            return redirect('new_vigilancia')

        messages.success(request, "Vigilancia creada correctamente.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    context = {
        'cuantitativas': cuantitativas_visibles,
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/new_vigilancia.html', context)


@login_required
@prevencion_riesgo_required
def vigilancia_trabajadores(request, vigilancia_id):
    # Vista principal por vigilancia: mezcla catalogo base + estado de ficha por trabajador.
    vigilancia = get_object_or_404(_vigilancias_visibles_para_usuario(request.user), id=vigilancia_id)
    _inicializar_snapshot_vigilancia(
        vigilancia=vigilancia,
        user=request.user,
        creado_por=request.user.username,
    )
    cuantitativa = vigilancia.cuantitativa
    trabajadores_catalogo = _catalogo_trabajadores_para_cuantitativa(request.user, cuantitativa)
    resumen_fichas = _resumen_fichas_por_usuario_en_vigilancia(vigilancia)
    snapshot_fijo = bool(getattr(vigilancia, 'snapshot_inicializado', False))

    filas = []
    trabajadores_dropdown = []
    user_ids_en_tabla = set()
    dropdown_user_ids = set()

    def _agregar_a_dropdown(trabajador_data):
        usuario_dropdown = trabajador_data.get('user')
        if not usuario_dropdown or not usuario_dropdown.id:
            return
        if usuario_dropdown.id in dropdown_user_ids or usuario_dropdown.id in user_ids_en_tabla:
            return
        trabajadores_dropdown.append(trabajador_data)
        dropdown_user_ids.add(usuario_dropdown.id)

    for trabajador in trabajadores_catalogo:
        user = trabajador['user']
        resumen_usuario = resumen_fichas.get(user.id)
        ficha = (resumen_usuario or {}).get('activa')
        coincide_automatico = _trabajador_coincide_filtro_automatico(vigilancia, trabajador)
        eliminado = _trabajador_eliminado_segun_resumen(resumen_usuario)
        puede_egresar, faltantes_egreso = (False, [])
        if ficha:
            puede_egresar, faltantes_egreso = _validar_ficha_para_egreso(ficha)
        if _trabajador_oculto_segun_resumen(resumen_usuario) and not ficha:
            # Egresados no deben reaparecer ni en tabla ni en "Agregar Trabajador".
            continue

        if eliminado and not ficha:
            _agregar_a_dropdown(trabajador)
            continue

        # Snapshot fijo: solo muestra fichas guardadas.
        # Snapshot no inicializado: mantiene comportamiento automatico por area+ges.
        if ficha or (not snapshot_fijo and coincide_automatico):
            fila = {
                **trabajador,
                'ficha': ficha,
                'puede_egresar': puede_egresar,
                'faltantes_egreso': faltantes_egreso,
            }
            if ficha:
                # Si ya existe ficha, la tabla debe reflejar el snapshot guardado en esa ficha.
                fila.update({
                    'username': ficha.rut_trabajador or trabajador.get('username'),
                    'nombre': ficha.nombre_completo or trabajador.get('nombre'),
                    'cargo': ficha.cargo or trabajador.get('cargo'),
                    'faena_id': ficha.faena_id or trabajador.get('faena_id'),
                    'faena_nombre': (
                        ficha.faena.faena
                        if getattr(ficha, 'faena', None)
                        else trabajador.get('faena_nombre')
                    ),
                    'area': ficha.area or trabajador.get('area'),
                    'ges': ficha.ges or trabajador.get('ges'),
                })
            filas.append(fila)
            user_ids_en_tabla.add(user.id)
        else:
            _agregar_a_dropdown(trabajador)

    # Mantiene visibles fichas activas aunque el trabajador ya no coincida con el catalogo actual.
    fichas_activas = (
        VigilanciaMedica.objects.filter(vigilancia=vigilancia, status=True)
        .select_related('faena', 'user')
        .order_by('nombre_completo', 'rut_trabajador', 'id')
    )
    for ficha in fichas_activas:
        if ficha.user_id and ficha.user_id in user_ids_en_tabla:
            continue
        puede_egresar, faltantes_egreso = _validar_ficha_para_egreso(ficha)
        nombre_usuario = '-'
        if ficha.user:
            nombre_usuario = ficha.user.get_full_name().strip() or ficha.user.username
        filas.append({
            'user': ficha.user,
            'username': ficha.rut_trabajador or (ficha.user.username if ficha.user else '-'),
            'nombre': ficha.nombre_completo or nombre_usuario,
            'cargo': ficha.cargo or '-',
            'faena_id': ficha.faena_id or '',
            'faena_nombre': ficha.faena.faena if ficha.faena else '-',
            'area': ficha.area or '-',
            'ges': ficha.ges or '-',
            'ficha': ficha,
            'puede_egresar': puede_egresar,
            'faltantes_egreso': faltantes_egreso,
        })
        if ficha.user_id:
            user_ids_en_tabla.add(ficha.user_id)

    # Si un trabajador eliminado ya no esta en el catalogo actual, igual debe quedar disponible para re-agregar.
    for user_id, resumen_usuario in resumen_fichas.items():
        if user_id in user_ids_en_tabla or user_id in dropdown_user_ids:
            continue
        if not _trabajador_eliminado_segun_resumen(resumen_usuario):
            continue

        ficha_ultima = resumen_usuario.get('ultima')
        if not ficha_ultima or not ficha_ultima.user_id:
            continue

        nombre_usuario = ficha_ultima.user.get_full_name().strip() if ficha_ultima.user else ''
        _agregar_a_dropdown({
            'user': ficha_ultima.user,
            'username': ficha_ultima.rut_trabajador or (ficha_ultima.user.username if ficha_ultima.user else '-'),
            'nombre': ficha_ultima.nombre_completo or nombre_usuario or (ficha_ultima.user.username if ficha_ultima.user else '-'),
            'cargo': ficha_ultima.cargo or '',
            'faena_id': ficha_ultima.faena_id or '',
            'faena_nombre': ficha_ultima.faena.faena if ficha_ultima.faena else '',
            'area': ficha_ultima.area or '',
            'ges': ficha_ultima.ges or '',
        })

    # Regla explícita: en "Agregar Trabajador" deben aparecer todos los usuarios
    # (excepto id=1), aunque no pertenezcan a la faena de la vigilancia.
    for trabajador in _catalogo_todos_los_trabajadores_para_agregar():
        usuario_catalogo = trabajador.get('user')
        if not usuario_catalogo or not usuario_catalogo.id:
            continue
        _agregar_a_dropdown(trabajador)

    context = {
        'vigilancia': vigilancia,
        'cuantitativa': cuantitativa,
        'filas_trabajadores': filas,
        'trabajadores_disponibles': trabajadores_dropdown,
        'total_fichas_activas': sum(1 for fila in filas if fila['ficha'] and fila['ficha'].status),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/vigilancia_trabajadores.html', context)


@login_required
@prevencion_riesgo_required
def add_vigilancia_trabajadores(request, vigilancia_id):
    # Agrega un trabajador seleccionado a la lista sin redirigir a la ficha.
    if request.method != 'POST':
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia_id)

    vigilancia = get_object_or_404(_vigilancias_visibles_para_usuario(request.user), id=vigilancia_id)
    cuantitativa = vigilancia.cuantitativa

    user_id = str(request.POST.get('user_id') or '').strip()
    if not user_id:
        legacy_ids = [str(item).strip() for item in request.POST.getlist('user_ids') if str(item).strip()]
        user_id = legacy_ids[0] if legacy_ids else ''

    if not user_id:
        messages.info(request, "Selecciona un trabajador para agregar.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    usuario = User.objects.filter(pk=user_id).first()
    if not usuario:
        messages.info(request, "El trabajador seleccionado no existe.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    ya_activo = VigilanciaMedica.objects.filter(
        vigilancia=vigilancia,
        user=usuario,
        status=True,
    ).exists()
    if ya_activo:
        messages.info(request, "El trabajador ya está agregado en la lista.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    ya_egresado = Egreso.objects.filter(vigilancia=vigilancia, user=usuario).exists()
    if ya_egresado:
        messages.info(request, "El trabajador ya fue egresado y no puede volver a agregarse.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    ficha_inactiva = (
        VigilanciaMedica.objects.filter(
            vigilancia=vigilancia,
            user=usuario,
            status=False,
            egresado=False,
        )
        .order_by('-id')
        .first()
    )
    if ficha_inactiva:
        ficha_inactiva.status = True
        ficha_inactiva.fecha_retiro = None
        ficha_inactiva.save(update_fields=['status', 'fecha_retiro'])
        messages.success(request, "Trabajador agregado a la lista correctamente.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    trabajadores_catalogo = _catalogo_todos_los_trabajadores_para_agregar()
    trabajadores_por_id = {
        str(item['user'].id): item
        for item in trabajadores_catalogo
        if item.get('user') and item['user'].id
    }
    trabajador = trabajadores_por_id.get(user_id)
    if not trabajador:
        messages.info(request, "El trabajador seleccionado no está disponible para agregar.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    _crear_ficha_vigilancia_desde_trabajador(
        vigilancia=vigilancia,
        trabajador=trabajador,
        creado_por=request.user.username,
        status=True,
        prefill_agente=_prefill_agente_desde_cuantitativa(cuantitativa),
    )

    messages.success(request, "Trabajador agregado a la lista correctamente.")

    return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)


@login_required
def egresar_vigilancia_trabajador(request, vigilancia_id):
    # Egreso: valida ficha, crea snapshot historico y marca ficha activa como egresada.
    if request.method != 'POST':
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia_id)

    vigilancia = get_object_or_404(_vigilancias_visibles_para_usuario(request.user), id=vigilancia_id)
    ficha_id = request.POST.get('ficha_id')
    if not ficha_id:
        messages.error(request, "No se recibió la ficha para egresar.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    ficha = get_object_or_404(
        VigilanciaMedica.objects.select_related('cuantitativa', 'vigilancia'),
        id=ficha_id,
        vigilancia=vigilancia,
    )

    if ficha.egresado:
        messages.info(request, "La ficha ya estaba egresada.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)
    if not ficha.status:
        messages.info(request, "La ficha seleccionada no está activa.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    ficha_valida, faltantes = _validar_ficha_para_egreso(ficha)
    if not ficha_valida:
        if faltantes:
            resumen = ', '.join(faltantes[:3])
            sufijo = '...' if len(faltantes) > 3 else ''
            messages.error(
                request,
                f"No se puede egresar: faltan {len(faltantes)} campos obligatorios ({resumen}{sufijo})."
            )
        else:
            messages.error(request, "No se puede egresar: la ficha no está completa.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    with transaction.atomic():
        fecha_egreso = timezone.now()
        _crear_egreso_desde_ficha(ficha, fecha_egreso=fecha_egreso)

        update_fields = ['egresado', 'fecha_egreso', 'status']
        ficha.egresado = True
        ficha.fecha_egreso = fecha_egreso
        ficha.status = False
        if not ficha.fecha_retiro:
            ficha.fecha_retiro = timezone.localdate()
            update_fields.append('fecha_retiro')
        ficha.save(update_fields=update_fields)

    messages.success(request, "Trabajador egresado correctamente.")
    user_name = request.user.get_full_name() or request.user.username
    notify_group('prevencion_general', f'Trabajador egresado: {ficha.nombre_completo}', f'Vigilancia Médica: {ficha.nombre_completo}\nEgresado por: {user_name}')
    return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)


@login_required
def remove_vigilancia_trabajador(request, vigilancia_id):
    # Quita al trabajador de la lista actual, sin borrar el usuario del sistema.
    if request.method != 'POST':
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia_id)

    vigilancia = get_object_or_404(_vigilancias_visibles_para_usuario(request.user), id=vigilancia_id)
    cuantitativa = vigilancia.cuantitativa
    user_id = str(request.POST.get('user_id') or '').strip()
    if not user_id:
        messages.error(request, "No se recibió el trabajador a remover.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    trabajadores_catalogo = _catalogo_trabajadores_para_cuantitativa(request.user, cuantitativa)
    trabajador = next((t for t in trabajadores_catalogo if str(t['user'].id) == user_id), None)
    if not trabajador:
        messages.error(request, "El trabajador no corresponde al filtro vigente.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    usuario = trabajador['user']
    ficha_activa = (
        VigilanciaMedica.objects.filter(vigilancia=vigilancia, user=usuario, status=True)
        .order_by('-id')
        .first()
    )
    if ficha_activa:
        ficha_activa.status = False
        ficha_activa.save(update_fields=['status'])
        messages.success(request, "Trabajador quitado de la lista.")
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)

    existe_oculto = VigilanciaMedica.objects.filter(
        vigilancia=vigilancia,
        user=usuario,
        status=False
    ).exists()
    if not existe_oculto:
        _crear_ficha_vigilancia_desde_trabajador(
            vigilancia=vigilancia,
            trabajador=trabajador,
            creado_por=request.user.username,
            status=False,
            prefill_agente=_prefill_agente_desde_cuantitativa(cuantitativa),
        )

    messages.success(request, "Trabajador quitado de la lista.")
    return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.id)


@login_required
def new_vigilancia_trabajador(request, vigilancia_id):
    # Alta de ficha por trabajador dentro de una vigilancia especifica.
    vigilancia_obj = get_object_or_404(_vigilancias_visibles_para_usuario(request.user), id=vigilancia_id)
    cuantitativa = vigilancia_obj.cuantitativa
    trabajadores_catalogo = _catalogo_trabajadores_para_cuantitativa(request.user, cuantitativa)
    trabajadores_ids_permitidos = {item['user'].id for item in trabajadores_catalogo}
    rut_preselect = (request.GET.get('rut') or '').strip()
    redirect_base = reverse('new_vigilancia_trabajador', args=[vigilancia_obj.id])

    if request.method == 'POST':
        try:
            archivos_validos, msg_archivos = _validar_archivos_subidos(request)
            if not archivos_validos:
                messages.error(request, msg_archivos)
                return redirect('new_vigilancia_trabajador', vigilancia_id=vigilancia_obj.id)
            archivos_documento = request.FILES.getlist('documento')
            if len(archivos_documento) > 1:
                messages.error(request, "Solo puedes subir 1 archivo en la ficha de vigilancia.")
                return redirect('new_vigilancia_trabajador', vigilancia_id=vigilancia_obj.id)

            faena_id_raw = str(cuantitativa.faena_id or request.POST.get('faena') or '').strip()
            try:
                faena_id = int(faena_id_raw) if faena_id_raw else None
            except (TypeError, ValueError):
                faena_id = None
            if not faena_id:
                messages.error(request, "Debes seleccionar una faena válida.")
                return redirect('new_vigilancia_trabajador', vigilancia_id=vigilancia_obj.id)
            if not _faena_permitida_para_usuario(request.user, faena_id):
                messages.error(request, "No tienes permisos para registrar una ficha en esa faena.")
                return redirect('new_vigilancia_trabajador', vigilancia_id=vigilancia_obj.id)

            rut_raw = (request.POST.get('rut') or '').strip()
            if not rut_raw:
                messages.error(request, "Debes seleccionar un trabajador.")
                return redirect('new_vigilancia_trabajador', vigilancia_id=vigilancia_obj.id)
            redirect_con_rut = f"{redirect_base}?rut={rut_raw}"

            rut_limpio = rut_raw.replace('.', '').replace('-', '').upper()
            rut_guion = f"{rut_limpio[:-1]}-{rut_limpio[-1]}" if len(rut_limpio) > 1 else rut_limpio
            usuario_asociado = User.objects.filter(
                Q(username__iexact=rut_limpio) | Q(username__iexact=rut_guion) | Q(username__iexact=rut_raw)
            ).first()
            if not usuario_asociado:
                messages.error(request, "No se encontró el trabajador seleccionado.")
                return redirect('new_vigilancia_trabajador', vigilancia_id=vigilancia_obj.id)

            if usuario_asociado.id not in trabajadores_ids_permitidos:
                messages.error(request, "El trabajador no corresponde a la faena de la cuantitativa.")
                return redirect('new_vigilancia_trabajador', vigilancia_id=vigilancia_obj.id)

            existe_activa = VigilanciaMedica.objects.filter(
                vigilancia=vigilancia_obj,
                user=usuario_asociado,
                status=True
            ).exists()
            if existe_activa:
                messages.error(request, "Este trabajador ya tiene una ficha activa para esta cuantitativa.")
                return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia_obj.id)

            # Regla de negocio: si ya fue egresado para esta cuantitativa no se puede volver a agregar.
            ya_egresado = Egreso.objects.filter(
                vigilancia=vigilancia_obj,
                user=usuario_asociado,
            ).exists()
            if ya_egresado:
                messages.error(request, "Este trabajador ya fue egresado y solo debe aparecer en la seccion Egreso.")
                return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia_obj.id)

            nombre_completo = (request.POST.get('nombre_completo') or '').strip()
            if not nombre_completo:
                messages.error(request, "Debes esperar la carga del trabajador antes de guardar.")
                return redirect(redirect_con_rut)

            area_valor = (request.POST.get('area') or '').strip() or (cuantitativa.area or '')
            ges_valor = (request.POST.get('ges') or '').strip() or (cuantitativa.ges or '')
            if not area_valor or not ges_valor:
                messages.error(request, "No se puede guardar: faltan Area o GES en la ficha.")
                return redirect(redirect_con_rut)
            prefill_agente = _prefill_agente_desde_cuantitativa(cuantitativa)

            def _valor_post_o_default(nombre_campo, default=''):
                valor = (request.POST.get(nombre_campo) or '').strip()
                return valor or (default or '')

            genero_id_raw = str(request.POST.get('genero') or '').strip()
            try:
                genero_id_int = int(genero_id_raw) if genero_id_raw else None
            except (TypeError, ValueError):
                genero_id_int = None
            genero_obj = Genero.objects.filter(pk=genero_id_int).first() if genero_id_int else None
            genero_id = genero_obj.id if genero_obj else None
            genero_txt = genero_obj.genero if genero_obj else None
            fecha_nacimiento_raw = request.POST.get('fecha_nacimiento') or None
            edad_calculada = _calcular_edad_desde_fecha(fecha_nacimiento_raw)
            vibra_fecha_vence = request.POST.get('vibra_fecha_vence') or None
            rad_fecha_vence = request.POST.get('rad_fecha_vence') or None
            humos_fecha_vence = request.POST.get('humos_fecha_vence') or None

            ficha_vigilancia = VigilanciaMedica(
                user=usuario_asociado,
                cuantitativa=cuantitativa,
                vigilancia=vigilancia_obj,
                rut_trabajador=rut_raw,
                nombre_completo=nombre_completo,
                area=area_valor,
                cargo=request.POST.get('cargo'),
                faena_id=faena_id,
                genero_texto=genero_txt,
                fecha_nacimiento=fecha_nacimiento_raw,
                edad=edad_calculada if edad_calculada is not None else 0,
                fecha_ingreso=request.POST.get('fecha_ingreso') or None,
                antiguedad_anos=request.POST.get('antiguedad_anos') or 0,
                antiguedad_meses=request.POST.get('antiguedad_meses') or 0,
                antiguedad_dias=request.POST.get('antiguedad_dias') or 0,
                tipo_contrato=request.POST.get('tipo_contrato'),
                fecha_retiro=request.POST.get('fecha_retiro') or None,
                ges=ges_valor,
                silice_eval_riesgo=_valor_post_o_default('silice_eval_riesgo', prefill_agente.get('silice_eval_riesgo')),
                silice_nivel_riesgo=_valor_post_o_default('silice_nivel_riesgo', prefill_agente.get('silice_nivel_riesgo')),
                silice_grado_expo=_valor_post_o_default('silice_grado_expo', prefill_agente.get('silice_grado_expo')),
                silice_fecha_radio=request.POST.get('silice_fecha_radio') or None,
                silice_fecha_vence=request.POST.get('silice_fecha_vence') or None,
                silice_vigencia=request.POST.get('silice_vigencia'),
                silice_obs=request.POST.get('silice_obs'),
                ruido_eval_riesgo=_valor_post_o_default('ruido_eval_riesgo', prefill_agente.get('ruido_eval_riesgo')),
                ruido_nivel_seguimiento=_valor_post_o_default('ruido_nivel_seguimiento', prefill_agente.get('ruido_nivel_seguimiento')),
                ruido_grado_expo=_valor_post_o_default('ruido_grado_expo', prefill_agente.get('ruido_grado_expo')),
                ruido_fecha_audio=request.POST.get('ruido_fecha_audio') or None,
                ruido_fecha_vence=request.POST.get('ruido_fecha_vence') or None,
                ruido_vigencia=request.POST.get('ruido_vigencia'),
                ruido_obs=request.POST.get('ruido_obs'),
                hipo_exposicion=_valor_post_o_default('hipo_exposicion', prefill_agente.get('hipo_exposicion')),
                hipo_fecha_hemo=request.POST.get('hipo_fecha_hemo') or None,
                hipo_fecha_vence=request.POST.get('hipo_fecha_vence') or None,
                hipo_vigencia=request.POST.get('hipo_vigencia'),
                hipo_obs=request.POST.get('hipo_obs'),
                vibra_eval_riesgo=_valor_post_o_default('vibra_eval_riesgo', prefill_agente.get('vibra_eval_riesgo')),
                vibra_exposicion=_valor_post_o_default('vibra_exposicion', prefill_agente.get('vibra_exposicion')),
                vibra_fecha_vence=vibra_fecha_vence,
                vibra_vigencia=_calcular_vigencia_desde_fecha(vibra_fecha_vence),
                rad_exposicion=_valor_post_o_default('rad_exposicion', prefill_agente.get('rad_exposicion')),
                rad_fecha_vence=rad_fecha_vence,
                rad_vigencia=_calcular_vigencia_desde_fecha(rad_fecha_vence),
                humos_eval_riesgo=_valor_post_o_default('humos_eval_riesgo', prefill_agente.get('humos_eval_riesgo')),
                humos_exposicion=_valor_post_o_default('humos_exposicion', prefill_agente.get('humos_exposicion')),
                humos_fecha_vence=humos_fecha_vence,
                humos_vigencia=_calcular_vigencia_desde_fecha(humos_fecha_vence),
                fecha_incidente=request.POST.get('fecha_incidente') or None,
                creado_por=request.user.username
            )

            with transaction.atomic():
                ficha_vigilancia.save()
                _sincronizar_datos_personales_usuario(
                    usuario_asociado,
                    genero_id=genero_id,
                    fecha_nacimiento_raw=fecha_nacimiento_raw,
                    sobrescribir=True
                )
                if archivos_documento:
                    VigilanciaAdjunto.objects.create(vigilancia=ficha_vigilancia, archivo=archivos_documento[0])

            messages.success(request, "Ficha de vigilancia guardada correctamente.")
            return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia_obj.id)

        except Exception as e:
            messages.error(request, f"Error al guardar: {e}")

    context = {
        'faenas': _faenas_seleccionables_para_usuario(request.user),
        'vigilancia': vigilancia_obj,
        'cuantitativas': [cuantitativa],
        'trabajadores': trabajadores_catalogo,
        'rut_preselect': rut_preselect,
        'back_url': reverse('vigilancia_trabajadores', args=[vigilancia_obj.id]),
        'generos': Genero.objects.filter(status=True),
        **_catalogos_vigilancia_formulario(),
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/new_vigilancia_trabajador.html', context)


@login_required
@prevencion_riesgo_required
def edit_vigilancia(request, pk):
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.prefetch_related('adjuntos_vigilancia'),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)
    if vigilancia.egresado:
        return _redireccion_egreso_desde_ficha(
            request,
            vigilancia,
            mensaje="La ficha egresada es solo lectura y no se puede editar.",
        )

    datos_personales_sincronizados = _obtener_datos_personales_sincronizados(vigilancia)
    
    if request.method == 'POST':
        try:
            archivos_validos, msg_archivos = _validar_archivos_subidos(request)
            if not archivos_validos:
                messages.error(request, msg_archivos)
                return redirect('edit_vigilancia', pk=pk)

            total_adjuntos_valido, total_adjuntos_msg = _validar_total_adjuntos_vigilancia_en_edicion(vigilancia, request)
            if not total_adjuntos_valido:
                messages.error(request, total_adjuntos_msg)
                return redirect('edit_vigilancia', pk=pk)

            # 1. PERSONALES
            vigilancia.rut_trabajador = request.POST.get('rut')
            vigilancia.nombre_completo = request.POST.get('nombre_completo')
            vigilancia.area = request.POST.get('area')
            vigilancia.cargo = request.POST.get('cargo')
            nueva_faena_id_raw = str(request.POST.get('faena') or '').strip()
            try:
                nueva_faena_id = int(nueva_faena_id_raw) if nueva_faena_id_raw else None
            except (TypeError, ValueError):
                nueva_faena_id = None
            if not nueva_faena_id:
                messages.error(request, "Debes seleccionar una faena válida.")
                return redirect('edit_vigilancia', pk=pk)
            if not _faena_permitida_para_usuario(request.user, nueva_faena_id):
                messages.error(request, "No tienes permisos para asignar esa faena.")
                return redirect('edit_vigilancia', pk=pk)
            vigilancia.faena_id = nueva_faena_id
            
            # 2. Género: guardar texto legible desde el ID seleccionado.
            genero_id_raw = str(request.POST.get('genero') or '').strip()
            try:
                genero_id_int = int(genero_id_raw) if genero_id_raw else None
            except (TypeError, ValueError):
                genero_id_int = None
            genero_obj = Genero.objects.filter(pk=genero_id_int).first() if genero_id_int else None
            genero_id = genero_obj.id if genero_obj else None
            vigilancia.genero_texto = genero_obj.genero if genero_obj else None  # Aquí se guarda el texto del género.
            
            fecha_nacimiento_raw = request.POST.get('fecha_nacimiento') or None
            edad_calculada = _calcular_edad_desde_fecha(fecha_nacimiento_raw)
            vigilancia.fecha_nacimiento = fecha_nacimiento_raw
            vigilancia.edad = edad_calculada if edad_calculada is not None else 0
            
            vigilancia.fecha_ingreso = request.POST.get('fecha_ingreso') or None
            vigilancia.antiguedad_anos = request.POST.get('antiguedad_anos') or 0
            vigilancia.antiguedad_meses = request.POST.get('antiguedad_meses') or 0
            vigilancia.antiguedad_dias = request.POST.get('antiguedad_dias') or 0
            vigilancia.tipo_contrato = request.POST.get('tipo_contrato')
            vigilancia.fecha_retiro = request.POST.get('fecha_retiro') or None
            vigilancia.ges = request.POST.get('ges')

            # 2. Sílice
            vigilancia.silice_eval_riesgo = request.POST.get('silice_eval_riesgo')
            vigilancia.silice_nivel_riesgo = request.POST.get('silice_nivel_riesgo')
            vigilancia.silice_grado_expo = request.POST.get('silice_grado_expo')
            vigilancia.silice_fecha_radio = request.POST.get('silice_fecha_radio') or None
            vigilancia.silice_fecha_vence = request.POST.get('silice_fecha_vence') or None
            vigilancia.silice_vigencia = request.POST.get('silice_vigencia')
            vigilancia.silice_obs = request.POST.get('silice_obs')

            # 3. RUIDO
            vigilancia.ruido_eval_riesgo = request.POST.get('ruido_eval_riesgo')
            vigilancia.ruido_nivel_seguimiento = request.POST.get('ruido_nivel_seguimiento')
            vigilancia.ruido_grado_expo = request.POST.get('ruido_grado_expo')
            vigilancia.ruido_fecha_audio = request.POST.get('ruido_fecha_audio') or None
            vigilancia.ruido_fecha_vence = request.POST.get('ruido_fecha_vence') or None
            vigilancia.ruido_vigencia = request.POST.get('ruido_vigencia')
            vigilancia.ruido_obs = request.POST.get('ruido_obs')

            # 4. HIPOBARIA
            vigilancia.hipo_exposicion = request.POST.get('hipo_exposicion')
            vigilancia.hipo_fecha_hemo = request.POST.get('hipo_fecha_hemo') or None
            vigilancia.hipo_fecha_vence = request.POST.get('hipo_fecha_vence') or None
            vigilancia.hipo_vigencia = request.POST.get('hipo_vigencia')
            vigilancia.hipo_obs = request.POST.get('hipo_obs')

            # 5. OTROS
            vibra_fecha_vence = request.POST.get('vibra_fecha_vence') or None
            rad_fecha_vence = request.POST.get('rad_fecha_vence') or None
            humos_fecha_vence = request.POST.get('humos_fecha_vence') or None
            vigilancia.vibra_eval_riesgo = request.POST.get('vibra_eval_riesgo')
            vigilancia.vibra_exposicion = request.POST.get('vibra_exposicion')
            vigilancia.vibra_fecha_vence = vibra_fecha_vence
            vigilancia.vibra_vigencia = _calcular_vigencia_desde_fecha(vibra_fecha_vence)
            vigilancia.rad_exposicion = request.POST.get('rad_exposicion')
            vigilancia.rad_fecha_vence = rad_fecha_vence
            vigilancia.rad_vigencia = _calcular_vigencia_desde_fecha(rad_fecha_vence)
            vigilancia.humos_eval_riesgo = request.POST.get('humos_eval_riesgo')
            vigilancia.humos_exposicion = request.POST.get('humos_exposicion')
            vigilancia.humos_fecha_vence = humos_fecha_vence
            vigilancia.humos_vigencia = _calcular_vigencia_desde_fecha(humos_fecha_vence)

            # META (Usamos 'fecha_incidente' para que coincida con el HTML)
            vigilancia.fecha_incidente = request.POST.get('fecha_incidente') or None

            with transaction.atomic():
                vigilancia.save()
                _sincronizar_datos_personales_usuario(
                    vigilancia.user,
                    genero_id=genero_id,
                    fecha_nacimiento_raw=fecha_nacimiento_raw,
                    sobrescribir=True
                )
                archivos_documento = request.FILES.getlist('documento')
                if archivos_documento:
                    _reemplazar_archivo_vigilancia(vigilancia, archivos_documento[0])
            messages.success(request, "Vigilancia actualizada correctamente.")
            if vigilancia.vigilancia_id:
                return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.vigilancia_id)
            return redirect('manage_vigilancia')
            
        except Exception as e:
            messages.error(request, f"Error al actualizar: {e}")

    context = {
        'vigilancia': vigilancia,
        'archivos_adjuntos': _serializar_adjuntos_vigilancia(vigilancia),
        'genero_seleccionado_id': datos_personales_sincronizados['genero_id'],
        'fecha_nacimiento_mostrada': datos_personales_sincronizados['fecha_nacimiento'],
        'faenas': _faenas_seleccionables_para_usuario(request.user),
        'generos': Genero.objects.filter(status=True),
        **_catalogos_vigilancia_formulario(),
        'back_url': (
            reverse('vigilancia_trabajadores', args=[vigilancia.vigilancia_id])
            if vigilancia.vigilancia_id else reverse('manage_vigilancia')
        ),
        'mostrar_filtro_faena': _usuario_puede_ver_todas_faenas(request.user),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/edit_vigilancia.html', context)

@login_required
@prevencion_riesgo_required
def view_vigilancia(request, pk):
    # "Ver" operativo: permite actualizar solo datos de seguimiento del agente en ficha activa.
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.select_related('cuantitativa', 'faena').prefetch_related('adjuntos_vigilancia'),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)
    if vigilancia.egresado:
        return _redireccion_egreso_desde_ficha(request, vigilancia)
    agente_actual = (vigilancia.cuantitativa.agente if vigilancia.cuantitativa else '')

    if request.method == 'POST':
        try:
            archivos_validos, msg_archivos = _validar_archivos_subidos(request)
            if not archivos_validos:
                messages.error(request, msg_archivos)
                return redirect('view_vigilancia', pk=pk)

            total_adjuntos_valido, total_adjuntos_msg = _validar_total_adjuntos_vigilancia_en_edicion(vigilancia, request)
            if not total_adjuntos_valido:
                messages.error(request, total_adjuntos_msg)
                return redirect('view_vigilancia', pk=pk)

            update_fields = []
            fecha_incidente_raw = request.POST.get('fecha_incidente')
            if fecha_incidente_raw is not None:
                vigilancia.fecha_incidente = fecha_incidente_raw or None
                update_fields.append('fecha_incidente')

            if agente_actual == 'silice':
                silice_fecha_radio = request.POST.get('silice_fecha_radio') or None
                silice_fecha_vence = request.POST.get('silice_fecha_vence') or None
                vigilancia.silice_fecha_radio = silice_fecha_radio
                vigilancia.silice_fecha_vence = silice_fecha_vence
                vigilancia.silice_vigencia = _calcular_vigencia_desde_fecha(silice_fecha_vence)
                vigilancia.silice_obs = request.POST.get('silice_obs')
                update_fields.extend([
                    'silice_fecha_radio',
                    'silice_fecha_vence',
                    'silice_vigencia',
                    'silice_obs',
                ])
            elif agente_actual == 'ruido':
                ruido_fecha_audio = request.POST.get('ruido_fecha_audio') or None
                ruido_fecha_vence = request.POST.get('ruido_fecha_vence') or None
                vigilancia.ruido_fecha_audio = ruido_fecha_audio
                vigilancia.ruido_fecha_vence = ruido_fecha_vence
                vigilancia.ruido_vigencia = _calcular_vigencia_desde_fecha(ruido_fecha_vence)
                vigilancia.ruido_obs = request.POST.get('ruido_obs')
                update_fields.extend([
                    'ruido_fecha_audio',
                    'ruido_fecha_vence',
                    'ruido_vigencia',
                    'ruido_obs',
                ])
            elif agente_actual == 'hipobaria':
                hipo_fecha_hemo = request.POST.get('hipo_fecha_hemo') or None
                hipo_fecha_vence = request.POST.get('hipo_fecha_vence') or None
                vigilancia.hipo_fecha_hemo = hipo_fecha_hemo
                vigilancia.hipo_fecha_vence = hipo_fecha_vence
                vigilancia.hipo_vigencia = _calcular_vigencia_desde_fecha(hipo_fecha_vence)
                vigilancia.hipo_obs = request.POST.get('hipo_obs')
                update_fields.extend([
                    'hipo_fecha_hemo',
                    'hipo_fecha_vence',
                    'hipo_vigencia',
                    'hipo_obs',
                ])
            elif agente_actual == 'vibraciones':
                vibra_fecha_vence = request.POST.get('vibra_fecha_vence') or None
                vigilancia.vibra_fecha_vence = vibra_fecha_vence
                vigilancia.vibra_vigencia = _calcular_vigencia_desde_fecha(vibra_fecha_vence)
                update_fields.extend([
                    'vibra_fecha_vence',
                    'vibra_vigencia',
                ])
            elif agente_actual == 'radiaciones':
                rad_fecha_vence = request.POST.get('rad_fecha_vence') or None
                vigilancia.rad_fecha_vence = rad_fecha_vence
                vigilancia.rad_vigencia = _calcular_vigencia_desde_fecha(rad_fecha_vence)
                update_fields.extend([
                    'rad_fecha_vence',
                    'rad_vigencia',
                ])
            elif agente_actual == 'humos':
                humos_fecha_vence = request.POST.get('humos_fecha_vence') or None
                vigilancia.humos_fecha_vence = humos_fecha_vence
                vigilancia.humos_vigencia = _calcular_vigencia_desde_fecha(humos_fecha_vence)
                update_fields.extend([
                    'humos_fecha_vence',
                    'humos_vigencia',
                ])

            archivos_documento = request.FILES.getlist('documento')
            if not update_fields and not archivos_documento:
                messages.info(request, "No se detectaron cambios para guardar.")
                return redirect('view_vigilancia', pk=pk)

            with transaction.atomic():
                if update_fields:
                    vigilancia.save(update_fields=list(dict.fromkeys(update_fields)))
                if archivos_documento:
                    _reemplazar_archivo_vigilancia(vigilancia, archivos_documento[0])

            messages.success(request, "Registro actualizado correctamente.")
            if vigilancia.vigilancia_id:
                return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.vigilancia_id)
            return redirect('manage_vigilancia')
        except Exception as e:
            messages.error(request, f"No se pudo actualizar el registro: {e}")
            return redirect('view_vigilancia', pk=pk)

    context = {
        'vigilancia': vigilancia,
        'agente_actual': agente_actual,
        'archivos_adjuntos': _serializar_adjuntos_vigilancia(vigilancia),
        'back_url': (
            reverse('vigilancia_trabajadores', args=[vigilancia.vigilancia_id])
            if vigilancia.vigilancia_id else reverse('manage_vigilancia')
        ),
        'sidebarmain': 'prevencion_riesgo', 
        'sidebar': 'vigilancia_medica'
    }
    return render(request, 'pages/prevencion/view_vigilancia.html', context)


@login_required
def view_vigilancia_egreso(request, pk):
    # "Ver" en egreso: reutiliza plantilla de vigilancia en modo lectura con actualizacion controlada.
    egresos = _egresos_visibles_para_usuario(request.user)
    vigilancia = egresos.filter(pk=pk).first()
    if vigilancia is None:
        # Compatibilidad: algunos flujos antiguos envian vigilancia_id en esta URL.
        vigilancia = egresos.filter(vigilancia_id=pk).order_by('-fecha_egreso', '-id').first()
    if vigilancia is None:
        raise Http404("No Egreso matches the given query.")
    agente_actual = (vigilancia.cuantitativa.agente if vigilancia.cuantitativa else '')

    if request.method == 'POST':
        archivos_validos, msg_archivos = _validar_archivos_subidos(request)
        if not archivos_validos:
            messages.error(request, msg_archivos)
            return redirect('view_vigilancia_egreso', pk=vigilancia.id)

        fecha_incidente = request.POST.get('fecha_incidente') or None
        archivos_documento = request.FILES.getlist('documento')
        if len(archivos_documento) > 1:
            messages.error(request, "Solo puedes subir 1 archivo en egreso.")
            return redirect('view_vigilancia_egreso', pk=vigilancia.id)

        with transaction.atomic():
            vigilancia.fecha_incidente = fecha_incidente
            vigilancia.save(update_fields=['fecha_incidente'])
            if archivos_documento:
                _reemplazar_archivo_egreso(vigilancia, archivos_documento[0])

        messages.success(request, "Datos de egreso actualizados correctamente.")
        return redirect('view_vigilancia_egreso', pk=vigilancia.id)

    archivos_egreso = _serializar_archivos_egreso(vigilancia)
    archivo_actual = _archivo_actual_egreso(vigilancia)

    context = {
        'vigilancia': vigilancia,
        'agente_actual': agente_actual,
        'archivos_egreso': archivos_egreso,
        'archivos_historial': archivos_egreso,
        'archivo_actual': archivo_actual,
        'solo_lectura': True,
        'permitir_subir_archivo_egreso': True,
        'back_url': (
            reverse('egreso_detalle', args=[vigilancia.vigilancia_id])
            if vigilancia.vigilancia_id else reverse('egreso')
        ),
        'sidebarmain': 'prevencion_riesgo',
        'sidebar': 'egreso'
    }
    return render(request, 'pages/prevencion/view_vigilancia.html', context)


@login_required
def status_vigilancia(request, pk):
    if request.method != 'POST':
        return redirect('manage_vigilancia')

    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.select_related('cuantitativa', 'faena'),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)
    if vigilancia.egresado:
        return _redireccion_egreso_desde_ficha(
            request,
            vigilancia,
            mensaje="No se puede cambiar el estado de una ficha egresada.",
        )
    vigilancia.status = not vigilancia.status
    vigilancia.save(update_fields=['status'])

    accion = "habilitada" if vigilancia.status else "deshabilitada"
    messages.success(request, f"Ficha {accion} correctamente.")
    user_name = request.user.get_full_name() or request.user.username
    notify_group('prevencion_general', f'Ficha vigilancia {accion}: {vigilancia.nombre_completo}', f'Vigilancia Médica: {vigilancia.nombre_completo}\n{accion.capitalize()} por: {user_name}')
    if vigilancia.vigilancia_id:
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.vigilancia_id)
    return redirect('manage_vigilancia')


@login_required
def delete_vigilancia(request, pk):
    if request.method != 'POST':
        return redirect('manage_vigilancia')

    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.select_related('cuantitativa', 'faena'),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)
    if vigilancia.egresado:
        return _redireccion_egreso_desde_ficha(
            request,
            vigilancia,
            mensaje="No se puede eliminar una ficha egresada.",
        )

    if vigilancia.status:
        vigilancia.status = False
        vigilancia.save(update_fields=['status'])
        user_name = request.user.get_full_name() or request.user.username
        notify_group('prevencion_general', f'Ficha vigilancia eliminada: {vigilancia.nombre_completo}', f'Vigilancia Médica: {vigilancia.nombre_completo}\nEliminada por: {user_name}')
        messages.success(request, "Ficha eliminada correctamente.")
    else:
        messages.info(request, "La ficha ya estaba eliminada.")

    if vigilancia.vigilancia_id:
        return redirect('vigilancia_trabajadores', vigilancia_id=vigilancia.vigilancia_id)
    return redirect('manage_vigilancia')


@login_required
@prevencion_riesgo_required
def vigilancia_pdf_view(request, pk):
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.prefetch_related('adjuntos_vigilancia'),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)
    # Prepara cada adjunto con metadatos de visualización en PDF (imagen/pdf/otro).
    archivos_adjuntos_raw = _serializar_adjuntos_vigilancia(vigilancia)
    archivos_adjuntos = []
    for archivo in archivos_adjuntos_raw:
        nombre = str((archivo or {}).get('name', '') or '')
        url = str((archivo or {}).get('url', '') or '')
        nombre_lower = nombre.lower()
        url_lower = url.lower().split('?', 1)[0]
        registro = {
            'id': (archivo or {}).get('id'),
            'name': nombre,
            'url': url,
            'paginas_pdf': [],
        }

        if nombre_lower.endswith(('.png', '.jpg', '.jpeg', '.gif', '.webp')) or url_lower.endswith(('.png', '.jpg', '.jpeg', '.gif', '.webp')):
            registro['tipo'] = 'imagen'
        elif nombre_lower.endswith('.pdf') or url_lower.endswith('.pdf'):
            # Si el adjunto es PDF, lo convierte a imágenes para respetar el orden de adjuntos.
            registro['tipo'] = 'pdf'
            ruta_local = _resolver_ruta_local_adjunto_pdf(url)
            if ruta_local:
                try:
                    registro['paginas_pdf'] = check_and_convert_pdf(ruta_local) or []
                except Exception:
                    registro['paginas_pdf'] = []
        else:
            registro['tipo'] = 'otro'

        archivos_adjuntos.append(registro)

    template = get_template('pages/pdfs/vigilancia_medica_pdf.html')
    html = template.render({
        'vigilancia': vigilancia,
        'archivos_adjuntos': archivos_adjuntos,
        'fecha_descarga': timezone.localtime(timezone.now()),
    })

    filename = f"vigilancia-medica-{vigilancia.id}.pdf"
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename=\"{filename}\"'
    pisa_status = pisa.CreatePDF(html, dest=response, link_callback=_resolver_src_pdf)

    if pisa_status.err:
        return HttpResponse('Error al generar el PDF.', status=500)

    return response

@login_required
def ajax_search_user_by_rut(request):
    rut_raw = request.GET.get('rut', '').strip()
    if not rut_raw: return JsonResponse({'found': False})
    
    rut_limpio = rut_raw.replace('.', '').replace('-', '').upper()
    rut_guion = f"{rut_limpio[:-1]}-{rut_limpio[-1]}" if len(rut_limpio) > 1 else ""

    try:
        user = User.objects.filter(Q(username__iexact=rut_limpio) | Q(username__iexact=rut_guion) | Q(username__iexact=rut_raw)).first()
        if user:
            nombre = f"{user.first_name} {user.last_name}".strip()
            cargo = ''
            genero_id = ''
            faena_id = ''
            fecha_nacimiento = ''
            area = ''
            ges = ''
            tipo_contrato = ''
            fecha_ingreso = ''
            fecha_retiro = ''
            
            if hasattr(user, 'usuarioprofile'):
                p = user.usuarioprofile
                if p.genero: genero_id = p.genero.id
                if p.faena and _faena_permitida_para_usuario(request.user, p.faena.id):
                    faena_id = p.faena.id
                if p.fechaNacimiento:
                    fecha_nacimiento = p.fechaNacimiento.strftime('%Y-%m-%d')

            # Trae datos laborales desde la nueva tabla user_informacion_laboral.
            info_laboral = UserInformacionLaboral.objects.filter(user=user).select_related(
                'area',
                'ges',
                'cargo',
                'tipo_contrato'
            ).first()
            if info_laboral:
                area = info_laboral.area.valor if info_laboral.area else ''
                ges = info_laboral.ges.valor if info_laboral.ges else ''
                cargo = info_laboral.cargo.valor if info_laboral.cargo else ''
                tipo_contrato = info_laboral.tipo_contrato.valor if info_laboral.tipo_contrato else ''
                fecha_ingreso = info_laboral.fechaIngreso.strftime('%Y-%m-%d') if info_laboral.fechaIngreso else ''
                fecha_retiro = (
                    info_laboral.fechaDesvinculacion.strftime('%Y-%m-%d')
                    if info_laboral.fechaDesvinculacion else ''
                )

            data = {
                'found': True,
                'username': user.username,
                'nombre': nombre,
                'cargo': cargo,
                'genero_id': genero_id,
                'faena_id': faena_id,
                'fecha_nacimiento': fecha_nacimiento,
                'area': area,
                'ges': ges,
                'tipo_contrato': tipo_contrato,
                'fecha_ingreso': fecha_ingreso,
                'fecha_retiro': fecha_retiro,
            }
        else:
            data = {'found': False}
    except Exception as e:
        data = {'found': False, 'error': str(e)}
    return JsonResponse(data)

@login_required
@prevencion_riesgo_required
def delete_vigilancia_doc(request, pk):
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.all(),
        request.user
    )
    vigilancia = get_object_or_404(vigilancias, pk=pk)
    if vigilancia.egresado:
        return _redireccion_egreso_desde_ficha(
            request,
            vigilancia,
            mensaje="No se puede modificar una ficha egresada.",
        )
    
    if vigilancia.documento_adjunto:
        # Borra el archivo del sistema de archivos y limpia el campo
        vigilancia.documento_adjunto.delete(save=False)
        vigilancia.documento_adjunto = None
        vigilancia.save()
        messages.success(request, "Archivo adjunto eliminado correctamente.")
        
    return redirect('edit_vigilancia', pk=pk)


@login_required
@prevencion_riesgo_required
def delete_vigilancia_adjunto(request, pk):
    vigilancias = _filtrar_queryset_por_faena_usuario(
        VigilanciaMedica.objects.all(),
        request.user
    )
    adjunto = get_object_or_404(
        VigilanciaAdjunto.objects.select_related('vigilancia'),
        pk=pk,
        vigilancia__in=vigilancias
    )

    if adjunto.vigilancia.egresado:
        return _redireccion_egreso_desde_ficha(
            request,
            adjunto.vigilancia,
            mensaje="No se puede modificar una ficha egresada.",
        )

    vigilancia_id = adjunto.vigilancia_id
    adjunto.archivo.delete(save=False)
    adjunto.delete()
    messages.success(request, "Archivo adjunto eliminado correctamente.")
    return redirect('edit_vigilancia', pk=vigilancia_id)

def limpiar_rut(rut_str):
    """
    Función auxiliar para normalizar el RUT.
    Quita puntos, guiones, espacios y lo deja en mayúsculas.
    Ejemplo: "12.345.678-k" -> "12345678K"
    """
    if not rut_str:
        return ""
    return str(rut_str).replace(".", "").replace("-", "").replace(" ", "").strip().upper()


@login_required
def generar_qr_validacion(request, doc_id):
    rut_usuario = request.user.username 
    
    validacion = ValidacionDobleFactor.objects.create(
        usuario=request.user,
        documento_id=doc_id,
        rut_esperado=rut_usuario
    )
    url_movil = request.build_absolute_uri(f'/prevencion/escanear-carnet/{validacion.token}/')
    
    return JsonResponse({'token': str(validacion.token), 'url_qr': url_movil})


def chequear_estado_validacion(request, token):
    validacion = get_object_or_404(ValidacionDobleFactor, token=token)
    
    if validacion.rut_esperado == "FAILED":
        return JsonResponse({'validado': False, 'rechazado': True})
        
    return JsonResponse({'validado': validacion.validado, 'rechazado': False})

def vista_escaner_movil(request, token):
    validacion = get_object_or_404(ValidacionDobleFactor, token=token)
    
    estilo_movil = (
        "<meta name='viewport' content='width=device-width, initial-scale=1.0, user-scalable=no'>"
        "<style>"
        "  body { font-family: sans-serif; background-color: #1a1a1a; color: #ffffff; margin: 0; padding: 0; }"
        "  .container { display: flex; flex-direction: column; justify-content: center; align-items: center; min-height: 100vh; padding: 24px; box-sizing: border-box; text-align: center; }"
        "  h2 { font-size: 2.2rem; margin-bottom: 16px; line-height: 1.2; }"
        "  p { font-size: 1.3rem; opacity: 0.85; line-height: 1.5; margin: 0; }"
        "</style>"
    )

    if validacion.validado or validacion.rut_esperado == "FAILED":
        return HttpResponseForbidden(
            f"{estilo_movil}"
            "<div class='container'>"
            "<h2 style='color: #ffc107;'>Código Expirado</h2>"
            "<p>Este código QR ya caducó o fue utilizado.</p>"
            "</div>"
        )
        
    session_key = f"qr_locked_{token}"

    if validacion.en_uso:
        if request.session.get(session_key) != True:
            return HttpResponseForbidden(
                f"{estilo_movil}"
                "<div class='container'>"
                "<h2 style='color: #dc3545;'>Acceso Denegado</h2>"
                "<p>Este código QR ya está siendo abierto por otro dispositivo.</p>"
                "</div>"
            )
    else:
        validacion.en_uso = True
        validacion.save(update_fields=['en_uso'])
        request.session[session_key] = True
        request.session.save()

    return render(request, 'pages/prevencion/escaner_carnet.html', {'token': validacion.token})

@csrf_exempt
def procesar_carnet_movil(request, token):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            rut_escaneado_raw = data.get('rut', '')
            validacion = get_object_or_404(ValidacionDobleFactor, token=token)

            rut_escaneado = limpiar_rut(rut_escaneado_raw)
            rut_esperado = limpiar_rut(validacion.rut_esperado)

            if rut_escaneado == rut_esperado:
                validacion.validado = True
                validacion.save()
                
                difusion = PrevencionDocumentoDifusionTrabajador.objects.filter(
                    documento_id=validacion.documento_id,
                    user=validacion.usuario
                ).first()

                if difusion and difusion.motivo_rechazo:
                    difusion.motivo_rechazo = None
                    difusion.save()
                
                return JsonResponse({'status': 'success', 'message': 'Identidad verificada.'})
            else:
                validacion.rut_esperado = "FAILED"
                validacion.save()

                difusion = PrevencionDocumentoDifusionTrabajador.objects.filter(
                    documento_id=validacion.documento_id,
                    user=validacion.usuario
                ).first()

                if difusion:
                    difusion.motivo_rechazo = "Identidad rechazada. El carnet escaneado no coincide con el usuario autenticado."
                    difusion.save()
                
                return JsonResponse({
                    'status': 'error',
                    'message': 'El RUT del carnet no coincide con el usuario autenticado.'
                }, status=400)

        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=500)
            
    return JsonResponse({'status': 'error', 'message': 'Método no permitido.'}, status=405)

@login_required
@prevencion_riesgo_required
def reiniciar_validacion_trabajador(request, documento_id, user_id):
    if request.method == 'POST':
        try:
            difusion = PrevencionDocumentoDifusionTrabajador.objects.filter(
                documento_id=documento_id,
                user_id=user_id
            ).first()
            
            if difusion:
                difusion.motivo_rechazo = None
                difusion.save(update_fields=['motivo_rechazo', 'updated_at'])
                return JsonResponse({'status': 'success', 'message': 'Validación reiniciada correctamente.'})
            
            return JsonResponse({'status': 'error', 'message': 'Registro de trabajador no encontrado.'}, status=404)
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=500)
            
    return JsonResponse({'status': 'error', 'message': 'Método no permitido.'}, status=405)

@login_required
@prevencion_riesgo_required
def reiniciar_curso_intentos(request, documento_id, user_id):
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': 'Método no permitido.'}, status=405)

    documentos = _filtrar_queryset_por_faena_usuario(
        PrevencionDocumento.objects.filter(status=True).select_related('faena', 'plantilla_base'),
        request.user
    )
    documento = get_object_or_404(documentos, pk=documento_id)

    tipos_permitidos = [
        PrevencionPlantilla.TipoRiesgo.CURSOS.value,
        PrevencionPlantilla.TipoRiesgo.DIFUSIONES.value,
        PrevencionPlantilla.TipoRiesgo.INFORMATIVOS.value,
        PrevencionPlantilla.TipoRiesgo.CHARLAS.value,
    ]

    tipo_documento = str(getattr(getattr(documento, 'plantilla_base', None), 'tipo', '') or '').strip().lower()

    if tipo_documento not in tipos_permitidos:
        return JsonResponse({'status': 'error', 'message': 'El documento no permite reiniciar intentos.'}, status=400)

    _reiniciar_resultado_curso_trabajador(documento, user_id)
    PrevencionResultadoCurso.objects.filter(documento=documento, trabajador_id=user_id).delete()
    ValidacionDobleFactor.objects.filter(documento_id=documento_id, usuario_id=user_id).delete()

    difusion = PrevencionDocumentoDifusionTrabajador.objects.filter(documento_id=documento_id, user_id=user_id).first()
    if difusion:
        difusion.motivo_rechazo = None
        difusion.save()

    return JsonResponse({'status': 'success', 'message': 'Curso reiniciado correctamente.'})