from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill, NamedStyle
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import locale
import pytz
import requests
from pathlib import Path
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn
from collections import defaultdict
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule
from pprint import pprint
from openpyxl import load_workbook
from ..reportes import helpers
import copy
import re

def _safe_sheet_title(title: str) -> str:

    if title is None:
        title = "SIN_DATO"
    title = re.sub(r"[:\\/?*\[\]]", "-", str(title))
    title = re.sub(r"\s+", " ", title.strip())
    if len(title) > 31:
        title = title[:31]
    return title


def _campana_nombre(campanas, campana_val):

    if campana_val in (None, "", 0, "0"):
        return "SIN DATO"

    if isinstance(campana_val, str) and not campana_val.strip().isdigit():
        return campana_val.strip()

    try:
        campana_id = int(campana_val)
    except Exception:
        return str(campana_val)

    for c in campanas:
        if c.get("id") == campana_id:
            return (c.get("campana") or c.get("nombre") or f"Campaña {campana_id}").strip()

    return f"Campaña {campana_id}"


def _build_meta_por_reco_id(all_recomendaciones_final):

    meta = {}

    if not all_recomendaciones_final:
        return meta

    for r in all_recomendaciones_final:

        rid = (
            r.get("id") or
            r.get("ID") or
            r.get("recomendacion") or
            r.get("recomendacion_id") or
            r.get("recomendacionFinal")  
        )
        try:
            rid_int = int(rid)
        except Exception:
            continue

        campana_id = (
            r.get("campana") or
            r.get("campana_id") or
            r.get("CAMPANA")
        )

        programa = r.get("programa") or r.get("PROGRAMA")
        programa_nombre = None
        if isinstance(programa, dict):
            programa_nombre = programa.get("programa") or programa.get("nombre")
        elif programa:
            programa_nombre = str(programa)

        meta[rid_int] = {
            "campana_id": campana_id,
            "programa_nombre": programa_nombre,
        }

    return meta


def fusionar_celdas_con_formato(hoja, celda_inicio, celda_fin, valor, color_fondo="FFC000"):
    rango = f"{celda_inicio}:{celda_fin}"
    hoja.merge_cells(rango)

    celda = hoja[celda_inicio]
    celda.value = valor

    celda.fill = PatternFill(start_color=color_fondo, end_color=color_fondo, fill_type="solid")
    celda.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    celda.alignment = Alignment(horizontal="center", vertical="center")

    return hoja


def agregar_datos_a_celdas(hoja, inicio_celda, datos, primer_color, segundo_color):
    fila_inicio = int(inicio_celda[1:])

    color_fondo_1 = PatternFill(start_color=primer_color, end_color=primer_color, fill_type="solid")
    color_fondo_2 = PatternFill(start_color=segundo_color, end_color=segundo_color, fill_type="solid")

    ultima_fila = fila_inicio

    mapeo_columnas = {
        "ID": "B",
        "Tipo de Sondaje": "C",
        "Sondaje": "D",
        "Sector": "E",
        "Este": "F",
        "Norte": "G",
        "Cota": "H",
        "Azimut": "I",
        "Inclinación": "J",
        "Largo (m)": "K",
        "Fecha Inicio": "L",
        "Fecha Termino": "M",
        "Por Perforar (m)": "N",
        "Avance Actual (m)": "O",
        "Mts. Faltantes": "P",
        "%Avance": "Q",
        "Estatus Perforación (m)": "R",
        "Largo Final (m)": "S",
        "Certificado Collar": "T",
        "Fecha Medición de Trayectoria": "U",
        "Observación": "V"
    }

    for idx, diccionario in enumerate(datos):
        fila_actual = fila_inicio + idx
        ultima_fila = fila_actual
        color_fondo = color_fondo_1 if idx % 2 == 0 else color_fondo_2

        for key, value in diccionario.items():
            if key in mapeo_columnas:
                columna = mapeo_columnas[key]
                celda = hoja[f"{columna}{fila_actual}"]
                celda.value = value
                celda.font = Font(name="Arial", size=8, bold=False)
                celda.alignment = Alignment(horizontal="center", vertical="center")
                celda.fill = color_fondo

    return hoja, ultima_fila


def obtener_collar(id_recomendacion, all_recomendaciones_final):
    for rec in all_recomendaciones_final:
        if rec.get('recomendacionFinal') == id_recomendacion:
            return "SI"
    return "NO"

def construir_data(reportes_agrupados, all_recomendaciones_final, campanas):
    opciones_estado = {
        '1': 'Abortado',
        '2': 'En Avance',
        '3': 'En Espera',
        '4': 'Finalizado',
    }

    meta_por_reco_id = _build_meta_por_reco_id(all_recomendaciones_final)

    resultado = defaultdict(list)

    for reporte, detalles in reportes_agrupados.items():
        max_id = max(d["id"] for d in detalles)
        ultimo_registro = next(d for d in detalles if d["id"] == max_id)

        if 'detalle_perforaciones' not in ultimo_registro:
            continue

        rec = ultimo_registro.get('recomendacion') or {}
        programa = rec.get('programa')
        campana_id = rec.get('campana') 

        rec_id = rec.get('id')
        rec_id_int = None
        try:
            rec_id_int = int(rec_id)
        except Exception:
            rec_id_int = None

        if (campana_id is None or campana_id == "" or campana_id == 0) and rec_id_int and rec_id_int in meta_por_reco_id:
            campana_id = meta_por_reco_id[rec_id_int].get("campana_id") or campana_id
            if not programa:
                programa = meta_por_reco_id[rec_id_int].get("programa_nombre") or programa

        if not programa:
            programa = "SIN DATO"

        campana_nombre = _campana_nombre(campanas, campana_id)
        key = (str(programa), int(campana_id) if str(campana_id).isdigit() else 0, str(campana_nombre))

        ultimo_detalle_perforacion = ultimo_registro['detalle_perforaciones'][-1]

        try:
            hasta = float(ultimo_detalle_perforacion.get('HASTA') or 0)
        except Exception:
            hasta = 0.0

        try:
            largo_programado = float(rec.get('largo_programado') or 0)
        except Exception:
            largo_programado = 0.0

        id_estado = rec.get('estado')
        if id_estado:
            estado = opciones_estado.get(str(id_estado), "SIN ESTADO")
            if estado == "Finalizado":
                fecha_termino = rec.get('fechaupdateestado') or rec.get('fechaUpdateEstado') or " "
            else:
                fecha_termino = " "
        else:
            estado = "SIN ESTADO"
            fecha_termino = " "

        collar = obtener_collar(rec.get('id'), all_recomendaciones_final)

        if largo_programado and largo_programado != 0:
            porcentaje = round((hasta / largo_programado) * 100, 2)
        else:
            porcentaje = 0.0

        resultado[key].append({
            "ID": str(programa), 
            "Tipo de Sondaje": "SIN DATO",
            "Sondaje": rec.get('pozo') or "SIN DATO",
            "Sector": rec.get('sector') or "SIN DATO",
            "Este": rec.get('este') or 1,
            "Norte": rec.get('norte') or 1,
            "Cota": rec.get('cota') or 1,
            "Azimut": rec.get('azimut') or 360,
            "Inclinación": rec.get('inclinacion') or 1,
            "Largo (m)": largo_programado,
            "Fecha Inicio": rec.get('fecha_inicio') or " ",
            "Fecha Termino": fecha_termino,
            "Por Perforar (m)": (largo_programado - hasta),
            "Avance Actual (m)": hasta,
            "Mts. Faltantes": (largo_programado - hasta),
            "%Avance": porcentaje,
            "Estatus Perforación (m)": estado,
            "Largo Final (m)": rec.get('largo_real') or largo_programado or 1,
            "Certificado Collar": collar,
            "Fecha Medición de Trayectoria": "SIN DATO",
            "Observación": "SIN DATO",
        })

    return resultado

def detalle_campanas(campanas, campana_actual):
    anio_inicial = "_SIN DATO"
    anio_final = "_SIN DATO"

    for c in campanas:
        if c['id'] == campana_actual:
            anio_inicial = f"_{c['anoInicial']}"
            anio_final = f"_{c['anoFinal']}"
            return anio_inicial, anio_final

    return anio_inicial, anio_final


def detalle_programas(programas, programa_actual, campanas):
    anio_inicial = "_SIN DATO"
    anio_final = "_SIN DATO"
    for p in programas:
        if p['programa'].upper() == programa_actual.upper():
            anio_inicial, anio_final = detalle_campanas(campanas, p['campana'])
            return anio_inicial, anio_final

    return anio_inicial, anio_final

def run(libro, reportes_agrupados, programas, campanas, all_recomendaciones_final):

    data_por_programa_y_campana = construir_data(reportes_agrupados, all_recomendaciones_final, campanas)

    lista_titulos_esperados = []


    anio_inicial_avance = 0
    anio_final_avance = 0

    for p in programas:
        anio_inicial, anio_final = detalle_campanas(campanas, p["campana"])

        if anio_inicial_avance == 0:
            anio_inicial_avance = anio_inicial
            anio_final_avance = anio_final

        if anio_inicial < anio_inicial_avance:
            anio_inicial_avance = anio_inicial

        if anio_final > anio_final_avance:
            anio_final_avance = anio_final

        campana_nom = _campana_nombre(campanas, p["campana"])
        titulo = _safe_sheet_title(f"{p['programa']} ({campana_nom})")
        lista_titulos_esperados.append(titulo)

    seen = set()
    lista_titulos_esperados = [x for x in lista_titulos_esperados if not (x in seen or seen.add(x))]

    titulos = [
        "ID", "Tipo de Sondaje", "Sondaje", "Sector", "Este", "Norte", "Cota", "Azimut",
        "Inclinación", "Largo (m)", "Fecha Inicio", "Fecha Termino", "Por Perforar (m)",
        "Avance Actual (m)", "Mts. Faltantes", "%Avance", "Estatus Perforación (m)",
        "Largo Final (m)", "Certificado Collar", "Fecha Medición de Trayectoria", "Observación"
    ]

    for (programa, campana_id, campana_nombre), filas in data_por_programa_y_campana.items():
        titulo = _safe_sheet_title(f"{programa} ({campana_nombre})")

        if titulo in lista_titulos_esperados:
            lista_titulos_esperados.remove(titulo)

        hoja = libro.create_sheet(title=titulo)

        hoja = helpers.obtener_fecha_documento(hoja, "B1", "B1:E1")
        hoja = fusionar_celdas_con_formato(hoja, "F3", "J3", "COORDENADAS", "70ad47")
        hoja = helpers.agregar_titulos(hoja, "B", 4, titulos, "70ad47", 3, 12)

        hoja, ultima_fila = agregar_datos_a_celdas(hoja, "B5", filas, "e2efd9", "b4c6e7")

        hoja.column_dimensions["A"].width = 1
        hoja.column_dimensions["B"].width = 13
        for col in "CDEFGHIJKLMNOPQRSTU":
            hoja.column_dimensions[col].width = 8
        hoja.column_dimensions["V"].width = 20

        libro._sheets.insert(0, libro._sheets.pop(-1))

    for titulo_esperado in lista_titulos_esperados:
        hoja = libro.create_sheet(title=titulo_esperado)

        hoja = helpers.obtener_fecha_documento(hoja, "B1", "B1:E1")
        hoja = fusionar_celdas_con_formato(hoja, "F3", "J3", "COORDENADAS", "70ad47")
        hoja = helpers.agregar_titulos(hoja, "B", 4, titulos, "70ad47", 3, 12)

        hoja.column_dimensions["A"].width = 1
        hoja.column_dimensions["B"].width = 13
        for col in "CDEFGHIJKLMNOPQRSTU":
            hoja.column_dimensions[col].width = 8
        hoja.column_dimensions["V"].width = 20

        libro._sheets.insert(0, libro._sheets.pop(-1))

    return libro, anio_inicial_avance, anio_final_avance
